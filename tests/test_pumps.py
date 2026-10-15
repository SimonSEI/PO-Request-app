"""
Tests for the Pumps app (pumps.py, pump_reports.py).

    python -m unittest tests.test_pumps -v

Runs against a throwaway database. Jobber and Claude are never called: Jobber
is replaced by a fake that records every request, and documents are "read"
from canned text. All names below are made up.
"""
import atexit
import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
import zipfile
import re
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix='pumps_test_')
os.environ.update(SECRET_KEY='test-secret-key-0123456789abcdef-pumps', DATA_DIR=TMP, PUMPS_AUTO_SCAN='false',
                  OPENCLAW_API_KEY='test-openclaw-key', WEBSITE_URL='https://office.example.com')
os.environ.pop('ANTHROPIC_API_KEY', None)
os.environ.pop('JOBBER_API_TOKEN', None)

# The preview's sheet rows, to check they become items when the app starts.
_db = sqlite3.connect(os.path.join(TMP, 'po_requests.db'))
_db.execute('''CREATE TABLE IF NOT EXISTS pump_invoices (id INTEGER PRIMARY KEY AUTOINCREMENT, po_number TEXT,
               entry_date TEXT, vendor TEXT, job_name TEXT, description TEXT, approved_by TEXT,
               jobber_request_made TEXT, vendor_quote_amount REAL, vendor_invoice_amount REAL,
               vendor_invoice_number TEXT, sei_invoice_number TEXT, amount REAL, notes TEXT, month_tab TEXT,
               source TEXT, email_uid TEXT, created_at TEXT, updated_at TEXT)''')
_db.execute('''INSERT INTO pump_invoices (po_number, entry_date, vendor, job_name, description, vendor_quote_amount,
               vendor_invoice_amount, vendor_invoice_number) VALUES ('0290', '2026-06-03', 'Wettech',
               'Old Sheet Lakes HOA', 'Replace transducer', 900, 900, '28001')''')
_db.commit()
_db.close()

import app as A  # noqa: E402
import pumps as P  # noqa: E402
import pump_reports as R  # noqa: E402
from docx import Document  # noqa: E402

P.ANTHROPIC_API_KEY = ''
A.app.config['TESTING'] = True


def make_report(customer='Stahlman-Lakeside Pines HOA', location='Pump #2', tech='Casey Rivera'):
    """A Wettech-style Word report, laid out like the real ones."""
    d = Document()
    for line in ('Water Equipment Technologies of Southwest Florida LLC',
                 'State of Florida Certified Plumbing Contractor #CFC1429137',
                 '451 interstate Court - Sarasota, FL  34240',
                 'Phone 941-232-4629   FAX 941-371-5151      EMAIL wettec@verizon.net', ''):
        d.add_paragraph(line)
    d.add_paragraph('IRRIGATION PUMP SYSTEM REPORT')
    t = d.add_table(rows=1, cols=4)
    for cell, text in zip(t.rows[0].cells, ('DATE', '9-12-2026', 'TIME', '0930')):
        cell.text = text
    t = d.add_table(rows=1, cols=4)
    for cell, text in zip(t.rows[0].cells, ('CUSTOMER', customer, 'LOCATION', location)):
        cell.text = text
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = '√', 'CHECK ALL ELECTRICAL CONNECTIONS'
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = '√', 'CLEAN SCREENS'
    d.add_paragraph('COMMENTS:')
    c = d.add_table(rows=1, cols=1)
    p = c.rows[0].cells[0].paragraphs[0]
    p.add_run(f'{tech.split()[0]} found the intake screen clogged. ')
    p.add_run('Wettech recommends a new ')
    p.add_run('check valve. Call 941-232-4629.')
    sig = d.add_table(rows=1, cols=2)
    sig.rows[0].cells[0].text, sig.rows[0].cells[1].text = 'Service Technician Signature', tech
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def extraction(kind, number, po='', client='Lakeside Pines', subtotal=1000.0, tax=65.0, wo='', items=None):
    """What Claude hands back for a quote or bill."""
    items = items if items is not None else [
        {'name': 'Replace check valve', 'description': 'Furnish and install 6" check valve. As per quotation.',
         'quantity': 1, 'unit_price': subtotal, 'amount': subtotal, 'taxable': True, 'is_tax': False},
        {'name': 'Sales Tax 6.5%', 'description': 'Sales Tax 6.5%', 'quantity': None, 'unit_price': None,
         'amount': tax, 'taxable': False, 'is_tax': True}]
    return {'kind': kind, 'vendor': 'Wettech', 'doc_number': number, 'doc_date': '2026-09-20', 'po_number': po,
            'ordered_by': '', 'wo_number': wo, 'quote_reference': '', 'client_name': client, 'site': '',
            'category': 'repair', 'description': 'Replace check valve', 'line_items': items,
            'subtotal': subtotal, 'tax': tax, 'total': round(subtotal + tax, 2), 'notes': ''}


class FakeJobber:
    """Stands in for Jobber: records every query and answers the ones Pumps sends."""

    def __init__(self, invoice_status='draft', quote_status='draft', properties=None, clients=None, quote_lines=None,
                 invoice_notes=False):
        self.calls = []
        self.invoice_notes = invoice_notes
        self.clients = clients
        self.quote_lines = quote_lines or []
        self.invoice_status = invoice_status
        self.quote_status = quote_status
        self.properties = properties if properties is not None else [
            {'id': 'P1', 'address': {'street1': '100 Lakeside Dr', 'city': 'Naples'}}]

    def __call__(self, query, variables=None):
        P.check_mutation_allowed(query)  # the real guard still runs
        if '__type(name' in query and (variables or {}).get('n') == 'NoteAttachmentAttributes':
            return {'__type': {'inputFields': [{'name': n} for n in ('url', 'fileName', 'contentType')]}}
        self.calls.append((query, variables))
        if 'mutationType' in query:
            return {'__schema': {'mutationType': {'fields': [{'name': n, 'type': {}, 'args': []} for n in
                                                             (['invoiceCreateNote'] if self.invoice_notes else [])]}}}
        if 'invoiceCreateNote' in query:
            return {'invoiceCreateNote': {'invoiceNote': {'id': 'IN1'}, 'userErrors': []}}
        if 'invoiceCreate' in query:
            return {'invoiceCreate': {'invoice': {'id': 'INV1', 'invoiceNumber': 5001,
                                                  'invoiceStatus': self.invoice_status,
                                                  'jobberWebUri': 'https://secure.getjobber.com/invoices/1'},
                                      'userErrors': []}}
        if 'quoteCreate(' in query:
            return {'quoteCreate': {'quote': {'id': 'Q1', 'quoteNumber': 812, 'quoteStatus': self.quote_status,
                                              'jobberWebUri': 'https://secure.getjobber.com/quotes/1'},
                                    'userErrors': []}}
        if 'quote(id' in query:
            return {'quote': {'quoteNumber': 812, 'lineItems': {'nodes': self.quote_lines}}}
        if 'quoteCreateNote' in query:
            return {'quoteCreateNote': {'quoteNote': {'id': 'QN1'}, 'userErrors': []}}
        if 'client(id' in query:
            return {'client': {'id': variables['id'], 'name': 'Lakeside Pines HOA',
                               'properties': self.properties, 'jobs': {'nodes': []}}}
        if 'jobCreateNote' in query:
            if 'attachments' in json.dumps(variables):
                raise P.JobberError("Argument 'attachments' is invalid")
            return {'jobCreateNote': {'jobNote': {'id': 'NOTE1'}, 'userErrors': []}}
        if 'clients(' in query and self.clients is not None:
            return {'clients': {'nodes': self.clients}}
        if 'clients(' in query:
            return {'clients': {'nodes': [
                {'id': 'C1', 'name': 'Lakeside Pines HOA c/o Example Management', 'isLead': False, 'isArchived': False},
                {'id': 'C2', 'name': 'Lakeside Pines - Work Orders', 'isLead': True, 'isArchived': False}]}}
        return {}


class PumpsTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.texts = {}
        cls.extracts = {}
        cls._orig = (P.extract_text, P._claude_extract, P.jobber_gql)
        real_extract = P.extract_text

        def fake_extract(name, data):
            if name in cls.texts:
                return cls.texts[name]
            if name in cls.extracts:
                return f'Document text for {name}'
            return real_extract(name, data)
        P.extract_text = fake_extract
        P._claude_extract = lambda text, sender='', subject='', filename='': cls.extracts.get(filename)
        # OpenClaw is switched off for the office (PUMPS_OPENCLAW); these tests still cover what it could do.
        cls._openclaw = P.OPENCLAW_ENABLED
        P.OPENCLAW_ENABLED = True

    @classmethod
    def tearDownClass(cls):
        P.extract_text, P._claude_extract, P.jobber_gql = cls._orig
        P.OPENCLAW_ENABLED = cls._openclaw
        # Removed when the whole run ends: other test files share the same app (and this folder).
        atexit.register(shutil.rmtree, TMP, True)

    def setUp(self):
        self.c = A.app.test_client()
        A.app.config['WTF_CSRF_ENABLED'] = False
        with self.c.session_transaction() as s:
            s['username'], s['role'], s['full_name'] = 'office1', 'office', 'Office Tester'

    def draft(self, doc_id, **body):
        """The office clicks Draft quote."""
        return self.c.post(f'/pumps/api/docs/{doc_id}/quote', json=body).get_json()

    def upload(self, name, data=b'%PDF-1.4 fake', kind=None, case_id=None):
        if not name.endswith('.docx'):
            data = data + name.encode()  # a distinct file per name
        form = {'file': (io.BytesIO(data), name)}
        if kind:
            form['kind'] = kind
        if case_id:
            form['case_id'] = str(case_id)
        r = self.c.post('/pumps/api/docs', data=form, content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        return r.get_json()['results'][0]

    def case(self, cid):
        return self.c.get(f'/pumps/api/cases/{cid}').get_json()['case']

    # ── reports ─────────────────────────────────────────────────────────────
    def test_rebrand_report(self):
        out, info = R.rebrand_report(make_report())
        text = '\n'.join(R.read_report(out)['lines'])
        for gone in ('Water Equipment', 'Wettech', 'CFC1429137', '941-232-4629', 'wettec', 'Casey', 'Rivera',
                     'Stahlman-Lakeside', 'Technician Signature'):
            self.assertNotIn(gone.lower(), text.lower(), gone)
        self.assertIn('STAHLMAN – ENGLAND, INC.', text)
        self.assertIn('IRRIGATION PUMP SYSTEM REPORT', text)
        self.assertIn('Lakeside Pines HOA', text)
        self.assertIn('Our technician found the intake screen clogged', text)
        self.assertIn('Stahlman-England Irrigation recommends a new check valve', text)
        self.assertEqual(info['technicians_removed'], ['Casey Rivera'])
        self.assertEqual(info['leftovers'], [])
        # The letterhead's logo came along.
        self.assertTrue(any('image' in rel.reltype for rel in Document(io.BytesIO(out)).part.rels.values()))

    def test_read_report_fields(self):
        f = R.read_report(make_report())
        self.assertEqual((f['customer'], f['location'], f['date']), ('Lakeside Pines HOA', 'Pump #2', '9-12-2026'))

    def test_report_upload_files_item_and_logs_to_jobber(self):
        res = self.upload('9-12 Lakeside Pines.docx', make_report(customer='Stahlman-Bayview Palms HOA'))
        self.assertEqual(res['kind'], 'report')
        case = self.case(res['case_id'])
        self.assertEqual(case['category'], 'maintenance')
        self.assertEqual(case['client_name'], 'Bayview Palms HOA')
        steps = {s['key']: s for s in case['step_list']}
        self.assertTrue(steps['work_done'].get('at'))
        self.assertTrue(steps['scheduled'].get('at'), 'work done implies it was scheduled')
        self.assertEqual(case['stage'], 'vendor_bill')
        # Branded file downloads with our name on it.
        r = self.c.get(f"/pumps/api/docs/{res['doc_id']}/file?version=branded")
        self.assertEqual(r.status_code, 200)
        self.assertIn('Stahlman-England', r.headers['Content-Disposition'])
        # Log it on a job; Jobber refuses the attachment, so it falls back to a link.
        fake = FakeJobber()
        P.jobber_gql = fake
        r = self.c.post(f"/pumps/api/docs/{res['doc_id']}/report_note",
                        json={'target_type': 'job', 'target_id': 'JOB9'})
        body = r.get_json()
        self.assertTrue(body['success'], body)
        self.assertFalse(body['note']['note_attached_file'])
        note_calls = [v for q, v in fake.calls if 'jobCreateNote' in q]
        self.assertEqual(len(note_calls), 2)
        self.assertNotIn('attachments', note_calls[-1]['input'])
        self.assertIn('Pump report - IRRIGATION PUMP SYSTEM REPORT', note_calls[-1]['input']['message'])
        message = note_calls[-1]['input']['message']
        for gone in ('Casey', 'Rivera', 'Wettech', 'Water Equipment', '941-232-4629', 'wettec', 'Stahlman-Bayview'):
            self.assertNotIn(gone, message)
        self.assertIn('Bayview Palms HOA', message)
        self.assertIn('Our technician found the intake screen clogged', message)
        steps = {s['key']: s for s in self.case(res['case_id'])['step_list']}
        self.assertTrue(steps['report_logged'].get('at'))

    def test_multi_page_report_like_spanish_wells(self):
        """Wettech's Spanish Wells report (Oct 2026): one page per pump, each with Wettech's letterhead,
        "CUSTOMER: Stahlman England", an empty comments box and the technician - rebranded on every page, then
        put on The Lake Club's "Quarterly Pump- 850" job as "October Pump Report"."""
        from docx import Document
        d = Document()
        for pump in ('Pump #1', 'Pump #2', 'Pump #3', 'Pump #5'):
            for t in ('Water Equipment Technologies of Southwest Florida LLC',
                      'State of Florida Certified Plumbing Contractor #CFC1429137',
                      '451 Interstate Court - Sarasota, FL  34240',
                      'Phone 941-232-4629   FAX 941-371-5151      EMAIL wettec@verizon.net', '',
                      '                   IRRIGATION PUMP SYSTEM REPORT'):
                d.add_paragraph(t)
            t = d.add_table(rows=2, cols=4)
            t.rows[0].cells[2].text, t.rows[0].cells[3].text = 'DATE', '10-1-2026'
            for cell, text in zip(t.rows[1].cells, ('CUSTOMER', 'Stahlman England', 'LOCATION', pump)):
                cell.text = text
            t = d.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = '√', 'GREASE MOTOR BEARINGS'
            d.add_paragraph('COMMENTS:')
            d.add_table(rows=1, cols=1)
            t = d.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = 'SERVICE TECHNICIAN', 'Justin Burnett'
            d.add_paragraph('')
        buf = io.BytesIO()
        d.save(buf)
        out, info = R.rebrand_report(buf.getvalue())
        self.assertEqual(info['pages_rebranded'], 4)
        self.assertEqual(info['technicians_removed'], ['Justin Burnett'])
        self.assertEqual(info['leftovers'], [])
        doc = Document(io.BytesIO(out))
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
        for gone in ('Water Equipment', 'wettec', '941-', 'Sarasota', 'CFC1429137', 'Justin', 'Burnett',
                     'Stahlman England', 'CUSTOMER', 'TECHNICIAN', 'COMMENTS'):
            self.assertNotIn(gone.lower(), text.lower(), gone)
        self.assertEqual(text.count('STAHLMAN – ENGLAND, INC.'), 4, 'our letterhead on every page')
        self.assertEqual(sum(1 for p in doc.paragraphs if p.paragraph_format.page_break_before), 3)
        for pump in ('Pump #1', 'Pump #2', 'Pump #3', 'Pump #5'):
            self.assertIn(pump, text)
        self.assertNotIn('   IRRIGATION', '\n'.join(p.text for p in doc.paragraphs))
        # Into the app and onto The Lake Club's job, with no one choosing anything.
        fake = FakeJobber()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            res = P.ingest_document('10-26  Spanish Wells.docx', buf.getvalue(), source='email', actor='email scan',
                                    email={'uid': 'sw-1', 'from': 'simon@stahlman-england.com',
                                           'subject': 'FW: Inv 29099', 'preview': ''})
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        self.assertEqual((res['kind'], res['review']), ('report', ''), res)
        got = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertEqual(got['client_name'], 'Spanish Wells')
        notes = [v for q, v in fake.calls if 'jobCreateNote' in q]
        self.assertTrue(notes, res)
        self.assertEqual(notes[0]['id'], 'Z2lkOi8vSm9iYmVyL0pvYi80MTQyMDAxOA==', 'job #1609 Quarterly Pump- 850')
        self.assertEqual(notes[0]['input']['message'], 'October Pump Report')
        att = notes[0]['input'].get('attachments') or []
        # The PDF where the server can make one (LibreOffice), else the Word file.
        if self.c.get(f"/pumps/api/docs/{res['doc_id']}/file?version=branded_pdf").status_code == 200:
            self.assertEqual((att[0]['fileName'], att[0]['contentType']), ('10-26  Spanish Wells.pdf', 'application/pdf'))
        else:
            self.assertEqual(att[0]['fileName'], '10-26  Spanish Wells.docx')
        # Jobber names the file from the link (it ignores fileName), so the link ends with the name.
        self.assertTrue(att[0]['url'].endswith('/10-26%20%20Spanish%20Wells' + os.path.splitext(att[0]['fileName'])[1]),
                        att[0]['url'])
        self.assertEqual(res['auto_note']['target'], 'The Lake Club #1609 Quarterly Pump- 850')
        # Not in Jobber after all (Lely, Oct 2026: the first try didn't take): the same email sent again tries again.
        conn = P._conn()
        conn.execute("UPDATE pump_docs SET jobber='{}' WHERE id=?", (res['doc_id'],))
        conn.commit()
        conn.close()
        fake.calls.clear()
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            again = P.ingest_document('10-26  Spanish Wells.docx', buf.getvalue(), source='email', actor='email scan',
                                      email={'uid': 'sw-2', 'from': 'simon@stahlman-england.com',
                                             'subject': 'FW: Inv 29099', 'preview': ''})
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        self.assertEqual((again['doc_id'], again.get('skipped')), (res['doc_id'], 'already have this file'), again)
        self.assertTrue(again['auto_note'].get('note_id'), again)
        self.assertTrue([v for q, v in fake.calls if 'jobCreateNote' in q])
        # Kept in the app by year.
        lib = self.c.get('/pumps/api/reports?year=2026').get_json()
        row = [r for r in lib['reports'] if r['id'] == res['doc_id']][0]
        self.assertEqual((row['date'], row['site'], row['jobber']['logged'], row['jobber']['title']),
                         ('2026-10-01', 'Spanish Wells', True, 'October Pump Report'))
        self.assertIn('2026', lib['years'])

    # ── quotes, bills, issues ───────────────────────────────────────────────
    def test_quote_then_matching_bill_checks_itself(self):
        self.extracts['q-0322.pdf'] = extraction('quote', 'Q-77', po='PO322', subtotal=1500)
        self.extracts['b-0322.pdf'] = extraction('bill', '29100', po='0322', subtotal=1500, tax=97.5)
        q = self.upload('q-0322.pdf')
        b = self.upload('b-0322.pdf')
        self.assertEqual(q['case_id'], b['case_id'], 'bill filed by PO number onto the quote\'s item')
        case = self.case(b['case_id'])
        self.assertEqual(case['compare']['state'], 'match')
        steps = {s['key']: s for s in case['step_list']}
        self.assertTrue(steps['bill_checked'].get('at'))
        self.assertEqual([i for i in case['issues'] if not i['resolved_at']], [])

    def test_bill_saying_a_quote_is_needed_puts_it_on_the_watchlist(self):
        """A bill that says more work needs a quote ("bearings worn - quote to
        follow") puts that work on the watchlist, waiting on the vendor's quote."""
        ex = extraction('bill', '29950', po='PO9950', client='Heron Bay', subtotal=300)
        ex['notes'] = 'Bearings on pump #2 are worn - quote to follow for replacement.'
        self.extracts['b-qneed.pdf'] = ex
        b = self.upload('b-qneed.pdf')
        conn = P._conn()
        try:
            new = [dict(r) for r in conn.execute("SELECT * FROM pump_cases WHERE source='bill' AND client_name=?",
                                                 ('Heron Bay',))]
            billed = dict(conn.execute('SELECT * FROM pump_cases WHERE id=?', (b['case_id'],)).fetchone())
            # Filing the same bill again doesn't add a second one.
            P.file_document(conn, b['doc_id'] if 'doc_id' in b else b['id'], 'Pumps')
            again = conn.execute("SELECT COUNT(*) FROM pump_cases WHERE source='bill' AND client_name=?",
                                 ('Heron Bay',)).fetchone()[0]
        finally:
            conn.close()
        self.assertEqual(len(new), 1)
        job = new[0]
        self.assertNotEqual(job['id'], billed['id'])
        self.assertEqual(job['stage'], 'vendor_quote')
        self.assertEqual(job['vendor'], 'Wettech')
        self.assertIn('quote to follow', job['title'])
        self.assertEqual(job['po_number'], '', "the billed job's PO stays with the billed job")
        self.assertEqual(again, 1)
        open_ids = [c['id'] for c in self.c.get('/pumps/api/cases?status=open').get_json()['cases']]
        self.assertIn(job['id'], open_ids)
        events = [e['action'] for e in self.case(b['case_id'])['events']]
        self.assertIn('quote needed', events)

    def test_bill_as_per_quotation_adds_nothing_to_the_watchlist(self):
        self.extracts['b-qnone.pdf'] = extraction('bill', '29951', po='PO9951', client='Osprey Point', subtotal=300)
        self.upload('b-qnone.pdf')
        conn = P._conn()
        try:
            n = conn.execute("SELECT COUNT(*) FROM pump_cases WHERE source='bill' AND client_name='Osprey Point'"
                             ).fetchone()[0]
        finally:
            conn.close()
        self.assertEqual(n, 0)

    def test_draft_invoice_todo_says_what_kind_of_bill_it_is(self):
        """Caymas (Oct 2026): the Draft invoice to-do says what the bill is for -
        the month's maintenance with no quote, quoted work, or no quote on file."""
        maint = extraction('bill', '29132', po='PO9132', client='Caymas Irrigation', subtotal=125, tax=0, items=[
            {'name': 'Maintenance', 'description': 'Service to perform irrigation pump maintenance OCTOBER 2026. '
             'As per quotation.', 'quantity': 1, 'unit_price': 125, 'amount': 125, 'taxable': False,
             'is_tax': False}])
        maint['description'] = 'Irrigation pump maintenance'
        self.extracts['b-kind-maint.pdf'] = maint
        self.extracts['q-kind.pdf'] = extraction('quote', 'Q-9133', po='PO9133', client='Pelican Marsh', subtotal=800)
        self.extracts['b-kind-quoted.pdf'] = extraction('bill', '29133', po='PO9133', client='Pelican Marsh',
                                                        subtotal=800)
        self.extracts['b-kind-none.pdf'] = extraction('bill', '29134', po='PO9134', client='Bay Colony', subtotal=400)
        m = self.upload('b-kind-maint.pdf')
        self.upload('q-kind.pdf')
        q = self.upload('b-kind-quoted.pdf')
        n = self.upload('b-kind-none.pdf')
        bills = {d['id']: d for d in self.c.get('/pumps/api/summary').get_json()['queue']['bills_to_draft']}
        self.assertEqual(bills[m['doc_id']]['bill_kind'], 'October maintenance · no quote')
        self.assertEqual(bills[q['doc_id']]['bill_kind'], 'Quoted work · quote #Q-9133')
        self.assertEqual(bills[n['doc_id']]['bill_kind'], 'No quote on file')
        page = self.c.get('/pumps').get_data(as_text=True)
        self.assertIn('See vendor invoice', page)

    def test_bill_over_quote_is_a_todo_to_check_with_the_vendor(self):
        """The client is invoiced what they approved; a bill that differs from
        Wettech's quote is a to-do to check with Tom, not a hold."""
        self.extracts['q-0400.pdf'] = extraction('quote', 'Q-80', po='PO400', client='Heron Bay', subtotal=2000)
        self.extracts['b-0400.pdf'] = extraction('bill', '29200', po='PO 400', client='Heron Bay', subtotal=2250)
        self.upload('q-0400.pdf')
        b = self.upload('b-0400.pdf')
        case = self.case(b['case_id'])
        self.assertEqual(case['compare']['state'], 'over')
        self.assertEqual([i for i in case['issues'] if not i['resolved_at']], [])
        self.assertTrue({s['key']: s for s in case['step_list']}['bill_checked'].get('at'))
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        t = [t for t in todos if t['link'].get('case_id') == b['case_id']]
        self.assertEqual(len(t), 1)
        self.assertIn('Check with Tom', t[0]['title'])
        self.assertIn('$250.00 more than their quote $2,000.00', t[0]['detail'])
        self.assertIn('tomm@wettec.biz', t[0]['detail'])

    def test_sheet_bill_amount_is_checked(self):
        cid = self.c.post('/pumps/api/sheet/save', json={'month': 'October 2026', 'client_name': 'Egret Run',
                                                          'vendor_quote_amount': '700',
                                                          'vendor_bill_amount': '760'}).get_json()['id']
        case = self.case(cid)
        self.assertEqual(case['compare']['state'], 'over')
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        self.assertTrue(any(t['link'].get('case_id') == cid and 'Check with' in t['title'] for t in todos))

    def test_bill_without_quote_raises_issue(self):
        self.extracts['b-noquote.pdf'] = extraction('bill', '29300', client='Coral Isles', subtotal=500)
        b = self.upload('b-noquote.pdf')
        case = self.case(b['case_id'])
        self.assertEqual(case['compare']['state'], 'no_quote')
        self.assertTrue(any(i['kind'] == 'no_quote' and not i['resolved_at'] for i in case['issues']))

    def test_bill_matched_by_client_name(self):
        self.extracts['q-name.pdf'] = extraction('quote', 'Q-90', client='The Reserve at Glenwood', subtotal=800)
        self.extracts['b-name.pdf'] = extraction('bill', '29400', client='Reserve at Glenwood - Lee County',
                                                 subtotal=800)
        q = self.upload('q-name.pdf')
        b = self.upload('b-name.pdf')
        self.assertEqual(q['case_id'], b['case_id'])

    def test_unread_bill_waits_in_inbox_for_review(self):
        self.texts['scan.pdf'] = ''
        res = self.upload('scan.pdf', kind='bill')
        self.assertIsNone(res['case_id'])
        self.assertIn('enter the details by hand', res['review'])
        self.c.patch(f"/pumps/api/docs/{res['doc_id']}", json={'client_name': 'Palm Cove', 'subtotal': 300,
                                                                'total': 319.5, 'doc_number': '29500', 'file': True})
        doc = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertTrue(doc['case_id'])

    def test_read_again_once_claude_is_back(self):
        # Claude is switched on but down: the basic reader marks both for review and leaves them in the Inbox.
        P.USE_CLAUDE, use = True, P.USE_CLAUDE
        self.addCleanup(setattr, P, 'USE_CLAUDE', use)
        for name in ('Inv_29041_from_Water_Equipment.pdf', 'Inv_29042_from_Water_Equipment.pdf'):
            self.texts[name] = 'INVOICE\nWater Equipment Technologies\nTotal $1,200.00'
        a = self.upload('Inv_29041_from_Water_Equipment.pdf')
        b = self.upload('Inv_29042_from_Water_Equipment.pdf')
        for res in (a, b):
            self.assertIsNone(res['case_id'])
            self.assertIn('without Claude', res['review'])
        # Claude is back.
        for name, num in (('Inv_29041_from_Water_Equipment.pdf', '29041'), ('Inv_29042_from_Water_Equipment.pdf', '29042')):
            self.extracts[name] = extraction('bill', num, po='PO950', client='Osprey Point', subtotal=1126.76)
            self.texts.pop(name)
        r = self.c.post(f"/pumps/api/docs/{a['doc_id']}/reread", json={})
        body = r.get_json()
        self.assertTrue(body['success'], body)
        self.assertEqual(body['kind'], 'bill')
        self.assertTrue(body['case_id'], 'filed onto an item once read cleanly')
        doc = self.c.get(f"/pumps/api/docs/{a['doc_id']}").get_json()['doc']
        self.assertEqual((doc['extracted_by'], doc['review_reason'], doc['doc_number']), ('claude', '', '29041'))
        self.assertEqual(self.case(body['case_id'])['vendor_bill_number'], '29041')
        # Already on an item: not read again over the office's work.
        self.assertEqual(self.c.post(f"/pumps/api/docs/{a['doc_id']}/reread", json={}).status_code, 400)
        # The next scan picks up the rest by itself.
        done = P.reread_pending('test')
        self.assertIn(b['doc_id'], [d['doc_id'] for d in done])
        self.assertTrue(self.c.get(f"/pumps/api/docs/{b['doc_id']}").get_json()['doc']['case_id'])

    def test_wettech_letter_quote(self):
        """Wettech quotes are Word letters: "RE: The Carlise", one "Your Cost" price that includes sales tax."""
        from docx import Document
        d = Document()
        for t in ('of Southwest Florida LLC', 'State of Florida Certified Plumbing Contractor', '#CFC1429137',
                  'Phone 941-232-4629   FAX 941-371-5151', 'Email: wettec@verizon.net', 'October 5, 2026',
                  'Stahlman England', 'Attn: Andrea',
                  'RE: The Carlise', 'We are pleased to quote you on the following services',
                  'Field service to check out pump station, found pipe had melted at the fitting going into the pump '
                  'suction.  Field service to pull and inspect suction line, clean screen and reinstall.',
                  'Your Cost --------------- $ 1158.99', 'Price includes Sales tax and in freight',
                  'Terms: Net 10 days', 'Thank You', 'H. H. (Tom) Morgan III'):
            d.add_paragraph(t)
        buf = io.BytesIO()
        d.save(buf)
        res = self.upload('Carslie Back Station .docx', data=buf.getvalue())
        self.assertEqual(res['kind'], 'quote', 'a quote letter, not a pump report')
        doc = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertEqual((doc['vendor'], doc['client_name'], doc['doc_date']), ('Wettech', 'The Carlise', '2026-10-05'))
        self.assertEqual((doc['subtotal'], doc['total']), (None, 1158.99))
        self.assertEqual(len(doc['line_items']), 1)
        line = doc['quote_suggestion']['line_items'][0]
        self.assertEqual(line['unit_price'], 1506.0, 'Simon quoted $1,506.69 in Jobber quote 9136 - whole dollars now')
        self.assertEqual(line['name'], 'Service Proposal Amount')
        self.assertFalse(line['taxable'], 'tax is already in the price - Jobber must not add it again')
        # Claude's reading of the same letter is cleaned up the same way.
        x = P._clean_extraction(extraction('quote', '', client='The Carlise', subtotal=1158.99, tax=0, items=[
            {'name': 'Suction line repair', 'description': 'Field service...', 'quantity': 1, 'unit_price': 1158.99,
             'amount': 1158.99, 'taxable': True, 'is_tax': False}]) | {'tax': None, 'subtotal': None, 'total': None,
                                                                       'tax_included': True})
        self.assertEqual((x['subtotal'], x['total'], x['line_items'][0]['taxable']), (None, 1158.99, False))

    def test_same_file_twice_is_skipped(self):
        self.extracts['dup.pdf'] = extraction('quote', 'Q-91', client='Dup Lakes', subtotal=100)
        first = self.upload('dup.pdf')
        # Not in Jobber yet: the same file again is read again, not filed twice.
        again = self.upload('dup.pdf')
        self.assertTrue(again.get('read_again'))
        self.assertEqual(again['doc_id'], first['doc_id'])
        # Once the office has corrected it by hand, a resend leaves it alone.
        self.c.patch(f"/pumps/api/docs/{first['doc_id']}", json={'total': 107.5})
        self.assertIn('already have', self.upload('dup.pdf').get('skipped', ''))

    # ── Jobber: drafts only, never sent ──────────────────────────────────────
    def test_guard_blocks_everything_but_drafts_and_notes(self):
        for q in ('mutation { invoiceSend(invoiceId: "1") { invoice { id } } }',
                  'mutation M($id: EncodedId!) { invoiceMarkAsSent(id: $id) { invoice { id } } }',
                  'mutation { x: invoiceMarkAsSent(id: "1") { invoice { id } } }',
                  'mutation { invoiceCreate(input: {}) { invoice { id } } quoteSend(quoteId: "1") { quote { id } } }',
                  'mutation { clientDelete(clientId: "1") { userErrors { message } } }'):
            with self.assertRaises(P.JobberError, msg=q):
                P.check_mutation_allowed(q)
        P.check_mutation_allowed(P.INVOICE_CREATE)
        P.check_mutation_allowed(P.QUOTE_CREATE)
        P.check_mutation_allowed('query { clients(first: 1) { nodes { id } } }')
        for target in P.NOTE_MUTATION.values():
            P.check_mutation_allowed(f'mutation N($id: EncodedId!, $input: {target[2]}!) {{ {target[0]}({target[1]}: $id, '
                                     f'input: $input) {{ {target[3]} {{ id }} userErrors {{ message }} }} }}')

    def test_no_send_mutation_anywhere_in_the_code(self):
        for name in ('pumps.py', 'pump_reports.py', 'pumps_page.py'):
            with open(os.path.join(ROOT, name)) as f:
                src = f.read()
            for word in ('invoiceSend', 'quoteSend', 'sendInvoice'):
                self.assertNotIn(word, src, f'{word} in {name}')
            # invoiceMarkAsSent may only be named in comments, never in a query.
            self.assertNotRegex(src, r"['\"][^'\"]*invoiceMarkAsSent", name)

    def test_draft_invoice_from_bill(self):
        self.extracts['q-0500.pdf'] = extraction('quote', 'Q-95', po='PO500', client='Lakeside Pines', subtotal=1803.81)
        self.extracts['b-0500.pdf'] = extraction('bill', '29600', po='PO500', client='Lakeside Pines',
                                                 subtotal=1803.81, tax=117.25)
        self.upload('q-0500.pdf')
        b = self.upload('b-0500.pdf')
        fake = FakeJobber()
        P.jobber_gql = fake
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={})
        body = r.get_json()
        self.assertTrue(body['success'], body)
        self.assertEqual(body['invoice']['invoice_status'], 'draft')
        creates = [v for q, v in fake.calls if 'invoiceCreate' in q]
        self.assertEqual(len(creates), 1)
        inp = creates[0]['input']
        self.assertEqual(inp['clientId'], 'C1', 'the clear best client, not the work-orders lead')
        self.assertEqual(len(inp['lineItems']), 1, 'sales tax line left off')
        # No Jobber quote to copy: one line, the bill before tax plus 30%.
        self.assertEqual((inp['lineItems'][0]['name'], inp['lineItems'][0]['unitPrice']), ('Service Proposal Amount', 2344.95))
        self.assertNotIn('quotation', inp['lineItems'][0]['description'].lower())
        self.assertEqual(inp['tax'], {'taxCalculationMethod': 'EXCLUSIVE'})
        self.assertFalse(any(word in q for q, _ in fake.calls for word in ('Send', 'MarkAsSent')))
        case = self.case(b['case_id'])
        self.assertEqual(case['sei_invoice_number'], '5001')
        self.assertTrue({s['key']: s for s in case['step_list']}['invoice_drafted'].get('at'))
        # A second draft for the same bill is refused.
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'C1'})
        self.assertEqual(r.status_code, 400)

    def test_sub_quote_goes_on_the_jobber_invoice(self):
        """Oct 2026: the sub's quote is noted on our Jobber invoice for later -
        on the invoice when Jobber takes invoice notes, else the client."""
        self.addCleanup(setattr, P, '_SCHEMA_CACHE', {})
        for n, notes in ((1, True), (2, False)):
            self.extracts[f'q-07{n}0.pdf'] = extraction('quote', f'Q-7{n}', po=f'PO7{n}', client='Lakeside Pines')
            self.extracts[f'b-07{n}0.pdf'] = extraction('bill', f'307{n}', po=f'PO7{n}', client='Lakeside Pines')
            self.upload(f'q-07{n}0.pdf')
            b = self.upload(f'b-07{n}0.pdf')
            fake = FakeJobber(invoice_notes=notes)
            P.jobber_gql, P._SCHEMA_CACHE = fake, {}
            body = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={}).get_json()
            self.assertTrue(body['success'], body)
            target = 'invoiceCreateNote' if notes else 'clientCreateNote'
            made = [v for q, v in fake.calls if target in q]
            self.assertEqual(len(made), 1, target)
            self.assertIn(f'Wettech quote #Q-7{n}', made[0]['input']['message'])
            self.assertEqual(made[0]['id'], 'INV1' if notes else 'C1')
            if not notes:
                self.assertIn('For invoice #5001', made[0]['input']['message'])
            self.assertIn('sub quote on invoice', [e['action'] for e in self.case(b['case_id'])['events']])

    def test_non_draft_status_from_jobber_raises_issue(self):
        self.extracts['b-0600.pdf'] = extraction('bill', '29700', po='PO600', client='Sunset Cove', subtotal=400)
        b = self.upload('b-0600.pdf')
        P.jobber_gql = FakeJobber(invoice_status='awaiting_payment')
        body = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'C9'}).get_json()
        self.assertTrue(body['success'], body)
        self.assertTrue(any(i['kind'] == 'not_draft' and not i['resolved_at'] for i in self.case(b['case_id'])['issues']))

    def test_draft_invoice_is_net_30(self):
        self.extracts['b-0630.pdf'] = extraction('bill', '29730', po='PO630', client='Sunset Cove', subtotal=400)
        b = self.upload('b-0630.pdf')

        class Net30Jobber(FakeJobber):
            reject_terms = False

            def __call__(self, query, variables=None):
                if '__type' in query:
                    n = variables['n']
                    if n == 'InvoiceCreateAttributes':
                        return {'__type': {'inputFields': [{'name': 'dueDetails', 'type': {
                            'kind': 'NON_NULL', 'ofType': {'kind': 'INPUT_OBJECT', 'name': 'InvoiceDueDetailsAttributes'}}}]}}
                    if n == 'InvoiceDueDetailsAttributes':
                        return {'__type': {'inputFields': [{'name': 'netTerms', 'type': {
                            'kind': 'ENUM', 'name': 'InvoiceNetTermsEnum'}}]}}
                    if n == 'InvoiceNetTermsEnum':
                        return {'__type': {'kind': 'ENUM', 'enumValues': [
                            {'name': 'DUE_ON_RECEIPT'}, {'name': 'NET_15'}, {'name': 'NET_30'}, {'name': 'NET_45'}]}}
                    return {'__type': {}}
                if 'invoiceCreate' in query and self.reject_terms and variables['input']['dueDetails']:
                    raise P.JobberError('Variable $input dueDetails is invalid')
                return super().__call__(query, variables)

        P._SCHEMA_CACHE.clear()
        fake = P.jobber_gql = Net30Jobber()
        try:
            body = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'C9'}).get_json()
            self.assertTrue(body['success'], body)
            sent = [v for q, v in fake.calls if 'invoiceCreate' in q][-1]['input']
            self.assertEqual(sent['dueDetails'], {'netTerms': 'NET_30'})
            # The invoice's amount in Jobber (with the client's tax) shows on the item once synced.
            inv_id = self.case(b['case_id'])['jobber']['invoice']['id']
            conn = P._conn()
            try:
                conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, total) "
                             "VALUES (?, 'invoice', '5001', 'draft', 556.4)", (inv_id,))
                P._follow_invoices(conn)
                conn.commit()
            finally:
                conn.close()
            self.assertEqual(self.case(b['case_id'])['jobber']['invoice']['total'], 556.4)
            # If Jobber won't take the terms, the draft is still made with its defaults.
            self.extracts['b-0631.pdf'] = extraction('bill', '29731', po='PO631', client='Sunset Cove', subtotal=400)
            b2 = self.upload('b-0631.pdf')
            fake.reject_terms = True
            body = self.c.post(f"/pumps/api/docs/{b2['doc_id']}/invoice", json={'client_id': 'C9'}).get_json()
            self.assertTrue(body['success'], body)
        finally:
            P._SCHEMA_CACHE.clear()

    def test_ambiguous_client_asks_a_person(self):
        self.extracts['b-amb.pdf'] = extraction('bill', '29800', client='Twin Oaks', subtotal=100)
        b = self.upload('b-amb.pdf')

        def two_equal(query, variables=None):
            return {'clients': {'nodes': [{'id': 'X1', 'name': 'Twin Oaks North'}, {'id': 'X2', 'name': 'Twin Oaks South'}]}}
        P.jobber_gql = two_equal
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={})
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.get_json()['needs_client'])

    def test_draft_quote_from_vendor_quote(self):
        self.extracts['q-0800.pdf'] = extraction('quote', 'Q-120', po='PO800', client='Lakeside Pines', subtotal=2450)
        q = self.upload('q-0800.pdf')
        cid = q['case_id']
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn(q['doc_id'], [d['id'] for d in queue['quotes_to_draft']])
        fake = FakeJobber()
        P.jobber_gql = fake
        body = self.c.post(f"/pumps/api/docs/{q['doc_id']}/quote", json={}).get_json()
        self.assertTrue(body['success'], body)
        self.assertEqual(body['quote']['quote_status'], 'draft')
        creates = [v for qq, v in fake.calls if 'quoteCreate(' in qq]
        self.assertEqual(len(creates), 1)
        attrs = creates[0]['attributes']
        self.assertEqual(attrs['clientId'], 'C1', 'the clear best client, not the work-orders lead')
        self.assertEqual(attrs['propertyId'], 'P1', "the client's only property")
        self.assertEqual(len(attrs['lineItems']), 1, 'sales tax line left off')
        self.assertEqual(attrs['lineItems'][0]['unitPrice'], 3185.0, "Wettech's price plus 30%")
        self.assertEqual(attrs['lineItems'][0]['name'], 'Service Proposal Amount')
        self.assertNotIn('quotation', attrs['lineItems'][0]['description'].lower())
        self.assertTrue(attrs['title'].startswith('Proposal'))
        notes = [v for qq, v in fake.calls if 'quoteCreateNote' in qq]
        self.assertEqual(notes[0]['id'], 'Q1', "Wettech's quote is saved as a note on our quote")
        self.assertFalse(any(word in qq for qq, _ in fake.calls for word in ('Send', 'MarkAsSent')))
        case = self.case(cid)
        self.assertEqual(case['jobber']['quote']['number'], '812')
        steps = {s['key']: s for s in case['step_list']}
        self.assertFalse(steps['client_quote'].get('at'), 'a draft is not "sent to client"')
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertNotIn(q['doc_id'], [d['id'] for d in queue['quotes_to_draft']])
        # A second draft for the same vendor quote is refused.
        r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/quote", json={'client_id': 'C1', 'property_id': 'P1'})
        self.assertEqual(r.status_code, 400)
        # When the office sends it from Jobber, and the client approves, the steps tick themselves.
        conn = P._conn()
        try:
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, updated_at) "
                         "VALUES ('Q1', 'quote', '812', 'Pump service', 'awaiting_response', '2026-10-03T10:00:00Z')")
            P._follow_quotes(conn)
            conn.commit()
        finally:
            conn.close()
        steps = {s['key']: s for s in self.case(cid)['step_list']}
        self.assertEqual(steps['client_quote'].get('at'), '2026-10-03')
        self.assertFalse(steps['client_approved'].get('at'))
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_jobber_items SET status='approved', updated_at='2026-10-06T09:00:00Z' "
                         "WHERE jobber_id='Q1'")
            P._follow_quotes(conn)
            conn.commit()
        finally:
            conn.close()
        steps = {s['key']: s for s in self.case(cid)['step_list']}
        self.assertEqual(steps['client_approved'].get('at'), '2026-10-06')

    def test_job_tracked_before_the_visit_step_waits_on_the_visit(self):
        """Homewood Suites was tracked from Jobber before the visit step existed:
        marked approved and scheduled, no vendor. Once, at start-up, it becomes
        'waiting on Wettech to assess' with the Jobber visit date."""
        conn = P._conn()
        try:
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, job_type, created_at, "
                         "start_at, category) VALUES ('J323', 'job', '14542', 'Homewood Suites - PO 323 Pump Service', "
                         "'upcoming', 'one_off', '2026-10-05T12:00:00Z', '2026-10-13T08:00:00Z', 'pump')")
            cid = P.create_case(conn, {'title': 'Homewood Suites - PO 323 Pump Service', 'client_name': 'Homewood Suites',
                                       'scheduled_for': '2026-10-13'}, 'test', source='jobber')
            conn.execute("UPDATE pump_cases SET jobber=? WHERE id=?", (json.dumps({'job': {'id': 'J323'}}), cid))
            P._set_step(conn, cid, 'client_approved', at='2026-10-05')
            P._set_step(conn, cid, 'scheduled', at='2026-10-05')
            conn.execute("DELETE FROM pump_state WHERE key='visit_step_backfill'")
            conn.commit()
            P._backfill_visit_step(conn)
        finally:
            conn.close()
        case = self.case(cid)
        self.assertEqual(case['stage'], 'assessment')
        self.assertEqual(case['vendor'], 'Wettech')
        steps = {s['key']: s for s in case['step_list']}
        self.assertEqual(steps['assessment'].get('due'), '2026-10-13')
        self.assertFalse(steps['scheduled'].get('at'))
        self.assertFalse(steps['client_approved'].get('at'))
        # And by hand: "Wettech needs to visit first" on a job waiting on a quote.
        other = P._conn()
        try:
            oid = P.create_case(other, {'title': 'Pump Service call - PO322', 'client_name': 'Greenscapes',
                                        'vendor': 'Wettech'}, 'test')
            other.commit()
        finally:
            other.close()
        self.assertEqual(self.case(oid)['stage'], 'vendor_quote')
        r = self.c.patch(f'/pumps/api/cases/{oid}', json={'steps': {'assessment': None}, 'assessment_due': '10/20/2026'})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        case = self.case(oid)
        self.assertEqual(case['stage'], 'assessment')
        self.assertEqual({s['key']: s for s in case['step_list']}['assessment'].get('due'), '2026-10-20')

    def test_every_huntington_lakes_quote_is_for_the_master_association(self):
        master = 'Z2lkOi8vSm9iYmVyL0NsaWVudC83NTE2ODg4Nw=='
        for name in ('Huntington Lakes', 'Huntington Lakes Residence Assoc.', 'Huntington Lakes Phase 3',
                     'HUNTINGTON LAKES PHASE 5 C/O ABILITY MANAGEMENT'):
            a = P.match_site_alias({'client_name': name, 'file_name': 'q.pdf'})
            self.assertEqual((a or {}).get('client_id'), master, name)
        for name in ('Huntington II HOA @ Legends', 'Huntington I at the Legends'):
            a = P.match_site_alias({'client_name': name, 'file_name': 'q.pdf'})
            self.assertNotEqual((a or {}).get('client_id'), master, name)
        # Other site names still don't match a name with words of its own.
        self.assertIsNone(P.match_site_alias({'client_name': 'Carlisle Golf Club', 'file_name': 'q.pdf'}))

    def test_read_a_quote_again(self):
        """A letter with two quotes was read as one; reading it again (the
        button, or the same file sent again) picks up both."""
        one = extraction('quote', 'Q-7701', po='PO7701', client='Zinnia Point', subtotal=1780.11)
        two = dict(one, line_items=[
            {'name': 'Control box', 'description': 'Furnish and install new control box', 'quantity': 1,
             'unit_price': 1780.11, 'amount': 1780.11, 'taxable': True, 'is_tax': False},
            {'name': 'Motor', 'description': 'Replace motor', 'quantity': 1, 'unit_price': 4836.9,
             'amount': 4836.9, 'taxable': True, 'is_tax': False, 'optional': True}], subtotal=None, tax=None)
        self.extracts['q-7701.pdf'] = one
        first = self.upload('q-7701.pdf')
        self.assertEqual(len(self.case(first['case_id'])['docs'][0]['line_items']), 2, 'one price + its tax line')
        self.extracts['q-7701.pdf'] = two
        r = self.c.post(f"/pumps/api/docs/{first['doc_id']}/read_again", json={})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.assertIn('1 optional', r.get_json()['summary'])
        doc = next(d for d in self.case(first['case_id'])['docs'] if d['id'] == first['doc_id'])
        self.assertEqual([bool(i.get('optional')) for i in doc['line_items']], [False, True])
        self.assertTrue(any(e['action'] == 'read again' for e in self.case(first['case_id'])['events']))
        # Sent again (same file): read again rather than ignored, while it isn't in Jobber.
        again = self.upload('q-7701.pdf')
        self.assertTrue(again.get('read_again'), again)
        self.assertEqual(again['doc_id'], first['doc_id'])
        # Reports aren't read this way.
        self.assertEqual(self.c.post('/pumps/api/docs/999999/read_again', json={}).status_code, 400)

    def test_quote_read_again_redoes_our_draft_in_jobber(self):
        """Our quote is already drafted in Jobber when the vendor's quote is
        read again: the draft's lines are redone - only while it's a draft."""
        T = lambda name, kind='INPUT_OBJECT': {'kind': kind, 'name': name, 'ofType': None}
        NN = lambda t: {'kind': 'NON_NULL', 'name': None, 'ofType': t}
        LST = lambda t: {'kind': 'LIST', 'name': None, 'ofType': t}

        class Fake(FakeJobber):
            status = 'draft'
            lines = [{'name': 'Service Proposal Amount', 'unitPrice': 650.0}]   # as the app made it

            def __call__(self, query, variables=None):
                if '__schema' in query:
                    return {'__schema': {'mutationType': {'fields': [
                        {'name': 'quoteCreateLineItems', 'type': T('QuoteCreateLineItemsPayload', 'OBJECT'),
                         'args': [{'name': 'quoteId', 'type': NN(T('EncodedId', 'SCALAR'))},
                                  {'name': 'lineItems', 'type': NN(LST(NN(T('QuoteCreateLineItemAttributes'))))}]},
                        {'name': 'quoteDeleteLineItems', 'type': T('QuoteDeleteLineItemsPayload', 'OBJECT'),
                         'args': [{'name': 'quoteId', 'type': NN(T('EncodedId', 'SCALAR'))},
                                  {'name': 'lineItemIds', 'type': NN(LST(NN(T('EncodedId', 'SCALAR'))))}]}]}}}
                if '__type' in query:
                    return {'__type': {'fields': [{'name': 'quote'}, {'name': 'userErrors'}]}}
                if 'quoteCreateLineItems(' in query or 'quoteDeleteLineItems(' in query:
                    P.check_mutation_allowed(query)
                    self.calls.append((query, variables))
                    name = 'quoteCreateLineItems' if 'Create' in query.split('{')[1] else 'quoteDeleteLineItems'
                    return {name: {'quote': {'id': 'Q1'}, 'userErrors': []}}
                if 'quoteStatus lineItems' in query:
                    return {'quote': {'quoteNumber': 812, 'quoteStatus': self.status,
                                      'lineItems': {'nodes': [dict(l, id='L-old') for l in self.lines]}}}
                return super().__call__(query, variables)
        P._SCHEMA_CACHE.clear()
        fake = Fake()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            self.extracts['q-0920.pdf'] = extraction('quote', 'Q-920', po='PO920', client='Lakeside Pines', subtotal=500)
            q = self.upload('q-0920.pdf', data=self._pdf())
            self.assertNotIn('quote', self.case(q['case_id'])['jobber'], 'nothing drafted on its own')
            self.draft(q['doc_id'])
            self.assertEqual(self.case(q['case_id'])['jobber']['quote']['number'], '812')
            self.extracts['q-0920.pdf'] = extraction('quote', 'Q-920', po='PO920', client='Lakeside Pines', subtotal=800)
            r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/read_again", json={}).get_json()
            self.assertTrue(r.get('updated'), r)
            added = [v for qq, v in fake.calls if 'quoteCreateLineItems(' in qq][0]
            self.assertEqual(added['quoteId'], 'Q1')
            self.assertEqual(sum(li['unitPrice'] for li in added['lineItems']), 1040.0, '800 + 30%')
            removed = [v for qq, v in fake.calls if 'quoteDeleteLineItems(' in qq][0]
            self.assertEqual(removed['lineItemIds'], ['L-old'])
            self.assertEqual(self.case(q['case_id'])['jobber']['quote']['total'], 1040.0)
            # Changed by hand in Jobber since: left alone.
            fake.lines = [{'name': 'Service Proposal Amount', 'unitPrice': 999.0}]
            fake.calls.clear()
            r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/read_again", json={}).get_json()
            self.assertFalse(r.get('updated'))
            self.assertIn('by hand', r['note'])
            self.assertFalse([1 for qq, _ in fake.calls if 'LineItems(' in qq])
            # Once it's sent, it is left alone.
            fake.lines = [{'name': 'Service Proposal Amount', 'unitPrice': 1040.0}]
            fake.status = 'awaiting_response'
            fake.calls.clear()
            r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/read_again", json={}).get_json()
            self.assertFalse(r.get('updated'))
            self.assertIn('already awaiting response', r['note'])
            self.assertFalse([1 for qq, _ in fake.calls if 'LineItems(' in qq])
            # Deleted in Jobber since: forgotten, and no new one made until Draft quote is clicked.
            orig = Fake.__call__

            def deleted(self_, query, variables=None):
                if 'quoteStatus lineItems' in query:
                    raise P.JobberError("Couldn't find Quote with 'id'=Q1")
                return orig(self_, query, variables)
            Fake.__call__ = deleted
            fake.calls.clear()
            r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/read_again", json={}).get_json()
            self.assertIn('#812 was deleted in Jobber - click Draft quote', r['note'])
            self.assertNotIn('quote', self.case(q['case_id'])['jobber'])
            self.assertFalse([1 for qq, _ in fake.calls if 'quoteCreate(' in qq])
        finally:
            P.JOBBER_STATIC_TOKEN = saved

    def test_letter_with_two_prices_is_one_quote_with_a_line_each(self):
        """Autumn Woods 2 (Oct 9 2026): a flow meter and a surge arrestor, each
        with its own "Your Cost" - ONE job and ONE quote with two lines."""
        self.texts['autumn2.pdf'] = (
            "October 9, 2026\nStahlman England\nAttn: Beatriz\nRE: Autumn Woods 2\n"
            "After preforming maintenance we found the following issues that should be addressed, We are pleased "
            "to quote on the following services and materials\n"
            "Item #1 ... Field service to furnish and install new 4 inch flow meter in place of flow meter that is "
            "unreadable in the ground\n                Your Cost ---------------- $ 2757.32\n"
            "Item #2 ... Field service to furnish and install new surge arrestor to replace failed surge arrestor\n"
            "                Your Cost ---------------- $ 1047.62\n"
            "Price does include sales tax and freight\nTerms: Net 10 days\n"
            "Delivery: 2-3 Weeks after receipt of order for completion\nPrices good for 30 days.\n")
        fake = FakeJobber()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            doc = self.upload('autumn2.pdf')
            self.assertFalse(doc.get('parts'), 'not split into two jobs')
            case = self.case(doc['case_id'])
            d = next(x for x in case['docs'] if x['id'] == doc['doc_id'])
            self.assertEqual([i['amount'] for i in d['line_items']], [2757.32, 1047.62])
            self.assertEqual(d['total'], 3804.94)
            self.assertFalse(d['line_items'][0]['description'].lower().startswith('item'))
            r = self.draft(doc['doc_id'], client_id='C1', property_id='P1')
            self.assertTrue(r['success'], r)
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        made = [v['attributes'] for qq, v in fake.calls if 'quoteCreate(' in qq]
        self.assertEqual(len(made), 1, 'one quote')
        lines = made[0]['lineItems']
        self.assertEqual([(l['name'], l['unitPrice'], l['taxable']) for l in lines],
                         [('Service Proposal Amount', 3584.0, False),
                          ('Service Proposal Amount', 1361.0, False)])
        self.assertIn('flow meter', lines[0]['description'])
        self.assertIn('surge arrestor', lines[1]['description'])
        self.assertFalse(any(l.get('optional') for l in lines))

    def test_letter_split_before_oct_9_is_one_quote_again(self):
        """Autumn Woods 2 came in while letters were split: the doc kept the
        flow meter and a copy made a job for the surge arrestor. Read again,
        it is one quote with both lines and the copy's job is cancelled."""
        self.texts['autumn2b.pdf'] = (
            "October 9, 2026\nStahlman England\nRE: Autumn Woods 2\n"
            "We are pleased to quote on the following services and materials\n"
            "Item #1 ... Field service to furnish and install new 4 inch flow meter\n"
            "                Your Cost ---------------- $ 2757.32\n"
            "Item #2 ... Field service to furnish and install new surge arrestor\n"
            "                Your Cost ---------------- $ 1047.62\n"
            "Price does include sales tax and freight\n")
        P.SPLIT_LETTERS = True
        try:
            doc = self.upload('autumn2b.pdf')
        finally:
            P.SPLIT_LETTERS = False
        self.assertEqual(len(doc.get('parts') or []), 1, 'split the old way')
        copy = doc['parts'][0]
        conn = P._conn()
        copy_case_id = conn.execute('SELECT case_id FROM pump_docs WHERE id=?', (copy,)).fetchone()[0]
        conn.close()
        d = next(x for x in self.case(doc['case_id'])['docs'] if x['id'] == doc['doc_id'])
        self.assertEqual([i['amount'] for i in d['line_items']], [2757.32])
        again = self.c.post(f"/pumps/api/docs/{doc['doc_id']}/read_again", json={}).get_json()
        self.assertTrue(again['success'], again)
        self.assertIn('one quote again', again['summary'])
        d = next(x for x in self.case(doc['case_id'])['docs'] if x['id'] == doc['doc_id'])
        self.assertEqual([i['amount'] for i in d['line_items']], [2757.32, 1047.62])
        self.assertNotIn('letter_parts', d.get('jobber') or {})
        conn = P._conn()
        self.assertEqual(conn.execute('SELECT status FROM pump_docs WHERE id=?', (copy,)).fetchone()[0], 'dismissed')
        self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (copy_case_id,)).fetchone()[0],
                         'cancelled')
        conn.close()
        # Reading it again once more doesn't split it or touch anything else.
        again = self.c.post(f"/pumps/api/docs/{doc['doc_id']}/read_again", json={}).get_json()
        self.assertTrue(again['success'], again)
        self.assertFalse(again.get('parts'))

    def test_duplicates_are_caught(self):
        # 1. Wettech sends invoice #40123 again as a different file: same amount = duplicate, never used again.
        self.extracts['b-40123.pdf'] = extraction('bill', '40123', po='PO777', client='Quail Hollow', subtotal=812.5)
        self.extracts['b-40123 resent.pdf'] = extraction('bill', '40123', po='PO777', client='Quail Hollow', subtotal=812.5)
        first = self.upload('b-40123.pdf')
        again = self.upload('b-40123 resent.pdf')
        self.assertIn('Duplicate', again.get('skipped', ''), again)
        self.assertEqual(again['case_id'], first['case_id'])
        q = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertEqual([d['id'] for d in q['bills_to_draft'] if d['doc_number'] == '40123'], [first['doc_id']],
                         'only one copy is offered for an invoice')
        # A different amount on the same number: a revised invoice? - a person decides.
        self.extracts['b-40123 v2.pdf'] = extraction('bill', '40123', po='PO777', client='Quail Hollow', subtotal=900)
        rev = self.upload('b-40123 v2.pdf')
        self.assertIn('different amount', rev['review'])
        self.assertTrue(any(i['kind'] == 'revised_doc' for i in self.case(first['case_id'])['issues']))

        # 2. Jobber already has a quote for this client at this price (made by hand): don't draft another.
        self.extracts['q-0901.pdf'] = extraction('quote', 'Q-901', po='PO901', client='Lakeside Pines', subtotal=1000)
        q1 = self.upload('q-0901.pdf')

        class HasQuote(FakeJobber):
            def __call__(self, query, variables=None):
                if 'quotes(first: 30)' in query:
                    return {'client': {'id': 'C1', 'quotes(first: 30)': {'nodes': [
                        {'id': 'QH', 'quoteNumber': 9200, 'quoteStatus': 'awaiting_response',
                         'createdAt': P._today().isoformat() + 'T10:00:00Z', 'amounts': {'total': 1300.0}}]}}}
                return super().__call__(query, variables)
        P.jobber_gql = HasQuote()
        r = self.c.post(f"/pumps/api/docs/{q1['doc_id']}/quote", json={'client_id': 'C1', 'property_id': 'P1'})
        self.assertEqual(r.status_code, 409, r.get_data(as_text=True))
        self.assertEqual(r.get_json()['possible_duplicate']['number'], '9200')
        self.assertFalse(any('quoteCreate(' in qq for qq, _ in P.jobber_gql.calls), 'nothing drafted')
        r = self.c.post(f"/pumps/api/docs/{q1['doc_id']}/quote", json={'client_id': 'C1', 'property_id': 'P1', 'force': True})
        self.assertEqual(r.status_code, 200, 'drafted anyway when asked')

        # 3. Two Jobber invoices for the same client and amount, close together: a to-do.
        conn = P._conn()
        try:
            for jid, num, day in (('IA', '7001', '2026-09-02'), ('IB', '7002', '2026-09-10'), ('IC', '7003', '2026-12-30')):
                conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, total, client_id, "
                             "client_name, created_at, web_uri) VALUES (?, 'invoice', ?, 'awaiting_payment', 450.0, 'CQ', "
                             "'Quail Run', ?, ?)", (jid, num, day + 'T12:00:00Z', 'https://jobber/i/' + num))
            found = P.duplicate_jobber_invoices(conn)
            conn.commit()
        finally:
            conn.close()
        self.assertEqual(found, ['dup-invoice:IA:IB'])
        titles = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertTrue(any('#7001 and #7002' in t for t in titles), titles)
        dup = [t for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos'] if '#7001 and #7002' in t['title']][0]
        self.assertEqual([u['uri'] for u in dup['link']['uris']], ['https://jobber/i/7001', 'https://jobber/i/7002'],
                         'both invoices can be opened from the to-do')

    def test_quote_sent_in_jobber_clears_send_quote(self):
        """A quote whose Jobber status was already noted as sent, but whose
        "Quote sent" step never got ticked, is ticked by the next sync."""
        j = self.c.post('/pumps/api/cases', json={'title': 'Sent already', 'client_name': 'Quail Run'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET jobber=? WHERE id=?", (P.json.dumps(
                {'quote': {'id': 'QS1', 'number': '9300', 'status': 'awaiting_response'}}), cid))
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, updated_at) "
                         "VALUES ('QS1', 'quote', '9300', 'awaiting_response', '2026-10-01T12:00:00Z')")
            P._follow_quotes(conn)
            conn.commit()
        finally:
            conn.close()
        self.assertTrue(self.case(cid)['steps']['client_quote'].get('at'))
        self.assertNotIn(cid, [c['id'] for c in self.c.get('/pumps/api/summary').get_json()['queue']['to_quote_client']])

    def test_quote_deleted_in_jobber_is_forgotten(self):
        j = self.c.post('/pumps/api/cases', json={'title': 'Deleted quote', 'client_name': 'Quail Run'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET jobber=? WHERE id=?", (P.json.dumps(
                {'quote': {'id': 'QDEL', 'number': '9149', 'status': 'draft'}}), cid))
            conn.execute("INSERT INTO pump_docs (kind, file_name, case_id, status, jobber) VALUES "
                         "('quote', 'q.pdf', ?, 'filed', ?)", (cid, P.json.dumps({'quote_id': 'QDEL', 'quote_number': '9149'})))
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, case_id) "
                         "VALUES ('QDEL', 'quote', '9149', 'draft', ?)", (cid,))
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, created_at) "
                         "VALUES ('QGONE2', 'quote', '9150', 'approved', ?)", (P._today().isoformat(),))
            conn.commit()

            class Gone(FakeJobber):
                def __call__(self, query, variables=None):
                    if variables and variables.get('id') == 'QDEL':
                        return {'quote': None}
                    if variables and variables.get('id') == 'QGONE2':
                        raise P.JobberError("Couldn't find Quote with 'id'=QGONE2")
                    return super().__call__(query, variables)
            P.jobber_gql = Gone()
            errors = P._refresh_linked_records(conn)
            conn.commit()
            self.assertEqual(errors, [])
            self.assertIsNone(conn.execute("SELECT 1 FROM pump_jobber_items WHERE jobber_id IN ('QDEL', 'QGONE2')").fetchone())
            doc_j = P.json.loads(conn.execute('SELECT jobber FROM pump_docs WHERE case_id=?', (cid,)).fetchone()[0])
            self.assertNotIn('quote_id', doc_j)
        finally:
            conn.close()
        case = self.case(cid)
        self.assertNotIn('quote', case['jobber'])
        self.assertTrue(any('#9149 was deleted' in (e.get('detail') or '') for e in case['events']))

    def test_undo_a_recent_change(self):
        j = self.c.post('/pumps/api/cases', json={'title': 'Undo me', 'client_name': 'Quail Run'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        self.c.post(f'/pumps/api/cases/{cid}', json={'title': 'Renamed'})
        self.assertEqual(self.case(cid)['title'], 'Renamed')
        acts = self.c.get('/pumps/api/undo').get_json()['actions']
        self.assertEqual([a['label'] for a in acts[:2]], [f'Changed job #{cid}', 'Added a job'])
        # Undoing the add first asks first - the rename came after it.
        r = self.c.post(f"/pumps/api/undo/{acts[1]['id']}", json={})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.get_json()['blocked_by'], [f'Changed job #{cid}'])
        self.assertTrue(self.c.post(f"/pumps/api/undo/{acts[0]['id']}", json={}).get_json()['success'])
        self.assertEqual(self.case(cid)['title'], 'Undo me')
        self.assertEqual(self.c.post(f"/pumps/api/undo/{acts[0]['id']}", json={}).status_code, 400, 'only once')
        # Cancelling a job, undone: it's back, and so is everything about it.
        self.c.post(f'/pumps/api/cases/{cid}/delete', json={})
        self.assertEqual(self.case(cid)['status'], 'cancelled')
        top = self.c.get('/pumps/api/undo').get_json()['actions'][0]
        self.assertEqual(top['label'], f'Deleted job #{cid}')
        self.c.post(f"/pumps/api/undo/{top['id']}", json={})
        self.assertNotEqual(self.case(cid)['status'], 'cancelled')
        # Undoing the add removes the job (and its history).
        r = self.c.post(f"/pumps/api/undo/{acts[1]['id']}", json={}).get_json(); self.assertTrue(r['success'], r)
        self.assertEqual(self.c.get(f'/pumps/api/cases/{cid}').status_code, 404)
        # A to-do ticked by mistake comes back; Jobber drafts carry a warning.
        todo = self.c.post('/pumps/api/todos', json={'title': 'Call Tom'}).get_json()
        tid = [t for t in todo['todos'] if t['title'] == 'Call Tom'][0]['id']
        self.c.post(f'/pumps/api/todos/{tid}/done', json={'done': True})
        top = self.c.get('/pumps/api/undo').get_json()['actions'][0]
        self.c.post(f"/pumps/api/undo/{top['id']}", json={})
        open_titles = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertIn('Call Tom', open_titles)
        self.assertIn('stays', P._undo_label('/pumps/api/docs/4/quote')[1])
        # A PATCH (the SCADA and job edit forms) is undoable too.
        j2 = self.c.post('/pumps/api/cases', json={'title': 'Patch me', 'client_name': 'Quail Run'}).get_json()
        c2 = j2.get('case_id') or j2['case']['id']
        self.c.patch(f'/pumps/api/cases/{c2}', json={'title': 'Patched'})
        self.c.post(f'/pumps/api/cases/{c2}', json={'title': 'Patched again'})
        acts = self.c.get('/pumps/api/undo').get_json()['actions']
        self.assertEqual(acts[1]['label'], f'Changed job #{c2}')
        # Undoing the PATCH and, when asked, the later change on the same job too.
        r = self.c.post(f"/pumps/api/undo/{acts[1]['id']}", json={'also_later': True}).get_json()
        self.assertTrue(r['success'], r)
        self.assertEqual(self.case(c2)['title'], 'Patch me')
        # Background work (scan, sync) is never put on the undo list.
        self.assertEqual(P._UNDO_ACTION.get(), 0)

    def test_follow_up_before_a_service_comes_due(self):
        def todos():
            return {t['title']: t for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']}
        # Mid-November: nothing of these four is due (quarterly Jan/Apr/Jul/Oct; Hodges Jan/May/Sep).
        self.assertEqual(P.service_follow_ups(P.datetime(2026, 11, 10).date()), [])
        # A week before January: Sopra, Hodges (SiteOne), Warm Springs (Naples Electric) and Kurt Biggs' dive.
        added = P.service_follow_ups(P.datetime(2026, 12, 26).date())
        self.assertEqual(len(added), 4, added)
        t = todos()
        self.assertIn('lmleczek@davisdevelopment.com', t["Follow up before Sopra Luxury Living's January service"]['detail'])
        self.assertIn('PSochar@siteone.com', t["Follow up before Hodges Funeral Home's January service"]['detail'])
        self.assertIn('Ramon', t["Follow up before Kurt Biggs's January dive"]['detail'])
        # Once per service: the hourly run doesn't add it again, even once it's done.
        self.assertEqual(P.service_follow_ups(P.datetime(2027, 1, 5).date()), [])
        # Set on the Accounts tab: a new one follows the same rule.
        aid = next(a['id'] for a in self.c.get('/pumps/api/maint').get_json()['accounts'] if a['name'] == 'Allura')
        self.c.post('/pumps/api/maint', json={'id': aid, 'follow_up': 'Call the manager for gate access.'})
        self.assertEqual(len(P.service_follow_ups(P.datetime(2027, 1, 5).date())), 1)
        self.c.post('/pumps/api/maint', json={'id': aid, 'follow_up': ''})

    def test_huntington_6_two_prices_one_optional(self):
        """Wettech's Huntington #6 letter (Oct 8 2026): the control box, and a
        'Possibly needed' motor replacement - two prices, the second optional."""
        text = ('October 8, 2026\nStahlman England\nAttn: Simon\nRE: Huntington #6\n'
                'We are pleased to quote you on the following services\n'
                'Field service to check out #6 pump station, found all of the capacitors and potential relay blown, '
                'suggest starting with a control box replacement.\nFurnish and install new 10 HP deluxe control box, '
                'wire up and test.\n\t\t\t\t\t\tYour Cost ---------------- $ 1780.11\n'
                'Possibly needed- Field service to pull and inspect pump and motor, replace motor with new 10 HP 230 '
                'volt 1 PH motor reusing pump end, reinstall, wire up and test\n'
                '        Your Cost --------------- $ 4836.90\nPrices include Sales tax\nTerms: Net 10 days\n')
        x = P._clean_extraction(P._regex_extract(text, 'tomm@wettec.biz', 'Huntington #6'))
        self.assertEqual(x['total'], 1780.11, 'the possibly-needed work is not in the total')
        self.assertEqual([i.get('optional', False) for i in x['line_items']], [False, True])
        self.assertTrue(x['line_items'][1]['description'].startswith('Field service to pull and inspect'))
        self.assertIn('Optional (not in the total)', x['notes'])
        # Two quotes to the client, not one with an optional line: the motor
        # job is its own quote, marked "If needed".
        first, motor = P._letter_part(x, 0), P._letter_part(x, 1)
        self.assertEqual([(l['name'], l['unit_price']) for l in P.one_line_items(first, 30)],
                         [('Service Proposal Amount', 2314.14)])
        lines = P.one_line_items(motor, 30)
        self.assertEqual([(l['name'], l['unit_price']) for l in lines], [('Service Proposal Amount', 6287.97)])
        self.assertTrue(lines[0]['description'].startswith('If needed: Field service to pull'))
        self.assertEqual((first['total'], motor['total']), (1780.11, 4836.9))
        # The client: the Master Association, whatever Wettech calls it.
        master = 'Z2lkOi8vSm9iYmVyL0NsaWVudC83NTE2ODg4Nw=='
        for name in ('Huntington #6', 'Huntington Lakes Residence Assoc.', 'HUNTINGTON LAKES PHASE 5 C/O ABILITY'):
            self.assertEqual((P.match_site_alias({'client_name': name, 'file_name': 'q.docx'}) or {}).get('client_id'),
                             master, name)
        for name in ('Huntington II HOA @ Legends', 'Huntington I at the Legends'):
            self.assertNotEqual((P.match_site_alias({'client_name': name, 'file_name': 'q.docx'}) or {}).get('client_id'),
                                master, name)
        self.assertIsNone(P.match_site_alias({'client_name': 'Carlisle Golf Club', 'file_name': 'q.pdf'}),
                          'other site names still need the name itself')
        # The property: the one Jobber labels with that pump number.
        props = [{'id': 'P4', 'label': '2437 Millcreek Lane Pump #4'}, {'id': 'P6', 'label': '2500 Oak Lane Pump #6'},
                 {'id': 'PL6', 'label': '6585 Huntington Lakes Circle Pool #6'}, {'id': 'P10', 'label': 'Pump 10-Phase 4'}]
        self.assertEqual(P._pump_number_property(props, 'RE: Huntington #6 check out #6 pump station')['id'], 'P6')
        self.assertIsNone(P._pump_number_property(props, 'Huntington #8'))
        # Jobber without optional lines: the option goes in the text, never the total.
        sent = []

        def jobber(query, variables=None):
            items = variables['attributes']['lineItems']
            if any('optional' in li for li in items):
                raise P.JobberError('Variable $attributes: lineItems.1.optional is not defined')
            sent.append(items)
            return {'quoteCreate': {'quote': {'id': 'Q1'}, 'userErrors': []}}
        real = P.jobber_gql
        P.jobber_gql = jobber
        try:
            # (an optional line still happens: a letter quoting one job and its options, read as one)
            P.gql_with_lines(P.QUOTE_CREATE, {'attributes': {'lineItems': P._jobber_lines(P.one_line_items(x, 30),
                                                                                         'quote')}}, 'attributes')
        finally:
            P.jobber_gql = real
        self.assertEqual(len(sent[0]), 1)
        self.assertEqual(sent[0][0]['unitPrice'], 2314.14)
        self.assertIn('Optional, if needed (not included)', sent[0][0]['description'])
        self.assertIn('$6,287.97', sent[0][0]['description'])

    def test_repairs_go_to_whoever_looks_after_the_account(self):
        conn = P._conn()
        try:
            by = dict(conn.execute("SELECT name, repairs_by FROM pump_maint_accounts WHERE name IN "
                                   "('Hodges Funeral Home', 'Warm Springs Comm. Assoc.', 'Autumn Woods', "
                                   "'The Reserve @ Estero', 'Allura')").fetchall())
            self.assertEqual(by, {'Hodges Funeral Home': 'SiteOne', 'Warm Springs Comm. Assoc.': 'Naples Electric',
                                  'Autumn Woods': P.IN_HOUSE, 'The Reserve @ Estero': P.IN_HOUSE, 'Allura': ''})
            self.assertIn('Oscar (our technician)', conn.execute(
                "SELECT notes FROM pump_maint_accounts WHERE name='Autumn Woods'").fetchone()[0])
            # A repair at Autumn Woods (tracked from Jobber, which says Wettech) is Oscar's: no Wettech quote or bill.
            oscar = P.create_case(conn, {'title': 'Pump tripping', 'client_name': 'Autumn Woods', 'vendor': 'Wettech'},
                                  'test', source='jobber')
            site_one = P.create_case(conn, {'title': 'Pump repair', 'client_name': 'Hodges Funeral Home'}, 'test')
            # Wettech's own document stays Wettech's; so does the maintenance.
            wettech = P.create_case(conn, {'title': 'Repair', 'client_name': 'Autumn Woods', 'vendor': 'Wettech'},
                                    'test', source='email')
            maint = P.create_case(conn, {'title': 'Quarterly', 'client_name': 'Autumn Woods', 'category': 'maintenance'},
                                  'test', source='jobber')
            conn.commit()
        finally:
            conn.close()
        c = self.case(oscar)
        self.assertEqual(c['vendor'], P.IN_HOUSE)
        self.assertEqual(c['stage'], 'client_quote', 'straight to our quote')
        self.assertEqual(c['account']['repairs_by'], P.IN_HOUSE)
        self.assertEqual(self.case(site_one)['vendor'], 'SiteOne')
        self.assertEqual(self.c.get(f'/pumps/api/cases/{site_one}/vendor_email').get_json()['email']['to'],
                         'PSochar@siteone.com')
        self.assertEqual(self.case(wettech)['vendor'], 'Wettech')
        self.assertNotEqual(self.case(maint)['vendor'], P.IN_HOUSE)
        # Set on the Accounts tab.
        aid = next(a['id'] for a in self.c.get('/pumps/api/maint').get_json()['accounts'] if a['name'] == 'Allura')
        r = self.c.post('/pumps/api/maint', json={'id': aid, 'repairs_by': 'Naples Electric'})
        self.assertEqual(next(a for a in r.get_json()['accounts'] if a['id'] == aid)['repairs_by'], 'Naples Electric')
        self.c.post('/pumps/api/maint', json={'id': aid, 'repairs_by': ''})

    def test_master_list_contacts_show_on_the_job(self):
        conn = P._conn()
        try:
            notes = conn.execute("SELECT notes FROM pump_maint_accounts WHERE name='Barrington Cove'").fetchone()[0]
            self.assertIn('quotes only to John Rodriguez', notes)
            self.assertIn('Ramon', conn.execute("SELECT diver_notes FROM pump_dive_sites WHERE name='Kurt Biggs'").fetchone()[0])
            heritage = json.loads(conn.execute("SELECT years FROM pump_scada_accounts WHERE client='Heritage Stations'")
                                  .fetchone()[0])
            self.assertEqual((heritage['2023'], heritage['2024'], heritage['2025']), ('18456', '23846', '28165'))
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM pump_scada_accounts WHERE client LIKE '%HERITAGE%'")
                             .fetchone()[0], 1, 'no duplicate Heritage')
            fgcu = conn.execute("SELECT active FROM pump_scada_accounts WHERE client='FGCU Intermural Pump Station'").fetchone()
            self.assertEqual(fgcu[0], 0, 'added switched off - no renewal quote drafted on its own')
            # Run again (the other worker): nothing doubles.
            P._apply_master_list(conn.cursor(), P._now_text())
            conn.commit()
            self.assertEqual(conn.execute("SELECT notes FROM pump_maint_accounts WHERE name='Barrington Cove'").fetchone()[0],
                             notes)
            cid = P.create_case(conn, {'title': 'Well at lake #4', 'client_name': 'Barrington Cove'}, 'test')
            conn.commit()
        finally:
            conn.close()
        acct = self.case(cid)['account']
        self.assertEqual(acct['name'], 'Barrington Cove')
        self.assertIn('Invoices only to Paramount', acct['notes'])

    def test_who_to_follow_up_with(self):
        """The account's contact from the master list; else the client's AP/billing email in Jobber."""
        self.assertEqual(P.billing_emails([{'address': 'Board.VP@example.com', 'primary': True},
                                           {'address': 'ap@mgmt-example.com', 'primary': False},
                                           {'address': 'AP@mgmt-example.com', 'primary': False}]),
                         ['ap@mgmt-example.com'])
        self.assertEqual(P.billing_emails([{'address': 'x@example.com'}, {'address': 'owner@example.com', 'primary': True}]),
                         ['owner@example.com'])
        self.assertEqual(P.billing_emails([{'address': 'accountspayable@gc-example.com', 'primary': True},
                                           {'address': 'pm@gc-example.com'}]), ['accountspayable@gc-example.com'])
        conn = P._conn()
        try:
            full = P.create_case(conn, {'title': 'Pump repair', 'client_name': 'Cypress Legends'}, 'test')
            part = P.create_case(conn, {'title': 'Pump repair', 'client_name': 'Huntington Lakes Residence Assoc.'}, 'test')
            none = P.create_case(conn, {'title': 'Pump repair', 'client_name': 'Nowhere Example HOA'}, 'test')
            conn.execute("UPDATE pump_cases SET jobber_client_id='CL-HL' WHERE id=?", (part,))
            conn.execute("INSERT OR REPLACE INTO pump_client_contacts (client_id, name, emails, fetched_at) VALUES "
                         "('CL-HL', 'Huntington', ?, ?)", (json.dumps([{'address': 'vp@example.com', 'primary': True},
                                                                      {'address': 'ap@condo-example.com'}]), P._now_text()))
            conn.commit()
            rows = {r['id']: r for r in P.list_cases(conn, 'open')}
            fu = P.follow_up_for(conn, rows[full])
            self.assertEqual((fu['name'], fu['email'], fu['source']), ('Bridget Wooten', 'bwooten@northland.com', 'account'))
            fu = P.follow_up_for(conn, rows[part])
            self.assertEqual((fu['name'], fu['phone'], fu['email'], fu['email_from']),
                             ('Michael', '973-615-1427', 'ap@condo-example.com', 'jobber_ap'))
            fu = P.follow_up_for(conn, rows[none])
            self.assertEqual((fu['name'], fu['email']), ('', ''))
        finally:
            conn.close()

    def test_follow_the_quote_already_in_jobber_instead_of_the_apps_draft(self):
        """The Carlise: the app drafted #9146 under 'The Carlise', but the real
        quote is #9136 under Greenscapes, already out with them."""
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'title': 'Yarrow Bend - Field service', 'client_name': 'Yarrow Bend',
                                       'vendor': 'Wettech'}, 'test')
            P._set_step(conn, cid, 'vendor_quote', at='2026-10-05')
            conn.execute('UPDATE pump_cases SET jobber=? WHERE id=?', (json.dumps(
                {'quote': {'id': 'QD9146', 'number': '9146', 'status': 'draft', 'uri': 'https://j/q/9146'}}), cid))
            conn.execute("INSERT INTO pump_docs (kind, status, vendor, doc_number, case_id, jobber) VALUES "
                         "('quote', 'filed', 'Wettech', 'Q-1', ?, ?)",
                         (cid, json.dumps({'quote_id': 'QD9146', 'quote_number': '9146'})))
            conn.commit()
        finally:
            conn.close()
        searched = []

        def jobber(query, variables=None):
            if 'quotes(searchTerm' in query:
                searched.append(variables['q'])
                return {'quotes': {'nodes': [
                    {'id': 'QX91360', 'quoteNumber': 91360, 'quoteStatus': 'draft'},
                    {'id': 'Q9136', 'quoteNumber': 9136, 'quoteStatus': 'awaiting_response', 'title': 'Proposal',
                     'jobberWebUri': 'https://j/q/9136', 'amounts': {'total': 1506.69},
                     'client': {'id': 'CGREEN', 'name': 'Greenscapes'}, 'property': {'id': 'PCARL'}}]}}
            return {}
        real, real_status = P.jobber_gql, P.jobber_status
        P.jobber_gql, P.jobber_status = jobber, lambda: {'connected': True}
        try:
            r = self.c.post(f'/pumps/api/cases/{cid}/jobber_quote', json={'number': '#9136'})
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
            self.assertEqual(self.c.post(f'/pumps/api/cases/{cid}/jobber_quote', json={'number': '7777'}).status_code, 400)
        finally:
            P.jobber_gql, P.jobber_status = real, real_status
        self.assertEqual(searched[0], '9136')
        case = self.case(cid)
        self.assertEqual(case['jobber']['quote']['number'], '9136')
        self.assertEqual(case['jobber']['client']['name'], 'Greenscapes')
        self.assertEqual(case['jobber_client_id'], 'CGREEN')
        self.assertEqual(case['stage'], 'client_approved', 'out with the client: waiting on their approval')
        doc = next(d for d in case['docs'] if d['kind'] == 'quote')
        self.assertEqual(doc['jobber']['quote_id'], 'Q9136')
        todos = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertIn('Delete the duplicate draft quote #9146 in Jobber', todos)

    def test_linked_quote_status_is_fetched_by_id(self):
        """Jobber's search can miss our own quote (it doesn't search titles), so
        a quote sent to the client still read 'draft'. Linked records are now
        fetched by id on each sync."""
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'title': 'Zephyr Pointe - Field service', 'client_name': 'Zephyr Pointe',
                                       'vendor': 'Wettech'}, 'test')
            P._set_step(conn, cid, 'vendor_quote', at='2026-10-01')
            conn.execute('UPDATE pump_cases SET jobber=? WHERE id=?', (json.dumps(
                {'quote': {'id': 'Q9146', 'number': '9146', 'status': 'draft', 'total': 1506.69}}), cid))
            conn.commit()
        finally:
            conn.close()
        calls = []

        def jobber(query, variables=None):
            calls.append(variables)
            if 'quote(id' in query:
                return {'quote': {'id': 'Q9146', 'quoteNumber': 9146, 'quoteStatus': 'awaiting_response',
                                  'amounts': {'total': 1612.16}, 'updatedAt': '2026-10-06T15:00:00Z'}}
            return {}
        real = P.jobber_gql
        P.jobber_gql = jobber
        conn = P._conn()
        try:
            self.assertEqual(P._refresh_linked_records(conn), [])
            P._follow_quotes(conn)
            conn.commit()
        finally:
            conn.close()
            P.jobber_gql = real
        self.assertIn({'id': 'Q9146'}, calls)
        case = self.case(cid)
        self.assertEqual(case['jobber']['quote']['status'], 'awaiting_response')
        self.assertEqual(case['jobber']['quote']['total'], 1612.16)
        self.assertEqual(case['stage'], 'client_approved', 'sent to the client: now waiting on their approval')

    def test_homeowners_pump_work_in_jobber_is_not_followed(self):
        """Only our accounts (the sheets) and companies show as new Jobber work;
        a homeowner (a person in Jobber, not on the sheets) doesn't."""
        rows = [('H1', 'Pump not starting', 'John Smith', 0),            # homeowner
                ('H2', 'Pump service call', 'Barrington Cove HOA', 0),   # on the sheets
                ('H3', 'PO 323 Pump Service', 'Homewood Suites', 1),     # a company
                ('H4', 'Pump check', 'Jane Doe', None)]                  # not known yet
        conn = P._conn()
        try:
            for jid, title, client, company in rows:
                conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, job_type, "
                             "client_name, created_at, category, client_company) VALUES "
                             "(?, 'job', ?, ?, 'upcoming', 'one_off', ?, '2026-10-07T12:00:00Z', 'pump', ?)",
                             (jid, jid, title, client, company))
            conn.commit()
            shown = {r['jobber_id'] for r in P.untracked_jobber_items(conn)}
            hidden = {r['jobber_id'] for r in P.untracked_jobber_items(conn, others=True)}
        finally:
            conn.close()
        self.assertEqual(shown & {'H1', 'H2', 'H3', 'H4'}, {'H2', 'H3', 'H4'})
        self.assertEqual(hidden & {'H1', 'H2', 'H3', 'H4'}, {'H1'})
        items = {i['jobber_id']: i for i in self.c.get('/pumps/api/jobber/items?open=1&kind=job').get_json()['items']}
        self.assertFalse(items['H1']['on_sheet'])
        self.assertTrue(items['H2']['on_sheet'])

    def test_email_wettech_asks_for_what_the_job_is_waiting_on(self):
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'title': 'Homewood Suites - PO 323 Pump Service', 'client_name': 'Homewood Suites',
                                       'po_number': '323', 'vendor': 'Wettech'}, 'test')
            P._set_assessment(conn, cid, '2026-10-13')
            conn.commit()
        finally:
            conn.close()

        def mail():
            return self.c.get(f'/pumps/api/cases/{cid}/vendor_email').get_json()['email']
        e = mail()
        self.assertEqual(e['to'], 'tomm@wettec.biz')
        self.assertTrue(e['body'].startswith('Hi Tom,'))
        self.assertIn('assess', e['body'])
        self.assertIn('Visit: 2026-10-13', e['body'])
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {'assessment': '2026-10-13'}})
        e = mail()
        self.assertTrue(e['subject'].startswith('Please quote: Homewood Suites - PO 323'), e['subject'])
        self.assertIn('send us your quote', e['body'])
        self.c.patch(f'/pumps/api/cases/{cid}', json={'vendor_bill_amount': 410, 'vendor_bill_number': '30501'})
        e = mail()
        self.assertIn('the quote that goes with it', e['body'])
        self.assertIn('Your invoice: #30501', e['body'])
        # Sent only when Send is clicked, from the PO mailbox, as edited.
        sent, real = [], P.graph_send
        P.graph_send = lambda to, subject, html, attachments=None, sender=None: sent.append((to, subject, html))
        try:
            self.assertEqual(sent, [])
            r = self.c.post(f'/pumps/api/cases/{cid}/vendor_email/send',
                            json={'to': 'tomm@wettec.biz', 'subject': 'Quote please', 'body': 'Hi Tom,\nQuote please.'})
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
            self.assertEqual(sent[0][:2], ('tomm@wettec.biz', 'Quote please'))
            self.assertIn('Quote please.', sent[0][2])
            self.assertEqual(self.c.post(f'/pumps/api/cases/{cid}/vendor_email/send', json={'to': ''}).status_code, 400)
        finally:
            P.graph_send = real
        self.assertTrue(any(ev['action'] == 'emailed the vendor' for ev in self.case(cid)['events']))

    def test_one_time_fixups_run_once_with_two_workers(self):
        conn = P._conn()
        try:
            conn.execute("DELETE FROM pump_state WHERE key='claim-test'")
            conn.commit()
            self.assertTrue(P._claim_once(conn, 'claim-test'))
            self.assertFalse(P._claim_once(conn, 'claim-test'), 'the second worker must not run it again')
            cid = P.create_case(conn, {'title': 'Twice noted', 'client_name': 'X'}, 'test')
            for _ in range(2):
                P._event(conn, 'Pumps', 'waiting on the vendor visit', 'Made in Jobber before any quote', case_id=cid)
            P._event(conn, 'Pumps', 'waiting on the vendor visit', 'A different note', case_id=cid)
            conn.execute("DELETE FROM pump_state WHERE key='startup_events_deduped'")
            conn.commit()
            P._dedupe_startup_events(conn)
            rows = conn.execute("SELECT detail FROM pump_events WHERE case_id=? AND action='waiting on the vendor visit'",
                                (cid,)).fetchall()
            self.assertEqual(sorted(r[0] for r in rows), ['A different note', 'Made in Jobber before any quote'])
        finally:
            conn.close()

    def test_the_apps_own_emails_are_not_service_calls(self):
        """The daily email ("Pumps today - 1 to do") names accounts and says
        "pump"; read back from the PO mailbox it must not open a service call."""
        for subject in ('Pumps today - 1 to do', 'FW: Pumps today - nothing urgent', 'Please pay Wettech invoice #29110',
                        'Diver schedule - October 2026 - Stahlman-England'):
            self.assertTrue(P.from_this_app('someone@example.com', subject), subject)
        self.assertFalse(P.from_this_app('manager@example.com', 'Pump at the front lake is down'))
        P.CFG['mail_from'], old = 'po@example.com', P.CFG.get('mail_from')
        try:
            self.assertTrue(P.from_this_app('PO <po@example.com>', 'anything'))
            self.assertIsNone(P.handle_service_call_email('u-own-1', 'po@example.com', 'Pumps today - 2 to do',
                                                          'Barrington Cove pump station - waiting on quote'))
        finally:
            P.CFG['mail_from'] = old
        # Ones opened before this was fixed are cancelled once, with their to-do.
        conn = P._conn()
        try:
            bad = P.create_case(conn, {'title': 'Service call - Pumps today - 1 to do', 'client_name': 'Barrington Cove'},
                                'email scan', source='service_call')
            real = P.create_case(conn, {'title': 'Service call - Front lake pump down', 'client_name': 'Barrington Cove'},
                                 'email scan', source='service_call')
            P.add_todo(conn, f'service-visit:{bad}', 'Add a "Service call - Pumps today - 1 to do" visit')
            conn.execute("DELETE FROM pump_state WHERE key='own_email_calls_removed_2'")
            conn.commit()
            P._remove_own_email_service_calls(conn)
            self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (bad,)).fetchone()[0], 'cancelled')
            self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (real,)).fetchone()[0], 'open')
            self.assertIsNotNone(conn.execute('SELECT done_at FROM pump_todos WHERE key=?', (f'service-visit:{bad}',)).fetchone()[0])
        finally:
            conn.close()

    def test_find_a_job_by_any_quote_or_invoice_number(self):
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'title': 'Lift station repair', 'client_name': 'Quail Run', 'vendor': 'Wettech',
                                       'vendor_quote_number': 'Q-7731', 'vendor_bill_number': '30455',
                                       'sei_invoice_number': '36999'}, 'test')
            conn.execute("UPDATE pump_cases SET jobber=? WHERE id=?",
                         (json.dumps({'quote': {'id': 'Z1', 'number': '9188', 'total': 1235.0}}), cid))
            conn.execute("INSERT INTO pump_docs (kind, status, vendor, doc_number, case_id) "
                         "VALUES ('report', 'filed', 'Wettech', 'WO-5512', ?)", (cid,))
            conn.commit()
        finally:
            conn.close()
        for q in ('Q-7731', '#30455', '36999', '9188', 'WO-5512'):
            ids = [c['id'] for c in self.c.get('/pumps/api/cases?status=all&q=' + q.replace('#', '%23')).get_json()['cases']]
            self.assertIn(cid, ids, q)
        ids = [c['id'] for c in self.c.get('/pumps/api/cases?status=all&q=Z1').get_json()['cases']]
        self.assertNotIn(cid, ids, "Jobber's internal ids don't match")
        self.assertEqual(self.case(cid)['jobber']['quote']['total'], 1235.0)

    def test_sync_left_running_by_a_restart_reads_as_interrupted(self):
        P._state_set('jobber_sync', {'state': 'running', 'started_at': P._now_text(), 'errors': []})
        st = self.c.get('/pumps/api/jobber/items').get_json()['sync']
        self.assertEqual(st['state'], 'interrupted')
        self.assertTrue(st['errors'])
        # While a sync really holds the lock, it still reads as running.
        P._state_set('jobber_sync', {'state': 'running', 'started_at': P._now_text(), 'errors': []})
        with P._FileLock('jobber_sync', blocking=False):
            self.assertEqual(P.jobber_sync_state()['state'], 'running')

    def test_failed_sync_does_not_stay_running(self):
        real = P._sync_jobber
        def boom(full=False, actor='system'):
            P._state_set('jobber_sync', {'state': 'running', 'started_at': P._now_text(), 'errors': []})
            raise RuntimeError('database is locked')
        P._sync_jobber = boom
        try:
            with self.assertRaises(RuntimeError):
                P.sync_jobber()
        finally:
            P._sync_jobber = real
        st = P._state_get('jobber_sync')
        self.assertEqual(st['state'], 'failed')
        self.assertIn('database is locked', st['errors'][-1])

    def test_job_made_in_jobber_shows_up_to_track(self):
        """A pump job made straight in Jobber (no request) is listed to track, and
        tracking it makes an item that is already scheduled."""
        conn = P._conn()
        try:
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, job_type, client_id, "
                         "client_name, property_id, created_at, start_at, category, po_number) VALUES "
                         "('J14542', 'job', '14542', 'Homewood Suites - PO 9323 Pump Service', 'upcoming', 'one_off', "
                         "'CHW', 'Homewood Suites', 'PHW', '2026-10-05T12:00:00Z', '2026-10-13T08:00:00Z', 'pump', '9323')")
            conn.commit()
        finally:
            conn.close()
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn('J14542', [i['jobber_id'] for i in queue['new_jobber_requests']])
        body = self.c.get('/pumps/api/jobber/items?open=1&kind=request,quote,job').get_json()
        items = body['items']
        self.assertIn('J14542', [i['jobber_id'] for i in items])
        cid = self.c.post('/pumps/api/jobber/items/J14542', json={'action': 'track'}).get_json()['case_id']
        case = self.case(cid)
        self.assertEqual(case['jobber']['job']['number'], '14542')
        # It waits on Wettech's visit to assess (Oct 13), not on a quote or a schedule.
        self.assertEqual(case['stage'], 'assessment')
        steps = {s['key']: s for s in case['step_list']}
        self.assertEqual(steps['assessment'].get('due'), '2026-10-13')
        self.assertFalse(steps['scheduled'].get('at'))
        self.assertFalse(steps['client_approved'].get('at'))
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn(cid, [c['id'] for c in queue['waiting_assessment']])
        # Wettech's quote coming in means the visit happened.
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {'vendor_quote': '2026-10-14'}})
        case = self.case(cid)
        self.assertEqual(case['stage'], 'client_quote')
        self.assertTrue({s['key']: s for s in case['step_list']}['assessment'].get('at'))
        # Items that never had the visit step don't get it.
        other = self.c.post('/pumps/api/cases', json={'title': 'Plain repair', 'client_name': 'X'}).get_json()
        self.assertEqual(self.case(other.get('case_id') or other['case']['id'])['stage'], 'vendor_quote')
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertNotIn('J14542', [i['jobber_id'] for i in queue['new_jobber_requests']])
        # A quote Jobber makes later for the same client and property joins that item by itself.
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET jobber_property_id='PHW' WHERE id=?", (cid,))
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, client_id, "
                         "property_id, created_at, category, po_number) VALUES ('Q14600', 'quote', '14600', "
                         "'Extra pump work', 'awaiting_response', 'CHW', 'PHW', '2026-10-07T12:00:00Z', 'pump', '')")
            P._link_jobber_items(conn)
            conn.commit()
        finally:
            conn.close()
        self.assertEqual(self.case(cid)['jobber']['quote']['number'], '14600')
        # A quote sent from Jobber and still out with the client isn't listed; once approved it is.
        conn = P._conn()
        try:
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, client_id, "
                         "created_at, category, po_number) VALUES ('Q14700', 'quote', '14700', 'Pump repair', "
                         "'awaiting_response', 'COTHER', '2026-10-07T12:00:00Z', 'pump', '')")
            conn.commit()
            self.assertNotIn('Q14700', [i['jobber_id'] for i in P.untracked_jobber_items(conn)])
            conn.execute("UPDATE pump_jobber_items SET status='approved' WHERE jobber_id='Q14700'")
            conn.commit()
            self.assertIn('Q14700', [i['jobber_id'] for i in P.untracked_jobber_items(conn)])
        finally:
            conn.close()

    def test_draft_quote_asks_which_property(self):
        self.extracts['q-0810.pdf'] = extraction('quote', 'Q-121', po='PO810', client='Lakeside Pines', subtotal=300)
        q = self.upload('q-0810.pdf')
        P.jobber_gql = FakeJobber(properties=[{'id': 'P1', 'address': {'street1': 'North gate'}},
                                              {'id': 'P2', 'address': {'street1': 'South gate'}}])
        r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/quote", json={'client_id': 'C1'})
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.get_json()['needs_property'])
        self.assertEqual([p['id'] for p in r.get_json()['properties']], ['P1', 'P2'])
        r = self.c.post(f"/pumps/api/docs/{q['doc_id']}/quote", json={'client_id': 'C1', 'property_id': 'P2'})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))

    def test_non_draft_quote_status_raises_issue(self):
        self.extracts['q-0820.pdf'] = extraction('quote', 'Q-122', po='PO820', client='Sunset Cove', subtotal=400)
        q = self.upload('q-0820.pdf')
        P.jobber_gql = FakeJobber(quote_status='awaiting_response')
        body = self.c.post(f"/pumps/api/docs/{q['doc_id']}/quote", json={'client_id': 'C9', 'property_id': 'P9'}).get_json()
        self.assertTrue(body['success'], body)
        self.assertTrue(any(i['kind'] == 'not_draft' and not i['resolved_at'] for i in self.case(q['case_id'])['issues']))

    def test_draft_quote_like_jobber_quote_9136(self):
        """Tom's Carlise quote (Oct 5 2026), as Simon then entered it by hand: Greenscapes, the Carlisle Pump #1
        exit ("back station"), "Proposal to inspect suction line", one Service Proposal Amount line of
        $1,506.69 (Wettech's $1,158.99 tax included, plus 30%), not taxable."""
        work = ('Field service to check out pump station, found pipe had melted at the fitting going into the pump '
                'suction. Field service to pull and inspect suction line, clean screen and reinstall, furnish ans '
                'install PVC parts needed to repair suction line, prime and test.')
        name = 'Carslie Back Station .docx'
        self.extracts[name] = extraction('quote', '', client='The Carlise', items=[
            {'name': 'Pump station repair', 'description': work, 'quantity': 1, 'unit_price': 1158.99,
             'amount': 1158.99, 'taxable': False, 'is_tax': False}]) | {
            'subtotal': None, 'tax': None, 'total': 1158.99, 'tax_included': True,
            'proposal_title': 'Proposal to inspect suction line', 'doc_date': '2026-10-05'}
        fake = FakeJobber(quote_lines=[{'name': 'Service Proposal Amount', 'description': work, 'quantity': 1,
                                         'unitPrice': 1506.69, 'taxable': False}])
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            res = self.upload(name, data=b'docx bytes')
            self.assertIsNone(res['auto_quote'], 'nothing drafted on its own')
            self.assertEqual(self.draft(res['doc_id'])['quote']['quote_number'], '812')
            attrs = [v for q, v in fake.calls if 'quoteCreate(' in q][0]['attributes']
            self.assertEqual(attrs['clientId'], 'Z2lkOi8vSm9iYmVyL0NsaWVudC80OTAwNzI1OA==', '"Carlisle" is the Greenscapes account')
            self.assertEqual(attrs['propertyId'], 'Z2lkOi8vSm9iYmVyL1Byb3BlcnR5LzUyOTc4MjYw', '"back station" is the Pump #1 exit')
            self.assertEqual(attrs['title'], 'Proposal to inspect suction line')
            self.assertEqual(len(attrs['lineItems']), 1)
            li = attrs['lineItems'][0]
            self.assertEqual((li['name'], li['unitPrice'], li['quantity'], li['taxable']),
                             ('Service Proposal Amount', 1506.0, 1.0, False))
            self.assertEqual(li['description'], work)
            note = [v for q, v in fake.calls if 'quoteCreateNote' in q][0]
            self.assertEqual(note['id'], 'Q1')
            self.assertIn('Wettech quote', note['input']['message'])
            self.assertIn('$1,158.99 (includes sales tax)', note['input']['message'])
            self.assertIn('$1,506.00 (30% markup)', note['input']['message'])
            self.assertFalse(any(w in q for q, _ in fake.calls for w in ('Send', 'MarkAsSent')))
            queue = self.c.get('/pumps/api/summary').get_json()['queue']
            self.assertNotIn(res['doc_id'], [d['id'] for d in queue['quotes_to_draft']])
            # Wettech's bill matches its quote: the client is invoiced what they were quoted.
            self.extracts['b-carlise.pdf'] = extraction('bill', '29150', client='The Carlise', subtotal=1088.25,
                                                        tax=70.74)
            b = self.upload('b-carlise.pdf')
            self.assertEqual(b['case_id'], res['case_id'])
            sugg = self.c.get(f"/pumps/api/docs/{b['doc_id']}").get_json()['doc']['invoice_suggestion']
            self.assertEqual(sugg['from_quote'], '812')
            self.assertEqual([(l['name'], l['unit_price']) for l in sugg['line_items']],
                             [('Service Proposal Amount', 1506.69)])
        finally:
            P.JOBBER_STATIC_TOKEN = saved

    def test_carlise_quote_without_claude(self):
        """No Anthropic key: the built-in reader alone turns Tom's emailed Word quote into the 9136 draft."""
        self.assertFalse(P.USE_CLAUDE)
        from docx import Document
        d = Document()
        for t in ('of Southwest Florida LLC', 'Email: wettec@verizon.net', 'October 5, 2026', 'Stahlman England',
                  'Attn: Andrea', 'RE: The Carlise', 'We are pleased to quote you on the following services',
                  'Field service to check out pump station, found pipe had melted at the fitting going into the pump '
                  'suction.  Field service to pull and inspect suction line, clean screen and reinstall, furnish ans '
                  'install PVC parts needed to repair suction line, prime and test.',
                  'Your Cost --------------- $ 1158.99', 'Price includes Sales tax and in freight', 'Thank You'):
            d.add_paragraph(t)
        buf = io.BytesIO()
        d.save(buf)
        fake = FakeJobber()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            form = {'file': (io.BytesIO(buf.getvalue()), 'Carslie Back Station .docx')}
            res = self.c.post('/pumps/api/docs', data=form, content_type='multipart/form-data').get_json()['results'][0]
            self.assertEqual(res['review'], '', 'a Wettech quote read cleanly is trusted')
            self.assertTrue(res['case_id'])
            self.assertEqual(self.draft(res['doc_id'])['quote']['quote_number'], '812')
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        attrs = [v for q, v in fake.calls if 'quoteCreate(' in q][0]['attributes']
        self.assertEqual((attrs['clientId'], attrs['propertyId']), ('Z2lkOi8vSm9iYmVyL0NsaWVudC80OTAwNzI1OA==', 'Z2lkOi8vSm9iYmVyL1Byb3BlcnR5LzUyOTc4MjYw'))
        self.assertEqual(attrs['title'], 'Proposal to pull and inspect suction line')
        li = attrs['lineItems'][0]
        self.assertEqual((li['name'], li['unitPrice'], li['taxable']), ('Service Proposal Amount', 1506.0, False))
        self.assertTrue([v for q, v in fake.calls if 'quoteCreateNote' in q])
        self.c.post(f"/pumps/api/cases/{res['case_id']}/delete", json={'reason': 'test done'})

    def test_forwarded_quote_with_letterhead_in_page_header(self):
        """Simon forwards Tom's quote to PO@: the sender is one of us and Wettech's details are only in the
        Word page header - it must still be read as a Wettech quote."""
        from docx import Document
        d = Document()
        hdr = d.sections[0].header.paragraphs[0]
        hdr.text = 'Water Equipment Technologies of Southwest Florida LLC - Email: wettec@verizon.net'
        for t in ('October 6, 2026', 'Stahlman England', 'RE: Egret Landing',
                  'We are pleased to quote you on the following services',
                  'Field service to replace check valve on pump #1, prime and test.',
                  'Your Cost --------------- $ 840.00', 'Price includes Sales tax and in freight'):
            d.add_paragraph(t)
        buf = io.BytesIO()
        d.save(buf)
        text = P.extract_text('Egret.docx', buf.getvalue())
        self.assertIn('wettec@verizon.net', text)
        res = P.ingest_document('Egret Landing pump.docx', buf.getvalue(), source='email', actor='email scan',
                                email={'uid': 'fw-1', 'from': 'simon@stahlman-england.com',
                                       'subject': 'FW: Egret Landing', 'preview': ''})
        self.assertEqual((res['kind'], res['review']), ('quote', ''), res)
        doc = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertEqual((doc['vendor'], doc['client_name'], doc['total']), ('Wettech', 'Egret Landing', 840.0))
        # A PDF forward whose only Wettech mark is the forwarded header in the email body.
        self.texts['fw-invoice.pdf'] = 'INVOICE\nBill To Stahlman-England\nRE: Egret Landing\nTotal $840.00'
        res = P.ingest_document('fw-invoice.pdf', b'%PDF fw', source='email', actor='email scan',
                                email={'uid': 'fw-2', 'from': 'simon@stahlman-england.com', 'subject': 'FW: Invoice',
                                       'preview': 'From: Tom Morgan <tomm@wettec.biz> Sent: Monday'})
        doc = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertEqual(doc['vendor'], 'Wettech')

    def test_client_found_through_its_property_despite_typos(self):
        props = [{'id': 'PX1', 'address': {'street1': '1 Osprey Ct', 'street2': 'Osprey Point - Pump #1', 'city': 'Naples'}},
                 {'id': 'PX2', 'address': {'street1': '2 Osprey Ct', 'street2': 'Osprey Point - Pump #2', 'city': 'Naples'}}]
        fake = FakeJobber(properties=props, clients=[
            {'id': 'LS1', 'name': 'Sunrise Landscaping', 'isLead': False, 'isArchived': False, 'properties': props},
            {'id': 'XX', 'name': 'Bayview Towers', 'isLead': False, 'isArchived': False, 'properties': []}])
        P.jobber_gql = fake
        cands = P.search_clients('The Ospray Pointe')
        self.assertEqual(cands[0]['id'], 'LS1', 'the community is a property of the landscape company')
        self.assertEqual(len(cands[0]['matching_properties']), 2)
        t = P.resolve_quote_target({'client_name': 'Ospray Pointe', 'site': 'Pump #2', 'file_name': 'q.pdf'})
        self.assertEqual((t['client_id'], t['property_id']), ('LS1', 'PX2'))
        t = P.resolve_quote_target({'client_name': 'Ospray Pointe', 'file_name': 'q.pdf'})
        self.assertEqual((t['client_id'], t['property_id']), ('LS1', None), 'two pumps, no hint: a person picks')

    def test_approved_client_quote_is_not_waiting_on_the_vendor_quote(self):
        """Carlisle (Oct 2026): our quote #9136 approved, no Wettech quote filed
        here - not "Waiting on quote from Wettech"."""
        self.assertEqual(P._stage({'client_quote': {'at': '2026-10-01'}, 'client_approved': {'at': '2026-10-02'}}),
                         'scheduled')
        self.assertEqual(P._stage({'client_quote': {'at': '2026-10-01'}}), 'client_approved')
        self.assertEqual(P._stage({}), 'vendor_quote')
        self.assertEqual(P._stage({'vendor_bill': {'at': '2026-10-01'}}), 'vendor_quote', 'bill before quote still waits')

    def test_parts_lead_waits_then_follow_up_after_14_days(self):
        """A quote saying "2-3 weeks": once approved the job waits on parts and
        comes up to follow up with Tom 14 days after the approval."""
        self.assertEqual(P.parts_lead_in('Delivery: 2-3 weeks after receipt of order'), '2-3 weeks')
        self.assertEqual(P.parts_lead_in('It will take 2-3 weeks to deliver after approval.'), '2-3 weeks')
        self.assertEqual(P.parts_lead_in('Delivery: 3-5 days after receipt of order'), '')
        case = {'id': 1, 'status': 'open', 'stage': 'scheduled', 'category': 'repair', 'vendor': 'Wettech',
                'parts_lead': '2-3 weeks', 'opened_on': '2026-09-20',
                'steps': {'client_approved': {'at': '2026-10-01'}}}
        ch = P.vendor_chase(case, today=P.date(2026, 10, 14))
        self.assertEqual((ch['state'], ch['check_on'], ch['parts']), ('waiting', '2026-10-15', '2-3 weeks'))
        ch = P.vendor_chase(case, today=P.date(2026, 10, 15))
        self.assertEqual(ch['state'], 'follow_up')
        self.assertIn('parts', ch['reason'])
        # No lead time on the quote: still waiting on the parts, followed up the same.
        ch = P.vendor_chase({**case, 'parts_lead': ''}, today=P.date(2026, 10, 2))
        self.assertEqual((ch['state'], ch['parts_wait'], ch['check_on']), ('waiting', True, '2026-10-15'))
        self.assertIn('waiting on parts', ch['reason'])
        # Our own technician has no parts to wait on: schedule straight away.
        ch = P.vendor_chase({**case, 'vendor': P.IN_HOUSE}, today=P.date(2026, 10, 2))
        self.assertEqual((ch['state'], ch['parts_wait']), ('follow_up', False))

    def test_approved_quote_sent_by_hand_waits_on_tom(self):
        """Huntington #6 (quote 9160, Oct 2026): the office emailed the approved
        quote to Tom themselves - the to-do is done and the job waits on Tom's
        reply about when it'll be scheduled."""
        self.extracts['q-0916.pdf'] = extraction('quote', 'Q-916', po='PO916', client='Huntington Lakes')
        q = self.upload('q-0916.pdf')
        cid = q['case_id']
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {'client_quote': '2026-10-09',
                                                                'client_approved': '2026-10-14'}})
        conn = P._conn()
        try:
            P.add_sign_todo(conn, dict(conn.execute('SELECT * FROM pump_cases WHERE id=?', (cid,)).fetchone()))
            conn.commit()
        finally:
            conn.close()
        case = self.case(cid)
        self.assertEqual(case['stage'], 'scheduled')
        self.assertFalse(case['chase']['last_at'], 'not sent yet')
        r = self.c.post(f'/pumps/api/cases/{cid}/quote_sent', json={}).get_json()
        self.assertTrue(r['success'], r)
        case = self.case(cid)
        ch = case['chase']
        self.assertTrue(ch['last_at'])
        self.assertEqual(ch['state'], 'waiting')
        self.assertIn("when it'll be scheduled", ch['said'])
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        self.assertFalse([t for t in todos if t['key'] == f'sign-quote:{cid}'], 'the send-it-to-Tom to-do is done')
        doc = self.c.get(f"/pumps/api/docs/{q['doc_id']}").get_json()['doc']
        self.assertTrue(doc['jobber']['sent_by_hand'])

    def test_waiting_on_quote_catches_up_with_a_quote_sent_in_jobber(self):
        """Magnolia Cove (Oct 2026): "Quote is needed" kept waiting on Wettech's
        quote after quote #9186 went to the client from another job - it's
        merged into that job. With no job behind the quote, it's linked."""
        conn = P._conn()
        try:
            b = P.create_case(conn, {'title': 'Quote is needed', 'client_name': 'Magnolia Cove',
                                     'opened_on': '2026-10-07'}, 'test')
            a = P.create_case(conn, {'title': 'Pressure transducer', 'client_name': 'Magnolia Cove',
                                     'opened_on': '2026-10-13',
                                     'jobber': {'quote': {'id': 'QM9186', 'number': '9186'}}}, 'test')
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, client_id, "
                         "client_name, total, created_at) VALUES ('QM9186', 'quote', '9186', 'Proposal - transducer', "
                         "'awaiting_response', 'CMAG', 'MAGNOLIA COVE', 1022, '2026-10-14T16:28:48Z')")
            # A job opened after the quoting job, at the same client, is new work: left alone.
            later = P.create_case(conn, {'title': 'Motor', 'client_name': 'Magnolia Cove',
                                         'opened_on': '2026-10-14'}, 'test')
            done = P.reconcile_quote_waits(conn)
            self.assertIn(('merged', b, a), done)
            self.assertNotIn(later, [x[1] for x in done])
            self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (b,)).fetchone()[0], 'cancelled')
            self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (a,)).fetchone()[0], 'open')
            # And Wettech's quote read as "Magnolia Cove Pressure Transducer" files on that job, not a new one.
            e = P.create_case(conn, {'title': 'Quote is needed', 'client_name': 'Lakewood Ranch',
                                     'opened_on': '2026-10-07'}, 'test')
            self.assertEqual(P._find_case_for(conn, {'kind': 'quote', 'client_name': 'Lakewood Ranch Pressure '
                                                                                     'Transducer'})[0], e)
            # A quote made by hand in Jobber, on no job: linked to the job waiting on it.
            c = P.create_case(conn, {'title': 'Quote is needed', 'client_name': 'Pelican Sound Estates',
                                     'opened_on': '2026-10-07'}, 'test')
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, client_id, "
                         "client_name, total, created_at) VALUES ('QB1', 'quote', '9190', 'Proposal - valve', "
                         "'awaiting_response', 'CPEL', 'Pelican Sound Estates', 900, '2026-10-14T10:00:00Z')")
            self.assertIn(('linked', c, '9190'), P.reconcile_quote_waits(conn))
            P._follow_quotes(conn)
            self.assertEqual(conn.execute('SELECT stage FROM pump_cases WHERE id=?', (c,)).fetchone()[0],
                             'client_approved', 'quote sent: waiting on the client, not on Wettech')
            # An older quote (before the job opened) is not this work.
            d = P.create_case(conn, {'title': 'Quote is needed', 'client_name': 'Sunset Cove',
                                     'opened_on': '2026-10-10'}, 'test')
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, client_id, "
                         "client_name, total, created_at) VALUES ('QSUN8000', 'quote', '8000', 'Old', "
                         "'awaiting_response', 'CSUN', 'Sunset Cove', 500, '2026-09-01T10:00:00Z')")
            self.assertNotIn(d, [x[1] for x in P.reconcile_quote_waits(conn)])
            conn.rollback()
        finally:
            conn.close()

    def test_approved_quote_is_emailed_to_the_office_once(self):
        """Quote 9160 (Oct 2026): an approval is emailed to the office straight
        away - send the approved quote to Tom - once per job."""
        self.extracts['q-0931.pdf'] = extraction('quote', 'Q-931', po='PO931', client='Huntington Lakes')
        q = self.upload('q-0931.pdf')
        cid = q['case_id']
        conn = P._conn()
        try:
            conn.execute('UPDATE pump_cases SET jobber=? WHERE id=?', (json.dumps(
                {'quote': {'id': 'Q9160', 'number': '9160', 'total': 1884.14, 'uri': 'https://j/q/9160'}}), cid))
            conn.commit()
        finally:
            conn.close()
        sent = []
        P.graph_send, saved = (lambda to, subject, html, atts=None, sender=None: sent.append((to, subject, html))), P.graph_send
        P.jobber_gql = FakeJobber()
        try:
            P.on_quote_approved(cid)
            P.on_quote_approved(cid)
        finally:
            P.graph_send = saved
        self.assertEqual(len(sent), 1, 'once per job')
        # The mailbox scan doesn't read the approval email as a service call (Lely, quote 9152).
        self.assertTrue(P.from_this_app('Simon <simon@stahlman-england.com>', sent[0][1]))
        self.assertIsNone(P.handle_service_call_email('m9152', 'simon@stahlman-england.com', sent[0][1],
                                                      'Lely approved our quote #9152. Next: send the approved pump quote'))
        to, subject, html = sent[0]
        self.assertEqual(to, P.DIGEST_TO)
        self.assertIn('Quote #9160 approved - Huntington Lakes', subject)
        self.assertIn('I already sent it', html)
        self.assertIn('https://j/q/9160', html)

    def test_quick_check_sees_an_approval_without_the_full_sync(self):
        """Huntington #5 (quote 9149, Oct 2026): approved in Jobber - the
        5-minute check reads our open quotes straight from Jobber, ticks the
        approval, puts "send it to Tom" on To do and emails the office."""
        self.extracts['q-0949.pdf'] = extraction('quote', 'Q-949', po='PO949', client='Huntington Five')
        q = self.upload('q-0949.pdf')
        cid = q['case_id']
        conn = P._conn()
        try:
            conn.execute('UPDATE pump_cases SET jobber=? WHERE id=?', (json.dumps(
                {'quote': {'id': 'Q9149', 'number': '9149', 'status': 'awaiting_response', 'total': 1760.02}}), cid))
            conn.commit()
        finally:
            conn.close()

        class Fake(FakeJobber):
            def __call__(self, query, variables=None):
                if 'quote(id' in query and (variables or {}).get('id') == 'Q9149':
                    self.calls.append((query, variables))
                    return {'quote': {'id': 'Q9149', 'quoteNumber': 9149, 'quoteStatus': 'approved',
                                      'amounts': {'total': 1760.02}, 'jobberWebUri': 'https://j/q/9149',
                                      'updatedAt': '2026-10-14T14:25:18Z'}}
                return super().__call__(query, variables)
        sent = []
        P.jobber_gql = Fake()
        P.JOBBER_STATIC_TOKEN, saved_tok = 'test-token', P.JOBBER_STATIC_TOKEN
        P.graph_send, saved = (lambda to, subject, html, atts=None, sender=None: sent.append(subject)), P.graph_send
        try:
            self.assertIn(cid, P.check_quote_approvals())
            self.assertEqual(P.check_quote_approvals(), [], 'seen once')
        finally:
            P.graph_send, P.JOBBER_STATIC_TOKEN = saved, saved_tok
        case = self.case(cid)
        self.assertTrue({s['key']: s for s in case['step_list']}['client_approved'].get('at'))
        self.assertEqual(case['stage'], 'scheduled')
        self.assertEqual(sent, ['Quote #9149 approved - Huntington Five: send it to Tom'])
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        self.assertTrue(any(t['key'] == f'sign-quote:{cid}' for t in todos))

    def test_approval_reopens_a_job_closed_while_its_quote_was_out(self):
        """Huntington #5 / #6 (Oct 2026): their jobs were off the watchlist when
        the client approved quotes 9149 and 9160 - the approval reopens them.
        A finished job whose approval was already ticked stays closed."""
        conn = P._conn()
        try:
            gone = P.create_case(conn, {'title': 'Check out #5 pump station', 'client_name': 'Huntington Reopen',
                                        'jobber': {'quote': {'id': 'QR9149', 'number': '9149',
                                                             'status': 'awaiting_response'}}}, 'test')
            conn.execute("UPDATE pump_cases SET status='cancelled', stage='cancelled' WHERE id=?", (gone,))
            done = P.create_case(conn, {'title': 'Old finished job', 'client_name': 'Huntington Reopen',
                                        'jobber': {'quote': {'id': 'QR8000', 'number': '8000',
                                                             'status': 'awaiting_response'}}}, 'test')
            for k in P.STEP_KEYS:
                P._set_step(conn, done, k, at='2026-10-01', by='test')
            self.assertEqual(conn.execute('SELECT status FROM pump_cases WHERE id=?', (done,)).fetchone()[0], 'closed')
            conn.commit()
        finally:
            conn.close()
        statuses = {'QR9149': 'approved', 'QR8000': 'converted'}

        class Fake(FakeJobber):
            def __call__(self, query, variables=None):
                jid = (variables or {}).get('id')
                if 'quote(id' in query and jid in statuses:
                    return {'quote': {'id': jid, 'quoteNumber': jid[2:], 'quoteStatus': statuses[jid],
                                      'amounts': {'total': 100}, 'jobberWebUri': '', 'updatedAt': '2026-10-14T14:25:00Z'}}
                return super().__call__(query, variables)
        P.jobber_gql = Fake()
        P.JOBBER_STATIC_TOKEN, saved_tok = 'test-token', P.JOBBER_STATIC_TOKEN
        P.graph_send, saved = (lambda *a, **k: None), P.graph_send
        try:
            self.assertIn(gone, P.check_quote_approvals())
        finally:
            P.graph_send, P.JOBBER_STATIC_TOKEN = saved, saved_tok
        case = self.case(gone)
        self.assertEqual((case['status'], case['stage']), ('open', 'scheduled'))
        self.assertIn('reopened', [e['action'] for e in case['events']])
        self.assertEqual(self.case(done)['status'], 'closed', 'a finished job stays closed')

    def test_jobs_show_days_since_the_last_action(self):
        """Oct 2026: each job says how long since anything was done on it, and what."""
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'title': 'Last action test', 'client_name': 'Idle Pines'}, 'test')
            conn.execute("UPDATE pump_events SET at='2026-01-02 09:00:00' WHERE case_id=?", (cid,))
            conn.commit()
        finally:
            conn.close()
        la = next(c for c in self.c.get('/pumps/api/cases?status=open').get_json()['cases'] if c['id'] == cid)['last_action']
        self.assertEqual((la['at'], la['action']), ('2026-01-02', 'created'))
        self.assertEqual(la['days'], (P._today() - P.date(2026, 1, 2)).days)
        self.c.patch(f'/pumps/api/cases/{cid}', json={'notes': 'Called the manager'})
        la = self.case(cid)['last_action']
        self.assertEqual(la['days'], 0, 'an edit is an action')

    def test_quote_drafted_from_an_unfiled_document_gets_a_job(self):
        """Homewood Suites (quote 9176, Oct 2026): drafted from a vendor quote that
        was on no job, so nothing followed it. Drafting now files it on a job;
        and those drafted before are filed once on start-up."""
        self.extracts['q-0976.pdf'] = extraction('quote', 'Q-976', po='PO976', client='Lakeside Pines')
        q = self.upload('q-0976.pdf')
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET status='cancelled' WHERE id=?", (q['case_id'],))
            conn.execute('UPDATE pump_docs SET case_id=NULL WHERE id=?', (q['doc_id'],))
            conn.commit()
        finally:
            conn.close()
        P.jobber_gql = FakeJobber()
        r = self.draft(q['doc_id'], client_id='C1', property_id='P1')
        self.assertTrue(r['success'], r)
        doc = self.c.get(f"/pumps/api/docs/{q['doc_id']}").get_json()['doc']
        self.assertTrue(doc['case_id'], 'filed on a job')
        case = self.case(doc['case_id'])
        self.assertEqual((case['status'], case['jobber']['quote']['number']), ('open', '812'))
        # Drafted before this: filed on start-up, the quote linked.
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET status='cancelled' WHERE id=?", (doc['case_id'],))
            conn.execute('UPDATE pump_docs SET case_id=NULL WHERE id=?', (q['doc_id'],))
            conn.execute("DELETE FROM pump_state WHERE key='file_quoted_docs'")
            conn.commit()
            P._file_quoted_docs(conn)
            cid = conn.execute('SELECT case_id FROM pump_docs WHERE id=?', (q['doc_id'],)).fetchone()[0]
        finally:
            conn.close()
        case = self.case(cid)
        self.assertEqual((case['status'], case['jobber']['quote']['number']), ('open', '812'))

    def test_short_job_names_to_skim(self):
        """Oct 2026: the watchlist names a job "Client · work" in a few words."""
        for title, client, short in (
                ('Magnolia Cove Pressure Transducer - Furnish and install new 0-150 Pressure transducer, wire up, '
                 'program and test', 'Magnolia Cove Pressure Transducer', 'Magnolia Cove · New 0-150 Pressure transducer'),
                ('Proposal to inspect suction line', 'Carlisle (The Carlisle)', 'Carlisle · Inspect suction line'),
                ('Proposal to check out #6 pump station', 'HUNTINGTON LAKES MASTER  ASSOCIATION  C/O Moore PM Services',
                 'Huntington Lakes · Check out #6 pump station'),
                ('Service call - Pump at the front lake is not running', 'Lely', 'Lely · Pump at the front lake'),
                ('Repairs as discussed', 'Quail Run', 'Quail Run · Repairs as discussed')):
            self.assertEqual(P.short_title({'title': title, 'client_name': client}), short)

    def test_homewood_optional_on_the_cost_line(self):
        """Wettech's Homewood Suites letter (Oct 14 2026): the option's work and
        its "Your Additional Cost" are on one line - an optional line, not in
        the total."""
        text = ('October 14, 2026\nStahlman England\nAttn: Andrea\nRE: Homewood Suites\n'
                'We are pleased to quote you on the following services\n'
                'Field service to check out well nor pumping, pulled with crane and crew, found pump had overheated '
                'and fell off pipe, we recovered all equipment from well and suggest the following\n'
                'Furnish and install new 5 HP pump, motor, 130 foot wire, 20 foot certalock pipe, drop nipple and '
                'coupling, wire up and test.\n                    Your Cost ---------------- $ 5427.63\n'
                'Optional ... replace all PVC pipe with galvanized pipe            Your Additional Cost ... $ 534.86\n'
                'Price includes Sales tax\nTerms: Net 10 days\n')
        x = P._clean_extraction(P._regex_extract(text, 'tomm@wettec.biz', 'Homewood Suites'))
        self.assertEqual(x['total'], 5427.63)
        self.assertEqual([(i['amount'], i.get('optional', False)) for i in x['line_items']],
                         [(5427.63, False), (534.86, True)])
        self.assertEqual(x['line_items'][1]['description'], 'Replace all PVC pipe with galvanized pipe')
        lines = P.one_line_items(x, 30)
        self.assertEqual([(l['unit_price'], bool(l.get('optional')), l['taxable']) for l in lines],
                         [(round(5427.63 * 1.3, 2), False, False), (round(534.86 * 1.3, 2), True, False)])
        # Our client quote in whole dollars, the cents dropped ($7,055.92 -> $7,055).
        self.assertEqual([l['unit_price'] for l in P.suggest_quote(x)['line_items']], [7055.0, 695.0])
        # Tom's typo ("well nor pumping") doesn't go to the client.
        s = P.suggest_quote(x)
        self.assertIn('check out well not pumping', s['line_items'][0]['description'])
        self.assertEqual(s['title'], 'Proposal to check out well not pumping')
        self.assertEqual(s['corrections'], ['"nor pumping" -> "not pumping"'])
        self.assertEqual(P.fix_typos('neither Wettech nor Gulfshore')[1], [], 'a real "nor" stays')

    def test_huntington_6_letter_is_one_quote_with_an_optional_line(self):
        """Wettech, Oct 8 2026: a control box at $1,780.11 and "Possibly
        needed" a new motor at $4,836.90 - one quote: the control box, and
        the motor as an optional line not in the total."""
        self.texts['huntington6.pdf'] = (
            "October 8, 2026\nStahlman England\nAttn: Simon\nRE: Huntington #6\n"
            "We are pleased to quote you on the following services\n"
            "Field service to check out #6 pump station, found all of the capacitors and potential relay blown, "
            "suggest starting with a control box replacement, and testing.\n\nFurnish and install new 10 HP deluxe "
            "control box, wire up and test.\n\n                Your Cost ---------------- $ 1780.11\n"
            "Possibly needed- Field service to pull and inspect pump and motor, replace motor with new 10 HP 230 "
            "volt 1 PH motor reusing pump end, reinstall, wire up and test\n"
            "                 Your Cost --------------- $ 4836.90\nPrices include Sales tax\n"
            "Terms: Net 10 days\nDelivery: 2-3 days after receipt of order\nPrices good for 30 days.\n")
        self.extracts['huntington6.pdf'] = extraction('quote', '', client='Huntington #6', subtotal=6616.01)
        fake = FakeJobber()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            doc = self.upload('huntington6.pdf')
            self.assertFalse(doc.get('parts'))
            self.assertEqual(self.case(doc['case_id'])['vendor_quote_total'], 1780.11)
            self.draft(doc['doc_id'])
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        made = [v['attributes'] for qq, v in fake.calls if 'quoteCreate(' in qq]
        self.assertEqual(len(made), 1, 'one quote')
        lines = made[0]['lineItems']
        self.assertEqual([(l['unitPrice'], bool(l.get('optional'))) for l in lines],
                         [(2314.0, False), (6287.0, True)])
        self.assertIn('motor', lines[1]['description'])

    def test_delete_from_to_do_list(self):
        """The ✕ on a To do row takes it off the list - the job is left as it
        is, Undo brings the row back, and the job's next step shows again."""
        self.extracts['q-0142.pdf'] = extraction('quote', 'Q-142', po='PO142', client='Lakeside Pines', subtotal=700)
        res = self.upload('q-0142.pdf', data=self._pdf())
        queue = lambda: self.c.get('/pumps/api/summary').get_json()['queue']
        row = next(d for d in queue()['quotes_to_draft'] if d['id'] == res['doc_id'])
        self.assertEqual(row['todo_key'], f"doc:{res['doc_id']}")
        self.assertTrue(self.c.post('/pumps/api/todos/hide', json={'key': row['todo_key']}).get_json()['success'])
        self.assertNotIn(res['doc_id'], [d['id'] for d in queue()['quotes_to_draft']])
        self.assertEqual(self.case(res['case_id'])['status'], 'open', 'the job itself is left alone')
        top = self.c.get('/pumps/api/undo').get_json()['actions'][0]
        self.assertEqual(top['label'], 'Took a row off the To do list')
        self.c.post(f"/pumps/api/undo/{top['id']}", json={})
        self.assertIn(res['doc_id'], [d['id'] for d in queue()['quotes_to_draft']])
        # An item's row: gone for this step only.
        j = self.c.post('/pumps/api/cases', json={'title': 'Hide me', 'client_name': 'Quail Run'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        rows = lambda: [c for k, v in queue().items() if isinstance(v, list) for c in v
                        if isinstance(c, dict) and c.get('todo_key', '').startswith(f'case:{cid}:')]
        key = rows()[0]['todo_key']
        self.c.post('/pumps/api/todos/hide', json={'key': key})
        self.assertEqual(rows(), [])
        conn = P._conn()
        conn.execute("UPDATE pump_cases SET stage='scheduled' WHERE id=?", (cid,))
        conn.commit()
        conn.close()
        self.assertEqual([r['todo_key'] for r in rows()], [f'case:{cid}:scheduled'])
        self.assertEqual(self.c.post('/pumps/api/todos/hide', json={'key': "x'; DROP"}).status_code, 400)
        # Deleted from the list, then the same letter uploaded again: it's back.
        self.c.post('/pumps/api/todos/hide', json={'key': row['todo_key']})
        self.assertNotIn(res['doc_id'], [d['id'] for d in queue()['quotes_to_draft']])
        self.upload('q-0142.pdf', data=self._pdf())
        self.assertIn(res['doc_id'], [d['id'] for d in queue()['quotes_to_draft']])

    def test_ready_client_approved(self):
        """Sopra called: ready for their maintenance. Off hold for the diver,
        followed under Waiting on others until both service reports are in."""
        acc = lambda: next(x for x in self.c.get('/pumps/api/accounts').get_json()['accounts']
                           if x['name'] == 'Sopra Luxury Living')
        self.assertEqual(acc()['lake']['status'], 'hold')
        r = self.c.post('/pumps/api/accounts/ready', json={'name': 'Sopra Luxury Living'}).get_json()
        self.assertEqual((r['vendors'], r['unheld']), (['Wettech', 'Gulfshore'], True))
        self.assertEqual(acc()['lake']['status'], 'active')
        self.assertEqual(acc()['visit']['link']['vendors'], ['Wettech', 'Gulfshore'])
        q = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertEqual([v['link']['account'] for v in q['waiting_visits']], ['Sopra Luxury Living'])
        self.assertNotIn('visit', [t['kind'] for t in q['todos']])
        self.assertEqual(self.c.get('/pumps/api/undo').get_json()['actions'][0]['label'], 'Client ready for maintenance')

        def report(vendor):
            conn = P._conn()
            conn.execute("INSERT INTO pump_docs (kind, status, vendor, client_name, file_name, created_at) "
                         "VALUES ('report', 'filed', ?, 'Sopra Luxury Living', 'r.docx', ?)", (vendor, P._now_text()))
            conn.commit()
            conn.close()
        report('Wettech')
        q = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertEqual(q['waiting_visits'][0]['link']['reported'], ['Wettech'])
        report('Gulfshore')
        self.assertEqual(self.c.get('/pumps/api/summary').get_json()['queue']['waiting_visits'], [])
        self.assertIsNone(acc()['visit'])
        self.assertEqual(self.c.post('/pumps/api/accounts/ready', json={'name': 'Nobody'}).status_code, 400)

    def test_dismissed_quote_sent_again_is_read_again(self):
        self.extracts['q-0143.pdf'] = extraction('quote', 'Q-143', po='PO143', client='Lakeside Pines', subtotal=700)
        res = self.upload('q-0143.pdf', data=self._pdf())
        self.c.post(f"/pumps/api/docs/{res['doc_id']}", json={'status': 'dismissed'})
        self.extracts['q-0143.pdf'] = extraction('quote', 'Q-143', po='PO143', client='Lakeside Pines', subtotal=750)
        again = self.upload('q-0143.pdf', data=self._pdf())
        self.assertTrue(again.get('read_again'), again)
        doc = self.c.get(f"/pumps/api/docs/{res['doc_id']}").get_json()['doc']
        self.assertEqual((doc['status'], doc['subtotal']), ('filed', 750))
        q = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn(res['doc_id'], [d['id'] for d in q['quotes_to_draft']])

    def test_page_knows_its_version(self):
        """An open page compares its version with the server's and reloads
        itself after a new deployment."""
        v = self.c.get('/pumps/api/summary').get_json()['version']
        self.assertEqual(v, P.APP_VERSION)
        self.assertIn(f'const APP_VERSION = "{v}"', self.c.get('/pumps').get_data(as_text=True))

    def test_pump_quote_goes_on_the_maintenance_jobs_property(self):
        """Anna's Place (Oct 2026): two properties - the billing P.O. Box and
        the one on the recurring "Quarterly 400" job. Our quote goes on the
        job's; a pump the vendor names is said in the title."""
        props = [{'id': 'PBILL', 'address': {'street1': 'C/O Advanced Property Management Services',
                                             'street2': 'P.O. Box 37296', 'city': 'Charlotte'}},
                 {'id': 'PJOB', 'address': {'street1': 'C/O Advanced Property Management Services',
                                            'street2': '1035 Collier Center Way, Suite #7', 'city': 'Naples'}}]
        jobs = [{'id': 'J475', 'jobNumber': 475, 'title': 'Quarterly 400', 'jobStatus': 'upcoming',
                 'jobType': 'RECURRING', 'property': {'id': 'PJOB'}},
                {'id': 'J7018', 'jobNumber': 7018, 'title': "Anna's Place- **WET**PO-0010", 'jobStatus': 'archived',
                 'jobType': 'ONE_OFF', 'property': {'id': 'PJOB'}}]

        class Fake(FakeJobber):
            def __call__(self, query, variables=None):
                if 'client(id' in query:
                    return {'client': {'id': variables['id'], 'name': "INC ANA'S PLACE HOMEOWNER'S ASSOCIATION",
                                       'properties': props, 'jobs(first: 40)': {'nodes': jobs}}}
                return super().__call__(query, variables)
        P.jobber_gql = Fake(clients=[{'id': 'CANA', 'name': "INC ANA'S PLACE HOMEOWNER'S ASSOCIATION",
                                      'isLead': False, 'isArchived': False, 'properties': props}])
        doc = {'client_name': "Anna's Place", 'file_name': 'Annas Place tank.docx',
               'description': 'Furnish and install new bladder tank, set to correct air pressure and test.'}
        t = P.resolve_quote_target(doc, {'jobber_client_id': 'CANA', 'jobber': '{}'})
        self.assertEqual((t['client_id'], t['property_id']), ('CANA', 'PJOB'))
        # The item remembers the billing address from an earlier quote: still the job's.
        t = P.resolve_quote_target(doc, {'jobber_client_id': 'CANA', 'jobber_property_id': 'PBILL', 'jobber': '{}'})
        self.assertEqual(t['property_id'], 'PJOB')
        # A big client (Autumn Woods HOA): the monthly job isn't in the first 40
        # jobs - it's asked for among the recurring ones.
        aw_props = [{'id': 'P6416', 'address': {'street1': '6416 Autumn Woods Boulevard', 'city': 'Naples'}},
                    {'id': 'P6721', 'address': {'street1': 'AUTUMN WOODS COMMUNITY', 'street2': '6721 AUTUMN WOODS BLVD',
                                                'city': 'Naples'}}]
        old = [{'id': f'JO{n}', 'jobNumber': n, 'title': 'Irrigation repair', 'jobStatus': 'archived',
                'jobType': 'ONE_OFF', 'property': {'id': 'P6416'}} for n in range(40)]

        class Big(FakeJobber):
            def __call__(self, query, variables=None):
                if 'jobType: RECURRING' in query:
                    return {'client': {'id': variables['id'], 'jobs': {'nodes': [
                        {'id': 'J9659', 'jobNumber': 9659, 'title': 'Monthly 3605', 'jobStatus': 'upcoming',
                         'jobType': 'RECURRING', 'property': {'id': 'P6721', 'address': {
                             'street1': 'AUTUMN WOODS COMMUNITY', 'street2': '6721 AUTUMN WOODS BLVD', 'city': 'Naples'}}}]}}}
                if 'client(id' in query:
                    return {'client': {'id': variables['id'], 'name': 'AUTUMN WOODS HOA COMMUNITY',
                                       'properties': aw_props, 'jobs(first: 40)': {'nodes': old}}}
                return super().__call__(query, variables)
        P.jobber_gql = Big()
        t = P.resolve_quote_target({'client_name': 'Autumn Woods 1', 'file_name': 'aw.docx'},
                                   {'jobber_client_id': 'CAW', 'jobber': '{}'})
        self.assertEqual(t['property_id'], 'P6721')
        # A home on Autumn Woods remembered for "Autumn Woods" from an earlier
        # quote: still the monthly job's property.
        conn = P._conn()
        conn.execute("INSERT INTO pump_site_aliases (place, area, client_id, client_name, property_id, property_label) "
                     "VALUES ('Autumn Woods', '', 'CAW', 'AUTUMN WOODS HOA COMMUNITY', 'P6416', '6416 Autumn Woods Boulevard')")
        conn.commit()
        conn.close()
        t = P.resolve_quote_target({'client_name': 'Autumn Woods', 'file_name': 'aw.docx'})
        self.assertEqual((t['client_id'], t['property_id']), ('CAW', 'P6721'))
        # The job's property isn't in the client's property list: the job says what it is.
        aw_props[1] = {'id': 'P6537', 'address': {'street1': '6537 Chestnut Circle', 'city': 'Naples'}}
        t = P.resolve_quote_target({'client_name': 'Autumn Woods 1', 'file_name': 'aw.docx'},
                                   {'jobber_client_id': 'CAW', 'jobber': '{}'})
        self.assertEqual((t['property_id'], t['property_label']),
                         ('P6721', 'AUTUMN WOODS COMMUNITY, 6721 AUTUMN WOODS BLVD, Naples'))
        P.jobber_gql = Fake()
        # Which pump goes in the title.
        self.assertEqual(P.suggest_quote({**doc, 'description': 'Field service to replace check valve, pump #4 '
                                                             'station', 'line_items': []})['title'],
                         'Proposal to replace check valve - Pump #4')
        self.assertEqual(P.suggest_quote({**doc, 'description': 'Furnish and install new check valve on pump #4',
                                          'line_items': []})['title'],
                         'Proposal - Furnish and install new check valve on pump #4', 'already says which pump')
        self.assertNotIn('Pump #', P.suggest_quote({**doc, 'line_items': []})['title'])

    def test_quote_is_never_drafted_on_its_own(self):
        """A vendor quote that reads cleanly waits in To do for the office to
        click Draft quote - nothing is made in Jobber until then."""
        self.extracts['q-0141.pdf'] = extraction('quote', 'Q-141', po='PO141', client='Lakeside Pines', subtotal=700)
        fake = FakeJobber()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            res = self.upload('q-0141.pdf', data=self._pdf())
            self.c.post(f"/pumps/api/docs/{res['doc_id']}/read_again", json={})
        finally:
            P.JOBBER_STATIC_TOKEN = saved
        self.assertTrue(res['case_id'])
        self.assertEqual([q for q, _ in fake.calls if 'quoteCreate(' in q], [])
        queue = self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn(res['doc_id'], [d['id'] for d in queue['quotes_to_draft']])

    def test_miramar_is_miromar(self):
        a = P.match_site_alias({'client_name': 'Miramar Lakes', 'file_name': 'Inv_29100.pdf'})
        self.assertEqual((a['client_name'], a['property_id']),
                         ('MIROMAR LAKES', 'Z2lkOi8vSm9iYmVyL1Byb3BlcnR5LzM4MTYzODEy'))
        self.assertEqual(P.match_site_alias({'client_name': 'Miromar Lakes', 'file_name': 'x.pdf'})['client_name'],
                         'MIROMAR LAKES')
        for other in ('Miramar Lakes Beach and Gulf Club', 'Miramar Lakes CDD', 'Miramar Lakes Golf & Country Club'):
            self.assertIsNone(P.match_site_alias({'client_name': other, 'file_name': 'x.pdf'}), other)
        # The Carlisle names still apply ("The Carlise" adds nothing of its own).
        self.assertEqual(P.match_site_alias({'client_name': 'The Carlise', 'file_name': 'Carslie Back Station .docx'})
                         ['area'], 'back station')

    def test_site_names(self):
        names = self.c.get('/pumps/api/site-names').get_json()['site_names']
        self.assertIn(('Carlisle', 'back station', 'Z2lkOi8vSm9iYmVyL1Byb3BlcnR5LzUyOTc4MjYw'), [(n['place'], n['area'], n['property_id']) for n in names])
        r = self.c.post('/pumps/api/site-names', json={'place': 'Heron Bay', 'area': 'front pump', 'client_id': 'C7',
                                                       'client_name': 'Heron Bay HOA', 'property_id': 'P7'})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        a = P.match_site_alias({'client_name': 'Heron Bay', 'file_name': 'Heron Bay front pump quote.pdf'})
        self.assertEqual(a['property_id'], 'P7')
        self.assertIsNone(P.match_site_alias({'client_name': 'Heron Glen', 'file_name': 'x.pdf'}))
        aid = [n['id'] for n in r.get_json()['site_names'] if n['place'] == 'Heron Bay'][0]
        bot = A.app.test_client()
        h = {'Authorization': 'Bearer test-openclaw-key'}
        self.assertEqual(bot.post('/api/pumps/site-names', json={'place': 'X', 'client_id': 'Y'}, headers=h).status_code, 403)
        self.assertEqual(bot.post(f'/api/pumps/site-names/{aid}/delete', json={}, headers=h).status_code, 403)
        self.c.post(f'/pumps/api/site-names/{aid}/delete', json={})
        self.assertIsNone(P.match_site_alias({'client_name': 'Heron Bay', 'file_name': 'Heron Bay front pump.pdf'}))

    def test_openclaw_cannot_draft_client_quotes(self):
        self.extracts['q-0830.pdf'] = extraction('quote', 'Q-123', po='PO830', client='Lakeside Pines', subtotal=500)
        q = self.upload('q-0830.pdf')
        fake = FakeJobber()
        P.jobber_gql = fake
        r = A.app.test_client().post(f"/api/pumps/docs/{q['doc_id']}/quote", json={'client_id': 'C1', 'property_id': 'P1'},
                                     headers={'Authorization': 'Bearer test-openclaw-key'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(fake.calls, [])

    def test_vendor_bill_due_once_client_pays(self):
        self.extracts['q-0900.pdf'] = extraction('quote', 'Q-130', po='PO900', client='Lakeside Pines', subtotal=900)
        self.extracts['b-0900.pdf'] = extraction('bill', '30100', po='PO900', client='Lakeside Pines', subtotal=900)
        self.upload('q-0900.pdf')
        b = self.upload('b-0900.pdf')
        cid = b['case_id']
        self.assertEqual(self.case(cid)['vendor_pay']['state'], 'unpaid')
        P.jobber_gql = FakeJobber()
        self.assertTrue(self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'C1'}).get_json()['success'])

        def jobber_says(status):
            conn = P._conn()
            try:
                conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status) "
                             "VALUES ('INV1', 'invoice', '5001', 'Pump service', ?) "
                             "ON CONFLICT(jobber_id) DO UPDATE SET status=excluded.status", (status,))
                P._follow_invoices(conn)
                conn.commit()
            finally:
                conn.close()

        def to_pay():
            return [c['id'] for c in self.c.get('/pumps/api/summary').get_json()['queue']['vendor_bills_to_pay']]
        jobber_says('awaiting_payment')
        self.assertEqual(self.case(cid)['vendor_pay']['state'], 'unpaid', 'sent, not paid yet')
        self.assertNotIn(cid, to_pay())
        # The item closes before the client pays; the alert must still come.
        conn = P._conn()
        try:
            conn.execute("UPDATE pump_cases SET status='closed' WHERE id=?", (cid,))
            conn.commit()
        finally:
            conn.close()
        jobber_says('paid')
        case = self.case(cid)
        self.assertEqual(case['vendor_pay']['state'], 'due')
        self.assertIn(cid, to_pay())
        self.assertTrue(any(e['action'] == 'pay vendor' for e in case['events']))
        # A ready-to-paste email asks Christian to pay it.
        conn = P._conn()
        try:
            conn.execute("CREATE TABLE IF NOT EXISTS app_settings (key TEXT PRIMARY KEY, value TEXT)")
            conn.execute("INSERT OR REPLACE INTO app_settings (key, value) VALUES ('christian_email', 'christian@example.com')")
            conn.commit()
        finally:
            conn.close()
        mail = self.c.get(f'/pumps/api/cases/{cid}/pay_email').get_json()['email']
        self.assertEqual(mail['to'], 'christian@example.com')
        self.assertIn('Wettech invoice #30100', mail['subject'])
        self.assertIn('Please pay Wettech invoice #30100 for $', mail['body'])
        self.assertIn('Jobber invoice #5001', mail['body'])
        self.assertEqual(mail['bill_doc_id'], b['doc_id'])
        # Clicking Send emails it from the PO mailbox with the bill attached - only when clicked.
        sent, real_send = [], P.graph_send
        P.graph_send = lambda to, subject, html, attachments=None, sender=None: sent.append((to, subject, html, attachments))
        try:
            self.assertEqual(sent, [], 'nothing goes out on its own')
            r = self.c.post(f'/pumps/api/cases/{cid}/pay_email/send', json={'body': 'Hi Christian,\nPlease pay it.'})
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
            self.assertEqual(len(sent), 1)
            self.assertEqual(sent[0][0], 'christian@example.com')
            self.assertIn('Please pay it.', sent[0][2])
            self.assertEqual(len(sent[0][3]), 1)
            q = self.c.get('/pumps/api/summary').get_json()['queue']['vendor_bills_to_pay']
            self.assertTrue(next(c for c in q if c['id'] == cid)['pay_email']['sent'])

            def refuse(*a, **k):
                raise RuntimeError('Microsoft 365 said HTTP 403: Access is denied.')
            P.graph_send = refuse
            r = self.c.post(f'/pumps/api/cases/{cid}/pay_email/send', json={})
            self.assertEqual(r.status_code, 502)
            self.assertIn('Access is denied', r.get_json()['error'])
        finally:
            P.graph_send = real_send
        # OpenClaw may not say a vendor was paid.
        r = A.app.test_client().patch(f'/api/pumps/cases/{cid}', json={'vendor_paid_on': '2026-10-08'},
                                      headers={'Authorization': 'Bearer test-openclaw-key'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.c.patch(f'/pumps/api/cases/{cid}', json={'vendor_paid_on': 'soon'}).status_code, 400)
        r = self.c.patch(f'/pumps/api/cases/{cid}', json={'vendor_paid_on': '2026-10-08'})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        pay = r.get_json()['case']['vendor_pay']
        self.assertEqual((pay['state'], pay['paid_on'], pay['paid_by']), ('paid', '2026-10-08', 'Office Tester'))
        self.assertNotIn(cid, to_pay())
        # Undo.
        self.c.patch(f'/pumps/api/cases/{cid}', json={'vendor_paid_on': ''})
        self.assertIn(cid, to_pay())

    # ── the diver ───────────────────────────────────────────────────────────
    def test_diver_sites_and_monthly_email(self):
        j = self.c.get('/pumps/api/dives').get_json()
        names = {x['name']: x for x in j['sites']}
        self.assertTrue({"Anna's Place", 'Spanish Wells Lake Club', 'Clubside', 'Riviera Golf Estates 2'} <= set(names))
        # From the lake sheet: quarterly, January only, April/October; former accounts kept but off the lists.
        self.assertEqual(names["Anna's Place"]['months'], '1,4,7,10')
        self.assertEqual((names["Anna's Place"]['diver_cost'], names["Anna's Place"]['our_bill']), ('$125.00', '$400.00'))
        self.assertEqual(names['Evergreen (Bradenton)']['months'], '1')
        self.assertEqual(names['Riviera Golf Estates 2']['months'], '4,10')
        self.assertEqual(names['Bentley Village']['active'], 0)
        self.assertEqual(names['Warm Springs Comm. Assoc.']['needs_dive'], 0, 'Naples Electric does it')
        self.assertEqual(names['Spanish Wells Lake Club']['status'], 'meet')
        self.assertEqual(names['Sopra Luxury Living']['status'], 'hold')
        self.assertEqual(j['settings']['to'], 'Gulfshoreyachts@gmail.com')
        # A site that doesn't need diving, one only in some months, and a one-time note.
        r = self.c.post('/pumps/api/dives/sites', json={'name': 'Lely Pump Station', 'equipment': '1 Pump',
                                                         'needs_dive': False})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.c.post('/pumps/api/dives/sites', json={'id': names['Morton Grove']['id'], 'months': '4,7,10'})
        self.c.post('/pumps/api/dives/sites', json={'id': names['Caymas']['id'],
                                                     'month_note': 'Gate is being replaced - call the office.'})
        self.assertEqual(self.c.get('/pumps/api/dives/preview?month=2026-11-01').get_json()['sites'], 0,
                         'nobody is due in November')
        p = self.c.get('/pumps/api/dives/preview?month=2027-01-01').get_json()
        self.assertEqual(p['subject'], 'Diver schedule - January 2027 - Stahlman-England')
        t = p['text']
        self.assertTrue(t.startswith('Hi Jordan!'))
        for name in ("Anna's Place", 'Clubside', 'Evergreen (Bradenton)', 'The Reserve @ Estero'):
            self.assertIn(name, t)
        for name in ('Riviera Golf Estates 2', 'Bentley Village', 'Warm Springs', 'Morton Grove'):
            self.assertNotIn(name, t)
        self.assertIn('Please see the attached list for January. Please note that Spanish Wells Lake Club has '
                      'requested to meet with you onsite.', t)
        self.assertIn('HOLD OFF - Called client to set up visit - waiting on approval. - Sopra Luxury Living', t)
        self.assertIn('Diver to check the rope that holds the float and filter.', t)
        self.assertIn('Gate code #0422', t)
        self.assertIn('Gate is being replaced', t)
        self.assertNotIn('Lely', t)
        self.assertIn('Andrea Mitchell', t)
        self.assertIn('<b>Anna&#x27;s Place</b>' if '&#x27;' in p['html'] else "<b>Anna's Place</b>", p['html'])
        october = self.c.get('/pumps/api/dives/preview?month=2026-10-01').get_json()['text']
        self.assertIn('Morton Grove', october)
        self.assertIn('Riviera Golf Estates 2', october)

        # Sending: once a month, from the PO mailbox, copied to Andrea, Word list attached.
        sent = []

        class Resp:
            status_code, content = 202, b''

        real_post = P.http_requests.post
        P.http_requests.post = lambda url, **kw: (sent.append((url, kw['json'])), Resp())[1]
        P.CFG['graph_token'], P.CFG['mail_from'] = (lambda: 'tok'), 'PO@stahlman-england.com'
        try:
            res = P.send_dive_email(P.datetime(2027, 1, 1).date())
            self.assertTrue(res['sent'])
            url, body = sent[0]
            self.assertIn('/users/PO@stahlman-england.com/sendMail', url)
            m = body['message']
            self.assertEqual([r['emailAddress']['address'] for r in m['toRecipients']], ['Gulfshoreyachts@gmail.com'])
            self.assertEqual([r['emailAddress']['address'] for r in m['ccRecipients']], ['Andrea@stahlman-england.com'])
            self.assertEqual([r['emailAddress']['address'] for r in m['replyTo']], ['Andrea@stahlman-england.com'])
            self.assertEqual(m['attachments'][0]['name'], 'Diver schedule January 2027.docx')
            self.assertGreater(len(m['attachments'][0]['contentBytes']), 1000)
            # The one-time note was for that email only; the month is not sent twice.
            caymas = [x for x in self.c.get('/pumps/api/dives').get_json()['sites'] if x['name'] == 'Caymas'][0]
            self.assertEqual(caymas['month_note'], '')
            self.assertEqual(P.send_dive_email(P.datetime(2027, 1, 1).date()), {'skipped': 'already sent this month'})
            self.assertEqual(len(sent), 1)
            # Switched off, the 1st only adds the to-do.
            P._now, real_now = (lambda: P.datetime(2027, 4, 1, 8, 5, tzinfo=P.TZ)), P._now
            try:
                self.c.post('/pumps/api/dives/settings', json={'auto': False})
                P._scheduled_dive_email()
                self.assertEqual(len(sent), 1)
                # Switched on (the default), the hourly check sends the month once; mid-month it does nothing.
                self.c.post('/pumps/api/dives/settings', json={'auto': True})
                P._scheduled_dive_email()
                P._scheduled_dive_email()
                self.assertEqual(len(sent), 2)
                self.assertIn('April 2027', sent[1][1]['message']['subject'])
                P._now = lambda: P.datetime(2027, 4, 15, 8, 5, tzinfo=P.TZ)
                P._scheduled_dive_email()
                self.assertEqual(len(sent), 2)
            finally:
                P._now = real_now
                self.c.post('/pumps/api/dives/settings', json={'auto': False})
            # A failed send is logged and tried again later.
            class Bad:
                status_code, content, text = 403, b'x', 'Access denied'

                def json(self):
                    return {'error': {'message': 'Access is denied'}}
            P.http_requests.post = lambda url, **kw: Bad()
            with self.assertRaises(RuntimeError):
                P.send_dive_email(P.datetime(2027, 7, 1).date())
            self.assertIsNone(P._state_get('dive_email:2027-07'))
            log = self.c.get('/pumps/api/dives').get_json()['sent']
            self.assertIn('Access is denied', log[0]['error'])
        finally:
            P.http_requests.post = real_post
            P.CFG['graph_token'] = None
        # OpenClaw can't change the list or email the diver.
        bot = A.app.test_client()
        h = {'Authorization': 'Bearer test-openclaw-key'}
        self.assertEqual(bot.post('/api/pumps/dives/send', json={}, headers=h).status_code, 403)
        self.assertEqual(bot.post('/api/pumps/dives/sites', json={'name': 'X'}, headers=h).status_code, 403)

    def test_todo_list_and_diver_email_todo(self):
        P.ensure_monthly_todos(P.datetime(2026, 11, 1).date())
        self.assertFalse([t for t in self.c.get('/pumps/api/todos').get_json()['todos']
                          if t['key'] == 'diver-email:2026-11'], 'no sites due in November - no to-do')
        P.ensure_monthly_todos(P.datetime(2027, 1, 1).date())
        P.ensure_monthly_todos(P.datetime(2027, 1, 1).date())
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        nov = [t for t in todos if t['key'] == 'diver-email:2027-01']
        self.assertEqual(len(nov), 1, 'one per month')
        self.assertEqual(nov[0]['title'], 'Email Jordan the January diver list')
        self.assertIn('Gulfshoreyachts@gmail.com', nov[0]['detail'])
        # The list is ready to attach.
        r = self.c.get('/pumps/api/dives/docx?month=2027-01-01')
        self.assertEqual(r.status_code, 200)
        self.assertIn('Diver schedule January 2027.docx', r.headers['Content-Disposition'])
        # Ticked off, and back again.
        j = self.c.post(f"/pumps/api/todos/{nov[0]['id']}/done", json={'done': True}).get_json()
        self.assertTrue([t for t in j['todos'] if t['id'] == nov[0]['id']][0]['done_at'])
        self.assertNotIn(nov[0]['id'], [t['id'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']])
        self.c.post(f"/pumps/api/todos/{nov[0]['id']}/done", json={'done': False})
        self.assertIn(nov[0]['id'], [t['id'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']])
        # The office's own to-dos.
        j = self.c.post('/pumps/api/todos', json={'title': 'Call Spanish Wells HOA about the dive'}).get_json()
        mine = [t for t in j['todos'] if t['title'] == 'Call Spanish Wells HOA about the dive'][0]
        self.assertEqual(self.c.post('/pumps/api/todos', json={'title': ' '}).status_code, 400)
        self.c.post(f"/pumps/api/todos/{mine['id']}/delete", json={})
        self.assertNotIn(mine['id'], [t['id'] for t in self.c.get('/pumps/api/todos').get_json()['todos']])

    def test_pipeline_dashboard(self):
        sample = self.c.get('/pumps/api/pipeline?sample=1').get_json()
        self.assertTrue(sample['sample'])
        self.assertEqual([st['key'] for st in sample['stages']],
                         ['quote', 'approval', 'scheduling', 'work', 'billing', 'payment'])
        self.assertEqual(sum(st['count'] for st in sample['stages']), sample['kpis']['open'])
        self.assertEqual(len(sample['weeks']), 8)
        # Live: a real item lands in its stage, with what's next.
        self.extracts['q-pl.pdf'] = extraction('quote', 'Q-150', po='PO990', client='Pipeline Pines', subtotal=700)
        q = self.upload('q-pl.pdf')
        live = self.c.get('/pumps/api/pipeline').get_json()
        self.assertFalse(live['sample'])
        quote = [st for st in live['stages'] if st['key'] == 'quote'][0]
        mine = [j for j in quote['jobs'] if j['id'] == q['case_id']][0]
        self.assertEqual(mine['next'], 'Quote sent to client')
        self.assertEqual(mine['jobber_uri'], '', 'not linked to Jobber yet')
        # Once linked, the job opens its Jobber record - the quote while quoting.
        self.c.patch(f"/pumps/api/cases/{q['case_id']}", json={'jobber': {
            'quote': {'id': 'QX', 'number': '9136', 'uri': 'https://secure.getjobber.com/quotes/66752875'},
            'job': {'id': 'JX', 'number': '1609', 'uri': 'https://secure.getjobber.com/work_orders/41420018'}}})
        live = self.c.get('/pumps/api/pipeline').get_json()
        mine = [j for st in live['stages'] for j in st['jobs'] if j['id'] == q['case_id']][0]
        self.assertEqual((mine['jobber_uri'], mine['jobber_label']),
                         ('https://secure.getjobber.com/quotes/66752875', 'quote #9136'))
        self.assertTrue(any(j['jobber_uri'] for st in sample['stages'] for j in st['jobs']))
        # Scheduled work shows up in the weeks ahead.
        when = (P._today() + P.timedelta(days=3)).isoformat()
        self.c.patch(f"/pumps/api/cases/{q['case_id']}", json={'scheduled_for': when})
        live = self.c.get('/pumps/api/pipeline').get_json()
        self.assertIn(when, [e['date'] for w in live['weeks'] for e in w['events']])
        ev = [e for w in live['weeks'] for e in w['events'] if e['date'] == when][0]
        self.assertEqual(ev['jobber_uri'], 'https://secure.getjobber.com/work_orders/41420018', 'the job, for a visit')

    # ── who can get in ───────────────────────────────────────────────────────
    def test_access(self):
        anon = A.app.test_client()
        self.assertEqual(anon.get('/pumps/api/summary').status_code, 401)
        self.assertEqual(anon.get('/api/pumps/summary').status_code, 401)
        self.assertEqual(anon.get('/api/pumps/summary', headers={'Authorization': 'Bearer wrong'}).status_code, 401)
        self.assertEqual(anon.get('/api/pumps/summary',
                                  headers={'Authorization': 'Bearer test-openclaw-key'}).status_code, 200)
        tech = A.app.test_client()
        with tech.session_transaction() as s:
            s['username'], s['role'] = 'tech1', 'technician'
        self.assertEqual(tech.get('/pumps/api/summary').status_code, 401)
        self.assertEqual(tech.get('/pumps').status_code, 302)
        self.assertEqual(self.c.get('/pumps').status_code, 200)
        # The job panel's script uses these; a page without them can't open a job.
        page = self.c.get('/pumps').get_data(as_text=True)
        for name in ('MONEY', 'CATS', 'STEPS'):
            self.assertRegex(page, rf'\bconst {name}\s*=', name)
        self.assertIn('function journeyHtml', page)

    def test_office_writes_need_csrf_but_openclaw_does_not(self):
        A.app.config['WTF_CSRF_ENABLED'] = True
        try:
            r = self.c.post('/pumps/api/cases', json={'title': 'CSRF probe'})
            self.assertIn(r.status_code, (400, 302), 'refused (the app redirects on a bad CSRF token)')
            A.app.config['WTF_CSRF_ENABLED'] = False
            self.assertEqual(self.c.get('/pumps/api/cases?status=all&q=CSRF probe').get_json()['cases'], [])
            A.app.config['WTF_CSRF_ENABLED'] = True
            bot = A.app.test_client()
            r = bot.post('/api/pumps/cases', json={'title': 'Pump #3 tripping', 'client_name': 'Gulf Harbor'},
                         headers={'Authorization': 'Bearer test-openclaw-key'})
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
            self.assertEqual(r.get_json()['case']['source'], 'openclaw')
        finally:
            A.app.config['WTF_CSRF_ENABLED'] = False

    def test_openclaw_is_off_unless_switched_on(self):
        """The office uses Pumps on its own: the bot door answers only with PUMPS_OPENCLAW on."""
        bot = A.app.test_client()
        h = {'Authorization': 'Bearer test-openclaw-key'}
        self.assertEqual(bot.get('/api/pumps/summary', headers=h).status_code, 200)
        P.OPENCLAW_ENABLED = False
        try:
            self.assertEqual(bot.get('/api/pumps/summary', headers=h).status_code, 401)
        finally:
            P.OPENCLAW_ENABLED = True

    def test_dashboard_tile_is_live(self):
        html = self.c.get('/dashboard').get_data(as_text=True)
        self.assertIn('href="/pumps"', html)
        folder = html.split('Features Coming Soon', 1)[1] if 'Features Coming Soon' in html else ''
        self.assertNotIn('/pumps"', folder)

    # ── sheet, checklist, SCADA ──────────────────────────────────────────────
    def test_old_sheet_rows_became_items(self):
        cases = self.c.get('/pumps/api/cases?status=all&q=Old Sheet').get_json()['cases']
        self.assertEqual(len(cases), 1)
        self.assertEqual(cases[0]['po_number'], '0290')
        self.assertEqual(cases[0]['compare']['state'], 'match')

    def test_sheet_edit_and_steps(self):
        r = self.c.post('/pumps/api/sheet/save', json={'month': 'October 2026', 'po_number': '0333',
                                                       'client_name': 'Mangrove Bay', 'vendor': 'Wettech',
                                                       'vendor_quote_amount': '$1,200.00'})
        cid = r.get_json()['id']
        case = self.case(cid)
        self.assertEqual(case['vendor_quote_amount'], 1200.0)
        self.assertTrue({s['key']: s for s in case['step_list']}['vendor_quote'].get('at'),
                        'an amount in the sheet ticks the step')
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {'vendor_quote': '2026-10-01', 'client_quote': 'na',
                                                                'client_approved': '2026-10-02'}})
        case = self.case(cid)
        self.assertEqual(case['stage'], 'scheduled')
        self.assertIn('mailto:', case['schedule_email'])
        summary = self.c.get('/pumps/api/summary').get_json()
        # Approved: waiting on the vendor's parts, under Waiting on others.
        self.assertTrue(any(c['id'] == cid for c in summary['queue']['scheduling_asked']))
        r = self.c.get('/pumps/api/export.xlsx?month=October 2026')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data[:2] == b'PK')

    def test_scada_sheet(self):
        """The office's SCADA sheet: renewal date = when it's next due; a year's cell holds the invoice # or a note."""
        today = P._today
        P._today = lambda: P.datetime(2026, 10, 5).date()
        try:
            j = self.c.get('/pumps/api/scada').get_json()
            rows = {r['client']: r for r in j['scada']}
            self.assertEqual(len(rows), 21)  # the sheet's 16 + 3 found in Jobber + 2 FGCU from the master list
            self.assertEqual(j['years'], ['2023', '2024', '2025', '2026'])
            allure = rows['ALLURE']
            self.assertEqual((allure['vendor_cost'], allure['our_bill'], allure['years']['2025']), (428.0, '$482.00', '29448'))
            # The sheet stopped at 2025; Jobber's invoices fill 2026.
            self.assertEqual((allure['years']['2026'], allure['next_due_on'], allure['state']), ('35944', '2027-08-28', 'current'))
            self.assertEqual(rows['LELY']['years'], {'2024': '19236', '2025': '25555', '2026': '32638'})
            self.assertEqual((rows['LELY']['next_due_on'], rows['LELY']['state']), ('2027-01-09', 'current'))
            self.assertEqual(rows['Heritage Stations']['years']['2026'], '33405')
            self.assertEqual(rows['OLD COLLIER']['next_due_on'], '2025-10-04')
            self.assertEqual(rows['AUTUMN WOODS']['years']['2024'], 'SEI pays for SCADA')
            self.assertEqual(rows['Tuscany Point']['vendor_cost'], 855.99)
            states = [r['state'] for r in j['scada']]
            # Old Collier and Camas Willows are off SCADA; Autumn Woods and Reserve are complimentary.
            self.assertEqual((states.count('overdue'), states.count('due_soon'), states.count('inactive')), (2, 0, 4))
            self.assertEqual((rows['AUTUMN WOODS']['complimentary'], rows['RESERVE AT ESTERO']['complimentary']), (1, 1))
            self.assertEqual(rows['CORSA (formerly Estero Crossing)']['state'], 'current')  # 46 days out
            # Renewed: the invoice goes in the year it was due, and the date moves on a year.
            sid = rows['CORSA (formerly Estero Crossing)']['id']
            j = self.c.post(f'/pumps/api/scada/{sid}', json={'action': 'renewed', 'value': '36800'}).get_json()
            corsa = [r for r in j['scada'] if r['id'] == sid][0]
            self.assertEqual((corsa['years']['2026'], corsa['next_due_on'], corsa['state']),
                             ('36800', '2027-11-20', 'current'))
            # A Jobber sync finds next year's renewal invoices by themselves.
            P.jobber_gql = lambda q, v=None: {'invoices': {'nodes': [
                {'invoiceNumber': 40001, 'invoiceStatus': 'paid', 'issuedDate': '2027-01-20T00:00:00Z',
                 'client': {'name': 'LELY CDD'}, 'lineItems': {'nodes': [{'name': 'PUMP SERVICE', 'description':
                 '2027 Renewal Annual Cellular and cloud subscription for SCADA system on irrigation pump station'}]}},
                {'invoiceNumber': 40002, 'invoiceStatus': 'awaiting_payment', 'issuedDate': '2027-03-02T00:00:00Z',
                 'client': {'name': 'MEDALLION HOME'}, 'lineItems': {'nodes': [{'name': 'PUMP SERVICE', 'description':
                 'Renewal of Annual Cellular and cloud subscription for SCADA system on irrigation pump station for '
                 'Cross Creek completed. 2027-2028'}]}},
                {'invoiceNumber': 40003, 'invoiceStatus': 'draft', 'issuedDate': '2027-03-02T00:00:00Z',
                 'client': {'name': 'CLUB CARE'}, 'lineItems': {'nodes': [{'description': 'SCADA subscription renewal'}]}},
                {'invoiceNumber': 40004, 'invoiceStatus': 'paid', 'issuedDate': '2027-02-02T00:00:00Z',
                 'client': {'name': 'Pebblebrook HOA'}, 'amounts': {'total': 600}, 'lineItems': {'nodes': [
                     {'description': 'Renewal of Annual Cellular and cloud subscription for SCADA system'}]}},
                {'invoiceNumber': 40005, 'invoiceStatus': 'paid', 'issuedDate': '2027-02-02T00:00:00Z',
                 'client': {'name': 'Barrington Cove'}, 'lineItems': {'nodes': [{'description': 'Pump Maintenance Complete'}]}},
            ], 'pageInfo': {'hasNextPage': False}}}
            conn = P._conn()
            self.assertEqual(P.scan_scada_invoices(conn), {'added': 1, 'recorded': 2})
            self.assertEqual(P.scan_scada_invoices(conn), {'added': 0, 'recorded': 0})  # once only
            conn.commit(); conn.close()
            rows = {r['client']: r for r in self.c.get('/pumps/api/scada').get_json()['scada']}
            self.assertEqual(rows['LELY']['years']['2027'], '40001')
            self.assertEqual(rows['CROSS CREEK']['years']['2027'], '40002')
            self.assertNotIn('2027', rows['CLUBCARE']['years'])  # drafts don't count
            self.assertEqual((rows['Pebblebrook HOA']['years'], rows['Pebblebrook HOA']['our_bill']), ({'2027': '40004'}, '$600.00'))
            # Edit and switch off.
            self.c.patch(f"/pumps/api/scada/{rows['CLUBCARE']['id']}", json={'active': False, 'notes': 'Left'})
            j = self.c.get('/pumps/api/scada').get_json()
            self.assertEqual([r for r in j['scada'] if r['client'] == 'CLUBCARE'][0]['state'], 'inactive')
            # A renewal item, one at a time.
            oc = rows['AUTUMN WOODS']['id']
            cid = self.c.post(f'/pumps/api/scada/{oc}', json={'action': 'renewal_item'}).get_json()['case_id']
            self.assertEqual(self.case(cid)['category'], 'scada')
            self.assertEqual(cid, self.c.post(f'/pumps/api/scada/{oc}', json={'action': 'renewal_item'}).get_json()['case_id'])
            self.assertIn('AUTUMN WOODS', [s['client_name'] for s in self.c.get('/pumps/api/summary').get_json()['queue']['scada_attention']])
        finally:
            P._today = today
            P.jobber_gql = self._orig[2]

    # ── the office's flow (Oct 2026) ────────────────────────────────────────
    def _pdf(self):
        from PyPDF2 import PdfWriter
        w, out = PdfWriter(), io.BytesIO()
        w.add_blank_page(width=612, height=792)
        w.write(out)
        return out.getvalue()

    def test_approved_quote_is_signed_sent_to_the_vendor_and_invoiced_at_the_quoted_price(self):
        self.extracts['q-0950.pdf'] = extraction('quote', 'Q-950', po='PO951', client='Lakeside Pines', subtotal=1000)
        self.extracts['b-0950.pdf'] = extraction('bill', '30950', po='PO951', client='Lakeside Pines', subtotal=1100)
        fake = FakeJobber(quote_lines=[{'name': 'Service Proposal Amount', 'description': 'Replace check valve',
                                        'quantity': 1, 'unitPrice': 1300, 'taxable': False}])
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            q = self.upload('q-0950.pdf', data=self._pdf())
            self.assertEqual(self.draft(q['doc_id'])['quote']['quote_number'], '812')
            cid = q['case_id']
            # The client approves quote #812 in Jobber (the sync calls the same thing).
            r = self.c.post(f'/pumps/api/cases/{cid}/approved', json={}).get_json()
            self.assertTrue(r['approval']['todo'], r)
            todos = {t['kind']: t for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']
                     if t['link'].get('case_id') == cid}
            self.assertEqual(todos['approved_quote']['title'], 'Sign the approved quote and send it to Tom - Lakeside Pines')
            self.assertEqual(todos['approved_quote']['link']['to'], 'tomm@wettec.biz')
            # Sign quote: signed (a PDF quote is stamped), on the Jobber quote, emailed to Tom.
            sent = []
            P.graph_send, saved_send = (lambda to, subject, html, atts=None, sender=None: sent.append((to, subject, html, atts))), P.graph_send
            try:
                prev = self.c.get(f'/pumps/api/cases/{cid}/sign_quote').get_json()
                self.assertIn(f'[P{cid}]', prev['email']['subject'])
                self.assertIn('When can we expect it to be done?', prev['email']['body'])
                r = self.c.post(f'/pumps/api/cases/{cid}/sign_quote', json={}).get_json()
            finally:
                P.graph_send = saved_send
            self.assertEqual(r['emailed'], 'tomm@wettec.biz', r)
            self.assertEqual(sent[0][0], 'tomm@wettec.biz')
            self.assertTrue(sent[0][3][0]['name'].endswith('SIGNED.pdf'))
            doc = self.c.get(f"/pumps/api/docs/{q['doc_id']}").get_json()['doc']
            self.assertEqual(doc['jobber']['approved_by'], 'Simon Weardon')
            self.assertTrue(doc['jobber']['signed_sent_at'])
            pdf = self.c.get(f"/pumps/api/docs/{q['doc_id']}/file?version=approved")
            self.assertEqual(pdf.status_code, 200)
            import pdfplumber
            with pdfplumber.open(io.BytesIO(pdf.data)) as f:
                text = f.pages[0].extract_text()
            self.assertIn('Simon Weardon', text)
            notes = [v for qq, v in fake.calls if 'quoteCreateNote' in qq]
            self.assertTrue(any('SIGNED.pdf' in json.dumps(v) for v in notes), 'signed copy on the Jobber quote')
            todos = {t['kind']: t for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']
                     if t['link'].get('case_id') == cid}
            self.assertNotIn('approved_quote', todos)
            # Tom answers with a date: on the item, and noted in Jobber.
            rep = P.handle_signed_quote_reply('m77', 'Tom <tomm@wettec.biz>', f'RE: {sent[0][1]}',
                                              'We can have it done by Oct 20th.\nFrom: PO <po@example.com>')
            self.assertEqual(rep['expect_by'][5:], '10-20', rep)
            case = self.c.get(f'/pumps/api/cases/{cid}').get_json()['case']
            self.assertEqual(case['scheduled_for'], rep['expect_by'])
            self.assertTrue(any('expects it done by 10/20' in json.dumps(v) for qq, v in fake.calls if 'CreateNote' in qq))
            self.assertIsNone(P.handle_signed_quote_reply('m77', 'tomm@wettec.biz', f'RE: {sent[0][1]}', 'Oct 20'))
            # This Jobber has no way to make a job from a quote: a to-do instead.
            self.assertTrue(any(t['title'].startswith('Convert quote for ') and t['title'].endswith(' (#812) to a job') for t in todos.values()), todos)
            # Wettech bills $1,100 against a $1,000 quote: the client is invoiced the quoted $1,300.
            b = self.upload('b-0950.pdf')
            self.assertEqual(b['case_id'], cid)
            inv = [v for qq, v in fake.calls if 'invoiceCreate' in qq][0]['input']
            self.assertEqual([(l['name'], l['unitPrice']) for l in inv['lineItems']], [('Service Proposal Amount', 1300)])
            todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
            self.assertTrue(any(t['link'].get('case_id') == cid and t['title'].startswith('Check with Tom') for t in todos))
        finally:
            P.JOBBER_STATIC_TOKEN = saved

    def test_approved_quote_typed_by_hand_joins_the_vendor_quotes_job(self):
        j = self.c.post('/pumps/api/cases', json={'title': 'Carlisle back station', 'client_name': 'Greenscapes',
                                                  'vendor_quote_total': 1158.99}).get_json()
        cid = j.get('case_id') or j['case']['id']
        today = P._today().isoformat()
        conn = P._conn()
        try:
            for jid, title, total, client in (('Q9136', 'Proposal to inspect suction line', 1506.69, 'Greenscapes'),
                                              ('Q9200', 'Irrigation repairs as well', 800.0, 'Greenscapes'),
                                              ('Q9201', 'Replace pump motor', 2600.0, 'Greenscapes')):
                conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, total, client_name, "
                             "client_company, created_at, updated_at, category) VALUES (?, 'quote', ?, ?, 'approved', ?, ?, 1, ?, ?, ?)",
                             (jid, jid[1:], title, total, client, today, today, P._category(title)))
            got = P.auto_track_approved(conn)
            conn.commit()
            linked = conn.execute("SELECT case_id FROM pump_jobber_items WHERE jobber_id='Q9136'").fetchone()[0]
            other = conn.execute("SELECT case_id FROM pump_jobber_items WHERE jobber_id='Q9200'").fetchone()[0]
        finally:
            conn.close()
        self.assertEqual(linked, cid, 'joins the Wettech quote job by price')
        self.assertIsNone(other, 'not pump work, no vendor quote, not an account we follow - left alone')
        self.assertEqual(len(got), 2)
        # The same kind of quote for an account we follow is followed.
        conn = P._conn()
        try:
            conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, title, status, total, client_name, "
                         "client_company, created_at, updated_at, category) VALUES ('Q9300', 'quote', '9300', "
                         "'Repairs as discussed', 'approved', 450, 'Quail Run', 1, ?, ?, '')", (today, today))
            P.auto_track_approved(conn)
            conn.commit()
            self.assertTrue(conn.execute("SELECT case_id FROM pump_jobber_items WHERE jobber_id='Q9300'").fetchone()[0])
        finally:
            conn.close()
        self.assertTrue(self.case(cid)['steps']['client_approved']['at'])
        # Tracked from Jobber's Greenscapes: on the job it's the pump account, The Carlisle.
        conn = P._conn()
        try:
            gid = conn.execute("SELECT client_id FROM pump_site_aliases WHERE place='Carlisle' LIMIT 1").fetchone()[0]
            self.assertEqual(P.account_for_jobber(conn, {'client_id': gid, 'client_name': 'Greenscapes'}),
                             'Carlisle (The Carlisle)')
        finally:
            conn.close()
        # A quote whose title says nothing: Wettech in a note on it is enough.
        saved = (P.jobber_status, P.fetch_quote_lines, P.jobber_gql)
        try:
            P.jobber_status = lambda: {'connected': True}
            P.fetch_quote_lines = lambda q: {'line_items': [{'name': 'Service Proposal Amount', 'description': 'Inspect suction line'}]}
            P.jobber_gql = lambda q, v=None: {'quote': {'notes': {'nodes': [{'message': 'Wettech quote dated 10/05/2026'}]}}}
            self.assertTrue(P._quote_is_pump_work('Q1'))
            P.jobber_gql = lambda q, v=None: {'quote': {'notes': {'nodes': [{'message': 'Client asked for Tuesday'}]}}}
            self.assertFalse(P._quote_is_pump_work('Q1'))
        finally:
            P.jobber_status, P.fetch_quote_lines, P.jobber_gql = saved
        self.assertFalse(P._category('Irrigation repairs as well'))
        self.assertEqual(P._category('Well pump motor replacement'), 'pump')

    def test_new_pump_account_linked_to_jobber(self):
        r = self.c.post('/pumps/api/maint', json={'name': 'XYZ Community', 'kind': 'Pump', 'months': '1,4,7,10',
                                                  'our_bill': '$300.00', 'active': True,
                                                  'jobber_client_id': 'C77', 'jobber_client_name': 'XYZ HOA',
                                                  'jobber_property_id': 'P77', 'jobber_property_label': '1 XYZ Way, Naples'}).get_json()
        self.assertTrue(r['success'], r)
        acct = [a for a in r['accounts'] if a['name'] == 'XYZ Community'][0]
        self.assertEqual(acct['jobber_client_id'], 'C77')
        conn = P._conn()
        try:
            alias = conn.execute("SELECT client_id, property_id FROM pump_site_aliases WHERE place='XYZ Community'").fetchone()
            self.assertEqual(tuple(alias), ('C77', 'P77'))
            self.assertTrue(P._tracked_account(conn, {'client_id': 'C77', 'client_name': 'XYZ HOA'}))
        finally:
            conn.close()
        todos = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertIn("Set up XYZ Community's pump maintenance job in Jobber", todos)
        # Picking the job later clears that to-do, and the account page links it.
        self.c.post('/pumps/api/maint', json={'id': acct['id'], 'jobber_client_id': 'C77', 'jobber_client_name': 'XYZ HOA',
                                              'jobber_job_id': 'J77', 'jobber_job_label': '#1700 Quarterly Pump'})
        todos = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertNotIn("Set up XYZ Community's pump maintenance job in Jobber", todos)
        h = self.c.get('/pumps/api/accounts/history?name=XYZ Community').get_json()
        self.assertEqual(h['jobber']['client']['id'], 'C77')
        self.assertEqual(h['jobber']['recurring'][0]['id'], 'J77')

    def test_edit_a_todo_and_stale_todos_close_themselves(self):
        todo = self.c.post('/pumps/api/todos', json={'title': 'Call Tom'}).get_json()
        tid = [t for t in todo['todos'] if t['title'] == 'Call Tom'][0]['id']
        r = self.c.post(f'/pumps/api/todos/{tid}/edit', json={'title': 'Call Tom about Lely', 'due_on': '2026-10-20'})
        self.assertEqual(r.status_code, 200)
        t = [t for t in self.c.get('/pumps/api/todos').get_json()['todos'] if t['id'] == tid][0]
        self.assertEqual((t['title'], t['due_on']), ('Call Tom about Lely', '2026-10-20'))
        self.assertEqual(self.c.post(f'/pumps/api/todos/{tid}/edit', json={'title': ''}).status_code, 400)
        # The county's WO/PO to-do goes once both are on the job.
        j = self.c.post('/pumps/api/cases', json={'title': 'Twin Lakes motor', 'client_name': 'SARASOTA COUNTY'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        P._gov_refs_todo({'id': cid, 'title': 'Twin Lakes motor'}, 'SARASOTA COUNTY')
        titles = lambda: [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
        self.assertTrue(any('WO # and PO #' in x for x in titles()))
        self.c.post(f'/pumps/api/cases/{cid}', json={'po_number': '260904', 'wo_number': 'WO2143416'})
        self.assertFalse(any('WO # and PO #' in x for x in titles()))
        # Quote converted to a job by hand in Jobber: the sync sees it, the to-do goes.
        conn = P._conn()
        conn.execute('UPDATE pump_cases SET jobber=? WHERE id=?', (json.dumps({'quote': {'id': 'Q9186', 'number': '9186'}}), cid))
        P.add_todo(conn, f'quote-to-job:{cid}', P.quote_to_job_title('Magnolia Cove', '9186'), '', {'case_id': cid})
        conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status) VALUES ('Q9186', 'quote', '9186', 'approved')")
        conn.commit()
        conn.close()
        self.assertIn('Convert quote for Magnolia Cove (#9186) to a job', titles())
        conn = P._conn()
        conn.execute("UPDATE pump_jobber_items SET status='converted' WHERE jobber_id='Q9186'")
        conn.commit()
        conn.close()
        self.assertNotIn('Convert quote for Magnolia Cove (#9186) to a job', titles())

    def test_sync_finds_approved_quotes_whose_title_says_nothing(self):
        today = P._today().isoformat()
        node = lambda i, title, desc, client: {
            'id': i, 'quoteNumber': i[1:], 'title': title, 'quoteStatus': 'APPROVED', 'createdAt': today, 'updatedAt': today,
            'jobberWebUri': '', 'amounts': {'total': 1506.69}, 'client': {'id': 'CX' + i, 'name': client, 'isCompany': True},
            'property': None, 'lineItems': {'nodes': [{'name': 'Service Proposal Amount', 'description': desc}]}}
        calls = []

        def fake(q, v=None):
            calls.append(q)
            if 'filter: { status: approved }' in q:
                return {'quotes': {'nodes': [node('Q9136', 'Proposal to inspect suction line',
                                                  'Field service to check out pump station, found pipe had melted', 'Greenscapes'),
                                             node('Q9500', 'Sod and irrigation repairs', 'Replace sod by the clubhouse', 'Some HOA')],
                                   'pageInfo': {'hasNextPage': False}}}
            return {'quotes': {'nodes': [], 'pageInfo': {'hasNextPage': False}}}
        saved = P.jobber_gql
        P.jobber_gql = fake
        try:
            seen = {}
            errs = P._sync_recent_quotes(seen)
        finally:
            P.jobber_gql = saved
        self.assertEqual(errs, [])
        self.assertIn('Q9136', seen, 'pump station in the line - kept')
        self.assertEqual((seen['Q9136']['status'], seen['Q9136']['category']), ('approved', 'pump'))
        self.assertNotIn('Q9500', seen, 'sod work for someone else - left out')

    def test_sarasota_county_is_billed_the_county_way(self):
        doc = {'line_items': [{'description': 'Field service to furnish and install new 15 HP vertical hollow shaft '
                                              'motor on pump #1 at Twin Lakes, wired up and tested.', 'amount': 5226.24}],
               'subtotal': 5226.24, 'client_name': 'Twin Lakes Park', 'description': ''}
        case = {'id': 0, 'client_name': 'SARASOTA COUNTY', 'po_number': '260904', 'wo_number': 'WO2143416'}
        s = P.suggest_invoice(doc, case)
        self.assertEqual(s['subject'], 'Invoice-PO260904')
        li = s['line_items']
        self.assertEqual(len(li), 1)
        self.assertEqual((li[0]['name'], li[0]['unit_price'], li[0]['taxable']), ('PUMP SERVICE', 5748.86, False))
        self.assertIn('WO2143416  PO260904', li[0]['description'])
        self.assertIn('Subcontractor Cost: $5,226.24\nVendor 10% Markup: $522.62\nTotal Cost: $5,748.86', li[0]['description'])
        q = P.suggest_quote(doc, case)
        self.assertEqual((q['line_items'][0]['unit_price'], q['markup_pct']), (5748.86, 10))
        # Everyone else: still one Service Proposal Amount line at +30%.
        other = P.suggest_invoice(doc, {'client_name': 'Quail Run'})
        self.assertEqual((other['line_items'][0]['name'], other['line_items'][0]['unit_price']),
                         ('Service Proposal Amount', 6794.11))
        self.assertEqual(P.gov_missing_refs({'po_number': '1'}), ['WO #'])

    def test_status_updates_on_an_approved_job(self):
        j = self.c.post('/pumps/api/cases', json={'title': 'Motor', 'client_name': 'Quail Run'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        r = self.c.post(f'/pumps/api/cases/{cid}/status', json={'kind': 'parts_ordered', 'text': '2 weeks out',
                                                                 'on': '2026-10-27'}).get_json()
        self.assertTrue(r['success'], r)
        self.assertEqual((r['case']['job_status']['label'], r['case']['job_status']['on']), ('Parts ordered', '2026-10-27'))
        self.c.post(f'/pumps/api/cases/{cid}/status', json={'kind': 'client_notified', 'text': 'emailed Lauren'})
        r = self.c.post(f'/pumps/api/cases/{cid}/status', json={'kind': 'scheduled', 'on': '2026-10-29'}).get_json()
        c = r['case']
        self.assertEqual(c['scheduled_for'], '2026-10-29')
        self.assertTrue(c['steps']['scheduled']['at'])
        hist = [e['detail'] for e in c['events'] if e['action'] == 'status']
        self.assertEqual(hist, ['Scheduled · for 2026-10-29', 'Client notified · emailed Lauren',
                                'Parts ordered · 2 weeks out · parts due 2026-10-27'])
        self.assertEqual(self.c.post(f'/pumps/api/cases/{cid}/status', json={'kind': 'scheduled'}).status_code, 400)
        self.assertEqual(self.c.get('/pumps/api/undo').get_json()['actions'][0]['label'], f'Status update on job #{cid}')

    def test_sign_docx_types_the_name_and_date_under_thank_you(self):
        body = ('<w:p><w:r><w:t>Your Cost ---- $ 1158.99</w:t></w:r></w:p>'
                '<w:p><w:pPr><w:pStyle w:val="NoSpacing"/></w:pPr><w:r><w:t xml:space="preserve">      Thank You </w:t></w:r></w:p>'
                '<w:p><w:pPr><w:pStyle w:val="NoSpacing"/></w:pPr><w:r><w:rPr><w:bCs/></w:rPr><w:tab/></w:r>'
                '<w:r><w:drawing/></w:r></w:p>'
                '<w:p><w:r><w:t>H. H. (Tom) Morgan III</w:t></w:r></w:p>')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            z.writestr('[Content_Types].xml', '<Types/>')
            z.writestr('word/document.xml', f'<w:document xmlns:w="w"><w:body>{body}</w:body></w:document>')
        out = P.sign_docx(buf.getvalue(), 'Simon Weardon 10/12/2026')
        xml = zipfile.ZipFile(io.BytesIO(out)).read('word/document.xml').decode()
        paras = re.findall(r'<w:p>.*?</w:p>', xml)
        self.assertIn('<w:rPr><w:bCs/></w:rPr><w:t xml:space="preserve">Simon Weardon 10/12/2026</w:t>', paras[2])
        self.assertLess(paras[2].index('Simon Weardon'), paras[2].index('<w:drawing/>'))
        self.assertIn('Thank You', paras[1])
        self.assertEqual(P.expected_date('Should be done by 10/20', date(2026, 10, 12)), '2026-10-20')
        self.assertEqual(P.expected_date('We will be out January 5th', date(2026, 12, 20)), '2027-01-05')
        self.assertEqual(P.expected_date('Will check and let you know', date(2026, 10, 12)), '')
        # This Jobber's note attachments take a url only (no fileName / contentType): send just that.
        saved = dict(P._SCHEMA_CACHE)
        try:
            P._SCHEMA_CACHE['type:NoteAttachmentAttributes'] = {'inputFields': [{'name': 'url'}]}
            self.assertEqual(P._note_attachment('https://x/f', 'a.pdf', 'application/pdf'), {'url': 'https://x/f'})
        finally:
            P._SCHEMA_CACHE.clear(); P._SCHEMA_CACHE.update(saved)

    def test_job_from_quote_and_service_visit_when_jobber_has_them(self):
        T = lambda name, kind='INPUT_OBJECT': {'kind': kind, 'name': name, 'ofType': None}
        NN = lambda t: {'kind': 'NON_NULL', 'name': None, 'ofType': t}

        class Fake(FakeJobber):
            def __call__(self, query, variables=None):
                if '__schema' in query:
                    return {'__schema': {'mutationType': {'fields': [
                        {'name': 'jobCreate', 'type': T('JobCreatePayload', 'OBJECT'),
                         'args': [{'name': 'input', 'type': NN(T('JobCreateAttributes'))}]},
                        {'name': 'visitCreate', 'type': T('VisitCreatePayload', 'OBJECT'),
                         'args': [{'name': 'jobId', 'type': NN(T('EncodedId', 'SCALAR'))},
                                  {'name': 'input', 'type': NN(T('VisitCreateAttributes'))}]}]}}}
                if '__type' in query:
                    n = variables['n']
                    return {'__type': {
                        'JobCreateAttributes': {'inputFields': [{'name': 'quoteId', 'type': NN(T('EncodedId', 'SCALAR'))},
                                                                {'name': 'title', 'type': T('String', 'SCALAR')}]},
                        'VisitCreateAttributes': {'inputFields': [{'name': 'title', 'type': NN(T('String', 'SCALAR'))},
                                                                  {'name': 'instructions', 'type': T('String', 'SCALAR')}]},
                        'JobCreatePayload': {'fields': [{'name': 'job'}, {'name': 'userErrors'}]},
                        'VisitCreatePayload': {'fields': [{'name': 'visit'}, {'name': 'userErrors'}]},
                        'EncodedId': {}}.get(n, {})}
                if 'jobCreate(' in query:
                    self.calls.append((query, variables))
                    return {'jobCreate': {'job': {'id': 'J9', 'jobNumber': 14500, 'jobberWebUri': 'https://secure.getjobber.com/work_orders/9'},
                                          'userErrors': []}}
                if 'visitCreate(' in query:
                    self.calls.append((query, variables))
                    return {'visitCreate': {'visit': {'id': 'V1'}, 'userErrors': []}}
                if 'client(id' in query:
                    return {'client': {'id': variables['id'], 'name': 'Banyan Bay', 'properties': self.properties,
                                       'jobs(first: 40)': {'nodes': [
                                           {'id': 'JM', 'jobNumber': 13196, 'title': 'Quarterly Pump Maintenance',
                                            'jobStatus': 'active', 'jobType': 'RECURRING', 'property': {'id': 'P1'}}]}}}
                return super().__call__(query, variables)
        P._SCHEMA_CACHE.clear()
        fake = Fake()
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            self.extracts['q-0910.pdf'] = extraction('quote', 'Q-910', po='PO910', client='Lakeside Pines', subtotal=500)
            q = self.upload('q-0910.pdf', data=self._pdf())
            self.draft(q['doc_id'])
            self.c.post(f"/pumps/api/cases/{q['case_id']}/approved", json={})
            self.assertEqual(self.case(q['case_id'])['jobber']['job']['number'], '14500')
            made = [v for qq, v in fake.calls if 'jobCreate(' in qq][0]
            self.assertEqual(made['input']['quoteId'], 'Q1')
            # A plain email about a pump at one of our accounts: an item and a visit on its maintenance job.
            fake.clients = [{'id': 'CB', 'name': 'Banyan Bay C/O SAK & Associated Mgmt, Inc', 'isLead': False,
                             'isArchived': False}]
            res = P.handle_service_call_email('m-1', 'manager@banyanbay.org', 'RE: Pump not running at Banyan Bay',
                                              'The pump by the clubhouse tripped again this morning.')
            self.assertEqual(res['account'], 'Banyan Bay')
            self.assertTrue(res['visit'])
            v = [v for qq, v in fake.calls if 'visitCreate(' in qq][0]
            self.assertEqual((v['jobId'], v['input']['title']), ('JM', 'Service call - Pump not running at Banyan Bay'))
            self.assertEqual(self.case(res['case_id'])['title'], 'Service call - Pump not running at Banyan Bay')
            self.assertIsNone(P.handle_service_call_email('m-1', 'manager@banyanbay.org', 'RE: Pump not running at Banyan Bay', ''))
        finally:
            P.JOBBER_STATIC_TOKEN = saved
            P._SCHEMA_CACHE.clear()

    def test_service_call_emails(self):
        # Not connected to Jobber: the visit becomes a to-do.
        res = P.handle_service_call_email('m-2', 'board@caymas.com', 'Caymas lake fountain is off', 'Please send someone.')
        self.assertEqual((res['account'], res['visit']), ('Caymas', False))
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['todos']
        self.assertIn("Add a \"Service call - Caymas lake fountain is off\" visit to Caymas's maintenance job",
                      [t['title'] for t in todos])
        self.assertEqual(self.case(res['case_id'])['vendor'], 'Gulfshore')
        # Not ours to act on: a vendor, nothing about pumps, or no account named.
        self.assertIsNone(P.handle_service_call_email('m-3', 'tomm@wettec.biz', 'Pump at Caymas', ''))
        self.assertIsNone(P.handle_service_call_email('m-4', 'a@b.com', 'Caymas invoice question', 'When is it due?'))
        self.assertIsNone(P.handle_service_call_email('m-5', 'a@b.com', 'Pump order for the shop', 'PO 5512'))

    def test_scada_coming_due_is_quoted_at_their_price(self):
        today = P._today
        P._today = lambda: P.datetime(2026, 12, 20).date()
        fake = FakeJobber(clients=[{'id': 'CL', 'name': 'LELY CDD', 'isLead': False, 'isArchived': False}])
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            done = P.scada_due_actions()
            self.assertIn('LELY', done)
            q = [v for qq, v in fake.calls if 'quoteCreate(' in qq and v['attributes']['clientId'] == 'CL'][0]['attributes']
            line = q['lineItems'][0]
            self.assertEqual((line['name'], line['unitPrice'], line['taxable']), ('Service Proposal Amount', 600.0, False))
            self.assertIn('for 2027-2028', line['description'])
            n = len([1 for qq, _ in fake.calls if 'quoteCreate(' in qq])
            P.scada_due_actions()
            self.assertEqual(len([1 for qq, _ in fake.calls if 'quoteCreate(' in qq]), n, 'once per renewal')
            todos = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
            self.assertIn('Send the SCADA renewal quote #812 - LELY', todos)
            self.assertIn('SCADA renewal due - AUTUMN WOODS (complimentary)', todos)
            self.assertFalse(any('Quote the SCADA renewal - AUTUMN' in t for t in todos))
        finally:
            P.JOBBER_STATIC_TOKEN = saved
            P._today = today

    def test_scada_vendor_quote_is_not_marked_up(self):
        doc = {'description': 'Renewal of Annual Cellular and cloud subscription for SCADA system', 'client_name': 'Banyan Bay',
               'line_items': [{'name': 'SCADA', 'description': 'Annual subscription - SCADA', 'unit_price': 428.0,
                               'quantity': 1}]}
        s = P.suggest_quote(doc)
        self.assertEqual((s['line_items'][0]['unit_price'], s['markup_pct']), (600.0, 0))
        doc['description'] = 'Replace pressure switch'
        doc['line_items'][0].update(name='Pressure switch', description='Replace pressure switch')
        self.assertEqual(P.suggest_quote(doc)['line_items'][0]['unit_price'], 556.0)

    def test_daily_email(self):
        sent = []

        class Resp:
            status_code, content = 202, b''
        real_post = P.http_requests.post
        P.http_requests.post = lambda url, **kw: (sent.append((url, kw['json'])), Resp())[1]
        P.CFG['graph_token'], P.CFG['mail_from'] = (lambda: 'tok'), 'PO@stahlman-england.com'
        try:
            self.c.post('/pumps/api/todos', json={'title': 'Call the Caymas board'})
            r = self.c.post('/pumps/api/digest/send', json={}).get_json()
            self.assertTrue(r['sent'], r)
            m = sent[-1][1]['message']
            self.assertEqual(m['toRecipients'][0]['emailAddress']['address'], 'simon@stahlman-england.com')
            self.assertIn('Call the Caymas board', m['body']['content'])
            self.assertTrue(m['subject'].startswith('Pumps today'))
            # The 8am send happens once a weekday.
            P._now, real_now = (lambda: P.datetime(2026, 10, 6, 8, 2, tzinfo=P.TZ)), P._now
            P._today, real_today = (lambda: P.datetime(2026, 10, 6).date()), P._today
            try:
                n = len(sent)
                P._scheduled_hourly()
                P._scheduled_hourly()
                self.assertEqual(len([u for u, _ in sent[n:] if 'sendMail' in u and 'Pumps today' in str(_)]), 1)
            finally:
                P._now, P._today = real_now, real_today
        finally:
            P.http_requests.post = real_post
            P.CFG['graph_token'] = None

    def test_accounts_in_one_list(self):
        rows = {a['name']: a for a in self.c.get('/pumps/api/accounts').get_json()['accounts']}
        bb = rows['Banyan Bay']
        self.assertTrue(bb['pump'] and bb['lake'] and bb['scada'])
        self.assertEqual(bb['scada']['client'], 'BANYAN BAY')
        car = rows['Carlisle (The Carlisle)']
        self.assertTrue(car['pump'] and car['lake'])
        self.assertFalse(any(r['pump'] is None and r['name'] == 'Carlisle (The Carlisle)' for r in rows.values()))

    def test_quotes_and_invoices_are_one_service_proposal_line(self):
        doc = {'subtotal': 600.0, 'total': 639.0, 'description': 'Pump repair', 'line_items': [
            {'name': 'After making repairs to get pump running', 'description': 'After making repairs to get pump running, replace the contactor.',
             'quantity': 1, 'unit_price': 450.0, 'amount': 450.0, 'taxable': True},
            {'name': 'Labor', 'description': 'Labor 2 hours', 'quantity': 2, 'unit_price': 75.0, 'amount': 150.0, 'taxable': True},
            {'name': 'Sales Tax', 'description': 'Sales Tax 6.5%', 'amount': 39.0, 'is_tax': True}]}
        for sugg in (P.suggest_quote(doc), P.suggest_invoice(doc)):
            self.assertEqual(len(sugg['line_items']), 1)
            line = sugg['line_items'][0]
            self.assertEqual((line['name'], line['quantity'], line['unit_price']), ('Service Proposal Amount', 1, 780.0))
            self.assertIn('replace the contactor', line['description'])
            self.assertIn('Labor 2 hours', line['description'])
            self.assertNotIn('Sales Tax', line['description'])

    def test_bill_without_quote_can_be_billed_plus_30(self):
        self.extracts['b-noq2.pdf'] = extraction('bill', '30990', client='Coral Isles', subtotal=500)
        b = self.upload('b-noq2.pdf')
        issue = [i for i in self.case(b['case_id'])['issues'] if i['kind'] == 'no_quote' and not i['resolved_at']][0]
        r = self.c.post(f"/pumps/api/issues/{issue['id']}/bill-anyway", json={}).get_json()
        self.assertEqual(r['doc_id'], b['doc_id'])
        case = self.case(b['case_id'])
        self.assertEqual([i for i in case['issues'] if not i['resolved_at']], [])
        self.assertTrue({s['key']: s for s in case['step_list']}['bill_checked'].get('at'))
        doc = self.c.get(f"/pumps/api/docs/{b['doc_id']}").get_json()['doc']
        self.assertEqual(doc['invoice_suggestion']['line_items'][0]['unit_price'], 650.0)

    def test_maintenance_bill_needs_no_quote_and_bills_the_recurring_job(self):
        """Wettech's bill for Allura's quarterly maintenance (Oct 2026): no quote is waited on, and our invoice
        is the recurring job's price - not the bill plus 30%."""
        items = [{'name': 'Maintenance', 'description': 'Service to perform quarterly maintenance on irrigation system '
                  'for OCTOBER 2026. Maintenance Report Attached!', 'quantity': 1, 'unit_price': 125.0,
                  'amount': 125.0, 'taxable': False, 'is_tax': False}]
        ex = extraction('bill', '29139', client='Allura', subtotal=125, tax=0, items=items)
        ex['description'] = items[0]['description']
        self.extracts['b-allura.pdf'] = ex
        b = self.upload('b-allura.pdf')
        case = self.case(b['case_id'])
        self.assertEqual(case['category'], 'maintenance')
        self.assertNotIn(case['stage'], ('vendor_quote', 'client_quote', 'client_approved'))
        self.assertEqual([i for i in case['issues'] if not i['resolved_at']], [])
        job = {'id': 'JR', 'number': '13196', 'title': 'Quarterly Pump Maintenance', 'uri': '', 'total': 320.0,
               'status': 'active', 'type': 'recurring', 'instructions': '', 'property_label': '',
               'line_items': [{'name': 'Quarterly Pump Maintenance', 'description': '', 'quantity': 1,
                               'unit_price': 320.0, 'taxable': False}]}
        saved = (P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail)
        P.jobber_status = lambda: {'connected': True}
        P.invoice_client = lambda doc, case=None: 'CA'
        P.recurring_jobs = lambda cid: [{'id': 'JR', 'number': '13196', 'title': job['title'], 'status': 'active'}]
        P.job_detail = lambda jid: job
        try:
            doc = P._doc_dict(P._conn().execute('SELECT * FROM pump_docs WHERE id=?', (b['doc_id'],)).fetchone())
            sugg = P.invoice_suggestion(doc, dict(P._conn().execute('SELECT * FROM pump_cases WHERE id=?',
                                                                    (b['case_id'],)).fetchone()))
        finally:
            P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail = saved
        self.assertEqual([li['unit_price'] for li in sugg['line_items']], [320.0])
        self.assertEqual(sugg['job']['id'], 'JR')
        # Wettech's payment terms next to the P.O. box aren't a PO number.
        self.extracts['b-enclave.pdf'] = {**ex, 'doc_number': '29140', 'client_name': 'Enclave @ Palmira',
                                          'po_number': '1.5% Finance Chrg ov'}
        e = self.upload('b-enclave.pdf')
        self.assertEqual(self.case(e['case_id'])['po_number'], '')
        self.assertEqual(P.clean_po('PO322'), 'PO322')
        conn = P._conn()
        conn.execute("UPDATE pump_cases SET po_number='1.5% Finance Chrg ov' WHERE id=?", (e['case_id'],))
        conn.commit()
        conn.close()
        self.c.get('/pumps/api/summary')
        self.assertEqual(self.case(e['case_id'])['po_number'], '')
        # A bill that names a quote still waits on it.
        self.assertFalse(P.maintenance_bill({'kind': 'bill', 'description': 'Pump maintenance repairs as per quote'}))

    def test_report_is_saved_in_jobber_under_the_vendors_name(self):
        """Jobber names a note's file after the link's last part (Oct 2026: a hash), so the link ends with the
        name the vendor gave the report."""
        b = self.upload('b-0950.pdf')
        saved = P.CFG.get('website_url')
        P.CFG['website_url'] = 'https://example.com'
        try:
            url = P._signed_file_url(b['doc_id'], 'original', name='10-26 Evergreen Pump Report.pdf')
        finally:
            P.CFG['website_url'] = saved
        self.assertTrue(url.endswith('/10-26%20Evergreen%20Pump%20Report.pdf'), url)
        r = self.c.get(url.replace('https://example.com', ''))
        self.assertEqual(r.status_code, 200)

    def test_report_only_accounts_get_no_maintenance_invoice(self):
        """The office (Oct 15 2026): Autumn Woods, The Reserve @ Estero and Wild Blue only get their pump report
        logged - Wettech's maintenance bill there is never invoiced."""
        items = [{'name': 'Maintenance', 'description': 'Service to perform monthly maintenance on pump station '
                  'for OCTOBER 2026.', 'quantity': 1, 'unit_price': 250.0, 'amount': 250.0, 'taxable': False,
                  'is_tax': False}]
        ex = extraction('bill', '29201', client='Wild Blue', subtotal=250, tax=0, items=items)
        ex['description'] = items[0]['description']
        self.extracts['b-wildblue.pdf'] = ex
        saved = P.jobber_status
        P.jobber_status = lambda: {'connected': True}
        try:
            b = self.upload('b-wildblue.pdf')
        finally:
            P.jobber_status = saved
        self.assertFalse(b.get('auto_invoice'), b)
        case = self.case(b['case_id'])
        self.assertEqual(case['category'], 'maintenance')
        self.assertTrue(case['steps']['invoice_drafted'].get('na'), case['steps'])
        self.assertTrue(case['steps']['closed'].get('na'), case['steps'])
        self.assertEqual(case['stage'], 'report_logged')
        drafts = self.c.get('/pumps/api/summary').get_json()['queue']['bills_to_draft']
        self.assertNotIn(b['case_id'], [d['case_id'] for d in drafts])
        # Wettech is paid once the account's monthly is billed - from the end of the bill's month.
        pay = case['vendor_pay']
        self.assertTrue(pay['report_only'], pay)
        saved = P._today
        P._today = lambda: date.fromisoformat(pay['month_end']).replace(day=1)
        try:
            self.assertEqual(self.case(b['case_id'])['vendor_pay']['state'], 'unpaid')
            q = self.c.get('/pumps/api/summary').get_json()['queue']
            self.assertNotIn(b['case_id'], [c['id'] for c in q['vendor_bills_to_pay']])
        finally:
            P._today = saved
        P._today = lambda: date.fromisoformat(pay['month_end'])
        try:
            q = self.c.get('/pumps/api/summary').get_json()['queue']
            self.assertIn(b['case_id'], [c['id'] for c in q['vendor_bills_to_pay']])
            conn = P._conn()
            self.assertIn('monthly has been billed', P.vendor_pay_email(conn, b['case_id'])['body'])
            conn.close()
        finally:
            P._today = saved
        # Anyone else's maintenance bill still gets invoiced.
        self.extracts['b-allura2.pdf'] = {**ex, 'doc_number': '29202', 'client_name': 'Allura'}
        a = self.upload('b-allura2.pdf')
        self.assertFalse((self.case(a['case_id'])['steps'].get('invoice_drafted') or {}).get('na'))

    def test_report_client_choices_are_clients_with_pump_work(self):
        """Rosewood (Oct 2026): logging a report offered everyone with "Rosewood" in an address or note. Only
        clients we do pump work for are offered, the rest behind "show all"."""
        cands = [{'id': 'C1', 'name': 'Rosewood HOA c/o Guardian', 'score': 1.0, 'is_lead': False, 'address': '',
                  'uri': '', 'matching_properties': []},
                 {'id': 'C2', 'name': 'Ken Kessler', 'score': 0.95, 'is_lead': False, 'address': '301 Rosewood Court',
                  'uri': '', 'matching_properties': []}]
        conn = P._conn()
        conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, client_id, category) "
                     "VALUES ('J1', 'job', '77', 'active', 'C1', 'pump')")
        conn.commit()
        conn.close()
        saved = P.search_clients
        P.search_clients = lambda name, limit=8: [dict(c) for c in cands]
        try:
            j = self.c.get('/pumps/api/jobber/clients?pump=1&q=Rosewood').get_json()
            self.assertEqual(([c['id'] for c in j['candidates']], j['hidden']), (['C1'], 1))
            j = self.c.get('/pumps/api/jobber/clients?pump=1&all=1&q=Rosewood').get_json()
            self.assertEqual([c['id'] for c in j['candidates']], ['C1', 'C2'])
            # Other searches (bills, quotes) are unchanged.
            j = self.c.get('/pumps/api/jobber/clients?q=Rosewood').get_json()
            self.assertEqual(len(j['candidates']), 2)
        finally:
            P.search_clients = saved

    def test_log_report_remembers_the_client_and_job(self):
        """The client and pump job a report was logged on are remembered for the vendor's name: the next
        Log report window starts on them, and the next report goes there by itself."""
        conn = P._conn()
        conn.execute("INSERT INTO pump_docs (kind, client_name, site, file_name, status) "
                     "VALUES ('report', 'Brookside Lakes', 'Brookside Lakes', 'r.docx', 'filed')")
        did = conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.commit()
        conn.close()
        saved = (P.log_report_to_jobber, P.search_clients)
        P.log_report_to_jobber = lambda doc_id, t, tid, actor, msg=None: {'note_id': 'N1', 'note_attached_file': True}
        P.search_clients = lambda name, limit=8: []
        try:
            r = self.c.post(f'/pumps/api/docs/{did}/report_note', json={
                'target_type': 'job', 'target_id': 'J1609', 'client_id': 'C1', 'client_name': 'Brookside Lakes HOA',
                'job_label': '#1609 Monthly Pump'})
            self.assertEqual(r.status_code, 200, r.get_json())
            j = self.c.get('/pumps/api/jobber/clients?pump=1&q=Brookside%20Lakes').get_json()
        finally:
            P.log_report_to_jobber, P.search_clients = saved
        self.assertEqual((j['pick']['id'], j['pick']['job_id']), ('C1', 'J1609'))
        t = P.resolve_report_target({'client_name': 'Brookside Lakes', 'site': 'Brookside Lakes'})
        self.assertEqual((t['type'], t['id']), ('job', 'J1609'))

    def test_maintenance_bill_as_per_quotation_takes_the_picked_clients_job_price(self):
        """Rosewood (Oct 2026): "Service to perform irrigation maintenance ... OCTOBER 2026. As per quotation."
        is the month's visit under the contract, not quoted work - and once the office picks the client, the
        invoice is that client's recurring job price, not $150 + 30%."""
        desc = ('Service to perform irrigation maintenance to ensure equipment is operating properly. '
                'OCTOBER 2026. As per quotation.')
        bill = {'kind': 'bill', 'client_name': 'Rosewood Irrigation Maintenance', 'subtotal': 150.0,
                'description': desc, 'line_items': [{'name': 'Rosewood Irrigation Maintenance', 'description': desc}]}
        self.assertTrue(P.maintenance_bill(bill))
        self.assertFalse(P.maintenance_bill({'kind': 'bill', 'description': 'Pump maintenance repairs as per quote'}))
        self.assertFalse(P.maintenance_bill({'kind': 'bill', 'description': 'Maintenance OCTOBER 2026 per quote #812'}))
        job = {'id': 'JRW', 'number': '1700', 'title': 'Quarterly Pump', 'uri': '', 'total': 195.0,
               'status': 'active', 'type': 'recurring', 'instructions': '', 'property_label': '',
               'line_items': [{'name': 'QUARTERLY PUMP PROGRAM', 'description': '', 'quantity': 1,
                               'unit_price': 195.0, 'taxable': False}]}
        saved = (P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail)
        P.jobber_status = lambda: {'connected': True}
        P.invoice_client = lambda doc, case=None: None   # the app can't tell which Rosewood
        P.recurring_jobs = lambda cid: [{'id': 'JRW', 'number': '1700', 'title': job['title'], 'status': 'active'}] \
            if cid == 'CRW' else []
        P.job_detail = lambda jid: job
        try:
            self.assertFalse(P.invoice_suggestion(bill, {}).get('job'))
            sugg = P.invoice_suggestion(bill, {}, 'CRW')
        finally:
            P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail = saved
        self.assertEqual(([li['unit_price'] for li in sugg['line_items']], sugg['job']['id']), ([195.0], 'JRW'))
        # The bill's notes for the visit go on the line, without "As per quotation" or the report note.
        self.assertEqual(sugg['line_items'][0]['description'],
                         'Service to perform irrigation maintenance to ensure equipment is operating properly. '
                         'OCTOBER 2026.')
        self.assertEqual(P._with_bill_notes([{'name': 'MONTHLY PUMP PROGRAM', 'description': 'SITE LOCATION: Evergreen'}],
                                            [{'description': 'Service to perform maintenance on irrigation system and '
                                                             'VFD for OCTOBER 2026. Maintenance Report Attached!'}])[0]['description'],
                         'SITE LOCATION: Evergreen\nService to perform maintenance on irrigation system and VFD for OCTOBER 2026.')

    def test_move_a_job_up_to_where_it_really_is(self):
        """Carlisle (Oct 2026): Wettech emailed that the work is done while the job still read "send the quote
        to Tom" - the office moves it up to waiting on the bill in one go, and can move it back."""
        j = self.c.post('/pumps/api/cases', json={'title': 'Inspect suction line', 'client_name': 'Carlisle'}).get_json()
        cid = j.get('case_id') or j['case']['id']
        r = self.c.post(f'/pumps/api/cases/{cid}/move', json={'stage': 'vendor_bill'}).get_json()
        self.assertTrue(r['success'], r)
        c = r['case']
        self.assertEqual(c['stage'], 'vendor_bill')
        for k in ('vendor_quote', 'client_quote', 'client_approved', 'scheduled', 'work_done'):
            self.assertTrue((c['steps'].get(k) or {}).get('at'), k)
        self.assertFalse((c['steps'].get('vendor_bill') or {}).get('at'))
        self.assertFalse((c['steps'].get('assessment') or {}).get('at'), 'an optional step left off stays off')
        # Back a step: work_done is open again.
        c = self.c.post(f'/pumps/api/cases/{cid}/move', json={'stage': 'work_done'}).get_json()['case']
        self.assertEqual(c['stage'], 'work_done')
        self.assertFalse((c['steps'].get('work_done') or {}).get('at'))
        self.assertEqual(self.c.post(f'/pumps/api/cases/{cid}/move', json={'stage': 'nope'}).status_code, 400)

    def test_send_invoice_todo_links_the_jobber_invoice(self):
        """Caymas (Oct 2026): "Send invoice" opens the drafted invoice in Jobber to look over and send - found by
        its number when the app didn't draft it."""
        conn = P._conn()
        conn.execute("INSERT INTO pump_jobber_items (jobber_id, kind, number, status, web_uri) "
                     "VALUES ('I36832', 'invoice', '36832', 'draft', 'https://secure.getjobber.com/invoices/36832')")
        conn.commit()
        c = P._with_invoice_link(conn, {'id': 1, 'jobber': {}, 'sei_invoice_number': '36832'})
        self.assertEqual(c['jobber']['invoice']['uri'], 'https://secure.getjobber.com/invoices/36832')
        c = P._with_invoice_link(conn, {'id': 2, 'jobber': {'invoice': {'id': 'X', 'uri': 'https://j/inv'}}})
        self.assertEqual(c['jobber']['invoice']['uri'], 'https://j/inv')
        self.assertEqual(P._with_invoice_link(conn, {'id': 3, 'jobber': {}, 'sei_invoice_number': ''}),
                         {'id': 3, 'jobber': {}, 'sei_invoice_number': ''})
        conn.close()

    def test_maintenance_bill_does_not_land_on_a_repair_waiting_on_its_quote(self):
        """Lely (Oct 2026): Wettech's quarterly maintenance bill went on the open "New 10 in valve" repair (quote
        #9152 approved). A maintenance bill only goes on a maintenance item by name; one filed wrong can be moved."""
        rq = extraction('quote', '9152', client='Lely', subtotal=2000)
        rq['description'] = 'Furnish and install new 10 in valve'
        self.extracts['q-lely.pdf'] = rq
        q = self.upload('q-lely.pdf')
        repair = q['case_id']
        desc = 'Service to perform quarterly maintenance on pump station for OCTOBER 2026. As per quotation.'
        items = [{'name': 'Lely Maintenance', 'description': desc, 'quantity': 1, 'unit_price': 300.0, 'amount': 300.0}]
        ex = extraction('bill', '29129', client='Lely', subtotal=300, tax=0, items=items)
        ex['description'] = desc
        self.extracts['b-lely.pdf'] = ex
        b = self.upload('b-lely.pdf')
        self.assertNotEqual(b['case_id'], repair)
        self.assertEqual(self.case(b['case_id'])['category'], 'maintenance')
        self.assertFalse((self.case(repair)['steps'].get('vendor_bill') or {}).get('at'))
        # Filed on the repair before this fix: "Wrong job" takes it off, undoes what it ticked, and files it again.
        conn = P._conn()
        conn.execute('UPDATE pump_docs SET case_id=? WHERE id=?', (repair, b['doc_id']))
        P._file_on_case(conn, dict(conn.execute('SELECT * FROM pump_docs WHERE id=?', (b['doc_id'],)).fetchone()),
                        repair, 'test', 'test')
        conn.commit()
        conn.close()
        self.assertTrue(self.case(repair)['steps']['vendor_bill'].get('at'))
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/refile").get_json()
        self.assertTrue(r['success'], r)
        self.assertNotEqual(r['case_id'], repair)
        c = self.case(repair)
        self.assertFalse((c['steps'].get('vendor_bill') or {}).get('at'))
        self.assertEqual(c['vendor_bill_number'], '')
        # The same email forwarded again does it too.
        conn = P._conn()
        conn.execute('UPDATE pump_docs SET case_id=? WHERE id=?', (repair, b['doc_id']))
        P._file_on_case(conn, dict(conn.execute('SELECT * FROM pump_docs WHERE id=?', (b['doc_id'],)).fetchone()),
                        repair, 'test', 'test')
        conn.commit()
        conn.close()
        again = self.upload('b-lely.pdf')
        self.assertEqual(again['doc_id'], b['doc_id'], again)
        self.assertNotEqual(again['case_id'], repair, again)
        self.assertFalse((self.case(repair)['steps'].get('vendor_bill') or {}).get('at'))

    def test_report_comment_saying_a_quote_is_coming_goes_on_the_watchlist(self):
        """Lely's October report (Oct 2026): "Pump #4 check valve has a hole in the body ... - quote to follow" in
        the comments. That work goes on the watchlist as its own job, waiting on Wettech's quote."""
        from docx import Document
        d = Document()
        for t in ('Water Equipment Technologies of Southwest Florida LLC', 'IRRIGATION PUMP SYSTEM REPORT'):
            d.add_paragraph(t)
        t = d.add_table(rows=2, cols=4)
        t.rows[0].cells[2].text, t.rows[0].cells[3].text = 'DATE', '10-7-2026'
        for cell, text in zip(t.rows[1].cells, ('CUSTOMER', 'Lely', 'LOCATION', 'Pump station')):
            cell.text = text
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = '√', 'VERIFY PUMPS GO TO SLEEP/SHUT DOWN PROPERLY'
        d.add_paragraph('COMMENTS:')
        box = d.add_table(rows=1, cols=1)
        box.rows[0].cells[0].text = ('Pump #4 check valve has a hole in the body of the check valve and a cooling fan '
                                     'is bad in the first control panel – quote to follow')
        buf = io.BytesIO()
        d.save(buf)
        res = P.ingest_document('10-26 Lely.docx', buf.getvalue(), source='email', actor='email scan',
                                email={'uid': 'lely-1', 'from': 'office@wettec.biz', 'subject': 'Inv 29129'})
        self.assertEqual(res['kind'], 'report', res)
        cases = self.c.get('/pumps/api/cases?status=open').get_json()['cases']
        quoted = [c for c in cases if 'check valve has a hole' in c['title']]
        self.assertEqual(len(quoted), 1, [c['title'] for c in cases])
        self.assertEqual((quoted[0]['category'], quoted[0]['stage']), ('repair', 'vendor_quote'))
        self.assertEqual(quoted[0]['client_name'], 'Lely')
        # A report that says nothing of a quote adds nothing.
        self.assertEqual(P.quote_needed_note({'kind': 'report', 'report_fields': {}, 'text_excerpt': 'All good.'}), '')

    def test_watercrest_is_water_crest(self):
        """Oct 2026: Wettech's "Watercrest @ Falling Waters" scored lowest against Jobber's "WATER CREST AT FALLING
        WATERS", so its report went on Falling Waters II. Words run together still match; the wrong pick is undone."""
        wc = P.similarity('Watercrest @ Falling Waters', 'WATER CREST AT FALLING WATERS')
        fw2 = P.similarity('Watercrest @ Falling Waters', 'FALLING WATERS II AT FALLING WATERS c/o GUARDIAN PROPERTY MANAGEMENT')
        self.assertEqual(wc, 1.0)
        self.assertGreater(wc - fw2, 0.3)
        self.assertEqual(P.pick_client([{'id': 'W', 'score': wc}, {'id': 'F', 'score': fw2}])['id'], 'W')
        conn = P._conn()
        conn.execute("INSERT INTO pump_site_aliases (place, area, client_id, client_name, job_id) VALUES "
                     "('Watercrest @ Falling Waters', '', ?, 'FALLING WATERS II', ?)", (P._FW2_CLIENT, P._FW2_JOB))
        conn.execute("INSERT INTO pump_docs (kind, client_name, file_name, status, jobber) VALUES "
                     "('report', 'Watercrest @ Falling Waters', '10-26 Watercrest.docx', 'filed', ?)",
                     (json.dumps({'note_id': 'N9', 'note_target': 'job', 'note_target_id': P._FW2_JOB}),))
        did = conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.execute("DELETE FROM pump_state WHERE key='watercrest_fw2_oct26'")
        P._undo_watercrest_mismatch(conn.cursor(), P._now_text())
        conn.commit()
        self.assertIsNone(conn.execute("SELECT 1 FROM pump_site_aliases WHERE client_id=?", (P._FW2_CLIENT,)).fetchone())
        self.assertNotIn('note_id', json.loads(conn.execute('SELECT jobber FROM pump_docs WHERE id=?', (did,)).fetchone()[0]))
        conn.close()
        todos = self.c.get('/pumps/api/summary').get_json()['queue']['reports_to_log']
        self.assertIn(did, [d['id'] for d in todos])

    def test_done_on_the_to_do_list_ticks_the_job(self):
        """Oct 2026: done on the to-do list is done on the job - an invoice drafted straight in Jobber ticks
        "Draft invoice in Jobber", a report logged by hand ticks "Report logged in Jobber"."""
        self.extracts['b-sync.pdf'] = extraction('bill', '31901', client='Sync Lakes HOA', subtotal=400)
        b = self.upload('b-sync.pdf')
        q = lambda: self.c.get('/pumps/api/summary').get_json()['queue']
        self.assertIn(b['doc_id'], [d['id'] for d in q()['bills_to_draft']])
        self.c.patch(f"/pumps/api/cases/{b['case_id']}", json={'steps': {'invoice_drafted': '2026-10-15'}})
        self.assertNotIn(b['doc_id'], [d['id'] for d in q()['bills_to_draft']])
        conn = P._conn()
        conn.execute("INSERT INTO pump_docs (kind, client_name, file_name, status, case_id, jobber) "
                     "VALUES ('report', 'Sync Lakes HOA', 'r.docx', 'filed', ?, '{}')", (b['case_id'],))
        rid = conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.commit()
        conn.close()
        self.assertIn(rid, [d['id'] for d in q()['reports_to_log']])
        r = self.c.post(f'/pumps/api/docs/{rid}/report_done').get_json()
        self.assertTrue(r['success'], r)
        self.assertNotIn(rid, [d['id'] for d in q()['reports_to_log']])
        self.assertTrue(self.case(b['case_id'])['steps']['report_logged'].get('at'))

    def test_maintenance_accounts_usual_bill_is_the_monthly_visit(self):
        """Evergreen (Oct 2026): Wettech's $100 bill says nothing of maintenance, but $100 is the account's
        usual vendor cost - it's the monthly visit, invoiced at the recurring job's $250, not $130."""
        conn = P._conn()
        conn.execute("UPDATE pump_maint_accounts SET active=1, vendor_cost='$100.00' WHERE name='Evergreen (Bradenton)'")
        conn.commit()
        conn.close()
        bill = {'kind': 'bill', 'client_name': 'Evergreen c/o Access Difference', 'subtotal': 100.0,
                'description': 'Pump service', 'line_items': [{'name': 'Service', 'description': 'Pump service'}]}
        self.assertTrue(P.maint_visit_bill(bill, {}))
        self.assertFalse(P.maint_visit_bill({**bill, 'subtotal': 1665.0}, {}), 'a repair, not the visit')
        self.assertFalse(P.maint_visit_bill({**bill, 'description': 'Repairs as per quote'}, {}))
        self.assertFalse(P.maint_visit_bill({**bill, 'client_name': 'Somewhere Else HOA'}, {}))
        job = {'id': 'J1450', 'number': '1450', 'title': 'Monthly Pump 250', 'uri': '', 'total': 250.0,
               'status': 'requires_invoicing', 'type': 'recurring', 'instructions': '', 'property_label': '',
               'line_items': [{'name': 'MONTHLY PUMP PROGRAM', 'description': '', 'quantity': 1,
                               'unit_price': 250.0, 'taxable': False}]}
        saved = (P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail)
        P.jobber_status = lambda: {'connected': True}
        P.invoice_client = lambda doc, case=None: 'CE'
        P.recurring_jobs = lambda cid: [{'id': 'J1450', 'number': '1450', 'title': job['title'], 'status': 'active'}]
        P.job_detail = lambda jid: job
        try:
            sugg = P.invoice_suggestion(bill, {})
        finally:
            P.jobber_status, P.invoice_client, P.recurring_jobs, P.job_detail = saved
        self.assertEqual([li['unit_price'] for li in sugg['line_items']], [250.0])
        self.assertEqual(sugg['job']['id'], 'J1450')

    def test_maintenance_invoice_marks_the_months_visit_complete(self):
        """Oct 2026: invoicing a maintenance bill on the recurring job also completes that job's visit for the
        month the bill is for - the one nearest the bill's date - and anything that can't be done is a to-do."""
        self.addCleanup(setattr, P, '_SCHEMA_CACHE', {})
        visits = [{'id': 'V9', 'title': 'Monthly Pump 250', 'startAt': '2026-09-25T12:00:00Z', 'completedAt': None},
                  {'id': 'V10', 'title': 'Monthly Pump 250', 'startAt': '2026-10-23T12:00:00Z', 'completedAt': None},
                  {'id': 'V11', 'title': 'Monthly Pump 250', 'startAt': '2026-11-27T12:00:00Z', 'completedAt': None}]

        class Visits(FakeJobber):
            def __call__(self, query, variables=None):
                P.check_mutation_allowed(query)
                if 'mutationType' in query:
                    self.calls.append((query, variables))
                    return {'__schema': {'mutationType': {'fields': [{'name': 'visitComplete', 'type': {}, 'args': [
                        {'name': 'visitId', 'type': {'kind': 'NON_NULL', 'ofType': {'kind': 'SCALAR', 'name': 'EncodedId'}}}]}]}}}
                if 'visits(' in query:
                    self.calls.append((query, variables))
                    return {'job': {'visits': {'nodes': visits, 'pageInfo': {'hasNextPage': False}}}}
                if 'visitComplete' in query:
                    self.calls.append((query, variables))
                    return {'visitComplete': {'__typename': 'VisitCompletePayload'}}
                return super().__call__(query, variables)
        job = {'id': 'J1450', 'number': '1450', 'title': 'Monthly Pump 250', 'type': 'recurring', 'uri': 'u'}
        doc = {'id': None, 'doc_date': '2026-10-10', 'description': 'Monthly pump readings for OCTOBER 2026',
               'line_items': []}
        self.assertEqual(P.bill_service_month(doc), (2026, 10))
        self.assertEqual(P.bill_service_month({'doc_date': '2026-09-30'}), (2026, 9))
        fake = P.jobber_gql = Visits()
        r = P.complete_visit_for_bill(doc, {}, job, 'test')
        self.assertTrue(r['completed'], r)
        self.assertEqual(r['visit']['id'], 'V10')
        self.assertEqual([v for q, v in fake.calls if 'visitComplete(' in q], [{'visitId': 'V10'}])
        # Already done: nothing is sent again.
        visits[1]['completedAt'] = '2026-10-23T15:00:00Z'
        fake.calls.clear()
        self.assertTrue(P.complete_visit_for_bill(doc, {}, job, 'test')['already_complete'])
        self.assertFalse([q for q, _ in fake.calls if 'visitComplete(' in q])
        # No visit that month: a to-do, not a guess at another month.
        r = P.complete_visit_for_bill({**doc, 'description': 'for DECEMBER 2026'}, {}, job, 'test')
        self.assertIn('no December 2026 visit', r['pending'])
        todos = [row[0] for row in P._conn().execute("SELECT title FROM pump_todos")]
        self.assertIn('Mark the December 2026 visit complete on Jobber job #1450', todos)
        # A one-off job's visits are left alone.
        self.assertIsNone(P.complete_visit_for_bill(doc, {}, {**job, 'type': 'one_off'}, 'test'))

    def test_monthly_marks_the_visit_complete_then_invoices_it_as_a_draft(self):
        """Oct 2026: a monthly/quarterly is invoiced the Jobber way - the recurring job's visit is marked
        complete first, then the draft invoice is made for that visit, so Jobber doesn't still list it as
        completed and not invoiced."""
        self.extracts['b-mon.pdf'] = extraction('bill', '31501', client='Evergreen', subtotal=100)
        b = self.upload('b-mon.pdf')
        job = {'id': 'J1450', 'number': '1450', 'title': 'Monthly Pump 250', 'uri': '', 'total': 250.0,
               'status': 'active', 'type': 'recurring'}
        order, sent = [], {}
        saved = (P.invoice_suggestion, P.complete_visit_for_bill, P.gql_with_lines, P._invoice_visit_field,
                 P.existing_jobber_record)
        P.invoice_suggestion = lambda doc, case=None, client_id=None: {
            'subject': 'Pump service', 'client_id': 'CE', 'job': job, 'recurring': True,
            'line_items': [{'name': 'MONTHLY PUMP PROGRAM', 'description': '', 'quantity': 1, 'unit_price': 250.0}]}
        P.complete_visit_for_bill = lambda doc, case, j, actor='system': (
            order.append('visit'), {'completed': True, 'month': 'October 2026', 'visit': {'id': 'V10'}})[1]

        def create(query, variables, key):
            order.append('invoice')
            sent.update(variables['input'])
            return {'invoiceCreate': {'invoice': {'id': 'I1', 'invoiceNumber': 9001, 'invoiceStatus': 'draft'}}}
        P.gql_with_lines = create
        P._invoice_visit_field = lambda: 'visitIds'
        P.existing_jobber_record = lambda *a, **k: None
        try:
            r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'CE'}).get_json()
        finally:
            (P.invoice_suggestion, P.complete_visit_for_bill, P.gql_with_lines, P._invoice_visit_field,
             P.existing_jobber_record) = saved
        self.assertTrue(r['success'], r)
        self.assertEqual(order, ['visit', 'invoice'])
        self.assertEqual((sent['jobId'], sent['visitIds']), ('J1450', ['V10']))
        self.assertEqual(r['invoice']['invoice_status'], 'draft')

    def test_picking_a_client_is_remembered_for_the_vendors_name(self):
        """Wettech still writes "Lee Memorial"; the office picks Hodges Funeral Home once."""
        hodges = [{'id': 'CH', 'name': 'Hodges Funeral Home', 'isLead': False, 'isArchived': False}]
        fake = FakeJobber(clients=hodges)
        P.jobber_gql = fake
        self.extracts['b-lee1.pdf'] = extraction('bill', '31001', client='Lee Memorial', subtotal=200)
        b = self.upload('b-lee1.pdf')
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice",
                        json={'client_id': 'CH', 'client_name': 'Hodges Funeral Home'}).get_json()
        self.assertTrue(r['success'], r)
        names = {n['place']: n for n in self.c.get('/pumps/api/site-names').get_json()['site_names']}
        self.assertEqual(names['Lee Memorial']['client_name'], 'Hodges Funeral Home')
        # Searching the name now offers Hodges first, already chosen.
        j = self.c.get('/pumps/api/jobber/clients?q=Lee%20Memorial').get_json()
        self.assertEqual((j['pick']['id'], j['candidates'][0]['remembered'] != ''), ('CH', True))
        # The next Lee Memorial bill goes to Hodges without asking.
        fake.clients = []
        self.extracts['b-lee2.pdf'] = extraction('bill', '31002', client='Lee Memorial', subtotal=300)
        b2 = self.upload('b-lee2.pdf')
        r = self.c.post(f"/pumps/api/docs/{b2['doc_id']}/invoice", json={}).get_json()
        self.assertTrue(r['success'], r)
        self.assertEqual([v for q, v in fake.calls if 'invoiceCreate' in q][-1]['input']['clientId'], 'CH')
        # A client picked under its own name teaches nothing.
        self.assertIsNone(P.remember_choice({'client_name': 'Hodges Funeral Home'}, 'CH', 'Hodges Funeral Home'))

    def test_invoice_retries_without_fields_jobber_refuses(self):
        """Jobber's invoice lines don't take saveToProductsAndServices (quote lines do)."""
        class Picky(FakeJobber):
            def __call__(self, query, variables=None):
                if 'invoiceCreate' in query and 'saveToProductsAndServices' in json.dumps(variables):
                    self.calls.append((query, variables))
                    raise P.JobberError('Variable $input of type InvoiceCreateInput! was provided invalid value for '
                                        'lineItems.0.saveToProductsAndServices (Field is not defined on '
                                        'InvoiceCreationLineItemInput)')
                return super().__call__(query, variables)
        fake = Picky()
        P.jobber_gql = fake
        self.extracts['b-picky.pdf'] = extraction('bill', '31100', client='Lakeside Pines', subtotal=600)
        b = self.upload('b-picky.pdf')
        r = self.c.post(f"/pumps/api/docs/{b['doc_id']}/invoice", json={'client_id': 'C1'}).get_json()
        self.assertTrue(r['success'], r)
        sent = [v for q, v in fake.calls if 'invoiceCreate' in q][-1]['input']['lineItems'][0]
        self.assertNotIn('saveToProductsAndServices', sent)
        self.assertEqual((sent['name'], sent['unitPrice'], sent['taxable']), ('Service Proposal Amount', 780.0, True))

    def test_bill_is_invoiced_on_the_matching_jobber_job(self):
        """Wettech bills "Lee Memorial"; Hodges' job #14511 "Replace broken drop pipe" needs invoicing."""
        HODGES = 'Z2lkOi8vSm9iYmVyL0NsaWVudC8zODg2MDA3MQ=='

        class Fake(FakeJobber):
            def __call__(self, query, variables=None):
                if 'job(id:' in query:
                    self.calls.append((query, variables))
                    jobs = {
                        'J14511': {'id': 'J14511', 'jobNumber': 14511, 'title': 'Replace broken drop pipe',
                                   'instructions': 'Lee Memorial Well 6\n\nAfter making repairs to get pump running, '
                                                   'we found the pipe is broke below the well seal.',
                                   'jobStatus': 'requires_invoicing', 'jobType': 'ONE_OFF', 'total': 3471.43,
                                   'jobberWebUri': 'https://secure.getjobber.com/work_orders/158570173',
                                   'property': {'id': 'P9', 'address': {'street1': '12777 Florida 82', 'city': 'Lehigh Acres'}},
                                   'lineItems': {'nodes': [{'name': 'PROPOSAL AMOUNT SERVICE', 'description': 'Lee Memorial Well 6 ...',
                                                            'quantity': 1, 'unitPrice': 3471.43, 'taxable': False}]}},
                        'J200': {'id': 'J200', 'jobNumber': 14400, 'title': 'Install new zones at the mausoleum',
                                 'instructions': 'Four new zones', 'jobStatus': 'active', 'jobType': 'ONE_OFF', 'total': 16795,
                                 'lineItems': {'nodes': [{'name': 'Zones', 'description': 'Four zones', 'quantity': 1,
                                                          'unitPrice': 16795, 'taxable': True}]}}}
                    return {'job': jobs[variables['id']]}
                if 'client(id' in query:
                    return {'client': {'id': variables['id'], 'name': 'Hodges Funeral Home', 'properties': [],
                                       'jobs(first: 40)': {'nodes': [
                                           {'id': 'J14511', 'jobNumber': 14511, 'title': 'Replace broken drop pipe',
                                            'jobStatus': 'requires_invoicing', 'jobType': 'ONE_OFF'},
                                           {'id': 'J200', 'jobNumber': 14400, 'title': 'Install new zones at the mausoleum',
                                            'jobStatus': 'active', 'jobType': 'ONE_OFF'},
                                           {'id': 'JR', 'jobNumber': 2069, 'title': 'Quarterly 2500', 'jobStatus': 'upcoming',
                                            'jobType': 'RECURRING'}]}}}
                return super().__call__(query, variables)
        fake = Fake(clients=[])
        P.jobber_gql = fake
        P.JOBBER_STATIC_TOKEN, saved = 'test-token', P.JOBBER_STATIC_TOKEN
        try:
            items = [{'name': 'After making repairs to get pump running', 'description': 'Lee Memorial Well 6. After making '
                      'repairs to get pump running, found pump stuck in well, customer will need to call well driller.',
                      'quantity': 1, 'unit_price': 600, 'amount': 600, 'taxable': True, 'is_tax': False}]
            self.extracts['b-leemem.pdf'] = extraction('bill', '31500', client='Lee Memorial', subtotal=600, items=items)
            b = self.upload('b-leemem.pdf')
            inv = b['auto_invoice']
            self.assertEqual(inv['job']['number'], '14511', inv)
            sent = [v for q, v in fake.calls if 'invoiceCreate' in q][-1]['input']
            self.assertEqual((sent['clientId'], sent['jobId']), (HODGES, 'J14511'))
            self.assertEqual([(l['name'], l['unitPrice']) for l in sent['lineItems']], [('PROPOSAL AMOUNT SERVICE', 3471.43)])
            case = self.case(b['case_id'])
            self.assertEqual(case['jobber']['job']['number'], '14511')
            self.assertEqual([i for i in case['issues'] if not i['resolved_at']], [])
            todos = [t['title'] for t in self.c.get('/pumps/api/summary').get_json()['queue']['todos']]
            self.assertIn("Check with Tom - Wettech bill doesn't line up with job #14511", todos)
        finally:
            P.JOBBER_STATIC_TOKEN = saved

    def test_mark_done_remove_and_reopen(self):
        self.extracts['b-done.pdf'] = extraction('bill', '31700', client='Coral Isles', subtotal=500)
        b = self.upload('b-done.pdf')
        cid = b['case_id']
        self.assertTrue([i for i in self.case(cid)['issues'] if not i['resolved_at']])
        self.assertTrue(self.c.post(f'/pumps/api/cases/{cid}/done', json={'note': 'Handled by phone'}).get_json()['success'])
        case = self.case(cid)
        self.assertEqual((case['status'], [i for i in case['issues'] if not i['resolved_at']]), ('closed', []))
        # Step changes later don't reopen a job marked done by hand.
        self.c.patch(f'/pumps/api/cases/{cid}', json={'notes': 'x'})
        self.assertEqual(self.case(cid)['status'], 'closed')
        self.assertNotIn(cid, [c['id'] for c in self.c.get('/pumps/api/cases?status=open').get_json()['cases']])
        self.c.post(f'/pumps/api/cases/{cid}/done', json={'reopen': True})
        self.assertEqual(self.case(cid)['status'], 'open')
        self.c.post(f'/pumps/api/cases/{cid}/delete', json={'reason': 'Removed'})
        self.assertEqual(self.case(cid)['status'], 'cancelled')

    def test_maintenance_accounts(self):
        today = P._today
        P._today = lambda: P.datetime(2026, 10, 5).date()
        try:
            j = self.c.get('/pumps/api/maint').get_json()
        finally:
            P._today = today
        acc = {a['name']: a for a in j['accounts']}
        self.assertEqual(sum(1 for a in j['accounts'] if a['active']), 39)
        self.assertEqual(sum(1 for a in j['accounts'] if not a['active']), 13)
        lc = acc['Lake Club (Spanish Wells)']
        self.assertEqual((lc['months'], lc['vendor_cost'], lc['due_this_month']), ('1,4,7,10', '$200.00', True))
        self.assertTrue(acc['Quail Run']['monthly'])
        self.assertFalse(acc['Forum c/o LandQwest Commercial Property']['due_this_month'], 'January/July only')
        self.assertEqual(acc['Warm Springs Comm. Assoc.']['naples_electric'], '$250.00')
        self.assertFalse(acc['Sapphire Lakes']['active'])
        r = self.c.post('/pumps/api/maint', json={'id': acc['Wilshire']['id'], 'active': True, 'months': 'jan apr'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.c.post('/api/pumps/maint', json={'name': 'X'},
                                     headers={'Authorization': 'Bearer test-openclaw-key'}).status_code, 403)

    def test_account_history_lists_every_todo_and_job(self):
        name = 'Lake Club (Spanish Wells)'
        r = self.c.post('/pumps/api/accounts/todo', json={'name': name, 'title': 'Ask about the fountain light',
                                                          'due_on': '2026-10-20'}).get_json()
        self.assertTrue(r['success'], r)
        self.assertEqual(r['todos'][0]['title'], 'Ask about the fountain light')
        # A to-do written on Today that names the client counts too, done or not.
        self.c.post('/pumps/api/todos', json={'title': 'Call Spanish Wells Lake Club about the gate code'})
        tid = [t for t in self.c.get('/pumps/api/todos').get_json()['todos'] if 'gate code' in t['title']][0]['id']
        self.c.post(f'/pumps/api/todos/{tid}/done', json={'done': True})
        self.c.post('/pumps/api/todos', json={'title': 'Order printer paper'})
        h = self.c.get('/pumps/api/accounts/history?name=' + name).get_json()
        titles = {t['title']: t for t in h['todos']}
        self.assertIn('Ask about the fountain light', titles)
        self.assertTrue(titles['Call Spanish Wells Lake Club about the gate code']['done_at'])
        self.assertNotIn('Order printer paper', titles)
        # The account page carries its setup too: the pump schedule, the lake, its Jobber links.
        self.assertEqual(h['account']['pump']['months'], '1,4,7,10')
        self.assertIn('docs', h)
        self.assertEqual(set(h['jobber']), {'client', 'recurring'})
        self.assertEqual(P._jobber_client_uri('Z2lkOi8vSm9iYmVyL0NsaWVudC80OTAwNzI1OA=='),
                         'https://secure.getjobber.com/clients/49007258')
        self.assertEqual(self.c.get('/pumps/api/accounts/history?name=Nobody').status_code, 400)

    def test_month_maintenance_and_subcontractor_notified(self):
        today = P._today
        P._today = lambda: P.datetime(2026, 10, 5).date()
        try:
            j = self.c.get('/pumps/api/maint/schedule?month=2026-10').get_json()
            rows = {(r['account'], r['vendor']): r for r in j['rows']}
            lc = rows[('Lake Club (Spanish Wells)', 'Wettech')]
            self.assertIsNone(lc['notified'])
            self.assertIn('Lake%20Club', j['wettech_mailto'])
            self.assertFalse(any(a == 'Forum c/o LandQwest Commercial Property' for a, _ in rows), 'January/July only')
            # Only the subcontractor the account has, per the master list: Allura and
            # Wild Blue are Wettech (Justin) only, Clubside's lake is Gulfshore (Jordan) only.
            for acct in ('Allura', 'Wild Blue'):
                self.assertEqual({v for a, v in rows if a == acct}, {'Wettech'}, acct)
            self.assertNotIn(('Clubside', 'Wettech'), rows)
            r = self.c.post('/pumps/api/maint/notified', json={'month': '2026-10', 'vendor': 'Wettech',
                                                               'accounts': ['Lake Club (Spanish Wells)'], 'how': 'phone'})
            j = r.get_json()
            lc = {(x['account'], x['vendor']): x for x in j['rows']}[('Lake Club (Spanish Wells)', 'Wettech')]
            self.assertEqual(lc['notified']['how'], 'phone')
            self.assertEqual(j['notified'], 1 + sum(1 for x in j['rows'] if x['vendor'] == 'Gulfshore' and x['notified']))
            hist = self.c.get('/pumps/api/accounts/history?name=Lake Club (Spanish Wells)').get_json()
            self.assertEqual([(n['month'], n['vendor']) for n in hist['notices']], [('2026-10', 'Wettech')])
            # Next month is its own list; un-marking takes it back off.
            self.assertNotIn(('Lake Club (Spanish Wells)', 'Wettech'),
                             {(x['account'], x['vendor']) for x in
                              self.c.get('/pumps/api/maint/schedule?month=2026-11').get_json()['rows']})
            j = self.c.post('/pumps/api/maint/notified', json={'month': '2026-10', 'vendor': 'Wettech', 'notified': False,
                                                               'account': 'Lake Club (Spanish Wells)'}).get_json()
            self.assertIsNone({(x['account'], x['vendor']): x for x in j['rows']}[('Lake Club (Spanish Wells)', 'Wettech')]['notified'])
            self.assertEqual(self.c.post('/pumps/api/maint/notified', json={'vendor': 'Acme', 'account': 'X'}).status_code, 400)
        finally:
            P._today = today

    def test_matching_helpers(self):
        self.assertGreater(P.similarity('Reserve at Estero - Lee County',
                                        'THE RESERVE AT ESTERO c/o ALLIANT PROPERTY MANAGEMENT, LLC.'), 0.95)
        self.assertGreater(P.similarity('Moritz HOA', 'St. Moritz C/O Suitor Middleton Cox & Associates'), 0.8)
        self.assertLess(P.similarity('Moritz HOA', 'Pebblebrook HOA'), 0.3)
        self.assertEqual(P.po_key('PO#0153'), '153')
        self.assertEqual(P.po_from_title('Service Call - WET***SC PUMP PO#0153'), '153')
        self.assertEqual(P.po_from_title('Pump Service call - PO322 Wttech'), '322')
        self.assertEqual(P._scada_site('Estimate to Renew the SCADA system-Dec 2023-Dec 2024'), '')
        self.assertEqual(P._scada_site('Scada Annual Renewal -Athletic Field'), 'Athletic Field')
        self.assertEqual(P._category('Pump Service call - PO322 Wttech'), 'pump')
        self.assertEqual(P._category('Service call - Diver to clean filter'), 'diver')
        self.assertEqual(P._category('SC- IRRIGATION'), '')

    def test_wettech_invoice_layout_without_claude(self):
        text = '''Invoice
Date Invoice #
of Southwest Florida LLC 9/30/2026 29067
451 Interstate Court
Phone # 941-232-4629 Fax # 941-371-5151
License #CFC1429137 office@wettec.biz
Bill To
Stahlman-England Irrigation
P.O. No. Terms
Andrea
Quantity Description Rate Amount
Sample Lakes - Lee County
Field service to furnish and install new 6" victaulic butterfly 1,803.81 1,803.81T
valve and couplings, pressurized and tested system. As per
quotation.
Sales Tax 6.5% 6.50% 117.25
Total $1,921.06
W/O No. 41105'''
        x = P._clean_extraction(P._regex_extract(text, 'office@wettec.biz'))
        self.assertEqual((x['kind'], x['doc_number'], x['doc_date'], x['vendor']), ('bill', '29067', '2026-09-30', 'Wettech'))
        self.assertEqual((x['po_number'], x['ordered_by'], x['wo_number']), ('', 'Andrea', '41105'))
        self.assertEqual((x['subtotal'], x['tax'], x['total']), (1803.81, 117.25, 1921.06))
        self.assertEqual(x['client_name'], 'Sample Lakes')
        self.assertEqual(len(x['line_items']), 1)
        self.assertTrue(x['line_items'][0]['taxable'])
        self.assertIn('pressurized and tested system', x['line_items'][0]['description'])

    def test_mailbox_scan(self):
        report = make_report(customer='Stahlman-Osprey Point HOA')
        self.extracts['Inv_31000.pdf'] = extraction('bill', '31000', po='PO777', client='Osprey Point', subtotal=250)
        self.texts['Inv_31000.pdf'] = 'Invoice ... office@wettec.biz'
        self.texts['menu.pdf'] = 'Lunch menu for Friday'
        mail = [('m1', {'from': {'emailAddress': {'address': 'office@wettec.biz'}}, 'subject': 'Report and invoice',
                        'receivedDateTime': '2026-10-01T12:00:00Z'}),
                ('m2', {'from': {'emailAddress': {'address': 'someone@example.com'}}, 'subject': 'Lunch',
                        'receivedDateTime': '2026-10-01T12:00:00Z'})]
        atts = {'m1': [('9-30 Osprey Point.docx', report), ('Inv_31000.pdf', b'%PDF fake 31000')],
                'm2': [('menu.pdf', b'%PDF menu')]}
        seen = {}

        def fetch(log_table, extensions, since):
            seen.update(log_table=log_table, extensions=extensions, since=since)
            return {'emails': mail, 'source': 'graph_api', 'diagnostics': {}}
        P.CFG.update(fetch_emails=fetch, graph_attachments=lambda uid: atts[uid], email_enabled=True)
        out = P.scan_mailbox()
        self.assertEqual(out['documents_added'], 2, out)
        self.assertEqual(out['skipped'], 1)
        self.assertIn('.docx', seen['extensions'])
        self.assertEqual(seen['log_table'], 'pump_email_scan_log')
        conn = P._conn()
        logged = {r[0] for r in conn.execute('SELECT email_uid FROM pump_email_scan_log')}
        conn.close()
        self.assertTrue({'m1', 'm2'} <= logged, 'every email is only looked at once')
        # What happened to each file is kept for "Recent files", even after an automatic scan.
        recent = self.c.get('/pumps/api/scan').get_json()['recent_files']
        what = {f['file']: f['what'] for f in recent}
        self.assertIn('Added as a bill', what['Inv_31000.pdf'])
        self.assertIn('not pump related', what['menu.pdf'])


    # ── chasing the vendor ──────────────────────────────────────────────────
    def _chase_case(self, **fields):
        conn = P._conn()
        try:
            cid = P.create_case(conn, {'vendor': 'Wettech', 'opened_on': '2026-10-05', **fields}, 'test')
            conn.commit()
        finally:
            conn.close()
        return cid

    def _chase(self, cid, today):
        conn = P._conn()
        try:
            return P.vendor_chase(P._case_dict(conn.execute('SELECT * FROM pump_cases WHERE id=?', (cid,)).fetchone()),
                                  today=today)
        finally:
            conn.close()

    def test_follow_up_with_the_vendor_on_a_job(self):
        """Approved job: it waits on Wettech's parts, and comes up to follow up
        14 days after the approval (Oct 2026). Tom says "later next week" on Fri Oct 9 -> scheduled for Fri Oct 16, nothing to do until Mon Oct 19,
        then a follow-up if the work isn't in."""
        from datetime import date
        cid = self._chase_case(title='Chase Suites lift pump', client_name='Chase Suites')
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {k: '2026-10-08' for k in
                                                                 ('vendor_quote', 'client_quote', 'client_approved')}})
        self.assertEqual(self.case(cid)['stage'], 'scheduled')
        ch = self._chase(cid, date(2026, 10, 9))
        self.assertEqual((ch['phase'], ch['state'], ch['check_on']), ('job', 'waiting', '2026-10-22'))
        self.assertIn('waiting on parts', ch['reason'])
        self.assertEqual(self._chase(cid, date(2026, 10, 22))['state'], 'follow_up')
        conn = P._conn()
        try:
            P.log_vendor_follow(conn, cid, 'Simon', 'Tom', 'call', "He'll be out later next week", '2026-10-16',
                                today=date(2026, 10, 9))
            conn.commit()
        finally:
            conn.close()
        case = self.case(cid)
        self.assertEqual(case['stage'], 'work_done', 'a date from the vendor schedules it')
        self.assertEqual(case['scheduled_for'], '2026-10-16')
        ch = self._chase(cid, date(2026, 10, 12))
        self.assertEqual((ch['state'], ch['check_on'], ch['said']), ('waiting', '2026-10-19', "He'll be out later next week"))
        self.assertIn('Tom said by 2026-10-16', ch['reason'])
        ch = self._chase(cid, date(2026, 10, 19))
        self.assertEqual(ch['state'], 'follow_up')
        self.assertIn("isn't in yet", ch['reason'])
        # Once the work is in, there is nothing to chase.
        self.c.patch(f'/pumps/api/cases/{cid}', json={'steps': {'work_done': '2026-10-15'}})
        self.assertNotIn(self.case(cid)['stage'], P.CHASE_PHASES)
        self.assertIsNone(self._chase(cid, date(2026, 10, 19)))
        self.assertIn('followed up with the vendor', [e['action'] for e in self.case(cid)['events']])

    def test_follow_up_with_the_vendor_on_a_quote(self):
        """Waiting on Wettech's quote: follow up after 3 workdays with no word;
        a call with no date waits another 3; the To do list shows it when due."""
        from datetime import date
        cid = self._chase_case(title='Chase Pointe pump quote', client_name='Chase Pointe', opened_on='2026-10-08')
        self.assertEqual(self.case(cid)['stage'], 'vendor_quote')
        self.assertEqual(self._chase(cid, date(2026, 10, 9))['state'], 'waiting')
        ch = self._chase(cid, date(2026, 10, 13))
        self.assertEqual((ch['phase'], ch['state'], ch['check_on']), ('quote', 'follow_up', '2026-10-13'))
        real = P._today, P._now_text
        P._today, P._now_text = (lambda: date(2026, 10, 13)), (lambda: '2026-10-13 10:00:00')
        try:
            q = self.c.get('/pumps/api/summary').get_json()['queue']
            self.assertIn(cid, [c['id'] for c in q['vendor_follow_ups']])
            self.assertNotIn(cid, [c['id'] for c in q['waiting_vendor_quote']])
            r = self.c.post(f'/pumps/api/cases/{cid}/vendor_follow', json={'who': 'Tom', 'how': 'text',
                                                                          'said': 'Working on it'})
            self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
            ch = r.get_json()['case']['chase']
            self.assertEqual((ch['state'], ch['check_on'], ch['how']), ('waiting', '2026-10-16', 'text'))
            q = self.c.get('/pumps/api/summary').get_json()['queue']
            self.assertNotIn(cid, [c['id'] for c in q['vendor_follow_ups']])
            self.assertIn(cid, [c['id'] for c in q['waiting_vendor_quote']])
            self.assertIn('Chase Pointe', P.build_digest()['html'])
            bad = self.c.post(f'/pumps/api/cases/{cid}/vendor_follow', json={'expect_by': 'soon'})
            self.assertEqual(bad.status_code, 400)
        finally:
            P._today, P._now_text = real

    def test_homewood_follow_up_is_kept(self):
        """Oct 9: Tom (Wettech) will be out to Homewood Suites later next week."""
        from datetime import date
        conn = P._conn()
        try:
            todo = conn.execute("SELECT title, due_on FROM pump_todos WHERE key='homewood_follow_oct26'").fetchone()
            self.assertIsNotNone(todo, 'no Homewood job in a fresh database: a to-do instead')
            self.assertEqual(todo['due_on'], '2026-10-19')
            cid = P.create_case(conn, {'title': 'Homewood Suites pump repair', 'vendor': 'Wettech'}, 'test')
            for k in ('vendor_quote', 'client_quote', 'client_approved'):
                P._set_step(conn, cid, k, at='2026-10-07')
            conn.execute("DELETE FROM pump_state WHERE key='homewood_follow_oct26'")
            conn.commit()
            P._seed_homewood_follow_up(conn)
        finally:
            conn.close()
        case = self.case(cid)
        self.assertEqual((case['stage'], case['scheduled_for']), ('work_done', '2026-10-16'))
        ch = self._chase(cid, date(2026, 10, 12))
        self.assertEqual((ch['state'], ch['who'], ch['check_on']), ('waiting', 'Tom', '2026-10-19'))


    def test_write_a_to_do_from_an_explanation(self):
        """Someone explains what they need to do; the app (no AI) shortens it to a to-do to check first."""
        from datetime import date
        fri = date(2026, 10, 9)
        r = self.c.post('/pumps/api/todos/reword', json={'text': 'I need to call Tom at Wettech about Homewood '
                                                                 'Suites. He said later next week.'})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        t = r.get_json()['todo']
        self.assertEqual((t['title'], t['detail']), ('Call Tom at Wettech about Homewood Suites',
                                                     'He said later next week.'))
        self.assertEqual(self.c.post('/pumps/api/todos/reword', json={'text': ' '}).status_code, 400)
        t = P.reword_todo('so remind me to send Allura the pump quote by Friday', today=fri)
        self.assertEqual((t['title'], t['detail'], t['due_on']), ('Send Allura the pump quote by Friday', '', '2026-10-16'))
        due = lambda text, today=fri: P.due_from_text(text, today)
        self.assertEqual(due('tomorrow'), '2026-10-10')
        self.assertEqual(due('get it done this week'), '2026-10-09')
        self.assertEqual(due('next week'), '2026-10-16')
        self.assertEqual(due('before the end of the month'), '2026-10-31')
        self.assertEqual(due('due 10/20'), '2026-10-20')
        self.assertEqual(due('by 1/5'), '2027-01-05')
        self.assertEqual(due('order a 1/2 HP motor'), '', 'a fraction is not a date')
        self.assertEqual(due('next friday', date(2026, 10, 12)), '2026-10-23')
        self.assertEqual(due('call Quail Run'), '')
        long = P.reword_todo('we need to get the SCADA quote out to Cross Creek before the end of the month, they renew '
                             'in November and Medallion pays it', today=fri)
        self.assertTrue(long['title'].startswith('Get the SCADA quote out to Cross Creek') and len(long['title']) <= 81)
        self.assertIn('Medallion pays it', long['detail'])
        self.assertEqual(long['due_on'], '2026-10-31')
        r = self.c.post('/pumps/api/todos', json={'title': t['title'], 'detail': 'They asked on the phone.',
                                                  'due_on': t['due_on']})
        row = next(x for x in r.get_json()['todos'] if x['title'] == 'Send Allura the pump quote by Friday')
        self.assertEqual((row['detail'], row['due_on'], row['kind']), ('They asked on the phone.', '2026-10-16', 'manual'))

if __name__ == '__main__':
    unittest.main(verbosity=2)
