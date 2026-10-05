"""
Tests for the Work Orders app (workorders.py).

    python -m unittest tests.test_workorders -v

Runs against a throwaway database. Microsoft 365, Jobber and Claude are fakes
that record what they were asked. All names and addresses are made up.
"""
import json
import os
import sys
import tempfile
import types
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix='workorders_test_')
os.environ.update(SECRET_KEY='test-secret-key-0123456789abcdef-wo', DATA_DIR=TMP, PUMPS_AUTO_SCAN='false',
                  WORKORDERS_AUTO_RUN='false', WEBSITE_URL='https://office.example.com')
os.environ.pop('ANTHROPIC_API_KEY', None)

import app as A  # noqa: E402
import workorders as W  # noqa: E402

A.app.config['TESTING'] = True
A.app.config['WTF_CSRF_ENABLED'] = False


class FakeGraph:
    def __init__(self, messages, attachments=None):
        self.messages = messages
        self.attachments = attachments or {}
        self.calls = []

    def __call__(self, method, path, **kw):
        self.calls.append((method, path, kw))
        if method == 'GET' and path.endswith('/messages'):
            return {'value': self.messages}
        if method == 'GET' and path.endswith('/attachments'):
            mid = path.split('/messages/')[1].split('/')[0]
            return {'value': [{'name': name, 'contentType': 'application/pdf', 'contentBytes': 'eA=='}
                              for name in self.attachments.get(mid, [])]}
        return {}


# What a Vertilinc maintenance request PDF reads as (made-up resident).
FORM = '''VeronaWalk HOA
8090 Sorrento Lane
Maintenance Request
Request # 20999-123456
Date Received: 10/4/2026
Time Received: 9:15 AM
Received by: Jamie Q. Sample
Service Location: Maintenance
Unit #: 8100 Example Court
Resident: Jamie Q. Sample
Contact Phone: 5550100
Work Requested:
Service: Irrigation
Status: New
Notes: 10-04-2026 9:15 AM - Jamie Q. SampleThe rotor by the mailbox
sprays the driveway. Please adjust it.
Special Instructions:
Pet Info:
Completion Notes:'''


class FakeJobber:
    def __init__(self, notes=None):
        self.calls = []
        self.notes = notes or []
        self.jobs = None

    def __call__(self, query, variables=None):
        self.calls.append((query, variables))
        if 'jobs(searchTerm' in query and self.jobs is not None:
            return {'jobs': {'nodes': self.jobs}}
        if 'jobs(searchTerm' in query:
            return {'jobs': {'nodes': [{'id': 'JOB_ENC_1', 'jobNumber': 1234, 'title': 'Verona Walk Work Orders',
                                        'client': {'id': 'CLIENT_1'}, 'property': {'id': 'PROP_1'}}]}}
        if 'jobCreateNote' in query:
            return {'jobCreateNote': {'userErrors': []}}
        if 'notes(first' in query:
            return {'job': {'notes': {'nodes': self.notes}}}
        if 'quoteCreate' in query:
            return {'quoteCreate': {'quote': {'id': 'Q_ENC', 'quoteNumber': 777,
                                              'jobberWebUri': 'https://secure.getjobber.com/quotes/777'},
                                    'userErrors': []}}
        return {}


class FakeClaude:
    """Answers the extract call and the clean-and-match call."""

    def __init__(self, needs_quote=False):
        self.prompts = []
        self.needs_quote = needs_quote
        self.messages = types.SimpleNamespace(create=self.create)

    def create(self, **kw):
        prompt = kw['messages'][0]['content']
        self.prompts.append(prompt)
        assert kw['model'] == W.CLAUDE_MODEL
        if prompt.startswith('Community:'):
            out = {'orders': [{'wo_number': '5521', 'wo_date': '2026-10-05', 'address': '8810 Example Lane',
                               'homeowner': 'Pat Example', 'description': 'Broken sprinkler head by the front walk.'}]}
        else:
            oid = int(prompt.split('- id ')[1].split(':')[0])
            out = {'work_order_id': oid, 'cleaned_notes': 'Replaced the broken spray head by the front walk.',
                   'needs_quote': self.needs_quote, 'quote_scope': 'Replace the zone 4 valve' if self.needs_quote else ''}
        return types.SimpleNamespace(stop_reason='end_turn',
                                     content=[types.SimpleNamespace(type='text', text=json.dumps(out))])


EMAIL = {'id': 'MSG1', 'subject': 'Verona Walk work order 5521',
         'from': {'emailAddress': {'address': 'manager@example.org'}},
         'receivedDateTime': '2026-10-05T13:00:00Z',
         'body': {'contentType': 'html', 'content': '<p>Please fix the broken head at 8810 Example Lane.</p>'}}


class WorkOrdersTest(unittest.TestCase):

    def setUp(self):
        conn = W._conn()
        for t in ('wo_inbox_orders', 'wo_inbox_communities', 'wo_inbox_tech_notes', 'wo_inbox_emails', 'wo_inbox_activity'):
            conn.execute(f'DELETE FROM {t}')
        conn.execute("DELETE FROM app_settings WHERE key LIKE 'workorders_%'")
        conn.execute('''INSERT INTO wo_inbox_communities (name, keywords, manager_name, manager_email, jobber_job, active)
                        VALUES ('Verona Walk', 'verona walk', 'Erwin Example', 'erwin@example.org', '1234', 1)''')
        conn.commit()
        conn.close()
        W._save_settings(dict(W.DEFAULT_SETTINGS, forward_to='regino@example.com, fredy@example.com'))
        self.graph = FakeGraph([EMAIL])
        self.jobber = FakeJobber()
        self.claude = FakeClaude()
        self._orig = (W._graph, W._jobber, dict(W._cfg))
        W._graph, W._jobber = self.graph, self.jobber
        W._cfg.update(graph_enabled=True, jobber_connected=lambda: True, anthropic_client=self.claude)
        self.c = A.app.test_client()

    def tearDown(self):
        W._graph, W._jobber = self._orig[0], self._orig[1]
        W._cfg.clear()
        W._cfg.update(self._orig[2])

    def login(self, role='office'):
        with self.c.session_transaction() as s:
            s['username'], s['role'], s['full_name'] = 'user1', role, 'Test User'

    def orders(self):
        conn = W._conn()
        rows = [dict(r) for r in conn.execute('SELECT * FROM wo_inbox_orders')]
        conn.close()
        return rows

    def test_new_email_is_saved_forwarded_and_logged_in_jobber(self):
        W.run_cycle(manual=True)
        [o] = self.orders()
        self.assertEqual((o['wo_number'], o['address'], o['status']), ('5521', '8810 Example Lane', 'open'))
        fwd = [c for c in self.graph.calls if c[1].endswith('/forward')]
        self.assertEqual(len(fwd), 1)
        self.assertIn('/users/po@stahlman-england.com/messages/MSG1/forward', fwd[0][1])
        self.assertEqual([r['emailAddress']['address'] for r in fwd[0][2]['json']['toRecipients']],
                         ['regino@example.com', 'fredy@example.com'])
        self.assertTrue(o['forwarded_at'] and o['jobber_logged_at'])
        note = [v for q, v in self.jobber.calls if 'jobCreateNote' in q][0]
        self.assertEqual(note['jobId'], 'JOB_ENC_1')
        self.assertIn('OUTSTANDING - Work order 5521 - 8810 Example Lane (Pat Example), 2026-10-05', note['input']['message'])

    def test_maintenance_request_pdf_is_read_field_by_field(self):
        self.graph.messages = [dict(EMAIL, id='MSG9', subject='Maintenance Request', hasAttachments=True,
                                    body={'contentType': 'text', 'content': 'See attached.'})]
        self.graph.attachments = {'MSG9': ['20999-123456.pdf']}
        orig = W._pdf_text
        W._pdf_text = lambda data: FORM
        try:
            W.run_cycle(manual=True)
        finally:
            W._pdf_text = orig
        [o] = self.orders()
        self.assertEqual((o['wo_number'], o['wo_date'], o['address'], o['homeowner'], o['phone']),
                         ('20999-123456', '2026-10-04', '8100 Example Court', 'Jamie Q. Sample', '5550100'))
        self.assertEqual(o['description'], 'The rotor by the mailbox sprays the driveway. Please adjust it.')
        self.assertFalse([p for p in self.claude.prompts if p.startswith('Community:')], 'the form needs no Claude')
        self.assertTrue([c for c in self.graph.calls if c[1].endswith('/MSG9/forward')])

    def test_work_order_job_is_found_when_no_number_is_set(self):
        conn = W._conn()
        conn.execute("UPDATE wo_inbox_communities SET jobber_job=''")
        conn.commit()
        conn.close()
        self.jobber.jobs = [
            {'id': 'OLD', 'jobNumber': 1672, 'title': 'Verona Walk- Service Call', 'client': {'id': 'C', 'name': 'VERONA  WALK HOA'}},
            {'id': 'WO_JOB', 'jobNumber': 11351, 'title': 'WORK ORDER', 'client': {'id': 'C', 'name': 'VERONA  WALK HOA'},
             'property': {'id': 'P'}},
            {'id': 'OTHER', 'jobNumber': 12000, 'title': 'Work order', 'client': {'id': 'X', 'name': 'Someone Else'}}]
        W.run_cycle(manual=True)
        note = [v for q, v in self.jobber.calls if 'jobCreateNote' in q][0]
        self.assertEqual(note['jobId'], 'WO_JOB')
        conn = W._conn()
        self.assertEqual(conn.execute("SELECT jobber_job FROM wo_inbox_communities").fetchone()[0], '11351')
        conn.close()

    def test_spreadsheet_import_and_export_use_the_office_layout(self):
        import io
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'October 2026'
        ws.append(['Work Order Date', 'workorder #', 'Address/Unit', 'Mary Hobbs, Trustee', 'Service Request', 'Results', 'TECH'])
        ws.append(['10/2/2026', '20771-111111', '8300 Example Court', 'Alex Sample', 'Check the system please', '', ''])
        ws.append(['10/3/202', '20775-222222', '8500 Example Court', 'Sam Sample', 'Adjust the head\nby the walk', 'Adjusted the head.', 'Regino'])
        ws.append([None] * 7)
        raw = io.BytesIO()
        wb.save(raw)
        data = raw.getvalue()
        self.login()
        cid = W._conn().execute('SELECT id FROM wo_inbox_communities').fetchone()[0]
        r = self.c.post('/workorders/import', data={'community_id': str(cid), 'file': (io.BytesIO(data), 'sheet.xlsx')},
                        content_type='multipart/form-data')
        self.assertEqual(r.get_json(), {'success': True, 'added': 2, 'closed': 1})
        rows = {o['wo_number']: o for o in self.orders()}
        self.assertEqual(rows['20771-111111']['status'], 'open')
        self.assertEqual((rows['20775-222222']['status'], rows['20775-222222']['wo_date'],
                          rows['20775-222222']['description'], rows['20775-222222']['tech']),
                         ('closed', '2026-10-03', 'Adjust the head by the walk', 'Regino'))
        # Importing the same sheet again adds nothing.
        r = self.c.post('/workorders/import', data={'community_id': str(cid), 'file': (io.BytesIO(data), 'sheet.xlsx')},
                        content_type='multipart/form-data')
        self.assertEqual(r.get_json()['added'], 0)
        out = load_workbook(io.BytesIO(self.c.get(f'/workorders/export.xlsx?community_id={cid}').data))
        self.assertEqual(out.sheetnames, ['October 2026'])
        vals = list(out['October 2026'].iter_rows(values_only=True))
        self.assertEqual(vals[0], W.SHEET_HEADERS)
        self.assertEqual(vals[2][1:], ('20775-222222', '8500 Example Court', 'Sam Sample', 'Adjust the head by the walk',
                                       'Adjusted the head.', 'Regino'))

    def test_same_email_is_not_picked_up_twice_and_replies_are_ignored(self):
        W.run_cycle(manual=True)
        self.graph.messages = [EMAIL, dict(EMAIL, id='MSG2', subject='RE: Verona Walk work order 5521'),
                               dict(EMAIL, id='MSG3', **{'from': {'emailAddress': {'address': 'regino@example.com'}}})]
        W.run_cycle(manual=True)
        self.assertEqual(len(self.orders()), 1)

    def test_email_without_a_known_community_is_skipped(self):
        self.graph.messages = [dict(EMAIL, subject='Work order for Some Other Place', body={'contentType': 'text', 'content': 'x'})]
        W.run_cycle(manual=True)
        self.assertEqual(self.orders(), [])

    def test_tech_note_is_cleaned_added_and_emailed_to_manager(self):
        W.run_cycle(manual=True)
        self.jobber.notes = [
            {'id': 'N1', 'message': 'replacd brokn head frnt walk', 'createdAt': '2099-01-01T10:00:00Z',
             'createdBy': {'name': {'full': 'Regino Example'}}},
            {'id': 'N2', 'message': 'not mine', 'createdAt': '2099-01-01T10:00:00Z',
             'createdBy': {'name': {'full': 'Someone Else'}}},
            {'id': 'N3', 'message': W.NOTE_MARK + ' our own note', 'createdAt': '2099-01-01T10:00:00Z',
             'createdBy': {'name': {'full': 'Regino Example'}}}]
        W.run_cycle(manual=True)
        [o] = self.orders()
        self.assertIn('Replaced the broken spray head by the front walk.', o['tech_notes'])
        self.assertEqual(len([p for p in self.claude.prompts if 'Note:' in p]), 1)
        conn = W._conn()
        [e] = [dict(r) for r in conn.execute('SELECT * FROM wo_inbox_emails')]
        conn.close()
        self.assertEqual((e['status'], e['to_email']), ('draft', 'erwin@example.org'))
        self.assertTrue(e['body'].startswith('Hi Erwin,'))
        self.assertIn("Technician's notes:\n2099-01-01: Replaced the broken spray head", e['body'])
        # Not sent until someone presses Send (auto-send is off by default).
        self.assertFalse([c for c in self.graph.calls if c[1].endswith('/sendMail')])
        self.login()
        r = self.c.post('/workorders/api/email', json={'id': e['id'], 'to_email': e['to_email'], 'subject': e['subject'],
                                                        'body': e['body'], 'send': True})
        self.assertTrue(r.get_json()['success'])
        sent = [c for c in self.graph.calls if c[1].endswith('/sendMail')]
        self.assertEqual(sent[0][2]['json']['message']['toRecipients'][0]['emailAddress']['address'], 'erwin@example.org')
        self.assertTrue(self.orders()[0]['emailed_at'])
        # A later cycle with nothing new does not draft again.
        W.run_cycle(manual=True)
        conn = W._conn()
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM wo_inbox_emails").fetchone()[0], 1)
        conn.close()

    def test_note_needing_a_quote_starts_a_jobber_quote_with_tech_notes(self):
        self.claude.needs_quote = True
        W.run_cycle(manual=True)
        self.jobber.notes = [{'id': 'N1', 'message': 'valve bad zone 4 need quote', 'createdAt': '2099-01-01T10:00:00Z',
                              'createdBy': {'name': {'full': 'Regino Example'}}}]
        W.run_cycle(manual=True)
        [o] = self.orders()
        self.assertEqual((o['status'], o['quote_status'], o['jobber_quote_number']), ('quote', 'Quote being prepared', '777'))
        attrs = [v for q, v in self.jobber.calls if 'quoteCreate' in q][0]['attributes']
        self.assertEqual((attrs['clientId'], attrs['propertyId']), ('CLIENT_1', 'PROP_1'))
        self.assertTrue(attrs['message'].startswith("Technician's notes:\n"))
        self.assertIn('Replaced the broken spray head', attrs['lineItems'][0]['description'])
        conn = W._conn()
        body = conn.execute('SELECT body FROM wo_inbox_emails').fetchone()[0]
        conn.close()
        self.assertIn('A quote is being prepared for this work order', body)

    def test_only_note_and_quote_changes_reach_jobber(self):
        W._jobber = self._orig[1]
        W._cfg['jobber_token'] = lambda force=False: self.fail('should be blocked before any token is used')
        with self.assertRaises(RuntimeError):
            W._jobber('mutation { invoiceCreate(input: {}) { invoice { id } } }')

    def test_office_only(self):
        r = self.c.get('/workorders')
        self.assertEqual(r.status_code, 302)
        for role in ('technician', 'property_manager', 'admin'):
            self.login(role)
            self.assertEqual(self.c.get('/workorders/api/data').status_code, 403)
            self.assertEqual(self.c.get('/workorders').status_code, 302)
        self.login('office')
        self.assertEqual(self.c.get('/workorders').status_code, 200)
        self.assertTrue(self.c.get('/workorders/api/data').get_json()['success'])

    def test_home_screen_shows_card_to_office_only(self):
        self.login('office')
        self.assertIn(b'/workorders', self.c.get('/dashboard').data)
        self.login('technician')
        self.assertNotIn(b'/workorders', self.c.get('/dashboard').data)


if __name__ == '__main__':
    unittest.main()
