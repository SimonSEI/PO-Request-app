"""
Receivables: QuickBooks A/R follow-ups, retainage and lien deadlines.

Office logins only (optionally narrowed by AR_USERS). The office uploads the A/R
sheet exported from QuickBooks; that upload is the list of what is owed. Then,
every hour (or "Run now"):

  1. Jobber: each open invoice is found in Jobber by its number for the client's
     email addresses, the online (client hub) link, line items, the job it belongs
     to, whether that job is complete, and whether Jobber already shows it paid.
  2. Replies: the sending mailbox is read for answers to our follow-ups. Each is
     read (Claude, or simple rules without it): "payment is coming" stops
     follow-ups until the invoice is very past due; "we paid" waits for the next
     upload; a dispute or question goes to a person.
  3. Follow-ups: invoices due for a reminder are grouped (one email per install
     job, one per customer for service calls), written in a tone that gets firmer
     the later the invoice is, and sent from simon@stahlman-england.com (Settings)
     with the invoice PDFs and any uploaded documents attached, or left as drafts
     until "Send emails automatically" is on.
  4. Retainage and liens: retainage (an invoice paid all but 5% or 10%) is held
     until the job is complete, then asked for. Install jobs carry a lien
     deadline (last day furnished + 90 days); ahead of it a Notice of Nonpayment
     is prepared for the owner, saying a claim of lien will follow.

Email goes through the Microsoft 365 app the PO app already uses (Graph: Mail.Send,
Mail.Read). Jobber goes through the Jobber app connected for Pumps, read only.

Hooked into app.py with init_receivables(...).
"""
import base64
import html
import io
import json
import os
import re
import sqlite3
import threading
import uuid
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import requests as http_requests
from flask import Blueprint, jsonify, redirect, render_template_string, request, send_file, session, url_for
from werkzeug.utils import secure_filename

try:
    import fcntl
except ImportError:  # Windows dev machines
    fcntl = None

TZ = ZoneInfo('America/New_York')
GRAPH = 'https://graph.microsoft.com/v1.0'
JOBBER_GQL = 'https://api.getjobber.com/api/graphql'
CLAUDE_MODEL = 'claude-opus-5-5'
# Optional: exactly these usernames get access (e.g. "simon,beatriz"). Unset = every office login.
AR_USERS = {u.strip().lower() for u in os.environ.get('AR_USERS', '').split(',') if u.strip()}

# An invoice's own state. 'open' is followed up on its schedule; the others hold it.
STATUSES = {
    'open': 'Open',
    'promised': 'Payment promised',
    'reported_paid': 'Customer says paid',
    'needs_person': 'Needs a person',
    'paused': 'Paused',
    'paid': 'Paid / closed',
}
NONP_STEPS = ('', 'prepared', 'sent', 'mailed', 'lien_recorded')

DEFAULT_STAGES = [
    {'key': 'reminder', 'name': 'Friendly reminder', 'from_days': 1, 'every_days': 14},
    {'key': 'past_due', 'name': 'Past due', 'from_days': 31, 'every_days': 10},
    {'key': 'very_past_due', 'name': 'Very past due', 'from_days': 61, 'every_days': 7},
    {'key': 'final', 'name': 'Final notice', 'from_days': 91, 'every_days': 5},
]

DEFAULT_SETTINGS = {
    'from_email': os.environ.get('AR_FROM_EMAIL', 'simon@stahlman-england.com'),
    'cc': '',                       # copied on every follow-up (e.g. the office)
    'signature': 'Thank you,\nSimon Weardon\nStahlman-England Irrigation',
    'auto_send': False,             # send follow-ups on their own (else they wait as drafts)
    'enabled': True,                # run the hourly cycle
    'read_replies': True,
    'stages': DEFAULT_STAGES,
    'promise_grace_days': 7,        # after a promised date passes, wait this long before following up
    'reported_paid_wait_days': 14,  # "we paid": wait this long for it to show up before following up again
    'min_balance': 1.0,
    'retainage_pcts': '5, 10',
    'retainage_tolerance': 0.6,     # percentage points either side
    'retainage_every_days': 30,     # once the job is complete
    'install_words': 'pay app, install, installation, proposal, retainage, retention',
    'lien_days': 90,                # Florida: claim of lien within 90 days of the last day furnished
    'lien_kinds': 'install',        # which invoices carry lien deadlines: install, or install,service
    'nonp_lead_days': 30,           # prepare the Notice of Nonpayment this many days before the lien deadline
    'nonp_auto': False,             # email the Notice of Nonpayment on its own
    'company_name': 'Stahlman-England Irrigation, Inc.',
    'company_address': '',
    'company_phone': '',
    'max_attach_mb': 20,
    'send_hours': '8-17',           # automatic emails go out Monday-Friday within these hours
}

bp = Blueprint('receivables', __name__)
_cfg = {}
_run_lock = threading.Lock()


# ── Plumbing ─────────────────────────────────────────────────────────────────

def _now():
    return datetime.now(TZ)


def _today():
    return _now().date()


def _stamp():
    return _now().strftime('%Y-%m-%d %H:%M:%S')


def _conn():
    conn = sqlite3.connect(_cfg['db_path'], timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def _files_dir():
    d = os.path.join(_cfg['data_dir'], 'receivables_files')
    os.makedirs(d, exist_ok=True)
    return d


def init_db():
    conn = _conn()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS ar_invoices (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  number TEXT UNIQUE NOT NULL,
                  customer TEXT DEFAULT '',
                  customer_key TEXT DEFAULT '',
                  qb_job TEXT DEFAULT '',
                  txn_date TEXT DEFAULT '',
                  due_date TEXT DEFAULT '',
                  terms TEXT DEFAULT '',
                  memo TEXT DEFAULT '',
                  po TEXT DEFAULT '',
                  amount REAL,
                  open_balance REAL DEFAULT 0,
                  kind TEXT DEFAULT 'service',
                  kind_locked INTEGER DEFAULT 0,
                  job_key TEXT DEFAULT '',
                  retainage INTEGER DEFAULT 0,
                  retainage_pct REAL,
                  retainage_locked INTEGER DEFAULT 0,
                  status TEXT DEFAULT 'open',
                  needs_reason TEXT DEFAULT '',
                  promised_date TEXT DEFAULT '',
                  promised_on TEXT DEFAULT '',
                  snooze_until TEXT DEFAULT '',
                  last_followup_at TEXT,
                  last_stage TEXT DEFAULT '',
                  followup_count INTEGER DEFAULT 0,
                  last_reply_at TEXT,
                  last_reply_summary TEXT DEFAULT '',
                  jobber_id TEXT DEFAULT '',
                  jobber_uri TEXT DEFAULT '',
                  client_hub_uri TEXT DEFAULT '',
                  jobber_status TEXT DEFAULT '',
                  jobber_subject TEXT DEFAULT '',
                  jobber_total REAL,
                  jobber_balance REAL,
                  jobber_paid REAL,
                  jobber_client_id TEXT DEFAULT '',
                  jobber_client_name TEXT DEFAULT '',
                  jobber_lines TEXT DEFAULT '[]',
                  billing_address TEXT DEFAULT '',
                  jobber_checked_at TEXT,
                  jobber_error TEXT DEFAULT '',
                  in_last_upload INTEGER DEFAULT 1,
                  first_seen_at TEXT,
                  closed_at TEXT,
                  updated_at TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_customers (
                  key TEXT PRIMARY KEY,
                  name TEXT DEFAULT '',
                  contact_name TEXT DEFAULT '',
                  emails TEXT DEFAULT '',
                  jobber_emails TEXT DEFAULT '',
                  jobber_client_id TEXT DEFAULT '',
                  kind_default TEXT DEFAULT '',
                  do_not_contact INTEGER DEFAULT 0,
                  notes TEXT DEFAULT '',
                  updated_at TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_jobs (
                  key TEXT PRIMARY KEY,
                  customer_key TEXT DEFAULT '',
                  name TEXT DEFAULT '',
                  kind TEXT DEFAULT 'install',
                  jobber_job_id TEXT DEFAULT '',
                  jobber_job_uri TEXT DEFAULT '',
                  jobber_status TEXT DEFAULT '',
                  jobber_completed_at TEXT DEFAULT '',
                  complete INTEGER,
                  property_address TEXT DEFAULT '',
                  contact_name TEXT DEFAULT '',
                  contact_emails TEXT DEFAULT '',
                  owner_name TEXT DEFAULT '',
                  owner_email TEXT DEFAULT '',
                  owner_address TEXT DEFAULT '',
                  first_furnished TEXT DEFAULT '',
                  last_furnished TEXT DEFAULT '',
                  nonp_status TEXT DEFAULT '',
                  nonp_at TEXT,
                  nonp_doc_id INTEGER,
                  notes TEXT DEFAULT '',
                  updated_at TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_notes (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  invoice_id INTEGER,
                  job_key TEXT DEFAULT '',
                  customer_key TEXT DEFAULT '',
                  at TEXT,
                  by TEXT DEFAULT '',
                  kind TEXT DEFAULT 'note',
                  text TEXT DEFAULT '')''')
    c.execute("CREATE INDEX IF NOT EXISTS ar_notes_inv ON ar_notes(invoice_id)")
    c.execute('''CREATE TABLE IF NOT EXISTS ar_documents (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  scope TEXT NOT NULL,
                  scope_key TEXT DEFAULT '',
                  kind TEXT DEFAULT 'document',
                  filename TEXT DEFAULT '',
                  stored_name TEXT DEFAULT '',
                  content_type TEXT DEFAULT '',
                  size INTEGER DEFAULT 0,
                  uploaded_by TEXT DEFAULT '',
                  uploaded_at TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_emails (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  kind TEXT DEFAULT 'followup',
                  group_key TEXT DEFAULT '',
                  customer_key TEXT DEFAULT '',
                  job_key TEXT DEFAULT '',
                  invoice_ids TEXT DEFAULT '[]',
                  stage TEXT DEFAULT '',
                  to_email TEXT DEFAULT '',
                  cc TEXT DEFAULT '',
                  subject TEXT DEFAULT '',
                  body TEXT DEFAULT '',
                  attachments TEXT DEFAULT '[]',
                  status TEXT DEFAULT 'draft',
                  manual INTEGER DEFAULT 0,
                  error TEXT DEFAULT '',
                  graph_id TEXT DEFAULT '',
                  conversation_id TEXT DEFAULT '',
                  created_at TEXT,
                  sent_at TEXT,
                  sent_by TEXT DEFAULT '')''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_replies (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  graph_id TEXT UNIQUE,
                  conversation_id TEXT DEFAULT '',
                  from_email TEXT DEFAULT '',
                  subject TEXT DEFAULT '',
                  received_at TEXT,
                  text TEXT DEFAULT '',
                  intent TEXT DEFAULT '',
                  promised_date TEXT DEFAULT '',
                  summary TEXT DEFAULT '',
                  customer_key TEXT DEFAULT '',
                  invoice_ids TEXT DEFAULT '[]',
                  created_at TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_uploads (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  filename TEXT DEFAULT '',
                  at TEXT,
                  by TEXT DEFAULT '',
                  info TEXT DEFAULT '{}')''')
    c.execute('''CREATE TABLE IF NOT EXISTS ar_activity (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  at TEXT,
                  kind TEXT,
                  message TEXT)''')
    conn.commit()
    conn.close()


def settings():
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    try:
        s.update(json.loads(_cfg['get_setting']('receivables_settings') or '{}'))
    except (TypeError, ValueError):
        pass
    if not isinstance(s.get('stages'), list) or not s['stages']:
        s['stages'] = json.loads(json.dumps(DEFAULT_STAGES))
    s['stages'] = sorted(s['stages'], key=lambda x: int(x.get('from_days') or 0))
    return s


def _save_settings(s):
    _cfg['set_setting']('receivables_settings', json.dumps(s))


def _state(key, default=''):
    return _cfg['get_setting'](f'receivables_{key}') or default


def _set_state(key, value):
    _cfg['set_setting'](f'receivables_{key}', value)


def _log(kind, message):
    print(f"  💵 Receivables [{kind}]: {message}")
    try:
        conn = _conn()
        conn.execute("INSERT INTO ar_activity (at, kind, message) VALUES (?,?,?)", (_stamp(), kind, message[:2000]))
        conn.execute("DELETE FROM ar_activity WHERE id NOT IN (SELECT id FROM ar_activity ORDER BY id DESC LIMIT 1000)")
        conn.commit()
        conn.close()
    except Exception:
        pass


def _note(conn, text, invoice_id=None, job_key='', customer_key='', kind='system', by='Office App'):
    conn.execute("INSERT INTO ar_notes (invoice_id, job_key, customer_key, at, by, kind, text) VALUES (?,?,?,?,?,?,?)",
                 (invoice_id, job_key or '', customer_key or '', _stamp(), by, kind, text[:5000]))


def _emails(text):
    return [e.strip().strip('<>') for e in re.split(r'[,;\s]+', text or '') if '@' in e]


def _squash(text):
    return re.sub(r'[^a-z0-9]', '', (text or '').lower())


def _ckey(name):
    return _squash((name or '').split(':')[0]) or 'unknown'


def _d(value):
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def _fmt_date(value):
    d = _d(value)
    return d.strftime('%m/%d/%Y') if d else ''


def _money_text(v):
    return f"${(v or 0):,.2f}"


def _words(text):
    return [k.strip().lower() for k in (text or '').split(',') if k.strip()]


def _nice_name(name):
    """QuickBooks and Jobber often hold names in capitals: MATTHEW LANGE -> Matthew Lange."""
    name = (name or '').strip()
    if name and name == name.upper() and any(ch.isalpha() for ch in name):
        return ' '.join(w if len(w) <= 3 and w in ('LLC', 'INC', 'HOA', 'POA', 'CDD', 'II', 'III') else w.capitalize()
                        for w in name.split())
    return name


def _html_to_text(body):
    text = re.sub(r'(?is)<(script|style).*?</\1>', ' ', body or '')
    text = re.sub(r'(?i)<br\s*/?>|</p>|</div>|</tr>|</li>', '\n', text)
    text = html.unescape(re.sub(r'<[^>]+>', ' ', text))
    text = re.sub(r'[ \t\r\f\v]+', ' ', text)
    return re.sub(r'\n\s*\n+', '\n\n', text).strip()


def receivables_allowed():
    """Office logins only (optionally narrowed by AR_USERS)."""
    if 'username' not in session or session.get('role') != 'office':
        return False
    if AR_USERS:
        return session['username'].lower() in AR_USERS
    return True


# ── Connections ──────────────────────────────────────────────────────────────

def graph_ready():
    return bool(_cfg.get('graph_enabled'))


def jobber_ready():
    try:
        return bool(_cfg['jobber_connected']())
    except Exception:
        return False


def claude_ready():
    return _cfg.get('anthropic_client') is not None


def _graph(method, path, **kw):
    token = _cfg['graph_token']()
    r = http_requests.request(method, path if path.startswith('http') else GRAPH + path, timeout=90,
                              headers={'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}, **kw)
    if r.status_code >= 400:
        try:
            detail = (r.json().get('error') or {}).get('message') or r.text[:300]
        except ValueError:
            detail = r.text[:300]
        raise RuntimeError(f'Microsoft 365 said HTTP {r.status_code}: {detail}')
    try:
        return r.json() if r.content else {}
    except ValueError:
        return {}


def _graph_upload(url, data):
    """Upload one large attachment through a Graph upload session (no auth header: the URL carries it)."""
    chunk = 320 * 1024 * 10
    for start in range(0, len(data), chunk):
        part = data[start:start + chunk]
        r = http_requests.put(url, data=part, timeout=120, headers={
            'Content-Length': str(len(part)),
            'Content-Range': f'bytes {start}-{start + len(part) - 1}/{len(data)}'})
        if r.status_code >= 400:
            raise RuntimeError(f'Microsoft 365 refused an attachment upload (HTTP {r.status_code})')


# Receivables only reads Jobber. Any mutation is refused here, whatever the token allows.
def _jobber(query, variables=None):
    if re.match(r'\s*mutation\b', query):
        raise RuntimeError('Blocked: Receivables only reads Jobber.')
    force = False
    for _ in range(4):
        token = _cfg['jobber_token'](force)
        r = http_requests.post(JOBBER_GQL, timeout=60, json={'query': query, 'variables': variables or {}},
                               headers={'Authorization': f'Bearer {token}', 'Content-Type': 'application/json',
                                        'X-JOBBER-GRAPHQL-VERSION': _cfg['jobber_version']})
        if r.status_code == 401 and not force:
            force = True
            continue
        if r.status_code == 429:
            threading.Event().wait(20)
            continue
        if r.status_code >= 400:
            raise RuntimeError(f'Jobber API HTTP {r.status_code}')
        out = r.json()
        errors = out.get('errors') or []
        if errors:
            if any((e.get('extensions') or {}).get('code') == 'THROTTLED' for e in errors):
                threading.Event().wait(15)
                continue
            raise RuntimeError(' | '.join(e.get('message', '?') for e in errors))
        return out.get('data') or {}
    raise RuntimeError('Jobber kept refusing the request (signed out or rate limited)')


def _claude_json(system, prompt, schema):
    """One Claude call that returns JSON matching schema, or None if Claude is unavailable or declines."""
    client = _cfg.get('anthropic_client')
    if client is None:
        return None
    resp = client.messages.create(
        model=CLAUDE_MODEL, max_tokens=2000, system=system,
        messages=[{'role': 'user', 'content': prompt}],
        extra_body={'output_config': {'effort': 'low', 'format': {'type': 'json_schema', 'schema': schema}}})
    if resp.stop_reason == 'refusal':
        return None
    text = ''.join(b.text for b in resp.content if b.type == 'text')
    try:
        return json.loads(text)
    except ValueError:
        return None


# ── 0. The QuickBooks A/R upload ─────────────────────────────────────────────

_COLS = {
    'type': ('transaction type', 'type', 'txn type', 'trans type'),
    'date': ('date', 'invoice date', 'txn date', 'transaction date'),
    'num': ('num', 'no.', 'no', 'invoice #', 'invoice no', 'invoice no.', 'invoice number', 'number', 'ref no',
            'ref no.', 'doc num', 'invoice'),
    'customer': ('customer', 'name', 'customer full name', 'customer name', 'client', 'customer:job', 'customer/job',
                 'customer:project', 'client name', 'customer/project'),
    'due': ('due date', 'due'),
    'terms': ('terms',),
    'amount': ('amount', 'total', 'orig. amount', 'original amount', 'invoice amount', 'total amount', 'orig amount'),
    'open': ('open balance', 'balance', 'amount due', 'open amount', 'balance due', 'open bal', 'open_balance'),
    'memo': ('memo/description', 'memo', 'description', 'message'),
    'po': ('p. o. #', 'p.o. #', 'po #', 'po number', 'p.o. number', 'po', 'p.o.'),
}
# Section rows QuickBooks puts in the A/R Aging Detail ("31 - 60 days past due") - not customers.
_BUCKET = re.compile(r'(?i)^(current|\d+\s*-\s*\d+\s*days?.*|\d+\s*(or|\+)\s*more.*|.*days past due.*|over\s+\d+.*|older.*|not due.*)$')


def _norm(v):
    return re.sub(r'\s+', ' ', str(v or '').strip().lower())


def _money(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace('$', '').replace(',', '')
    neg = t.startswith('(') and t.endswith(')')
    t = t.strip('()')
    try:
        n = float(t)
    except ValueError:
        return None
    return -n if neg else n


def _date_text(v):
    if v is None or v == '':
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v).strip()
    for fmt in ('%m/%d/%Y', '%m/%d/%y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d-%b-%Y', '%b %d, %Y', '%m-%d-%Y'):
        try:
            return datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            continue
    return ''


def _read_table(fileobj, filename):
    name = (filename or '').lower()
    data = fileobj.read()
    if name.endswith('.csv') or name.endswith('.txt'):
        import csv
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('latin-1')
        return [list(r) for r in csv.reader(io.StringIO(text))]
    if name.endswith('.xls'):
        raise ValueError('That is an old .xls file. In QuickBooks export to Excel (.xlsx) or CSV instead.')
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if len(rows) > len(best):
            best = rows
    return best


def _map_header(row):
    cells = [_norm(c) for c in row]
    m = {}
    for field, names in _COLS.items():
        for i, c in enumerate(cells):
            if c in names and i not in m.values():
                m[field] = i
                break
    return m


def parse_ar_table(rows):
    """Invoices from a QuickBooks A/R Aging Detail, Open Invoices or Invoice List export.
    Returns (invoices, info). Customer comes from its column or, without one, the section heading."""
    hdr, m = None, {}
    for i, row in enumerate(rows[:40]):
        got = _map_header(row)
        if 'num' in got and ({'open', 'amount'} & got.keys()) and len(got) > len(m):
            hdr, m = i, got
    if hdr is None:
        raise ValueError("Couldn't find the invoice number (Num) and Open Balance columns. In QuickBooks run "
                         "Reports > A/R Aging Detail (or Open Invoices) and export it to Excel.")
    out, seen = [], set()
    info = {'columns': {f: str(rows[hdr][j]) for f, j in m.items()}, 'skipped_types': 0, 'skipped_zero': 0,
            'duplicates': 0, 'header_row': hdr + 1}
    section = ''

    def cell(row, f):
        j = m.get(f)
        return row[j] if j is not None and j < len(row) else None

    for row in rows[hdr + 1:]:
        filled = [c for c in row if c not in (None, '')]
        if not filled:
            continue
        num = str(cell(row, 'num') or '').strip()
        if re.fullmatch(r'\d+\.0', num):
            num = num[:-2]
        if not num:
            first = str(filled[0]).strip()
            if len(filled) == 1 and not first.lower().startswith('total') and not _BUCKET.match(first):
                section = first
            continue
        kind = _norm(cell(row, 'type'))
        if 'type' in m and kind and not any(w in kind for w in ('invoice', 'charge')):
            info['skipped_types'] += 1
            continue
        open_bal = _money(cell(row, 'open'))
        amount = _money(cell(row, 'amount')) if 'amount' in m else None
        if open_bal is None:
            open_bal = amount
        if open_bal is None or open_bal <= 0.004:
            info['skipped_zero'] += 1
            continue
        if num in seen:
            info['duplicates'] += 1
            continue
        seen.add(num)
        customer = str(cell(row, 'customer') or '').strip() or section or 'Unknown customer'
        parent, _, sub = customer.partition(':')
        out.append({'number': num, 'customer': parent.strip() or customer, 'qb_job': sub.strip(),
                    'txn_date': _date_text(cell(row, 'date')), 'due_date': _date_text(cell(row, 'due')),
                    'terms': str(cell(row, 'terms') or '').strip(), 'memo': str(cell(row, 'memo') or '').strip()[:1000],
                    'po': str(cell(row, 'po') or '').strip(), 'amount': amount, 'open_balance': round(open_bal, 2)})
    info['invoices'] = len(out)
    return out, info


def apply_upload(parsed, filename, by, confirm=False):
    """Make the open invoices match the upload. Anything open here that the upload no longer lists was paid
    (or credited) in QuickBooks and is closed. Asks first when that would close most of them at once."""
    conn = _conn()
    existing = {r['number']: dict(r) for r in conn.execute("SELECT * FROM ar_invoices")}
    open_now = {n for n, r in existing.items() if r['status'] != 'paid'}
    nums = {p['number'] for p in parsed}
    missing = sorted(open_now - nums)
    if not confirm and len(open_now) >= 10 and len(missing) > len(open_now) / 2:
        conn.close()
        return {'needs_confirm': True, 'missing': len(missing), 'open': len(open_now)}
    now = _stamp()
    new = updated = reopened = paid_down = 0
    for p in parsed:
        ck = _ckey(p['customer'])
        conn.execute("INSERT OR IGNORE INTO ar_customers (key, name, updated_at) VALUES (?,?,?)", (ck, p['customer'], now))
        old = existing.get(p['number'])
        fields = dict(p, customer_key=ck, in_last_upload=1, updated_at=now)
        if not old:
            fields.update(status='open', first_seen_at=now)
            cols = ', '.join(fields)
            cur = conn.execute(f"INSERT INTO ar_invoices ({cols}) VALUES ({','.join('?' * len(fields))})", list(fields.values()))
            _note(conn, f"On the QuickBooks A/R upload ({filename}): {_money_text(p['open_balance'])} open.", cur.lastrowid,
                  customer_key=ck)
            new += 1
            continue
        if old['status'] == 'paid':
            fields.update(status='open', closed_at=None)
            _note(conn, 'Back on the QuickBooks A/R list, so it is open again.', old['id'], customer_key=ck)
            reopened += 1
        prev = old['open_balance'] or 0
        if p['open_balance'] < prev - 0.004:
            _note(conn, f"Balance went from {_money_text(prev)} to {_money_text(p['open_balance'])} "
                        f"(payment or credit of {_money_text(prev - p['open_balance'])}).", old['id'], customer_key=ck)
            paid_down += 1
            if old['status'] == 'reported_paid':
                fields.update(status='open', snooze_until='')
        conn.execute(f"UPDATE ar_invoices SET {', '.join(k + '=?' for k in fields)} WHERE id=?",
                     list(fields.values()) + [old['id']])
        updated += 1
    for n in missing:
        r = existing[n]
        conn.execute("UPDATE ar_invoices SET status='paid', closed_at=?, in_last_upload=0, updated_at=? WHERE id=?",
                     (now, now, r['id']))
        _note(conn, f'No longer on the QuickBooks A/R list ({filename}): paid or credited. Follow-ups stopped.', r['id'],
              customer_key=r['customer_key'])
    conn.commit()
    conn.close()
    classify_all()
    _close_stale_drafts()
    summary = {'new': new, 'updated': updated, 'closed': len(missing), 'reopened': reopened, 'paid_down': paid_down,
               'open': len(parsed)}
    conn = _conn()
    conn.execute("INSERT INTO ar_uploads (filename, at, by, info) VALUES (?,?,?,?)", (filename, now, by, json.dumps(summary)))
    conn.commit()
    conn.close()
    _log('upload', f"{filename}: {len(parsed)} open invoice(s); {new} new, {len(missing)} paid since the last upload, "
                   f"{paid_down} partly paid.")
    return summary


# ── Classification: service or install, which job, retainage ─────────────────

def _retainage_pct(amount, open_bal, pcts, tol):
    if not amount or amount <= 0 or open_bal is None or open_bal <= 0.004 or open_bal >= amount - 0.004:
        return None
    pct = open_bal / amount * 100
    for p in pcts:
        if abs(pct - p) <= tol:
            return p
    return None


def _pcts(s):
    out = []
    for t in re.split(r'[,\s]+', str(s.get('retainage_pcts') or '')):
        try:
            out.append(float(t.strip('%')))
        except ValueError:
            pass
    return out or [5.0, 10.0]


_PROJECT = re.compile(r'^\s*(\d{4,6})\s*[-–]\s*(.+?)\s*$')
_PAY_APP = re.compile(r'(?i)pay\s*app(?:lication)?\s*#?\s*\d*\s*[-–:]\s*(.+)$')


def _job_for(inv, lines, jobber_jobs):
    """(job_key, job_name) for an invoice, or ('', '') when it stands on its own."""
    if jobber_jobs:
        j = jobber_jobs[0]
        title = (j.get('title') or '').strip()
        return f"jobber:{j['id']}", f"#{j.get('jobNumber')} {title}".strip()
    for ln in lines:
        first = (ln.get('description') or '').strip().split('\n')[0]
        m = _PROJECT.match(first)
        if m:
            return f"proj:{m.group(1)}", f"{m.group(1)} - {m.group(2).strip()}"
    for text in (inv.get('jobber_subject') or '', inv.get('memo') or ''):
        m = _PAY_APP.search(text)
        if m and _squash(m.group(1)):
            return f"name:{inv['customer_key']}:{_squash(m.group(1))}", _nice_name(m.group(1).strip())
    if inv.get('qb_job'):
        return f"qb:{inv['customer_key']}:{_squash(inv['qb_job'])}", inv['qb_job']
    return '', ''


def classify_all(only_ids=None):
    s = settings()
    pcts, tol = _pcts(s), float(s.get('retainage_tolerance') or 0.6)
    words = _words(s['install_words'])
    conn = _conn()
    custs = {r['key']: dict(r) for r in conn.execute("SELECT * FROM ar_customers")}
    q = "SELECT * FROM ar_invoices WHERE status != 'paid'"
    rows = [dict(r) for r in conn.execute(q)]
    if only_ids:
        rows = [r for r in rows if r['id'] in only_ids]
    now = _stamp()
    for inv in rows:
        lines = json.loads(inv.get('jobber_lines') or '[]')
        text = ' '.join([inv.get('jobber_subject') or '', inv.get('memo') or '', inv.get('qb_job') or ''] +
                        [(ln.get('name') or '') + ' ' + (ln.get('description') or '') for ln in lines]).lower()
        fields = {}
        # Retainage
        if not inv['retainage_locked']:
            amount = inv['amount'] if inv['amount'] else inv['jobber_total']
            p = _retainage_pct(amount, inv['open_balance'], pcts, tol)
            named = bool(re.search(r'\b(retainage|retention)\b', text))
            fields['retainage'] = 1 if (p is not None or named) else 0
            fields['retainage_pct'] = p
        retain = fields.get('retainage', inv['retainage'])
        # Service call or install
        if not inv['kind_locked']:
            kd = (custs.get(inv['customer_key']) or {}).get('kind_default') or ''
            if retain or any(w in text for w in words):
                fields['kind'] = 'install'
            elif kd in ('service', 'install'):
                fields['kind'] = kd
            else:
                fields['kind'] = 'service'
        kind = fields.get('kind', inv['kind'])
        # Which job (the Jobber job is set by the Jobber sync; keep it)
        if not inv['job_key'].startswith('jobber:') and not inv['job_key'].startswith('manual:'):
            key, name = _job_for(inv, lines, [])
            if not key and kind == 'install':
                key, name = f"inv:{inv['number']}", f"Invoice {inv['number']}"
            fields['job_key'] = key
            if key:
                conn.execute("INSERT OR IGNORE INTO ar_jobs (key, customer_key, name, kind, updated_at) VALUES (?,?,?,?,?)",
                             (key, inv['customer_key'], name, kind, now))
        if fields:
            conn.execute(f"UPDATE ar_invoices SET {', '.join(k + '=?' for k in fields)} WHERE id=?",
                         list(fields.values()) + [inv['id']])
    # A job is an install if any of its invoices is, and so are its other invoices.
    conn.execute('''UPDATE ar_invoices SET kind='install' WHERE kind_locked=0 AND status != 'paid' AND job_key != ''
                    AND job_key IN (SELECT job_key FROM ar_invoices WHERE kind='install' AND job_key != '')''')
    conn.execute('''UPDATE ar_jobs SET kind='install' WHERE key IN
                    (SELECT job_key FROM ar_invoices WHERE kind='install' AND job_key != '')''')
    conn.commit()
    conn.close()


# ── 1. Jobber ────────────────────────────────────────────────────────────────

_ADDR = 'street1 street2 city province postalCode'
_INV_TIERS = [
    f'''id invoiceNumber subject invoiceStatus issuedDate dueDate jobberWebUri clientHubUri
        amounts {{ total invoiceBalance paymentsTotal }}
        billingAddress {{ {_ADDR} }}
        client {{ id name emails {{ address primary }} }}
        lineItems(first: 30) {{ nodes {{ name description quantity unitPrice totalPrice }} }}
        jobs(first: 3) {{ nodes {{ id jobNumber title jobStatus completedAt jobberWebUri property {{ address {{ {_ADDR} }} }} }} }}''',
    '''id invoiceNumber subject invoiceStatus issuedDate dueDate jobberWebUri clientHubUri
        amounts { total invoiceBalance paymentsTotal }
        client { id name emails { address primary } }
        lineItems(first: 30) { nodes { name description quantity unitPrice totalPrice } }''',
    'id invoiceNumber subject invoiceStatus jobberWebUri amounts { total } client { id name }',
]


def _addr_text(a):
    a = a or {}
    street = ', '.join(x for x in (a.get('street1'), a.get('street2')) if x)
    tail = ' '.join(x for x in (a.get('province'), a.get('postalCode')) if x)
    return ', '.join(x for x in (street, a.get('city'), tail) if x)


def _find_invoice(number, jobber_id=''):
    """The Jobber invoice with this number, trying smaller queries if this API version lacks a field."""
    start = int(_cfg.get('tier') or 0)
    last = None
    for tier in range(start, len(_INV_TIERS)):
        fields = _INV_TIERS[tier]
        try:
            if jobber_id:
                data = _jobber(f'query($id: EncodedId!) {{ invoice(id: $id) {{ {fields} }} }}', {'id': jobber_id})
                inv = data.get('invoice')
            else:
                data = _jobber(f'query($q: String!) {{ invoices(searchTerm: $q, first: 10) {{ nodes {{ {fields} }} }} }}',
                               {'q': number})
                nodes = ((data.get('invoices') or {}).get('nodes')) or []
                inv = next((n for n in nodes if str(n.get('invoiceNumber')) == str(number)), None)
            _cfg['tier'] = tier
            return inv
        except RuntimeError as e:
            last = e
            if not re.search(r"(?i)field|doesn't exist|undefined|argument", str(e)):
                raise
    raise last


def sync_jobber(limit=150, only_ids=None):
    if not jobber_ready():
        return _log('jobber', 'Skipped: Jobber is not connected (connect it in the Pumps app).')
    cutoff = (_now() - timedelta(hours=12)).strftime('%Y-%m-%d %H:%M:%S')
    conn = _conn()
    rows = [dict(r) for r in conn.execute(
        '''SELECT * FROM ar_invoices WHERE status != 'paid' AND (jobber_checked_at IS NULL OR jobber_checked_at < ?)
           ORDER BY jobber_checked_at IS NOT NULL, jobber_checked_at LIMIT ?''', (cutoff, limit))]
    if only_ids:
        rows = [dict(r) for r in conn.execute(
            f"SELECT * FROM ar_invoices WHERE id IN ({','.join('?' * len(only_ids))})", list(only_ids))]
    conn.close()
    found = closed = 0
    for inv in rows:
        try:
            j = _find_invoice(inv['number'], inv['jobber_id'])
        except Exception as e:
            _set_invoice(inv['id'], jobber_error=str(e)[:500], jobber_checked_at=_stamp())
            _log('jobber', f"Invoice {inv['number']}: Jobber lookup failed: {e}")
            continue
        if not j:
            _set_invoice(inv['id'], jobber_error='Not found in Jobber', jobber_checked_at=_stamp())
            continue
        found += 1
        closed += _apply_jobber(inv, j)
    classify_all()
    if rows:
        _log('jobber', f'Checked {len(rows)} invoice(s) in Jobber: {found} found, {closed} already paid there.')


def _apply_jobber(inv, j):
    amounts = j.get('amounts') or {}
    client = j.get('client') or {}
    lines = [{k: n.get(k) for k in ('name', 'description', 'quantity', 'unitPrice', 'totalPrice')}
             for n in (((j.get('lineItems') or {}).get('nodes')) or []) if n]
    jobs = [n for n in (((j.get('jobs') or {}).get('nodes')) or []) if n and n.get('id')]
    status = (j.get('invoiceStatus') or '').lower()
    fields = dict(jobber_id=j.get('id', ''), jobber_uri=j.get('jobberWebUri') or '', client_hub_uri=j.get('clientHubUri') or '',
                  jobber_status=status, jobber_subject=j.get('subject') or '', jobber_total=amounts.get('total'),
                  jobber_balance=amounts.get('invoiceBalance'), jobber_paid=amounts.get('paymentsTotal'),
                  jobber_client_id=client.get('id', ''), jobber_client_name=client.get('name', ''),
                  jobber_lines=json.dumps(lines), billing_address=_addr_text(j.get('billingAddress')),
                  jobber_checked_at=_stamp(), jobber_error='')
    conn = _conn()
    now = _stamp()
    if jobs:
        jb = jobs[0]
        key = f"jobber:{jb['id']}"
        name = f"#{jb.get('jobNumber')} {(jb.get('title') or '').strip()}".strip()
        conn.execute("INSERT OR IGNORE INTO ar_jobs (key, customer_key, name, kind, updated_at) VALUES (?,?,?,?,?)",
                     (key, inv['customer_key'], name, inv['kind'], now))
        conn.execute('''UPDATE ar_jobs SET name=?, jobber_job_id=?, jobber_job_uri=?, jobber_status=?, jobber_completed_at=?,
                        property_address=CASE WHEN property_address='' THEN ? ELSE property_address END, updated_at=? WHERE key=?''',
                     (name, jb['id'], jb.get('jobberWebUri') or '', (jb.get('jobStatus') or '').lower(),
                      (jb.get('completedAt') or '')[:10], _addr_text((jb.get('property') or {}).get('address')), now, key))
        if not inv['job_key'].startswith('manual:'):
            fields['job_key'] = key
    emails = [e.get('address') for e in (client.get('emails') or []) if e and e.get('address')]
    primary = [e.get('address') for e in (client.get('emails') or []) if e and e.get('primary') and e.get('address')]
    if emails:
        ordered = primary + [e for e in emails if e not in primary]
        conn.execute("UPDATE ar_customers SET jobber_emails=?, jobber_client_id=?, updated_at=? WHERE key=?",
                     (', '.join(ordered), client.get('id', ''), now, inv['customer_key']))
    paid = 0
    if status in ('paid', 'voided', 'bad_debt') and inv['status'] != 'paid':
        fields.update(status='paid', closed_at=now)
        _note(conn, f"Jobber shows this invoice as {status.replace('_', ' ')}. Follow-ups stopped "
                    f"(it will drop off the next QuickBooks upload).", inv['id'], customer_key=inv['customer_key'])
        paid = 1
    conn.commit()
    conn.close()
    _set_invoice(inv['id'], **fields)
    return paid


def _set_invoice(inv_id, **fields):
    fields['updated_at'] = _stamp()
    conn = _conn()
    conn.execute(f"UPDATE ar_invoices SET {', '.join(k + '=?' for k in fields)} WHERE id=?", list(fields.values()) + [inv_id])
    conn.commit()
    conn.close()


# ── Where each invoice stands ────────────────────────────────────────────────

def _due(inv):
    d = _d(inv.get('due_date'))
    if d:
        return d
    t = _d(inv.get('txn_date'))
    return t + timedelta(days=30) if t else None


def _stage_for(dpd, s):
    cur = None
    for st in s['stages']:
        if dpd >= int(st.get('from_days') or 0):
            cur = st
    return cur


def _very_past_due_days(s):
    for st in s['stages']:
        if st.get('key') == 'very_past_due':
            return int(st['from_days'])
    return int(s['stages'][min(2, len(s['stages']) - 1)]['from_days'])


def job_complete(job):
    if not job:
        return False
    if job.get('complete') is not None:
        return bool(job['complete'])
    return bool(job.get('jobber_completed_at')) or (job.get('jobber_status') or '') in ('archived', 'requires_invoicing')


def contacts_for(inv, job, cust):
    """Who a follow-up goes to: the job's contact, else the customer's emails, else Jobber's."""
    if job and _emails(job.get('contact_emails')):
        return _emails(job['contact_emails'])
    if cust and _emails(cust.get('emails')):
        return _emails(cust['emails'])
    if cust and _emails(cust.get('jobber_emails')):
        return _emails(cust['jobber_emails'])[:2]
    return []


def plan(inv, job, cust, s, today=None):
    """What happens next for one invoice: {due_now, reason, next, stage, stage_name, dpd}."""
    today = today or _today()
    due = _due(inv)
    dpd = (today - due).days if due else 0
    out = {'due_now': False, 'reason': '', 'next': '', 'stage': '', 'stage_name': '', 'dpd': dpd}
    st = _stage_for(dpd, s)
    if st:
        out['stage'], out['stage_name'] = st['key'], st['name']
    if inv['status'] == 'paid':
        out['reason'] = 'Paid / closed'
        return out
    if cust and cust.get('do_not_contact'):
        out['reason'] = 'Customer is set to "do not contact"'
        return out
    if inv['status'] == 'paused':
        out['reason'] = 'Follow-ups paused'
        return out
    if inv['status'] == 'needs_person':
        out['reason'] = 'Waiting on a person: ' + (inv.get('needs_reason') or 'see notes')
        return out
    if (inv['open_balance'] or 0) < float(s.get('min_balance') or 0):
        out['reason'] = f"Balance under {_money_text(float(s.get('min_balance') or 0))}"
        return out
    last = _d(inv.get('last_followup_at'))
    if inv['retainage']:
        if not job_complete(job):
            out['reason'] = 'Retainage: held until the job is complete'
            return out
        out['stage'], out['stage_name'] = 'retainage', 'Retainage release'
        every = int(s.get('retainage_every_days') or 30)
    else:
        if not st:
            first = int(s['stages'][0]['from_days'])
            out['reason'] = 'Not past due yet'
            out['next'] = (due + timedelta(days=first)).isoformat() if due else ''
            return out
        every = int(st.get('every_days') or 7)
    snooze = _d(inv.get('snooze_until'))
    if snooze and today < snooze:
        out['next'] = snooze.isoformat()
        if inv['status'] == 'promised':
            pd = f" for {_fmt_date(inv['promised_date'])}" if inv.get('promised_date') else ''
            out['reason'] = f"Payment promised{pd}; next follow-up {_fmt_date(snooze)}"
        elif inv['status'] == 'reported_paid':
            out['reason'] = f"Customer says it's paid; waiting for it to show up until {_fmt_date(snooze)}"
        else:
            out['reason'] = f'Snoozed until {_fmt_date(snooze)}'
        return out
    if last and (today - last).days < every:
        nxt = last + timedelta(days=every)
        out['next'] = nxt.isoformat()
        out['reason'] = f"Followed up {_fmt_date(last)}; next {_fmt_date(nxt)}"
        return out
    out['due_now'] = True
    out['next'] = today.isoformat()
    out['reason'] = f"Due for a {out['stage_name'].lower() or 'follow-up'}"
    return out


def lien_info(job, invoices, s, today=None):
    """{applies, last_furnished, estimated, deadline, days_left, nonp_due} for a job."""
    today = today or _today()
    kinds = _words(s.get('lien_kinds') or 'install')
    if not job or job.get('kind') not in kinds:
        return {'applies': False}
    open_inv = [i for i in invoices if i['status'] != 'paid' and (i['open_balance'] or 0) > 0.004]
    last = _d(job.get('last_furnished'))
    estimated = False
    if not last:
        last = _d(job.get('jobber_completed_at'))
    if not last:
        dates = [_d(i['txn_date']) for i in invoices if _d(i['txn_date'])]
        last = max(dates) if dates else None
        estimated = True
    if not last:
        return {'applies': True, 'last_furnished': '', 'estimated': True, 'deadline': '', 'days_left': None,
                'nonp_due': False}
    deadline = last + timedelta(days=int(s.get('lien_days') or 90))
    days_left = (deadline - today).days
    past_due = any(_due(i) and _due(i) < today for i in open_inv)
    nonp_due = bool(open_inv) and past_due and days_left <= int(s.get('nonp_lead_days') or 30) and \
        not (job.get('nonp_status') or '')
    return {'applies': True, 'last_furnished': last.isoformat(), 'estimated': estimated, 'deadline': deadline.isoformat(),
            'days_left': days_left, 'nonp_due': nonp_due, 'open_balance': round(sum(i['open_balance'] or 0 for i in open_inv), 2)}


def _load_all():
    conn = _conn()
    invs = [dict(r) for r in conn.execute("SELECT * FROM ar_invoices")]
    jobs = {r['key']: dict(r) for r in conn.execute("SELECT * FROM ar_jobs")}
    custs = {r['key']: dict(r) for r in conn.execute("SELECT * FROM ar_customers")}
    conn.close()
    return invs, jobs, custs


# ── 2. Replies ───────────────────────────────────────────────────────────────

_REPLY_SCHEMA = {
    'type': 'object', 'additionalProperties': False,
    'required': ['intent', 'promised_date', 'summary', 'invoice_numbers'],
    'properties': {
        'intent': {'type': 'string', 'enum': ['promise', 'paid', 'dispute', 'question', 'wrong_contact',
                                              'auto_reply', 'other']},
        'promised_date': {'type': 'string', 'description': 'YYYY-MM-DD when they say payment will be made/arrive, else ""'},
        'summary': {'type': 'string', 'description': 'one short sentence for the office'},
        'invoice_numbers': {'type': 'array', 'items': {'type': 'string'}},
    },
}

_WEEKDAYS = ['monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday']
_MONTHS = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']


def _find_date(text, today):
    low = text.lower()
    m = re.search(r'\b(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\b', low)
    if m:
        mo, dy = int(m.group(1)), int(m.group(2))
        yr = int(m.group(3)) if m.group(3) else today.year
        yr = yr + 2000 if yr < 100 else yr
        try:
            d = date(yr, mo, dy)
            if not m.group(3) and d < today - timedelta(days=30):
                d = date(yr + 1, mo, dy)
            return d
        except ValueError:
            pass
    m = re.search(r'\b(' + '|'.join(_MONTHS) + r')[a-z]*\.?\s+(\d{1,2})(?:st|nd|rd|th)?\b', low)
    if m:
        try:
            d = date(today.year, _MONTHS.index(m.group(1)) + 1, int(m.group(2)))
            return d if d >= today - timedelta(days=30) else date(today.year + 1, d.month, d.day)
        except ValueError:
            pass
    if re.search(r'end of (the )?month', low):
        nxt = (today.replace(day=28) + timedelta(days=4)).replace(day=1)
        return nxt - timedelta(days=1)
    if 'next week' in low:
        return today + timedelta(days=7)
    if 'tomorrow' in low:
        return today + timedelta(days=1)
    for i, wd in enumerate(_WEEKDAYS):
        if re.search(rf'\b{wd}\b', low):
            ahead = (i - today.weekday()) % 7 or 7
            return today + timedelta(days=ahead)
    return None


def classify_reply(subject, text, invoices, today=None):
    """What a customer's reply means. Claude when available, else plain rules."""
    today = today or _today()
    listing = '\n'.join(f"- Invoice {i['number']}: {_money_text(i['open_balance'])} open, due {_fmt_date(i['due_date'])}"
                        for i in invoices) or '(none)'
    try:
        got = _claude_json(
            'You read replies to past-due invoice reminders sent by an irrigation company. Decide what the reply '
            'means. intent: "promise" = they say payment is coming, being processed, approved, scheduled or will be '
            'sent; "paid" = they say it was already paid, mailed or sent; "dispute" = they disagree with the charge '
            'or say it is not theirs; "question" = they ask for something (a copy, a W-9, lien waiver, details); '
            '"wrong_contact" = we wrote to the wrong person or they forward us elsewhere; "auto_reply" = out of '
            'office or automatic; "other" = anything else. promised_date: the date they say payment will be made '
            f'(today is {today.isoformat()}), else "". invoice_numbers: the invoices it is about, from the list.',
            f'Open invoices for this customer:\n{listing}\n\nSubject: {subject}\n\nReply:\n{text[:8000]}',
            _REPLY_SCHEMA)
    except Exception as e:
        _log('replies', f'Claude could not read a reply: {e}')
        got = None
    if got:
        return got
    low = f'{subject}\n{text}'.lower()
    found = _find_date(text, today)
    if re.search(r'automatic reply|out of (the )?office|auto-?reply|away from (the|my) office', low):
        intent = 'auto_reply'
    elif re.search(r'dispute|incorrect|not (ours|our responsibility)|never (received|ordered|authorized)|disagree|'
                   r'overcharg|should not have been|wrong (amount|invoice)', low):
        intent = 'dispute'
    elif re.search(r'(will|going to|should|to be|is being|are) (be )?(paid|pay|mail|sent|send|process|cut|release)|'
                   r'scheduled|in process|processing|next (check|pay|payment) run|on (its|it\'s|the) way|'
                   r'approved for payment|cut a check|payment is coming', low):
        intent = 'promise'
    elif re.search(r'\b(paid|mailed|sent (the |a )?(check|payment)|check (was )?(mailed|sent)|payment (was )?(sent|made|'
                   r'processed|submitted))\b', low):
        intent = 'paid'
    elif re.search(r'no longer|not the right|wrong (person|contact)|please contact|forwarded? (this|to)', low):
        intent = 'wrong_contact'
    elif '?' in text:
        intent = 'question'
    else:
        intent = 'other'
    nums = [i['number'] for i in invoices if re.search(rf"\b{re.escape(i['number'])}\b", low)]
    first = next((ln.strip() for ln in text.splitlines() if ln.strip()), '')
    return {'intent': intent, 'promised_date': found.isoformat() if found and intent == 'promise' else '',
            'summary': first[:200], 'invoice_numbers': nums}


def _apply_reply(conn, invoices, got, from_email, s, today):
    intent = got.get('intent') or 'other'
    summary = (got.get('summary') or '').strip()
    vpd = _very_past_due_days(s)
    for inv in invoices:
        note = f"Reply from {from_email}: {summary}" if summary else f"Reply from {from_email}."
        fields = {'last_reply_at': _stamp(), 'last_reply_summary': f'{intent}: {summary}'[:500]}
        if intent == 'promise':
            due = _due(inv) or today
            very = due + timedelta(days=vpd)
            pd = _d(got.get('promised_date'))
            if today < very:
                snooze = very  # payment is coming: leave them be until the invoice is very past due
            elif pd and pd + timedelta(days=int(s['promise_grace_days'])) > today:
                snooze = pd + timedelta(days=int(s['promise_grace_days']))
            else:
                snooze = today + timedelta(days=14)
            fields.update(status='promised', promised_date=pd.isoformat() if pd else '', promised_on=today.isoformat(),
                          snooze_until=snooze.isoformat(), needs_reason='')
            note += f" Payment promised{' for ' + _fmt_date(pd) if pd else ''}; no more follow-ups until {_fmt_date(snooze)}."
        elif intent == 'paid':
            snooze = today + timedelta(days=int(s['reported_paid_wait_days']))
            fields.update(status='reported_paid', snooze_until=snooze.isoformat(), needs_reason='')
            note += f" They say it's paid; waiting for it to show up (next follow-up {_fmt_date(snooze)} if it hasn't)."
        elif intent == 'auto_reply':
            note += ' (automatic reply, nothing changed)'
        else:
            label = {'dispute': 'Disputed', 'question': 'Customer asked a question', 'wrong_contact': 'Wrong contact',
                     'other': 'Customer replied'}.get(intent, 'Customer replied')
            fields.update(status='needs_person', needs_reason=f"{label}: {summary}"[:300])
            note += ' Follow-ups stopped until someone answers it.'
        if inv['status'] in ('paid', 'paused') and 'status' in fields:
            fields.pop('status')
        conn.execute(f"UPDATE ar_invoices SET {', '.join(k + '=?' for k in fields)} WHERE id=?",
                     list(fields.values()) + [inv['id']])
        _note(conn, note, inv['id'], customer_key=inv['customer_key'], kind='reply')


def scan_replies():
    s = settings()
    if not s.get('read_replies'):
        return
    if not graph_ready():
        return _log('replies', 'Skipped: Microsoft 365 is not connected.')
    mailbox = s['from_email'].strip()
    since = _state('last_reply_scan') or (datetime.now(timezone.utc) - timedelta(days=7)).strftime('%Y-%m-%dT%H:%M:%SZ')
    started = (datetime.now(timezone.utc) - timedelta(minutes=15)).strftime('%Y-%m-%dT%H:%M:%SZ')
    params = {'$filter': f'receivedDateTime ge {since}', '$orderby': 'receivedDateTime asc', '$top': 50,
              '$select': 'id,subject,from,receivedDateTime,conversationId,uniqueBody,body'}
    data = _graph('GET', f'/users/{mailbox}/mailFolders/inbox/messages', params=params)
    messages = list(data.get('value', []))
    while data.get('@odata.nextLink') and len(messages) < 1000:
        data = _graph('GET', data['@odata.nextLink'])
        messages.extend(data.get('value', []))

    conn = _conn()
    seen = {r[0] for r in conn.execute("SELECT graph_id FROM ar_replies")}
    convs = {}
    for r in conn.execute("SELECT conversation_id, invoice_ids, customer_key FROM ar_emails WHERE conversation_id != ''"):
        convs.setdefault(r['conversation_id'], set()).update(json.loads(r['invoice_ids'] or '[]'))
    open_invs = [dict(r) for r in conn.execute("SELECT * FROM ar_invoices WHERE status != 'paid'")]
    by_id = {i['id']: i for i in open_invs}
    by_cust = {}
    for i in open_invs:
        by_cust.setdefault(i['customer_key'], []).append(i)
    addr = {}
    for r in conn.execute("SELECT key, emails, jobber_emails FROM ar_customers"):
        for e in _emails(r['emails']) + _emails(r['jobber_emails']):
            addr.setdefault(e.lower(), set()).add(r['key'])
    for r in conn.execute("SELECT customer_key, contact_emails, owner_email FROM ar_jobs"):
        for e in _emails(r['contact_emails']) + _emails(r['owner_email']):
            addr.setdefault(e.lower(), set()).add(r['customer_key'])
    conn.close()

    today = _today()
    handled = 0
    for m in messages:
        if m['id'] in seen:
            continue
        sender = ((m.get('from') or {}).get('emailAddress') or {}).get('address', '').lower()
        if not sender or sender == mailbox.lower():
            continue
        subject = m.get('subject') or ''
        body = m.get('uniqueBody') or m.get('body') or {}
        text = _html_to_text(body.get('content', '')) if body.get('contentType') == 'html' else (body.get('content') or '')
        invs = [by_id[i] for i in convs.get(m.get('conversationId') or '', ()) if i in by_id]
        if not invs and sender in addr:
            for ck in addr[sender]:
                invs.extend(by_cust.get(ck, []))
            named = [i for i in invs if re.search(rf"\b{re.escape(i['number'])}\b", subject + ' ' + text)]
            if named:
                invs = named
            elif not re.search(r'(?i)invoice|payment|balance|statement|past due|retainage', subject + ' ' + text):
                invs = []
        if not invs:
            nums = set(re.findall(r'\b(\d{4,6})\b', subject))
            invs = [i for i in open_invs if i['number'] in nums] if re.search(r'(?i)\binv', subject) else []
        if re.search(r'(?i)undeliverable|delivery (status|has failed)|mail delivery|returned mail', subject):
            hit = [e for e in addr if e in text.lower()]
            bounced = [i for ck in {c for e in hit for c in addr[e]} for i in by_cust.get(ck, [])]
            if bounced:
                conn = _conn()
                for i in bounced:
                    conn.execute("UPDATE ar_invoices SET status='needs_person', needs_reason=? WHERE id=? AND status != 'paid'",
                                 (f"Email to {', '.join(hit)} bounced: check the address", i['id']))
                    _note(conn, f"Our email to {', '.join(hit)} bounced. Follow-ups stopped until the address is fixed.",
                          i['id'], customer_key=i['customer_key'])
                conn.execute("INSERT OR IGNORE INTO ar_replies (graph_id, from_email, subject, received_at, intent, summary, "
                             "customer_key, invoice_ids, created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                             (m['id'], sender, subject, m.get('receivedDateTime', ''), 'bounce', f"Bounced: {', '.join(hit)}",
                              bounced[0]['customer_key'], json.dumps([i['id'] for i in bounced]), _stamp()))
                conn.commit()
                conn.close()
                handled += 1
            continue
        if not invs:
            continue  # someone else's mail: not stored
        got = classify_reply(subject, text, invs, today)
        nums = set(got.get('invoice_numbers') or [])
        if nums and any(i['number'] in nums for i in invs):
            invs = [i for i in invs if i['number'] in nums]
        conn = _conn()
        conn.execute('''INSERT OR IGNORE INTO ar_replies (graph_id, conversation_id, from_email, subject, received_at, text,
                        intent, promised_date, summary, customer_key, invoice_ids, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''',
                     (m['id'], m.get('conversationId') or '', sender, subject, m.get('receivedDateTime', ''), text[:8000],
                      got.get('intent', ''), got.get('promised_date', ''), got.get('summary', ''), invs[0]['customer_key'],
                      json.dumps([i['id'] for i in invs]), _stamp()))
        _apply_reply(conn, invs, got, sender, s, today)
        conn.commit()
        conn.close()
        handled += 1
        _log('replies', f"{sender} replied about invoice(s) {', '.join(i['number'] for i in invs)}: {got.get('intent')}.")
    _set_state('last_reply_scan', started)
    if handled:
        _log('replies', f'{handled} repl(ies) read from {mailbox}.')


# ── 3. Follow-up emails ──────────────────────────────────────────────────────

def _group_key(inv):
    if inv['kind'] == 'install' and inv['job_key']:
        return f"job:{inv['job_key']}"
    return f"cust:{inv['customer_key']}"


def _stage_rank(key, s):
    keys = [st['key'] for st in s['stages']]
    return keys.index(key) if key in keys else -1


def _greeting(cust, job):
    name = (job or {}).get('contact_name') or (cust or {}).get('contact_name') or ''
    if name:
        return f"Hi {_nice_name(name).split(' ')[0]},"
    return f"Hello {_nice_name((cust or {}).get('name') or '')},".replace('Hello ,', 'Hello,')


def _inv_line(inv, today):
    due = _due(inv)
    dpd = (today - due).days if due else 0
    bits = [f"Invoice #{inv['number']}"]
    if inv['txn_date']:
        bits.append(f"dated {_fmt_date(inv['txn_date'])}")
    line = ' '.join(bits) + f" - {_money_text(inv['open_balance'])} open"
    if inv['retainage']:
        line += ' (retainage)'
    elif due:
        line += f", due {_fmt_date(due)}" + (f" ({dpd} days past due)" if dpd > 0 else '')
    if inv.get('client_hub_uri'):
        line += f"\n    View or pay online: {inv['client_hub_uri']}"
    return line


def compose(group_invoices, stage_key, cust, job, s, today=None, other_open=None):
    """(subject, body) for one follow-up email."""
    today = today or _today()
    total = sum(i['open_balance'] or 0 for i in group_invoices)
    nums = ', '.join(f"#{i['number']}" for i in group_invoices[:6]) + (' and more' if len(group_invoices) > 6 else '')
    where = f" - {job['name']}" if job and job.get('name') and job.get('kind') == 'install' else ''
    company = 'Stahlman-England Irrigation'
    promised = [i for i in group_invoices if i.get('promised_on')]
    install = bool(job) and job.get('kind') == 'install'
    lines = [_greeting(cust, job), '']
    if stage_key == 'retainage':
        subject = f"Retainage release{where} - {_money_text(total)}"
        lines += [f"Our work{(' on ' + job['name']) if job and job.get('name') else ''} is complete, so the retainage "
                  f"held on it is now due. The retainage outstanding is:", '']
    elif stage_key == 'reminder':
        subject = f"Past due invoice{'s' if len(group_invoices) > 1 else ''} {nums} - {company}"
        lines += ['This is a friendly reminder that the following is now past due:', '']
    elif stage_key == 'past_due':
        subject = f"Second notice: past due balance {_money_text(total)}{where} ({nums})"
        lines += ["We're following up on the past due balance below and haven't received payment yet:", '']
    elif stage_key == 'very_past_due':
        subject = f"Past due balance {_money_text(total)}{where} - payment needed ({nums})"
        lines += ['The balance below is now well past due:', '']
    else:
        subject = f"FINAL NOTICE: past due balance {_money_text(total)}{where} ({nums})"
        lines += ['This is our final notice for the balance below:', '']
    for i in group_invoices:
        lines.append('  ' + _inv_line(i, today))
    lines += ['', f"Total: {_money_text(total)}", '']
    if promised:
        p = promised[0]
        when = f" for {_fmt_date(p['promised_date'])}" if p.get('promised_date') else ''
        lines += [f"On {_fmt_date(p['promised_on'])} you let us know payment{when} was on the way, but we haven't "
                  f"received it yet.", '']
    if stage_key == 'retainage':
        lines += ['Please let us know when we can expect it, or what you need from us to release it (final lien waiver, '
                  'close-out documents, etc.).']
    elif stage_key == 'reminder':
        lines += ["If it's already on its way, thank you and please disregard this note. Otherwise, please let us know "
                  "when we can expect payment. The invoice is attached."]
    elif stage_key == 'past_due':
        lines += ['Please send payment, or reply with the date we can expect it. If anything on the invoice needs '
                  'attention, let us know and we will take care of it.']
    elif stage_key == 'very_past_due':
        lines += ['Please arrange payment within 7 days, or reply today with the date it will be paid.']
        if install:
            lines += ['', 'Please note that we have preserved our lien rights on this project.']
    else:
        if install:
            lines += ['Unless payment is received within 10 days, we will serve a Notice of Nonpayment on the property '
                      'owner and record a claim of lien against the property to protect our rights.']
        else:
            lines += ['Please contact us right away to bring this account current. Unless payment is received within '
                      '10 days, we will have to escalate collection of this balance.']
    if other_open:
        held = sum(i['open_balance'] or 0 for i in other_open)
        lines += ['', f"For your records, retainage of {_money_text(held)} is also held on this job until completion."]
    lines += ['', s.get('signature') or company]
    return subject[:250], '\n'.join(lines)


def _docs(scope, key):
    conn = _conn()
    rows = [dict(r) for r in conn.execute("SELECT * FROM ar_documents WHERE scope=? AND scope_key=? ORDER BY id",
                                          (scope, key or ''))]
    conn.close()
    return rows


def _doc_bytes(doc):
    with open(os.path.join(_files_dir(), doc['stored_name']), 'rb') as fh:
        return fh.read()


def attachments_for(invoices, job_key, customer_key, kind, s):
    """[(filename, bytes, content type)] for a follow-up: each invoice (uploaded PDF, else made from Jobber),
    then for installs the job's documents, then the customer's and the "every email" ones."""
    out, names = [], set()

    def add(name, data, ctype):
        if name in names:
            return
        names.add(name)
        out.append((name, data, ctype))

    for inv in invoices:
        uploaded = [d for d in _docs('invoice', str(inv['id'])) if d['kind'] == 'invoice_pdf']
        if uploaded:
            for d in uploaded:
                add(d['filename'], _doc_bytes(d), d['content_type'] or 'application/pdf')
        else:
            add(f"Invoice-{inv['number']}.pdf", invoice_pdf(inv, s), 'application/pdf')
    if kind == 'install' and job_key:
        for d in _docs('job', job_key):
            add(d['filename'], _doc_bytes(d), d['content_type'] or 'application/octet-stream')
    for d in _docs('customer', customer_key) + _docs('all', ''):
        add(d['filename'], _doc_bytes(d), d['content_type'] or 'application/octet-stream')
    cap = float(s.get('max_attach_mb') or 20) * 1024 * 1024
    kept, size, dropped = [], 0, []
    for a in out:
        if size + len(a[1]) > cap:
            dropped.append(a[0])
            continue
        kept.append(a)
        size += len(a[1])
    return kept, dropped


def _close_stale_drafts():
    """Drop automatic drafts whose invoices are no longer due (paid, promised, needs a person...)."""
    s = settings()
    invs, jobs, custs = _load_all()
    by_id = {i['id']: i for i in invs}
    conn = _conn()
    for e in conn.execute("SELECT * FROM ar_emails WHERE status='draft' AND kind='followup' AND manual=0").fetchall():
        ids = json.loads(e['invoice_ids'] or '[]')
        still = [i for i in ids if i in by_id and plan(by_id[i], jobs.get(by_id[i]['job_key']),
                                                         custs.get(by_id[i]['customer_key']), s)['due_now']]
        if not still:
            conn.execute("UPDATE ar_emails SET status='discarded', error='No longer needed' WHERE id=?", (e['id'],))
    conn.commit()
    conn.close()


def _in_send_window(s):
    now = _now()
    try:
        a, b = (int(x) for x in str(s.get('send_hours') or '8-17').split('-'))
    except ValueError:
        a, b = 8, 17
    return now.weekday() < 5 and a <= now.hour < b


def plan_followups(force_invoice_id=None):
    """Draft (and, when automatic, send) the follow-ups that are due. Returns the draft ids made or refreshed."""
    s = settings()
    today = _today()
    invs, jobs, custs = _load_all()
    conn = _conn()
    # Promises and "we paid" that have run out go back to normal follow-ups.
    for inv in invs:
        snooze = _d(inv.get('snooze_until'))
        if inv['status'] in ('promised', 'reported_paid') and snooze and snooze <= today:
            msg = ('The promised payment has not arrived; follow-ups resumed.' if inv['status'] == 'promised'
                   else "The payment they reported hasn't shown up; follow-ups resumed.")
            conn.execute("UPDATE ar_invoices SET status='open', snooze_until='' WHERE id=?", (inv['id'],))
            _note(conn, msg, inv['id'], customer_key=inv['customer_key'])
            inv['status'], inv['snooze_until'] = 'open', ''
    conn.commit()
    conn.close()

    groups = {}
    for inv in invs:
        job, cust = jobs.get(inv['job_key']), custs.get(inv['customer_key'])
        p = plan(inv, job, cust, s, today)
        if p['due_now'] or inv['id'] == force_invoice_id:
            inv['_stage'] = p['stage'] or (s['stages'][0]['key'] if not inv['retainage'] else 'retainage')
            groups.setdefault(_group_key(inv), []).append(inv)
    made = []
    for gk, members in groups.items():
        first = members[0]
        job = jobs.get(first['job_key']) if gk.startswith('job:') else None
        cust = custs.get(first['customer_key'])
        regular = [i for i in members if not i['retainage']]
        if regular:
            group_invs = regular
            stage = max((i['_stage'] for i in regular), key=lambda k: _stage_rank(k, s))
        else:
            group_invs, stage = members, 'retainage'
        held = []
        if job and stage != 'retainage':
            held = [i for i in invs if i['job_key'] == job['key'] and i['retainage'] and i['status'] != 'paid'
                    and i not in group_invs]
        to = contacts_for(first, job, cust)
        subject, body = compose(group_invs, stage, cust, job, s, today, held)
        ids = sorted(i['id'] for i in group_invs)
        conn = _conn()
        draft = conn.execute("SELECT * FROM ar_emails WHERE group_key=? AND status='draft' AND kind='followup' "
                             "ORDER BY id DESC LIMIT 1", (gk,)).fetchone()
        if draft and json.loads(draft['invoice_ids'] or '[]') == ids and draft['stage'] == stage:
            email_id = draft['id']  # keep any edits made to it
        elif draft:
            conn.execute("UPDATE ar_emails SET to_email=?, cc=?, subject=?, body=?, invoice_ids=?, stage=?, created_at=?, "
                         "manual=? WHERE id=?", (', '.join(to), s.get('cc') or '', subject, body, json.dumps(ids), stage,
                                                 _stamp(), 1 if force_invoice_id else draft['manual'], draft['id']))
            email_id = draft['id']
        else:
            email_id = conn.execute('''INSERT INTO ar_emails (kind, group_key, customer_key, job_key, invoice_ids, stage,
                                       to_email, cc, subject, body, status, manual, created_at)
                                       VALUES ('followup',?,?,?,?,?,?,?,?,?, 'draft', ?, ?)''',
                                    (gk, first['customer_key'], job['key'] if job else '', json.dumps(ids), stage,
                                     ', '.join(to), s.get('cc') or '', subject, body, 1 if force_invoice_id else 0,
                                     _stamp())).lastrowid
        if not to:
            conn.execute("UPDATE ar_emails SET error=? WHERE id=?",
                         (f"No email address for {(cust or {}).get('name') or first['customer']}: add one on the "
                          f"Customers tab", email_id))
        conn.commit()
        conn.close()
        made.append(email_id)
        if s.get('auto_send') and to and not force_invoice_id and _in_send_window(s):
            send_email(email_id, by='automatic')
    if made:
        _log('followups', f'{len(made)} follow-up email(s) ready' + (' (sending automatically)' if s.get('auto_send') else
                                                                      ' to review on the Today tab') + '.')
    return made


def _send_mail(mailbox, to, cc, subject, body, attachments):
    """Send through Graph as a draft first, so the conversation id is known for matching replies."""
    msg = {'subject': subject, 'body': {'contentType': 'Text', 'content': body},
           'toRecipients': [{'emailAddress': {'address': a}} for a in to]}
    if cc:
        msg['ccRecipients'] = [{'emailAddress': {'address': a}} for a in cc]
    draft = _graph('POST', f'/users/{mailbox}/messages', json=msg)
    mid = draft.get('id')
    try:
        for name, data, ctype in attachments:
            if len(data) < 3 * 1024 * 1024:
                _graph('POST', f'/users/{mailbox}/messages/{mid}/attachments', json={
                    '@odata.type': '#microsoft.graph.fileAttachment', 'name': name, 'contentType': ctype,
                    'contentBytes': base64.b64encode(data).decode()})
            else:
                sess = _graph('POST', f'/users/{mailbox}/messages/{mid}/attachments/createUploadSession', json={
                    'AttachmentItem': {'attachmentType': 'file', 'name': name, 'size': len(data), 'contentType': ctype}})
                _graph_upload(sess['uploadUrl'], data)
        _graph('POST', f'/users/{mailbox}/messages/{mid}/send')
    except Exception:
        try:
            _graph('DELETE', f'/users/{mailbox}/messages/{mid}')
        except Exception:
            pass
        raise
    return mid, draft.get('conversationId') or ''


def send_email(email_id, by=''):
    s = settings()
    conn = _conn()
    e = conn.execute("SELECT * FROM ar_emails WHERE id=?", (email_id,)).fetchone()
    conn.close()
    if not e or e['status'] == 'sent':
        return False, 'Already sent'
    e = dict(e)
    to, cc = _emails(e['to_email']), _emails(e['cc'])
    if not to:
        return False, 'No email address to send to'
    if not graph_ready():
        return False, 'Microsoft 365 is not connected'
    ids = json.loads(e['invoice_ids'] or '[]')
    conn = _conn()
    invs = [dict(r) for r in conn.execute(f"SELECT * FROM ar_invoices WHERE id IN ({','.join('?' * len(ids))})", ids)] if ids else []
    conn.close()
    try:
        if e['kind'] == 'nonp':
            job = _job(e['job_key'])
            atts = [(f"Notice-of-Nonpayment-{_squash(job['name'])[:30] or 'job'}.pdf", nonp_pdf(job, invs, s),
                     'application/pdf')] if job else []
            atts += [(f"Invoice-{i['number']}.pdf", invoice_pdf(i, s), 'application/pdf') for i in invs]
            dropped = []
        else:
            kind = 'install' if e['job_key'] else (invs[0]['kind'] if invs else 'service')
            atts, dropped = attachments_for(invs, e['job_key'], e['customer_key'], kind, s)
        mid, conv = _send_mail(s['from_email'].strip(), to, cc, e['subject'], e['body'], atts)
    except Exception as ex:
        conn = _conn()
        conn.execute("UPDATE ar_emails SET error=? WHERE id=?", (str(ex)[:500], email_id))
        conn.commit()
        conn.close()
        _log('email', f'Email #{email_id} not sent: {ex}')
        return False, str(ex)
    now = _stamp()
    conn = _conn()
    conn.execute('''UPDATE ar_emails SET status='sent', sent_at=?, sent_by=?, error=?, graph_id=?, conversation_id=?,
                    attachments=? WHERE id=?''',
                 (now, by or 'office', ('Too large to attach: ' + ', '.join(dropped)) if dropped else '', mid or '', conv,
                  json.dumps([a[0] for a in atts]), email_id))
    for inv in invs:
        if e['kind'] == 'nonp':
            _note(conn, f"Notice of Nonpayment emailed to {', '.join(to)}.", inv['id'], job_key=e['job_key'],
                  customer_key=inv['customer_key'], kind='email')
            continue
        conn.execute("UPDATE ar_invoices SET last_followup_at=?, last_stage=?, followup_count=followup_count+1 WHERE id=?",
                     (now, e['stage'], inv['id']))
        _note(conn, f"Follow-up ({e['stage'].replace('_', ' ')}) emailed to {', '.join(to)} with "
                    f"{len(atts)} attachment(s).", inv['id'], customer_key=inv['customer_key'], kind='email',
              by=by or 'office')
    if e['kind'] == 'nonp' and e['job_key']:
        conn.execute("UPDATE ar_jobs SET nonp_status='sent', nonp_at=? WHERE key=?", (now, e['job_key']))
        _note(conn, f"Notice of Nonpayment emailed to {', '.join(to)}. Serve it by certified mail as well and record the "
                    f"claim of lien before the deadline.", job_key=e['job_key'], kind='nonp')
    conn.commit()
    conn.close()
    _log('email', f"Sent to {', '.join(to)}: {e['subject']}")
    return True, ''


# ── 4. Lien deadlines and the Notice of Nonpayment ───────────────────────────

def _job(key):
    conn = _conn()
    r = conn.execute("SELECT * FROM ar_jobs WHERE key=?", (key,)).fetchone()
    conn.close()
    return dict(r) if r else None


_CONTRACTOR = re.compile(r'(?i)contract|construct|builder|homes\b|development|general|landscap|\bgc\b|pool')


def _owner(job, cust):
    """(name, email, address) of the property owner. A customer who looks like a contractor is not the owner."""
    if job.get('owner_name') or job.get('owner_email'):
        return job.get('owner_name') or '', job.get('owner_email') or '', job.get('owner_address') or ''
    if cust and not _CONTRACTOR.search(cust.get('name') or ''):
        em = _emails(cust.get('emails')) or _emails(cust.get('jobber_emails'))
        return _nice_name(cust.get('name')), (em[0] if em else ''), ''
    return '', '', ''


def prepare_nonp(job_key, by='Office App'):
    """Make the Notice of Nonpayment PDF and a draft email to the owner. Returns (ok, message)."""
    s = settings()
    invs, jobs, custs = _load_all()
    job = jobs.get(job_key)
    if not job:
        return False, 'No such job'
    members = [i for i in invs if i['job_key'] == job_key and i['status'] != 'paid' and (i['open_balance'] or 0) > 0.004]
    if not members:
        return False, 'Nothing is owed on this job'
    cust = custs.get(job['customer_key'])
    owner_name, owner_email, _addr = _owner(job, cust)
    if not owner_name:
        return False, 'Add the property owner on the job first (the customer looks like a contractor)'
    pdf = nonp_pdf(job, members, s)
    stored = f"{uuid.uuid4().hex}_nonp.pdf"
    with open(os.path.join(_files_dir(), stored), 'wb') as fh:
        fh.write(pdf)
    li = lien_info(job, [i for i in invs if i['job_key'] == job_key], s)
    total = sum(i['open_balance'] or 0 for i in members)
    conn = _conn()
    doc_id = conn.execute('''INSERT INTO ar_documents (scope, scope_key, kind, filename, stored_name, content_type, size,
                             uploaded_by, uploaded_at) VALUES ('job', ?, 'nonp', ?, ?, 'application/pdf', ?, ?, ?)''',
                          (job_key, f"Notice of Nonpayment - {job['name'][:40]}.pdf", stored, len(pdf), by,
                           _stamp())).lastrowid
    # The NONP is its own document; keep it out of ordinary follow-ups.
    conn.execute("UPDATE ar_documents SET scope='nonp' WHERE id=?", (doc_id,))
    contractor = _emails((cust or {}).get('emails')) or _emails((cust or {}).get('jobber_emails'))
    body = '\n'.join([
        f"Dear {owner_name},", '',
        f"{s['company_name']} furnished irrigation labor and materials for the improvement of the property at "
        f"{job.get('property_address') or job['name']}" + (f", under contract with {_nice_name(cust['name'])}" if cust and
                                                            _nice_name(cust['name']) != owner_name else '') + '.', '',
        f"We have not been paid {_money_text(total)} for that work. The attached Notice of Nonpayment is served on you as "
        f"owner of the property.", '',
        "Unless payment is received, it will be followed by a claim of lien recorded against the property"
        + (f" on or before {_fmt_date(li.get('deadline'))}" if li.get('deadline') else '') + '.', '',
        'Please contact us right away to resolve this.', '', s.get('signature') or s['company_name']])
    ids = json.dumps(sorted(i['id'] for i in members))
    conn.execute("UPDATE ar_emails SET status='discarded' WHERE kind='nonp' AND job_key=? AND status='draft'", (job_key,))
    email_id = conn.execute('''INSERT INTO ar_emails (kind, group_key, customer_key, job_key, invoice_ids, stage, to_email, cc,
                               subject, body, status, manual, created_at, error)
                               VALUES ('nonp', ?, ?, ?, ?, 'nonp', ?, ?, ?, ?, 'draft', 1, ?, ?)''',
                            (f"nonp:{job_key}", job['customer_key'], job_key, ids, owner_email, ', '.join(contractor),
                             f"Notice of Nonpayment - {job.get('property_address') or job['name']}", body, _stamp(),
                             '' if owner_email else 'No owner email: add it on the job, or mail the notice')).lastrowid
    conn.execute("UPDATE ar_jobs SET nonp_status='prepared', nonp_at=?, nonp_doc_id=? WHERE key=?", (_stamp(), doc_id, job_key))
    _note(conn, f"Notice of Nonpayment prepared for {owner_name} ({_money_text(total)}). Lien deadline "
                f"{_fmt_date(li.get('deadline'))}.", job_key=job_key, kind='nonp', by=by)
    conn.commit()
    conn.close()
    _log('lien', f"Notice of Nonpayment prepared for {job['name']} ({_money_text(total)}).")
    if s.get('nonp_auto') and owner_email and _in_send_window(s):
        send_email(email_id, by='automatic')
    return True, email_id


def check_liens():
    s = settings()
    invs, jobs, custs = _load_all()
    for key, job in jobs.items():
        members = [i for i in invs if i['job_key'] == key]
        li = lien_info(job, members, s)
        if li.get('applies') and li.get('nonp_due'):
            ok, msg = prepare_nonp(key)
            if not ok:
                _log('lien', f"{job['name']}: Notice of Nonpayment due but not prepared: {msg}")


# ── PDFs ─────────────────────────────────────────────────────────────────────

def _pdf_doc():
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
    return buf, doc, getSampleStyleSheet()


def _p(text, style):
    from reportlab.platypus import Paragraph
    return Paragraph(html.escape(str(text or '')).replace('\n', '<br/>'), style)


def invoice_pdf(inv, s):
    """The invoice as a PDF: Jobber's line items when known, else QuickBooks' totals."""
    from reportlab.lib import colors
    from reportlab.platypus import Spacer, Table, TableStyle
    buf, doc, st = _pdf_doc()
    story = [_p(s['company_name'], st['Title'])]
    head = '\n'.join(x for x in (s.get('company_address'), s.get('company_phone'), s.get('from_email')) if x)
    if head:
        story.append(_p(head, st['Normal']))
    story += [Spacer(1, 12), _p(f"INVOICE #{inv['number']}", st['Heading2'])]
    meta = [['Bill to', _nice_name(inv.get('jobber_client_name') or inv['customer'])]]
    if inv.get('billing_address'):
        meta.append(['', inv['billing_address']])
    if inv.get('txn_date'):
        meta.append(['Invoice date', _fmt_date(inv['txn_date'])])
    if _due(inv):
        meta.append(['Due date', _fmt_date(_due(inv))])
    if inv.get('terms'):
        meta.append(['Terms', inv['terms']])
    if inv.get('po'):
        meta.append(['PO #', inv['po']])
    if inv.get('jobber_subject'):
        meta.append(['Subject', inv['jobber_subject']])
    t = Table(meta, colWidths=[90, 400])
    t.setStyle(TableStyle([('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    story += [t, Spacer(1, 14)]
    lines = json.loads(inv.get('jobber_lines') or '[]')
    rows = [['Item', 'Qty', 'Unit price', 'Total']]
    for ln in lines:
        desc = (ln.get('name') or '') + (('\n' + ln['description'].strip()) if ln.get('description') else '')
        rows.append([_p(desc, st['Normal']), f"{ln.get('quantity') or ''}", _money_text(ln.get('unitPrice')),
                     _money_text(ln.get('totalPrice'))])
    if not lines:
        rows.append([_p(inv.get('memo') or 'Services as invoiced', st['Normal']), '1',
                     _money_text(inv.get('amount') or inv['open_balance']), _money_text(inv.get('amount') or inv['open_balance'])])
    t = Table(rows, colWidths=[300, 40, 75, 75], repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E2E8F0')),
                           ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                           ('ALIGN', (1, 0), (-1, -1), 'RIGHT'), ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.grey)]))
    story += [t, Spacer(1, 10)]
    total = inv.get('jobber_total') or inv.get('amount') or inv['open_balance']
    paid = max(0.0, (total or 0) - (inv['open_balance'] or 0))
    tot = [['Total', _money_text(total)], ['Payments', _money_text(paid)], ['Balance due', _money_text(inv['open_balance'])]]
    if inv.get('retainage'):
        tot.append(['', f"Balance is retainage ({inv['retainage_pct']:g}%)" if inv.get('retainage_pct') else 'Retainage'])
    t = Table(tot, colWidths=[415, 75])
    t.setStyle(TableStyle([('ALIGN', (0, 0), (-1, -1), 'RIGHT'), ('FONTNAME', (0, 2), (-1, 2), 'Helvetica-Bold')]))
    story.append(t)
    if inv.get('client_hub_uri'):
        story += [Spacer(1, 14), _p(f"View or pay online: {inv['client_hub_uri']}", st['Normal'])]
    doc.build(story)
    return buf.getvalue()


def nonp_pdf(job, invoices, s):
    from reportlab.platypus import Spacer
    buf, doc, st = _pdf_doc()
    conn = _conn()
    cust = conn.execute("SELECT * FROM ar_customers WHERE key=?", (job['customer_key'],)).fetchone()
    conn.close()
    cust = dict(cust) if cust else {}
    owner_name, owner_email, owner_addr = _owner(job, cust)
    allinv = invoices
    li = lien_info(job, allinv, s)
    total = sum(i['open_balance'] or 0 for i in invoices)
    dates = sorted(_d(i['txn_date']) for i in invoices if _d(i['txn_date']))
    first = _d(job.get('first_furnished')) or (dates[0] if dates else None)
    last = _d(li.get('last_furnished'))
    pay_by = _today() + timedelta(days=10)
    if li.get('deadline'):
        pay_by = min(pay_by, _d(li['deadline']) - timedelta(days=5))
    story = [_p('NOTICE OF NONPAYMENT', st['Title']), Spacer(1, 6),
             _p(f"Date: {_today().strftime('%B %d, %Y')}", st['Normal']), Spacer(1, 10),
             _p(f"To (Owner): {owner_name or '________________'}\n{owner_addr}", st['Normal']), Spacer(1, 6),
             _p(f"Property: {job.get('property_address') or job['name']}", st['Normal']), Spacer(1, 6)]
    if cust and _nice_name(cust.get('name')) != owner_name:
        story += [_p(f"Contractor: {_nice_name(cust.get('name'))}", st['Normal']), Spacer(1, 6)]
    story += [_p(f"Lienor: {s['company_name']}\n{s.get('company_address') or ''}\n{s.get('company_phone') or ''}",
                 st['Normal']), Spacer(1, 14),
              _p("The undersigned lienor notifies you that it has furnished labor, services and/or materials consisting of "
                 "irrigation work for the improvement of the real property identified above"
                 + (f", from {_fmt_date(first)}" if first else '') + (f" through {_fmt_date(last)}" if last else '')
                 + f", and that it has not been paid. The amount now due and unpaid is {_money_text(total)}, as "
                   f"shown on the invoices listed below.", st['Normal']), Spacer(1, 10)]
    for i in invoices:
        story.append(_p(f"Invoice #{i['number']} dated {_fmt_date(i['txn_date'])}: {_money_text(i['open_balance'])}"
                        + (' (retainage)' if i['retainage'] else ''), st['Normal']))
    story += [Spacer(1, 12),
              _p(f"Unless the amount due is paid in full by {_fmt_date(pay_by)}, the lienor intends to record a claim of "
                 f"lien against the property under Part I of Chapter 713, Florida Statutes"
                 + (f", on or before {_fmt_date(li['deadline'])}" if li.get('deadline') else '') + '.', st['Normal']),
              Spacer(1, 30), _p(f"{s['company_name']}\n\nBy: ______________________________\n\nTitle: ___________________",
                                st['Normal'])]
    doc.build(story)
    return buf.getvalue()


# ── The cycle ────────────────────────────────────────────────────────────────

class _FileLock:
    """Only one gunicorn worker runs the cycle at a time."""

    def __init__(self):
        self.fh = None

    def __enter__(self):
        if not _run_lock.acquire(blocking=False):
            return False
        if fcntl:
            self.fh = open(os.path.join(_cfg['data_dir'], '.receivables_cycle.lock'), 'w')
            try:
                fcntl.flock(self.fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError:
                self.fh.close()
                self.fh = None
                _run_lock.release()
                return False
        return True

    def __exit__(self, *exc):
        if self.fh:
            fcntl.flock(self.fh, fcntl.LOCK_UN)
            self.fh.close()
            self.fh = None
        if _run_lock.locked():
            _run_lock.release()
        return False


def run_cycle(manual=False):
    if not manual and not settings().get('enabled'):
        return False
    with _FileLock() as got:
        if not got:
            return False
        for step in (sync_jobber, scan_replies, _close_stale_drafts, plan_followups, check_liens):
            try:
                step()
            except Exception as e:
                _log(step.__name__, f'Failed: {e}')
        _set_state('last_run', _stamp())
    return True


# ── Pages and API ────────────────────────────────────────────────────────────

def _deny():
    if 'username' not in session:
        return redirect(url_for('login'))
    return redirect(url_for('dashboard'))


def _forbidden():
    return jsonify({'success': False, 'error': 'Office only'}), 403


def _who():
    return session.get('full_name') or session.get('username') or 'office'


@bp.route('/receivables')
def page():
    if not receivables_allowed():
        return _deny()
    return render_template_string(PAGE_TEMPLATE, full_name=_who())


@bp.route('/receivables/api/data')
def api_data():
    if not receivables_allowed():
        return _forbidden()
    s = settings()
    today = _today()
    invs, jobs, custs = _load_all()
    by_job = {}
    for i in invs:
        if i['job_key']:
            by_job.setdefault(i['job_key'], []).append(i)
    out_inv = []
    for i in invs:
        job, cust = jobs.get(i['job_key']), custs.get(i['customer_key'])
        p = plan(i, job, cust, s, today)
        i.pop('jobber_lines', None)
        i.update(plan=p, job_name=(job or {}).get('name', ''), contacts=contacts_for(i, job, cust),
                 customer_name=_nice_name((cust or {}).get('name') or i['customer']))
        out_inv.append(i)
    out_jobs = []
    for key, j in jobs.items():
        members = by_job.get(key, [])
        open_m = [i for i in members if i['status'] != 'paid']
        if not open_m and not j.get('nonp_status'):
            continue
        j['lien'] = lien_info(j, members, s, today)
        j['complete_now'] = job_complete(j)
        j['open_count'] = len(open_m)
        j['open_balance'] = round(sum(i['open_balance'] or 0 for i in open_m), 2)
        j['retainage_held'] = round(sum(i['open_balance'] or 0 for i in open_m if i['retainage']), 2)
        j['retainage_count'] = sum(1 for i in open_m if i['retainage'])
        j['customer_name'] = _nice_name((custs.get(j['customer_key']) or {}).get('name') or '')
        j['owner'] = list(_owner(j, custs.get(j['customer_key'])))
        out_jobs.append(j)
    out_cust = []
    for key, c in custs.items():
        mine = [i for i in invs if i['customer_key'] == key and i['status'] != 'paid']
        if not mine:
            continue
        c['open_count'] = len(mine)
        c['open_balance'] = round(sum(i['open_balance'] or 0 for i in mine), 2)
        c['nice_name'] = _nice_name(c['name'])
        out_cust.append(c)
    conn = _conn()
    emails = [dict(r) for r in conn.execute("SELECT * FROM ar_emails WHERE status != 'discarded' ORDER BY id DESC LIMIT 300")]
    replies = [dict(r) for r in conn.execute("SELECT id, from_email, subject, received_at, intent, promised_date, summary, "
                                             "customer_key, invoice_ids FROM ar_replies ORDER BY id DESC LIMIT 100")]
    uploads = [dict(r) for r in conn.execute("SELECT * FROM ar_uploads ORDER BY id DESC LIMIT 10")]
    docs = [dict(r) for r in conn.execute("SELECT id, scope, scope_key, kind, filename, size, uploaded_by, uploaded_at "
                                          "FROM ar_documents WHERE scope IN ('all','customer','job') ORDER BY id DESC")]
    activity = [dict(r) for r in conn.execute("SELECT at, kind, message FROM ar_activity ORDER BY id DESC LIMIT 150")]
    conn.close()
    return jsonify({'success': True, 'invoices': out_inv, 'jobs': out_jobs, 'customers': out_cust, 'emails': emails,
                    'replies': replies, 'uploads': uploads, 'docs': docs, 'activity': activity, 'settings': s,
                    'statuses': STATUSES, 'today': today.isoformat(),
                    'status': {'graph': graph_ready(), 'jobber': jobber_ready(), 'claude': claude_ready(),
                               'last_run': _state('last_run'), 'last_reply_scan': _state('last_reply_scan')}})


@bp.route('/receivables/api/invoice/<int:inv_id>')
def api_invoice(inv_id):
    if not receivables_allowed():
        return _forbidden()
    s = settings()
    conn = _conn()
    inv = conn.execute("SELECT * FROM ar_invoices WHERE id=?", (inv_id,)).fetchone()
    if not inv:
        conn.close()
        return jsonify({'success': False, 'error': 'Not found'}), 404
    inv = dict(inv)
    job = conn.execute("SELECT * FROM ar_jobs WHERE key=?", (inv['job_key'],)).fetchone()
    job = dict(job) if job else None
    cust = conn.execute("SELECT * FROM ar_customers WHERE key=?", (inv['customer_key'],)).fetchone()
    cust = dict(cust) if cust else None
    notes = [dict(r) for r in conn.execute("SELECT * FROM ar_notes WHERE invoice_id=? OR (job_key != '' AND job_key=?) "
                                           "ORDER BY id DESC", (inv_id, inv['job_key'] or '~'))]
    emails = [dict(r) for r in conn.execute("SELECT * FROM ar_emails WHERE status != 'discarded' ORDER BY id DESC")
              if inv_id in json.loads(r['invoice_ids'] or '[]')]
    job_invs = [dict(r) for r in conn.execute("SELECT * FROM ar_invoices WHERE job_key=? AND job_key != '' ORDER BY txn_date",
                                              (inv['job_key'],))]
    siblings = [{k: x[k] for k in ('id', 'number', 'txn_date', 'due_date', 'open_balance', 'amount', 'retainage', 'status')}
                for x in job_invs if x['status'] != 'paid']
    conn.close()
    inv['jobber_lines'] = json.loads(inv.get('jobber_lines') or '[]')
    docs = _docs('invoice', str(inv_id))
    if job:
        job['lien'] = lien_info(job, job_invs or [inv], s)
        job['complete_now'] = job_complete(job)
        docs += _docs('job', job['key'])
    if cust:
        docs += _docs('customer', cust['key'])
    docs += _docs('all', '')
    for d in docs:
        d.pop('stored_name', None)
    return jsonify({'success': True, 'invoice': inv, 'job': job, 'customer': cust, 'notes': notes, 'emails': emails,
                    'siblings': siblings, 'docs': docs, 'plan': plan(inv, job, cust, s),
                    'contacts': contacts_for(inv, job, cust)})


@bp.route('/receivables/api/upload', methods=['POST'])
def api_upload():
    if not receivables_allowed():
        return _forbidden()
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'success': False, 'error': 'Choose the A/R Excel file from QuickBooks'}), 400
    try:
        parsed, info = parse_ar_table(_read_table(f.stream, f.filename))
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'Could not read that file: {e}'}), 400
    if not parsed:
        return jsonify({'success': False, 'error': 'No open invoices found in that file'}), 400
    summary = apply_upload(parsed, f.filename, _who(), confirm=request.form.get('confirm') == '1')
    if summary.get('needs_confirm'):
        return jsonify({'success': True, 'needs_confirm': True, **summary})
    threading.Thread(target=run_cycle, kwargs={'manual': True}, daemon=True).start()
    return jsonify({'success': True, 'summary': summary, 'info': info})


@bp.route('/receivables/api/invoice', methods=['POST'])
def api_invoice_update():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    conn = _conn()
    inv = conn.execute("SELECT * FROM ar_invoices WHERE id=?", (d.get('id'),)).fetchone()
    if not inv:
        conn.close()
        return jsonify({'success': False, 'error': 'Not found'}), 404
    inv = dict(inv)
    who = _who()
    fields, notes = {}, []
    if d.get('kind') in ('service', 'install'):
        fields.update(kind=d['kind'], kind_locked=1)
        notes.append(f"Set to {d['kind']}.")
    if d.get('kind') == 'auto':
        fields.update(kind_locked=0)
    if 'retainage' in d:
        if d['retainage'] == 'auto':
            fields.update(retainage_locked=0)
        else:
            fields.update(retainage=1 if d['retainage'] else 0, retainage_locked=1)
            notes.append('Marked as retainage.' if d['retainage'] else 'Marked as not retainage.')
    if 'job_key' in d:
        key = str(d['job_key'] or '').strip()
        if key == 'new':
            name = str(d.get('job_name') or '').strip() or f"Invoice {inv['number']}"
            key = f"manual:{inv['customer_key']}:{_squash(name)}"
            conn.execute("INSERT OR IGNORE INTO ar_jobs (key, customer_key, name, kind, updated_at) VALUES (?,?,?,?,?)",
                         (key, inv['customer_key'], name, 'install', _stamp()))
        fields['job_key'] = key
        notes.append(f"Moved to job {key or '(none)'}.")
    action = d.get('action')
    if action == 'pause':
        fields.update(status='paused')
        notes.append('Follow-ups paused.')
    elif action == 'resume':
        fields.update(status='open', needs_reason='', snooze_until='')
        notes.append('Follow-ups resumed.')
    elif action == 'promise':
        pd = _d(d.get('date'))
        s = settings()
        due = _due(inv) or _today()
        very = due + timedelta(days=_very_past_due_days(s))
        snooze = very if _today() < very else ((pd + timedelta(days=int(s['promise_grace_days']))) if pd else
                                               _today() + timedelta(days=14))
        fields.update(status='promised', promised_date=pd.isoformat() if pd else '', promised_on=_today().isoformat(),
                      snooze_until=snooze.isoformat(), needs_reason='')
        notes.append(f"Payment promised{' for ' + _fmt_date(pd) if pd else ''} (entered by hand); next follow-up "
                     f"{_fmt_date(snooze)}.")
    elif action == 'snooze':
        sd = _d(d.get('date'))
        if sd:
            fields.update(snooze_until=sd.isoformat())
            notes.append(f"Snoozed until {_fmt_date(sd)}.")
    elif action == 'close':
        fields.update(status='paid', closed_at=_stamp())
        notes.append('Closed by hand (paid or written off).')
    if fields:
        fields['updated_at'] = _stamp()
        conn.execute(f"UPDATE ar_invoices SET {', '.join(k + '=?' for k in fields)} WHERE id=?",
                     list(fields.values()) + [inv['id']])
    if (d.get('note') or '').strip():
        notes.append(d['note'].strip())
    for n in notes:
        _note(conn, n, inv['id'], customer_key=inv['customer_key'], kind='note', by=who)
    conn.commit()
    conn.close()
    if 'kind' in d or 'retainage' in d:
        classify_all()
    if action == 'draft':
        made = plan_followups(force_invoice_id=inv['id'])
        return jsonify({'success': True, 'drafts': made})
    if action == 'jobber':
        threading.Thread(target=sync_jobber, kwargs={'only_ids': [inv['id']]}, daemon=True).start()
    return jsonify({'success': True})


@bp.route('/receivables/api/job', methods=['POST'])
def api_job():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    job = _job(d.get('key') or '')
    if not job:
        return jsonify({'success': False, 'error': 'Not found'}), 404
    fields = {}
    for k in ('name', 'property_address', 'contact_name', 'contact_emails', 'owner_name', 'owner_email', 'owner_address',
              'notes'):
        if k in d:
            fields[k] = str(d[k] or '').strip()[:2000]
    for k in ('first_furnished', 'last_furnished'):
        if k in d:
            fields[k] = _date_text(d[k]) if d[k] else ''
    if 'complete' in d:
        fields['complete'] = None if d['complete'] in (None, 'auto', '') else (1 if d['complete'] else 0)
    if d.get('nonp_status') in NONP_STEPS:
        fields['nonp_status'] = d['nonp_status']
    if fields:
        fields['updated_at'] = _stamp()
        conn = _conn()
        conn.execute(f"UPDATE ar_jobs SET {', '.join(k + '=?' for k in fields)} WHERE key=?",
                     list(fields.values()) + [job['key']])
        if 'nonp_status' in fields:
            label = {'mailed': 'Notice of Nonpayment served by certified mail.', 'lien_recorded': 'Claim of lien recorded.',
                     'sent': 'Notice of Nonpayment marked as sent.', '': 'Notice of Nonpayment cleared.',
                     'prepared': 'Notice of Nonpayment marked as prepared.'}[fields['nonp_status']]
            _note(conn, label, job_key=job['key'], kind='nonp', by=_who())
        conn.commit()
        conn.close()
    if (d.get('note') or '').strip():
        conn = _conn()
        _note(conn, d['note'].strip(), job_key=job['key'], kind='note', by=_who())
        conn.commit()
        conn.close()
    if d.get('action') == 'nonp':
        ok, msg = prepare_nonp(job['key'], by=_who())
        if not ok:
            return jsonify({'success': False, 'error': msg}), 400
        return jsonify({'success': True, 'email_id': msg})
    return jsonify({'success': True})


@bp.route('/receivables/api/job_notes')
def api_job_notes():
    if not receivables_allowed():
        return _forbidden()
    key = request.args.get('key', '')
    conn = _conn()
    notes = [dict(r) for r in conn.execute("SELECT * FROM ar_notes WHERE job_key=? OR invoice_id IN "
                                           "(SELECT id FROM ar_invoices WHERE job_key=?) ORDER BY id DESC LIMIT 200",
                                           (key, key))]
    conn.close()
    docs = _docs('job', key) + _docs('nonp', key)
    for x in docs:
        x.pop('stored_name', None)
    return jsonify({'success': True, 'notes': notes, 'docs': docs})


@bp.route('/receivables/api/customer', methods=['POST'])
def api_customer():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    conn = _conn()
    if not conn.execute("SELECT 1 FROM ar_customers WHERE key=?", (d.get('key'),)).fetchone():
        conn.close()
        return jsonify({'success': False, 'error': 'Not found'}), 404
    fields = {}
    for k in ('contact_name', 'emails', 'notes'):
        if k in d:
            fields[k] = str(d[k] or '').strip()[:2000]
    if d.get('kind_default') in ('', 'service', 'install'):
        fields['kind_default'] = d['kind_default']
    if 'do_not_contact' in d:
        fields['do_not_contact'] = 1 if d['do_not_contact'] else 0
    if fields:
        fields['updated_at'] = _stamp()
        conn.execute(f"UPDATE ar_customers SET {', '.join(k + '=?' for k in fields)} WHERE key=?",
                     list(fields.values()) + [d['key']])
        if 'emails' in fields:
            # A fixed address clears "bounced" holds for this customer.
            conn.execute("UPDATE ar_invoices SET status='open', needs_reason='' WHERE customer_key=? AND status='needs_person' "
                         "AND needs_reason LIKE 'Email to % bounced%'", (d['key'],))
            conn.execute("UPDATE ar_emails SET to_email=?, error='' WHERE customer_key=? AND status='draft' AND kind='followup' "
                         "AND job_key=''", (fields['emails'], d['key']))
    conn.commit()
    conn.close()
    if 'kind_default' in fields:
        classify_all()
    return jsonify({'success': True})


@bp.route('/receivables/api/email', methods=['POST'])
def api_email():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    eid = d.get('id')
    conn = _conn()
    e = conn.execute("SELECT status FROM ar_emails WHERE id=?", (eid,)).fetchone()
    if not e or e['status'] == 'sent':
        conn.close()
        return jsonify({'success': False, 'error': 'Not found or already sent'}), 400
    if d.get('discard'):
        conn.execute("UPDATE ar_emails SET status='discarded' WHERE id=?", (eid,))
        conn.commit()
        conn.close()
        return jsonify({'success': True})
    conn.execute("UPDATE ar_emails SET to_email=?, cc=?, subject=?, body=?, manual=1 WHERE id=?",
                 (str(d.get('to_email', '')).strip(), str(d.get('cc', '')).strip(), str(d.get('subject', ''))[:300],
                  str(d.get('body', ''))[:20000], eid))
    conn.commit()
    conn.close()
    if d.get('send'):
        ok, err = send_email(eid, by=_who())
        if not ok:
            return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True})


@bp.route('/receivables/api/email/<int:eid>/attachments')
def api_email_attachments(eid):
    """The files that would go with a draft (names and sizes), for the review screen."""
    if not receivables_allowed():
        return _forbidden()
    conn = _conn()
    e = conn.execute("SELECT * FROM ar_emails WHERE id=?", (eid,)).fetchone()
    if not e:
        conn.close()
        return jsonify({'success': False, 'error': 'Not found'}), 404
    if e['status'] == 'sent':
        conn.close()
        return jsonify({'success': True, 'files': [{'name': n} for n in json.loads(e['attachments'] or '[]')]})
    ids = json.loads(e['invoice_ids'] or '[]')
    invs = [dict(r) for r in conn.execute(f"SELECT * FROM ar_invoices WHERE id IN ({','.join('?' * len(ids))})", ids)] if ids else []
    conn.close()
    files, dropped = [], []
    if e['kind'] == 'nonp':
        files = [{'name': 'Notice of Nonpayment.pdf'}] + [{'name': f"Invoice-{i['number']}.pdf"} for i in invs]
    else:
        kind = 'install' if e['job_key'] else (invs[0]['kind'] if invs else 'service')
        atts, dropped = attachments_for(invs, e['job_key'], e['customer_key'], kind, settings())
        files = [{'name': a[0], 'size': len(a[1])} for a in atts]
    return jsonify({'success': True, 'files': files, 'dropped': dropped})


@bp.route('/receivables/api/settings', methods=['POST'])
def api_settings():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    s = settings()
    for k in ('from_email', 'cc', 'signature', 'retainage_pcts', 'install_words', 'lien_kinds', 'company_name',
              'company_address', 'company_phone', 'send_hours'):
        if k in d:
            s[k] = str(d[k]).strip()[:3000]
    for k in ('auto_send', 'enabled', 'read_replies', 'nonp_auto'):
        if k in d:
            s[k] = bool(d[k])
    for k in ('promise_grace_days', 'reported_paid_wait_days', 'retainage_every_days', 'lien_days', 'nonp_lead_days',
              'max_attach_mb'):
        if k in d:
            try:
                s[k] = max(0, int(d[k]))
            except (TypeError, ValueError):
                return jsonify({'success': False, 'error': f'{k} must be a whole number'}), 400
    for k in ('min_balance', 'retainage_tolerance'):
        if k in d:
            try:
                s[k] = max(0.0, float(d[k]))
            except (TypeError, ValueError):
                return jsonify({'success': False, 'error': f'{k} must be a number'}), 400
    if isinstance(d.get('stages'), list) and d['stages']:
        stages = []
        for st in d['stages']:
            try:
                stages.append({'key': str(st['key']), 'name': str(st.get('name') or st['key'])[:60],
                               'from_days': max(1, int(st['from_days'])), 'every_days': max(1, int(st['every_days']))})
            except (KeyError, TypeError, ValueError):
                return jsonify({'success': False, 'error': 'Each stage needs whole numbers of days'}), 400
        s['stages'] = sorted(stages, key=lambda x: x['from_days'])
    if '@' not in s['from_email']:
        return jsonify({'success': False, 'error': 'The "send from" address must be an email address'}), 400
    _save_settings(s)
    if 'retainage_pcts' in d or 'install_words' in d or 'retainage_tolerance' in d:
        classify_all()
    return jsonify({'success': True})


@bp.route('/receivables/api/run', methods=['POST'])
def api_run():
    if not receivables_allowed():
        return _forbidden()
    threading.Thread(target=run_cycle, kwargs={'manual': True}, daemon=True).start()
    return jsonify({'success': True})


@bp.route('/receivables/api/docs/upload', methods=['POST'])
def api_doc_upload():
    if not receivables_allowed():
        return _forbidden()
    f = request.files.get('file')
    scope = request.form.get('scope', '')
    key = request.form.get('scope_key', '')
    kind = request.form.get('kind', 'document')
    if scope not in ('all', 'customer', 'job', 'invoice') or not f or not f.filename:
        return jsonify({'success': False, 'error': 'Choose a file'}), 400
    if scope == 'all':
        key = ''
    data = f.read()
    if len(data) > 25 * 1024 * 1024:
        return jsonify({'success': False, 'error': 'Files must be under 25 MB'}), 400
    name = secure_filename(f.filename) or 'document'
    stored = f"{uuid.uuid4().hex}_{name}"
    with open(os.path.join(_files_dir(), stored), 'wb') as fh:
        fh.write(data)
    conn = _conn()
    conn.execute('''INSERT INTO ar_documents (scope, scope_key, kind, filename, stored_name, content_type, size, uploaded_by,
                    uploaded_at) VALUES (?,?,?,?,?,?,?,?,?)''',
                 (scope, key, 'invoice_pdf' if kind == 'invoice_pdf' else 'document', f.filename[:200], stored,
                  f.mimetype or 'application/octet-stream', len(data), _who(), _stamp()))
    if scope == 'invoice' and key.isdigit():
        _note(conn, f"Uploaded {f.filename}" + (' (sent in place of the invoice made from Jobber)' if kind == 'invoice_pdf'
                                                else ''), int(key), kind='note', by=_who())
    elif scope == 'job':
        _note(conn, f"Uploaded {f.filename}; it goes with every follow-up on this job.", job_key=key, kind='note', by=_who())
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@bp.route('/receivables/api/docs/delete', methods=['POST'])
def api_doc_delete():
    if not receivables_allowed():
        return _forbidden()
    d = request.get_json(silent=True) or {}
    conn = _conn()
    doc = conn.execute("SELECT * FROM ar_documents WHERE id=?", (d.get('id'),)).fetchone()
    if doc:
        conn.execute("DELETE FROM ar_documents WHERE id=?", (doc['id'],))
        conn.commit()
        try:
            os.remove(os.path.join(_files_dir(), doc['stored_name']))
        except OSError:
            pass
    conn.close()
    return jsonify({'success': True})


@bp.route('/receivables/docs/<int:doc_id>')
def doc_download(doc_id):
    if not receivables_allowed():
        return _deny()
    conn = _conn()
    doc = conn.execute("SELECT * FROM ar_documents WHERE id=?", (doc_id,)).fetchone()
    conn.close()
    if not doc:
        return 'Not found', 404
    return send_file(os.path.join(_files_dir(), doc['stored_name']), download_name=doc['filename'],
                     mimetype=doc['content_type'] or 'application/octet-stream')


@bp.route('/receivables/invoice/<int:inv_id>.pdf')
def invoice_pdf_view(inv_id):
    if not receivables_allowed():
        return _deny()
    conn = _conn()
    inv = conn.execute("SELECT * FROM ar_invoices WHERE id=?", (inv_id,)).fetchone()
    conn.close()
    if not inv:
        return 'Not found', 404
    return send_file(io.BytesIO(invoice_pdf(dict(inv), settings())), mimetype='application/pdf',
                     download_name=f"Invoice-{inv['number']}.pdf")


@bp.route('/receivables/nonp.pdf')
def nonp_pdf_view():
    """Preview the Notice of Nonpayment for a job before preparing it."""
    if not receivables_allowed():
        return _deny()
    job = _job(request.args.get('key', ''))
    if not job:
        return 'Not found', 404
    conn = _conn()
    invs = [dict(r) for r in conn.execute("SELECT * FROM ar_invoices WHERE job_key=? AND status != 'paid'", (job['key'],))]
    conn.close()
    return send_file(io.BytesIO(nonp_pdf(job, invs, settings())), mimetype='application/pdf',
                     download_name='Notice-of-Nonpayment.pdf')


def init_receivables(app, db_path, *, data_dir, get_setting, set_setting, graph_token, graph_enabled,
                     jobber_token, jobber_connected, jobber_version, anthropic_client=None, scheduler_cls=None):
    """Create the tables, register the routes and schedule the cycle every hour.
    jobber_token(force_refresh) hands over an access token from the Jobber app connected for Pumps."""
    _cfg.update(db_path=db_path, data_dir=data_dir, get_setting=get_setting, set_setting=set_setting,
                graph_token=graph_token, graph_enabled=graph_enabled, jobber_token=jobber_token,
                jobber_connected=jobber_connected, jobber_version=jobber_version, anthropic_client=anthropic_client)
    init_db()
    app.register_blueprint(bp)
    if scheduler_cls and os.environ.get('AR_AUTO_RUN', 'true').lower() not in ('0', 'false', 'no', 'off'):
        try:
            sched = scheduler_cls()
            sched.add_job(run_cycle, 'interval', hours=1, id='receivables_cycle',
                          next_run_time=datetime.now() + timedelta(minutes=4))
            sched.start()
            print("✓ Receivables: Jobber, replies and follow-up cycle scheduled every hour")
        except Exception as e:
            print(f"⚠ Could not start the receivables scheduler: {e}")


from receivables_page import PAGE_TEMPLATE  # noqa: E402
