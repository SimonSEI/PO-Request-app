"""
Vendor Invoices: the CURRENT VENDOR INVOICES workbook as an app, for every
non-pump vendor (sub contractor) invoice that comes into PO@Stahlman-England.

  The sheet   One tab per month, same columns, colors, bold/red text and merged
              cells as the workbook (OK TO PROCESS? ... EXPECTED PAYMENT DATE).
              Every row of the workbook is imported once, the first time the app
              starts (vendor_invoices_seed.json); more workbooks can be uploaded,
              and the whole thing downloads back out as an .xlsx.
  The inbox   Invoice PDFs that arrive at PO@ from vendors are listed until
              someone enters them on the sheet. Pump vendors (Wettech) are left
              to the Pumps app. A vendor invoice nobody has entered after a few
              days is emailed to Christian.
  Jobber      A person links a row to the SEI (Jobber) invoice it was billed on.
              When that Jobber invoice is paid, Christian gets an email, once.
              Jobber is only read here, never changed.

Hooked into app.py with init_vendor_invoices(...); tables start with vib_.
"""
import html
import io
import json
import os
import re
import sqlite3
import threading
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import requests as http_requests
from flask import Blueprint, jsonify, redirect, render_template_string, request, send_file, session, url_for

TZ = ZoneInfo('America/New_York')
GRAPH = 'https://graph.microsoft.com/v1.0'
JOBBER_GQL = 'https://api.getjobber.com/api/graphql'
SEED_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'vendor_invoices_seed.json')

VI_ROLES = {'office', 'admin'}
VI_USERS = {u.strip().lower() for u in os.environ.get('VI_USERS', '').split(',') if u.strip()}

bp = Blueprint('vendor_invoices', __name__)
_cfg = {}
_lock = threading.Lock()

# The workbook's columns, A to P, in order. kind: text / date / money / spacer (the thin gray columns).
# width is Excel's column width, so the page lays out like the sheet.
COLS = [
    ('A', 'ok', 'OK TO PROCESS?', 'text', 9.3),
    ('B', 'vendor', 'SUB CONTRACTOR', 'text', 22.4),
    ('C', 'inv_date', 'INVOICE DATE', 'date', 10.7),
    ('D', 'inv_no', 'INVOICE #', 'text', 10.0),
    ('E', 'inv_total', 'INVOICE TOTAL', 'money', 12.7),
    ('F', 'sp_f', '', 'spacer', 1.1),
    ('G', 'proposal', 'SEI PROPOSAL #', 'text', 16.6),
    ('H', 'date_billed', 'DATE BILLED', 'date', 10.7),
    ('I', 'amt_billed', 'AMOUNT SEI BILLED', 'money', 12.4),
    ('J', 'notes', 'NOTES', 'text', 96.1),
    ('K', 'sp_k', '', 'spacer', 3.3),
    ('L', 'pm', 'PROJECT MGR', 'text', 9.9),
    ('M', 'to_christian', 'INVOICE TO CHRISTIAN', 'date', 10.7),
    ('N', 'bill_com', 'BILL.COM', 'date', 11.9),
    ('O', 'sp_o', '', 'spacer', 2.3),
    ('P', 'exp_pay', 'EXPECTED PAYMENT DATE', 'date', 13.6),
]
KEYS = [c[1] for c in COLS]
KIND = {c[1]: c[3] for c in COLS}
SPACERS = [k for k in KEYS if KIND[k] == 'spacer']
HEADERS = {c[1]: c[2] for c in COLS}

# Header text in older tabs -> column. "INVOICE TO JOANNE" became "INVOICE TO CHRISTIAN" in 2024.
_ALIASES = {
    'ok': ('ok to process?', 'ok to process'),
    'vendor': ('sub contractor', 'subcontractor', 'vendor'),
    'inv_date': ('invoice date',),
    'inv_no': ('invoice #', 'invoice#', 'invoice no'),
    'inv_total': ('invoice total', 'cost'),
    'proposal': ('sei proposal #', 'sei proposal', 'proposal #'),
    'date_billed': ('date billed',),
    'amt_billed': ('amount sei billed', 'invoice billed', 'sei billed', 'amount billed'),
    'notes': ('notes',),
    'pm': ('project mgr', 'project manager'),
    'to_christian': ('invoice to christian', 'invoice to joanne'),
    'bill_com': ('bill.com', 'bill com'),
    'exp_pay': ('expected payment date', 'expected payment'),
}

# The sheet's fill colors, offered on the page.
COLORS = [('92D050', 'Green - OK to process'), ('66FFFF', 'Blue - needs a change order / waiting'),
          ('FFFF00', 'Yellow'), ('FF0000', 'Red - not OK / cancelled'), ('99FFCC', 'Mint'), ('FF99FF', 'Pink')]
# Office theme colors, for cells colored with a theme color instead of a plain one.
_THEME = ['FFFFFF', '000000', 'E7E6E6', '44546A', '4472C4', 'ED7D31', 'A5A5A5', 'FFC000', '5B9BD5', '70AD47']

DEFAULT_PUMP_WORDS = 'wettech, wettec, wet tec, wet tech, w.e.t., wttech, wetech, wet-tec'
# Words too common in vendor names to tell two vendors apart.
_GENERIC = set('''LAWN LAWNS SOD LANDSCAPE LANDSCAPES LANDSCAPING SERVICE SERVICES INC LLC CO COMPANY COMPANIES CORP
GROUP ELECTRIC ELECTRICAL WELL WELLS DRILLING DRILL TREE TREES TURF NAPLES GULF SHORE FLORIDA SWFL ASSOCIATES ASSOC
ENGINEERING MAPPING MAP IRRIGATION SUPPLY PUMP PUMPS THE AND NURSERY TRUCKING EXCAVATING PEST ENVIRONMENTAL SOUTHERN
MOTOR WORKS INVOICE OFFICE BILLING ACCOUNTS PAYABLE RECEIVABLE NOREPLY INFO MAIL EMAIL GMAIL OUTLOOK YAHOO HOTMAIL
COM NET ORG BIZ'''.split())


# ── Small helpers ────────────────────────────────────────────────────────────

def _now():
    return datetime.now(TZ)


def _today():
    return _now().date()


def _stamp():
    return _now().strftime('%Y-%m-%d %H:%M:%S')


def _conn():
    conn = sqlite3.connect(_cfg['db_path'], timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def _files_dir():
    d = os.path.join(_cfg.get('data_dir') or '.', 'vendor_invoice_files')
    os.makedirs(d, exist_ok=True)
    return d


def _who():
    return session.get('full_name') or session.get('username') or 'office'


def _log(kind, message):
    conn = _conn()
    conn.execute("INSERT INTO vib_log (at, kind, message) VALUES (?,?,?)", (_stamp(), kind, message))
    conn.commit()
    conn.close()


def _setting(key, default=''):
    conn = _conn()
    r = conn.execute("SELECT value FROM vib_settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return r['value'] if r and r['value'] is not None else default


def _set_setting(key, value):
    conn = _conn()
    conn.execute("INSERT OR REPLACE INTO vib_settings (key, value) VALUES (?,?)", (key, str(value)))
    conn.commit()
    conn.close()


def _int_setting(key, default):
    try:
        return max(0, int(_setting(key, str(default))))
    except ValueError:
        return default


def christian_email():
    get = _cfg.get('get_christian_email')
    try:
        return (get() if get else '') or ''
    except Exception:
        return ''


def graph_ready():
    return bool(_cfg.get('graph_enabled') and _cfg.get('mail_from'))


def jobber_ready():
    try:
        return bool(_cfg['jobber_connected']())
    except Exception:
        return False


def vendor_invoices_allowed():
    """Office (and admin) logins, or only the usernames in VI_USERS."""
    if 'username' not in session:
        return False
    if VI_USERS:
        return session['username'].lower() in VI_USERS
    return session.get('role') in VI_ROLES


def _money(v):
    if v in (None, ''):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace('$', '').replace(',', '').strip()
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()').strip()
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def _num_text(v):
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ('%.6f' % v).rstrip('0').rstrip('.')
    return str(v)


def _parse_date(text):
    """'2026-09-17', '9/17/26', '09-17-2026' -> date, else None."""
    s = (text or '').strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y', '%m-%d-%Y', '%m-%d-%y', '%B %d, %Y', '%b %d, %Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _norm_cell(key, value):
    """What a person typed into a cell -> the stored text. Dates that read as dates are kept as YYYY-MM-DD,
    amounts as plain numbers; anything else (e.g. 'TBD OCT 20') is kept exactly as typed."""
    s = '' if value is None else str(value).strip()
    if not s:
        return ''
    kind = KIND.get(key)
    if kind == 'date':
        d = _parse_date(s)
        return d.isoformat() if d else s
    if kind == 'money':
        n = _money(s)
        return _num_text(n) if n is not None else s
    if key == 'ok':
        return s.upper() if s.lower() in ('yes', 'no') else s
    return s


def _letters(text):
    return re.sub(r'[^A-Z0-9]+', ' ', (text or '').upper()).split()


def vendor_tokens(name):
    """The words that tell vendors apart: PELICAN, SANCHEZ, MONGOVEN, GRABER..."""
    return {w for w in _letters(name) if len(w) >= 4 and w not in _GENERIC and not w.isdigit()}


def _inv_key(text):
    """Invoice numbers compare without spaces, dashes or leading zeros ('INV-00461' == '461')."""
    s = re.sub(r'(?i)^(inv(oice)?\s*(no\.?|#)?)', '', (text or '').strip())
    s = re.sub(r'[^A-Za-z0-9]', '', s).upper().lstrip('0')
    return s


def _month_label(d):
    return d.strftime('%B %Y').upper()


# ── Database ─────────────────────────────────────────────────────────────────

_ROW_COLS = ', '.join(f'{k} TEXT DEFAULT ""' for k in KEYS)


def init_db():
    conn = _conn()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS vib_tabs (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   label TEXT NOT NULL,
                   position REAL NOT NULL,
                   headers TEXT DEFAULT '{}',
                   source TEXT DEFAULT 'manual',
                   created_at TEXT)''')
    c.execute(f'''CREATE TABLE IF NOT EXISTS vib_rows (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   tab_id INTEGER NOT NULL,
                   position REAL NOT NULL,
                   {_ROW_COLS},
                   style TEXT DEFAULT '{{}}',
                   spans TEXT DEFAULT '{{}}',
                   is_total INTEGER DEFAULT 0,
                   jobber_id TEXT DEFAULT '',
                   jobber_number TEXT DEFAULT '',
                   jobber_client TEXT DEFAULT '',
                   jobber_subject TEXT DEFAULT '',
                   jobber_total REAL,
                   jobber_status TEXT DEFAULT '',
                   jobber_url TEXT DEFAULT '',
                   jobber_linked_by TEXT DEFAULT '',
                   jobber_linked_at TEXT DEFAULT '',
                   jobber_checked_at TEXT DEFAULT '',
                   jobber_paid_at TEXT DEFAULT '',
                   paid_email_at TEXT DEFAULT '',
                   source TEXT DEFAULT 'manual',
                   created_at TEXT,
                   updated_at TEXT,
                   updated_by TEXT DEFAULT '')''')
    c.execute("CREATE INDEX IF NOT EXISTS vib_rows_tab ON vib_rows (tab_id, position)")
    # Invoice PDFs that came into PO@ from vendors.
    # status: received (a vendor invoice to enter) / other (a PDF from someone we don't know yet)
    #         / inputted (on the sheet) / ignored (not a vendor invoice)
    c.execute('''CREATE TABLE IF NOT EXISTS vib_received (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   email_uid TEXT,
                   filename TEXT,
                   file_path TEXT DEFAULT '',
                   sender TEXT DEFAULT '',
                   subject TEXT DEFAULT '',
                   received_at TEXT DEFAULT '',
                   vendor TEXT DEFAULT '',
                   inv_no TEXT DEFAULT '',
                   inv_date TEXT DEFAULT '',
                   inv_total REAL,
                   proposal TEXT DEFAULT '',
                   summary TEXT DEFAULT '',
                   status TEXT DEFAULT 'received',
                   row_id INTEGER,
                   alerted_at TEXT DEFAULT '',
                   alert_count INTEGER DEFAULT 0,
                   decided_by TEXT DEFAULT '',
                   decided_at TEXT DEFAULT '',
                   created_at TEXT)''')
    # Senders a person has told us about: a vendor's address, or one to ignore.
    c.execute('''CREATE TABLE IF NOT EXISTS vib_senders (
                   address TEXT PRIMARY KEY,
                   kind TEXT NOT NULL,
                   vendor TEXT DEFAULT '')''')
    # Every PO@ email the scanner has looked at (shared_mail.fetch_scan_emails reads email_uid).
    c.execute('''CREATE TABLE IF NOT EXISTS vib_email_log (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   email_uid TEXT UNIQUE,
                   email_sender TEXT,
                   email_subject TEXT,
                   email_date TEXT,
                   scanned_at TEXT,
                   invoices_found INTEGER DEFAULT 0,
                   results TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS vib_settings (key TEXT PRIMARY KEY, value TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS vib_log (
                   id INTEGER PRIMARY KEY AUTOINCREMENT, at TEXT, kind TEXT, message TEXT)''')
    conn.commit()
    if c.execute("SELECT value FROM vib_settings WHERE key='scan_since'").fetchone() is None:
        # Only mail from the day the app first ran: older invoices are already on the workbook.
        c.execute("INSERT INTO vib_settings (key, value) VALUES ('scan_since', ?)", (_today().isoformat(),))
        conn.commit()
    conn.close()


# ── 1. The workbook ──────────────────────────────────────────────────────────

def _rgb(color):
    """An openpyxl color -> 'RRGGBB', or '' for none / automatic / white."""
    if color is None:
        return ''
    try:
        if color.type == 'rgb' and isinstance(color.rgb, str):
            v = color.rgb[-6:].upper()
        elif color.type == 'theme' and color.theme is not None and color.theme < len(_THEME):
            base = _THEME[color.theme]
            tint = float(color.tint or 0)
            ch = [int(base[i:i + 2], 16) for i in (0, 2, 4)]
            ch = [round(x + (255 - x) * tint) if tint > 0 else round(x * (1 + tint)) for x in ch]
            v = ''.join('%02X' % max(0, min(255, x)) for x in ch)
        else:
            return ''
    except Exception:
        return ''
    return '' if v in ('FFFFFF', '000000') else v


def _cell_style(cell):
    st = {}
    try:
        if cell.fill is not None and cell.fill.fill_type == 'solid':
            bg = _rgb(cell.fill.fgColor)
            if bg:
                st['bg'] = bg
        if cell.font is not None:
            fg = _rgb(cell.font.color)
            if fg:
                st['fg'] = fg
            if cell.font.b:
                st['b'] = 1
        fmt = cell.number_format or ''
        if 'mmm' in fmt and 'yy' not in fmt:
            st['fmt'] = 'd-mmm'
        elif '$' in fmt and not isinstance(cell.value, str):
            st['fmt'] = 'money'
    except Exception:
        pass
    return st


def _header_map(header_cells):
    """Workbook column index -> our column key, and the tab's own header labels."""
    colmap, labels, blanks = {}, {}, []
    for i, v in enumerate(header_cells):
        text = re.sub(r'\s+', ' ', str(v or '')).strip()
        low = text.lower()
        key = next((k for k, al in _ALIASES.items() if low in al and k not in labels), None)
        if key:
            colmap[i] = key
            labels[key] = text
        elif not text:
            blanks.append(i)
    # Untitled columns are the thin spacer columns (F, K, O), in order; any more are left out.
    for i, key in zip(blanks, SPACERS):
        colmap[i] = key
    return colmap, labels


def parse_workbook(data):
    """The vendor invoices workbook (bytes) -> [{'label', 'headers', 'rows': [{key: text, 'style', 'spans',
    'is_total'}]}], in the workbook's order. The MASTER COPY BLANK template tab is left out."""
    import openpyxl
    wb_f = openpyxl.load_workbook(io.BytesIO(data))                  # formulas, styles and merges
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)  # the values Excel last showed
    tabs = []
    for ws in wb_f.worksheets:
        title = ws.title.strip()
        if 'MASTER' in title.upper() and 'BLANK' in title.upper():
            continue
        wv = wb_v[ws.title]
        header = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        colmap, labels = _header_map(header)
        if len(labels) < 3:
            continue  # not a vendor invoice tab
        order = {k: i for i, k in enumerate(KEYS)}
        # Merged cells below the header: anchor -> (rows, cols); covered cells are left blank.
        spans = {}
        for mr in ws.merged_cells.ranges:
            if mr.min_row < 2:
                continue
            k1, k2 = colmap.get(mr.min_col - 1), colmap.get(mr.max_col - 1)
            if not k1:
                continue
            cols = (order[k2] - order[k1] + 1) if k2 and order[k2] >= order[k1] else 1
            rows = mr.max_row - mr.min_row + 1
            if rows > 1 or cols > 1:
                spans[(mr.min_row, k1)] = [rows, cols]
        rows, last = [], 1
        for r in range(2, ws.max_row + 1):
            row = {k: '' for k in KEYS}
            style, has, is_total = {}, False, 0
            for ci, key in colmap.items():
                fcell = ws.cell(row=r, column=ci + 1)
                v = wv.cell(row=r, column=ci + 1).value
                if isinstance(fcell.value, str) and fcell.value.startswith('=SUM'):
                    is_total = 1
                if v is None:
                    text = ''
                elif isinstance(v, datetime):
                    text = v.date().isoformat()
                elif isinstance(v, date):
                    text = v.isoformat()
                elif isinstance(v, (int, float)):
                    text = _num_text(v)
                else:
                    text = str(v).strip()
                if text:
                    has = True
                row[key] = text
                st = _cell_style(fcell)
                if not text:
                    st.pop('fmt', None)
                    st.pop('b', None)
                    st.pop('fg', None)
                if st:
                    style[key] = st
            sp = {k: v for (rr, k), v in spans.items() if rr == r}
            if sp:
                has = True
            row['style'], row['spans'], row['is_total'] = style, sp, is_total
            rows.append(row)
            if has or any(rr >= r and rr - 1 + v[0] >= r for (rr, _), v in spans.items() if rr <= r):
                last = r
        rows = rows[:max(0, last - 1)]
        tabs.append({'label': title, 'headers': labels, 'rows': rows})
    return tabs


def import_tabs(tabs, by='workbook', replace=False):
    """Add parsed tabs. A tab already here (same name) is skipped, or replaced when replace=True.
    New tabs keep the workbook's order and go after the tabs already here."""
    now = _stamp()
    added, replaced, skipped, nrows = [], [], [], 0
    with _lock:
        conn = _conn()
        try:
            conn.execute('BEGIN IMMEDIATE')
            existing = {r['label'].upper(): r['id'] for r in conn.execute("SELECT id, label FROM vib_tabs")}
            pos = (conn.execute("SELECT MAX(position) FROM vib_tabs").fetchone()[0] or 0) + 1
            for t in tabs:
                tid = existing.get(t['label'].upper())
                if tid and not replace:
                    skipped.append(t['label'])
                    continue
                if tid:
                    conn.execute("DELETE FROM vib_rows WHERE tab_id=? AND source='workbook'", (tid,))
                    conn.execute("UPDATE vib_tabs SET headers=? WHERE id=?", (json.dumps(t['headers']), tid))
                    first = conn.execute("SELECT MIN(position) FROM vib_rows WHERE tab_id=?", (tid,)).fetchone()[0]
                    start = (first if first is not None else 0) - len(t['rows']) - 1
                    replaced.append(t['label'])
                else:
                    tid = conn.execute("INSERT INTO vib_tabs (label, position, headers, source, created_at) "
                                       "VALUES (?,?,?,?,?)",
                                       (t['label'], pos, json.dumps(t['headers']), 'workbook', now)).lastrowid
                    pos += 1
                    start = 0
                    added.append(t['label'])
                for i, r in enumerate(t['rows']):
                    conn.execute(f"INSERT INTO vib_rows (tab_id, position, {', '.join(KEYS)}, style, spans, is_total, "
                                 f"source, created_at, updated_at, updated_by) VALUES "
                                 f"(?,?,{','.join('?' * len(KEYS))},?,?,?,?,?,?,?)",
                                 [tid, start + i + 1] + [r.get(k, '') or '' for k in KEYS] +
                                 [json.dumps(r.get('style') or {}), json.dumps(r.get('spans') or {}),
                                  int(r.get('is_total') or 0), 'workbook', now, now, by])
                    nrows += 1
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
    if added or replaced:
        _log('import', f'{by}: {nrows} rows in {len(added) + len(replaced)} tab(s)'
                       + (f', {len(skipped)} tab(s) already here' if skipped else '') + '.')
    return {'added': added, 'replaced': replaced, 'skipped': skipped, 'rows': nrows}


def seed_from_workbook():
    """The first time the app runs, bring in every row of the CURRENT VENDOR INVOICES workbook."""
    if _setting('seeded') or not os.path.exists(_cfg.get('seed_file') or ''):
        return None
    conn = _conn()
    have = conn.execute("SELECT COUNT(*) FROM vib_tabs").fetchone()[0]
    conn.close()
    if have:
        _set_setting('seeded', _stamp())
        return None
    with open(_cfg['seed_file'], encoding='utf-8') as fh:
        tabs = json.load(fh)['tabs']
    out = import_tabs(tabs, by='CURRENT VENDOR INVOICES workbook')
    _set_setting('seeded', _stamp())
    print(f"✓ Vendor Invoices: imported {out['rows']} rows in {len(out['added'])} month tabs from the workbook")
    return out


def export_workbook():
    """The whole sheet back out as an .xlsx laid out like the original workbook."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    conn = _conn()
    tabs = conn.execute("SELECT * FROM vib_tabs ORDER BY position").fetchall()
    money_fmt = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
    used = set()
    for t in tabs:
        name = re.sub(r'[\[\]\*\?/\\:]', ' ', t['label'])[:31] or 'Sheet'
        while name.upper() in used:
            name = (name[:28] + ' ' + str(len(used)))[:31]
        used.add(name.upper())
        ws = wb.create_sheet(name)
        heads = json.loads(t['headers'] or '{}') if t['source'] == 'workbook' else HEADERS
        for i, (letter, key, _, kind, width) in enumerate(COLS, start=1):
            ws.column_dimensions[letter].width = width
            c = ws.cell(row=1, column=i, value=heads.get(key) or None)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.fill = PatternFill('solid', fgColor='FFFF00')
        ws.cell(row=1, column=len(COLS) + 1, value='JOBBER INVOICE').font = Font(bold=True)
        ws.column_dimensions[get_column_letter(len(COLS) + 1)].width = 16
        rows = conn.execute("SELECT * FROM vib_rows WHERE tab_id=? ORDER BY position", (t['id'],)).fetchall()
        for ri, r in enumerate(rows, start=2):
            style = json.loads(r['style'] or '{}')
            for ci, (letter, key, _, kind, _) in enumerate(COLS, start=1):
                text = r[key] or ''
                st = style.get(key) or {}
                val, fmt = text, None
                if text:
                    if kind == 'date' or st.get('fmt') == 'd-mmm':
                        d = _parse_date(text) if re.match(r'^\d{4}-\d{2}-\d{2}$', text) else None
                        if d:
                            val, fmt = datetime(d.year, d.month, d.day), ('d-mmm' if st.get('fmt') == 'd-mmm'
                                                                           else 'mm-dd-yy')
                    elif kind == 'money' or st.get('fmt') == 'money':
                        n = _money(text) if re.match(r'^-?\d+(\.\d+)?$', text) else None
                        if n is not None:
                            val, fmt = n, money_fmt
                    elif re.match(r'^-?\d{1,15}$', text) and not text.startswith('0'):
                        val = int(text)
                cell = ws.cell(row=ri, column=ci, value=val if val != '' else None)
                if fmt:
                    cell.number_format = fmt
                if st.get('bg'):
                    cell.fill = PatternFill('solid', fgColor=st['bg'])
                if st.get('fg') or st.get('b'):
                    cell.font = Font(bold=bool(st.get('b')), color=st.get('fg') or None)
            if r['jobber_number']:
                ws.cell(row=ri, column=len(COLS) + 1,
                        value=f"{r['jobber_number']}{' - PAID ' + r['jobber_paid_at'][:10] if r['jobber_paid_at'] else ''}")
            for key, (nr, nc) in (json.loads(r['spans'] or '{}')).items():
                if key in KEYS:
                    ci = KEYS.index(key) + 1
                    ws.merge_cells(start_row=ri, start_column=ci, end_row=ri + nr - 1,
                                   end_column=min(len(COLS), ci + nc - 1))
        ws.freeze_panes = 'A2'
    conn.close()
    if not wb.worksheets:
        wb.create_sheet('VENDOR INVOICES')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── 2. The sheet ─────────────────────────────────────────────────────────────

def _row_out(r):
    d = {k: r[k] or '' for k in KEYS}
    d.update(id=r['id'], tab_id=r['tab_id'], style=json.loads(r['style'] or '{}'),
             spans=json.loads(r['spans'] or '{}'), is_total=r['is_total'],
             jobber={'id': r['jobber_id'], 'number': r['jobber_number'], 'client': r['jobber_client'],
                     'subject': r['jobber_subject'], 'total': r['jobber_total'], 'status': r['jobber_status'],
                     'url': r['jobber_url'], 'linked_by': r['jobber_linked_by'], 'linked_at': r['jobber_linked_at'],
                     'paid_at': r['jobber_paid_at'], 'paid_email_at': r['paid_email_at']} if r['jobber_id'] else None)
    return d


def tabs_summary():
    conn = _conn()
    # A tab from the workbook shows its own header row (e.g. INVOICE TO JOANNE, or no OK TO PROCESS? in 2020).
    out = [dict(id=t['id'], label=t['label'], rows=t['n'],
                headers=json.loads(t['headers'] or '{}') if t['source'] == 'workbook' else dict(HEADERS))
           for t in conn.execute('''SELECT t.*, (SELECT COUNT(*) FROM vib_rows r WHERE r.tab_id=t.id) AS n
                                    FROM vib_tabs t ORDER BY t.position''')]
    conn.close()
    return out


def tab_rows(tab_id):
    conn = _conn()
    rows = [_row_out(r) for r in conn.execute("SELECT * FROM vib_rows WHERE tab_id=? ORDER BY position, id",
                                              (tab_id,))]
    conn.close()
    return rows


def add_tab(label=''):
    """A new month tab at the front, like copying MASTER COPY BLANK."""
    label = (label or '').strip().upper() or _month_label(_today())
    conn = _conn()
    have = conn.execute("SELECT id FROM vib_tabs WHERE UPPER(label)=?", (label,)).fetchone()
    if have:
        conn.close()
        return have['id'], False
    first = conn.execute("SELECT MIN(position) FROM vib_tabs").fetchone()[0]
    tid = conn.execute("INSERT INTO vib_tabs (label, position, headers, source, created_at) VALUES (?,?,?,?,?)",
                       (label, (first if first is not None else 1) - 1, '{}', 'manual', _stamp())).lastrowid
    conn.commit()
    conn.close()
    return tid, True


def current_tab_id():
    """This month's tab (made if it isn't there yet)."""
    return add_tab(_month_label(_today()))[0]


def add_row(tab_id, after_id=None, values=None, by='', source='manual'):
    conn = _conn()
    if not conn.execute("SELECT 1 FROM vib_tabs WHERE id=?", (tab_id,)).fetchone():
        conn.close()
        raise ValueError('No such tab')
    pos = None
    if after_id:
        a = conn.execute("SELECT position FROM vib_rows WHERE id=? AND tab_id=?", (after_id, tab_id)).fetchone()
        if a:
            nxt = conn.execute("SELECT MIN(position) FROM vib_rows WHERE tab_id=? AND position>?",
                               (tab_id, a['position'])).fetchone()[0]
            pos = (a['position'] + nxt) / 2 if nxt is not None else a['position'] + 1
    if pos is None:
        # New rows go under the last row that has something in it (above any totals row).
        filled = ' OR '.join(f"{k}<>''" for k in KEYS)
        last = conn.execute(f"SELECT MAX(position) FROM vib_rows WHERE tab_id=? AND is_total=0 AND ({filled})",
                            (tab_id,)).fetchone()[0]
        if last is None:
            first = conn.execute("SELECT MIN(position) FROM vib_rows WHERE tab_id=?", (tab_id,)).fetchone()[0]
            pos = (first - 1) if first is not None else 1
        else:
            nxt = conn.execute("SELECT MIN(position) FROM vib_rows WHERE tab_id=? AND position>?",
                               (tab_id, last)).fetchone()[0]
            pos = (last + nxt) / 2 if nxt is not None else last + 1
    vals = {k: _norm_cell(k, (values or {}).get(k)) for k in KEYS}
    style = {}
    if vals['ok'] in ('YES', 'NO'):
        style['ok'] = {'bg': '92D050' if vals['ok'] == 'YES' else 'FF0000'}
    now = _stamp()
    rid = conn.execute(f"INSERT INTO vib_rows (tab_id, position, {', '.join(KEYS)}, style, source, created_at, "
                       f"updated_at, updated_by) VALUES (?,?,{','.join('?' * len(KEYS))},?,?,?,?,?)",
                       [tab_id, pos] + [vals[k] for k in KEYS] + [json.dumps(style), source, now, now, by]).lastrowid
    conn.commit()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (rid,)).fetchone()
    conn.close()
    if vals['inv_no']:
        match_received(rid)
    return _row_out(r)


def set_cell(row_id, key, value, by=''):
    if key not in KEYS:
        raise ValueError('Unknown column')
    text = _norm_cell(key, value)
    conn = _conn()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    if not r:
        conn.close()
        raise ValueError('No such row')
    style = json.loads(r['style'] or '{}')
    if key == 'ok' and text in ('YES', 'NO') and text != (r['ok'] or '').upper():
        # The sheet's habit: YES is green, NO is red.
        style.setdefault('ok', {})['bg'] = '92D050' if text == 'YES' else 'FF0000'
    conn.execute(f"UPDATE vib_rows SET {key}=?, style=?, updated_at=?, updated_by=? WHERE id=?",
                 (text, json.dumps(style), _stamp(), by, row_id))
    conn.commit()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    conn.close()
    if key in ('inv_no', 'vendor'):
        match_received(row_id)
    return _row_out(r)


def set_style(row_id, keys, bg=None, fg=None, bold=None, by=''):
    conn = _conn()
    r = conn.execute("SELECT style FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    if not r:
        conn.close()
        raise ValueError('No such row')
    style = json.loads(r['style'] or '{}')
    for key in keys:
        if key not in KEYS:
            continue
        st = style.get(key) or {}
        for attr, val in (('bg', bg), ('fg', fg)):
            if val is not None:
                val = re.sub(r'[^0-9A-Fa-f]', '', val).upper()[:6]
                if len(val) == 6:
                    st[attr] = val
                else:
                    st.pop(attr, None)
        if bold is not None:
            if bold:
                st['b'] = 1
            else:
                st.pop('b', None)
        if st:
            style[key] = st
        else:
            style.pop(key, None)
    conn.execute("UPDATE vib_rows SET style=?, updated_at=?, updated_by=? WHERE id=?",
                 (json.dumps(style), _stamp(), by, row_id))
    conn.commit()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    conn.close()
    return _row_out(r)


def delete_row(row_id, by=''):
    conn = _conn()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    if not r:
        conn.close()
        return False
    # A merged cell that reached into this row now stops one row sooner.
    above = conn.execute("SELECT id, spans, position FROM vib_rows WHERE tab_id=? AND position<? AND spans<>'{}' "
                         "ORDER BY position", (r['tab_id'], r['position'])).fetchall()
    order = [x['id'] for x in conn.execute("SELECT id FROM vib_rows WHERE tab_id=? ORDER BY position, id",
                                           (r['tab_id'],))]
    idx = order.index(row_id)
    for a in above:
        sp = json.loads(a['spans'] or '{}')
        ai = order.index(a['id'])
        changed = False
        for k, (nr, nc) in list(sp.items()):
            if ai + nr - 1 >= idx:
                sp[k] = [nr - 1, nc]
                changed = True
                if nr - 1 <= 1 and nc <= 1:
                    sp.pop(k)
        if changed:
            conn.execute("UPDATE vib_rows SET spans=? WHERE id=?", (json.dumps(sp), a['id']))
    conn.execute("UPDATE vib_received SET status='received', row_id=NULL WHERE row_id=?", (row_id,))
    conn.execute("DELETE FROM vib_rows WHERE id=?", (row_id,))
    conn.commit()
    conn.close()
    _log('row', f"{by} deleted a row from the sheet ({r['vendor'] or '-'} #{r['inv_no'] or '-'}).")
    return True


def search_rows(q, limit=300):
    words = [w for w in re.split(r'\s+', (q or '').strip()) if w]
    if not words:
        return []
    where, args = [], []
    for w in words:
        where.append('(' + ' OR '.join(f'r.{k} LIKE ?' for k in KEYS if KIND[k] != 'spacer')
                     + ' OR r.jobber_number LIKE ?)')
        args += [f'%{w}%'] * (len([k for k in KEYS if KIND[k] != 'spacer']) + 1)
    conn = _conn()
    rows = conn.execute(f'''SELECT r.*, t.label AS tab_label FROM vib_rows r JOIN vib_tabs t ON t.id=r.tab_id
                            WHERE {' AND '.join(where)} ORDER BY t.position, r.position LIMIT ?''',
                        args + [limit]).fetchall()
    conn.close()
    out = []
    for r in rows:
        d = _row_out(r)
        d['tab_label'] = r['tab_label']
        out.append(d)
    return out


# ── 3. The PO@ inbox ─────────────────────────────────────────────────────────

def _pump_words():
    return [w.strip().lower() for w in _setting('pump_words', DEFAULT_PUMP_WORDS).split(',') if w.strip()]


def is_pump(*texts):
    """Pump work (Wettech) belongs to the Pumps app, not this sheet."""
    blob = ' '.join(t or '' for t in texts).lower()
    squashed = re.sub(r'[^a-z]', '', blob)
    for w in _pump_words():
        if w in blob or (len(re.sub(r'[^a-z]', '', w)) >= 5 and re.sub(r'[^a-z]', '', w) in squashed):
            return True
    return False


def _address(sender):
    m = re.search(r'[\w.+-]+@[\w.-]+', sender or '')
    return m.group(0).lower() if m else (sender or '').strip().lower()


def known_vendors():
    """Every vendor on the sheet: {distinctive word: a name as written on the sheet}."""
    conn = _conn()
    names = [r[0] for r in conn.execute("SELECT vendor, COUNT(*) n FROM vib_rows WHERE vendor<>'' "
                                        "GROUP BY UPPER(vendor) ORDER BY n DESC")]
    conn.close()
    out = {}
    for n in names:
        for w in vendor_tokens(n):
            out.setdefault(w, n)
    return out


def match_vendor(name, sender=''):
    """The vendor on the sheet this invoice is from, or ''."""
    addr = _address(sender)
    conn = _conn()
    s = conn.execute("SELECT kind, vendor FROM vib_senders WHERE address=?", (addr,)).fetchone()
    conn.close()
    if s and s['kind'] == 'vendor':
        return s['vendor'] or name or addr
    known = known_vendors()
    for w in vendor_tokens(name):
        if w in known:
            return known[w]
    domain = addr.split('@')[-1].split('.')[0] if '@' in addr else ''
    for w in vendor_tokens(domain) | {domain.upper()}:
        for k, v in known.items():
            if len(w) >= 5 and (w.startswith(k) or k.startswith(w)):
                return v
    return ''


def _sender_kind(sender):
    conn = _conn()
    s = conn.execute("SELECT kind FROM vib_senders WHERE address=?", (_address(sender),)).fetchone()
    conn.close()
    return s['kind'] if s else ''


def match_received(row_id):
    """A row typed on the sheet by hand counts as entering the PO@ invoice with the same number."""
    conn = _conn()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    if not r or not _inv_key(r['inv_no']):
        conn.close()
        return None
    key, toks = _inv_key(r['inv_no']), vendor_tokens(r['vendor'])
    hit = None
    for x in conn.execute("SELECT * FROM vib_received WHERE status IN ('received','other') AND row_id IS NULL"):
        if _inv_key(x['inv_no']) != key:
            continue
        xt = vendor_tokens(x['vendor']) | vendor_tokens(x['sender'].split('@')[-1] if x['sender'] else '')
        if not toks or not xt or toks & xt or any(a.startswith(b) or b.startswith(a) for a in toks for b in xt):
            hit = x
            break
    if hit:
        conn.execute("UPDATE vib_received SET status='inputted', row_id=?, decided_at=?, decided_by=? WHERE id=?",
                     (row_id, _stamp(), r['updated_by'] or 'sheet', hit['id']))
        conn.commit()
    conn.close()
    return hit['id'] if hit else None


def _pdf_text(data):
    import pdfplumber
    text = ''
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:6]:
            text += (page.extract_text() or '') + '\n'
    return text


def add_received(uid, filename, data, sender, subject, when, fields=None, text=''):
    """File one PDF from PO@. Returns (id, status) or (None, why it was skipped)."""
    fields = fields or {}
    vendor = (fields.get('sub_contractor') or '').strip()
    if is_pump(vendor, sender, subject, filename, text[:3000]):
        return None, 'pump'
    kind = _sender_kind(sender)
    if kind == 'ignore':
        return None, 'ignored sender'
    sheet_vendor = match_vendor(vendor, sender)
    status = 'received' if sheet_vendor else 'other'
    inv_no = str(fields.get('invoice_number') or '').strip()
    conn = _conn()
    if inv_no:
        for x in conn.execute("SELECT id, vendor, sender, inv_no FROM vib_received"):
            if _inv_key(x['inv_no']) == _inv_key(inv_no) and (
                    _address(x['sender']) == _address(sender) or vendor_tokens(x['vendor']) & vendor_tokens(vendor)):
                conn.close()
                return None, 'already in the inbox'
    total = fields.get('invoice_total')
    try:
        total = float(total) if total not in (None, '') else None
    except (TypeError, ValueError):
        total = _money(total)
    rid = conn.execute('''INSERT INTO vib_received (email_uid, filename, sender, subject, received_at, vendor, inv_no,
                            inv_date, inv_total, proposal, summary, status, created_at)
                          VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                       (str(uid), filename, sender, subject, _when_text(when), sheet_vendor or vendor, inv_no,
                        str(fields.get('invoice_date') or ''), total, str(fields.get('sei_proposal_number') or ''),
                        str(fields.get('notes') or '')[:500], status, _stamp())).lastrowid
    path = os.path.join(_files_dir(), f'{rid}.pdf')
    with open(path, 'wb') as fh:
        fh.write(data)
    conn.execute("UPDATE vib_received SET file_path=? WHERE id=?", (path, rid))
    conn.commit()
    # Already typed on the sheet by hand? Then it's entered.
    if inv_no:
        for r in conn.execute("SELECT id, vendor, inv_no FROM vib_rows WHERE inv_no<>''"):
            if _inv_key(r['inv_no']) == _inv_key(inv_no) and (
                    not vendor_tokens(r['vendor']) or vendor_tokens(r['vendor']) &
                    (vendor_tokens(sheet_vendor) | vendor_tokens(vendor))):
                conn.execute("UPDATE vib_received SET status='inputted', row_id=?, decided_by='sheet', decided_at=? "
                             "WHERE id=?", (r['id'], _stamp(), rid))
                conn.commit()
                status = 'inputted'
                break
    conn.close()
    return rid, status


def _when_text(when):
    """Graph ('2026-10-14T13:05:00Z') or IMAP ('Tue, 14 Oct 2026 09:05:00 -0400') -> local 'YYYY-MM-DD HH:MM'."""
    import email.utils
    s = (when or '').strip()
    try:
        if 'T' in s and s[:4].isdigit():
            d = datetime.fromisoformat(s.replace('Z', '+00:00'))
        else:
            d = email.utils.parsedate_to_datetime(s)
        if d.tzinfo:
            d = d.astimezone(TZ)
        return d.strftime('%Y-%m-%d %H:%M')
    except Exception:
        return _now().strftime('%Y-%m-%d %H:%M')


def scan_inbox():
    """Read new mail in PO@ for vendor invoice PDFs (each email once, vib_email_log)."""
    summary = {'emails_checked': 0, 'added': 0, 'other': 0, 'pump': 0, 'skipped': 0, 'errors': []}
    fetch = _cfg.get('fetch_emails')
    if not fetch:
        summary['errors'].append('The PO@ mailbox is not set up')
        return summary
    try:
        fetched = fetch(log_table='vib_email_log', extensions=('.pdf',), since=_setting('scan_since') or None)
    except Exception as e:
        summary['errors'].append(f'Mailbox: {e}')
        return summary
    diag = fetched.get('diagnostics') or {}
    if diag.get('error'):
        summary['errors'].append(diag['error'])
    source = fetched.get('source')
    for uid, msg in fetched.get('emails', []):
        summary['emails_checked'] += 1
        found = 0
        sender = subject = when = ''
        try:
            if source == 'graph_api':
                sender = ((msg.get('from') or {}).get('emailAddress') or {}).get('address', '')
                subject, when = msg.get('subject', ''), msg.get('receivedDateTime', '')
                attachments = _cfg['graph_attachments'](uid)
            else:
                sender, subject, when = msg.get('From', ''), msg.get('Subject', ''), msg.get('Date', '')
                attachments = _cfg['email_attachments'](msg)
            for filename, data in attachments:
                if not (filename or '').lower().endswith('.pdf') or not data:
                    continue
                if is_pump(sender, subject, filename):
                    summary['pump'] += 1
                    continue
                try:
                    text = _pdf_text(data)
                except Exception as e:
                    text = ''
                    summary['errors'].append(f'{filename}: could not read the PDF ({e})')
                fields = {}
                if text.strip() and _cfg.get('extract_fields'):
                    try:
                        fields = _cfg['extract_fields'](text, sender) or {}
                    except Exception as e:
                        summary['errors'].append(f'{filename}: {e}')
                rid, status = add_received(uid, filename, data, sender, subject, when, fields, text)
                if rid is None:
                    summary['pump' if status == 'pump' else 'skipped'] += 1
                    continue
                found += 1
                summary['added' if status in ('received', 'inputted') else 'other'] += 1
        except Exception as e:
            summary['errors'].append(f'Email {uid}: {e}')
        conn = _conn()
        conn.execute('''INSERT OR IGNORE INTO vib_email_log (email_uid, email_sender, email_subject, email_date,
                          scanned_at, invoices_found, results) VALUES (?,?,?,?,?,?,?)''',
                     (str(uid), sender, subject, when, _stamp(), found, json.dumps({'found': found})))
        conn.commit()
        conn.close()
    _set_setting('last_scan', json.dumps({'at': _stamp(), **{k: v for k, v in summary.items() if k != 'errors'},
                                          'errors': summary['errors'][:5]}))
    if summary['added'] or summary['other']:
        _log('scan', f"PO@: {summary['added']} vendor invoice(s) and {summary['other']} other PDF(s) came in.")
    return summary


def inbox():
    conn = _conn()
    rows = [dict(r) for r in conn.execute('''SELECT x.*, r.tab_id AS row_tab FROM vib_received x
                                             LEFT JOIN vib_rows r ON r.id=x.row_id
                                             WHERE x.status IN ('received','other')
                                                OR x.decided_at >= ? ORDER BY x.received_at DESC, x.id DESC''',
                                          ((_now() - timedelta(days=14)).strftime('%Y-%m-%d'),))]
    conn.close()
    days = _int_setting('alert_days', 3)
    today = _today()
    for r in rows:
        r.pop('file_path', None)
        try:
            got = datetime.strptime(r['received_at'][:10], '%Y-%m-%d').date()
            r['days_waiting'] = (today - got).days
        except ValueError:
            r['days_waiting'] = 0
        r['overdue'] = r['status'] == 'received' and r['days_waiting'] >= days
    return rows


def enter_received(rec_id, tab_id=None, by=''):
    """'Enter on the sheet': a new row on the open tab, filled in from the PDF."""
    conn = _conn()
    x = conn.execute("SELECT * FROM vib_received WHERE id=?", (rec_id,)).fetchone()
    conn.close()
    if not x:
        raise ValueError('No such invoice')
    if x['row_id']:
        raise ValueError('Already on the sheet')
    row = add_row(tab_id or current_tab_id(), values={
        'ok': 'YES', 'vendor': (x['vendor'] or '').upper(), 'inv_date': x['inv_date'], 'inv_no': x['inv_no'],
        'inv_total': x['inv_total'], 'proposal': x['proposal'], 'notes': (x['summary'] or '').upper()},
        by=by, source='inbox')
    conn = _conn()
    conn.execute("UPDATE vib_received SET status='inputted', row_id=?, decided_by=?, decided_at=? WHERE id=?",
                 (row['id'], by, _stamp(), rec_id))
    conn.commit()
    conn.close()
    _log('inbox', f"{by} entered {x['vendor'] or 'a vendor'} invoice #{x['inv_no'] or '-'} on the sheet.")
    return row


def decide_received(rec_id, action, by=''):
    """vendor: it is a vendor invoice (and this sender is a vendor) / ignore: not a vendor invoice
    (and this sender never is) / entered: it's on the sheet already / undo: back to waiting."""
    conn = _conn()
    x = conn.execute("SELECT * FROM vib_received WHERE id=?", (rec_id,)).fetchone()
    if not x:
        conn.close()
        raise ValueError('No such invoice')
    addr = _address(x['sender'])
    if action == 'vendor':
        conn.execute("UPDATE vib_received SET status='received' WHERE id=?", (rec_id,))
        if addr:
            conn.execute("INSERT OR REPLACE INTO vib_senders (address, kind, vendor) VALUES (?,?,?)",
                         (addr, 'vendor', x['vendor']))
    elif action == 'ignore':
        conn.execute("UPDATE vib_received SET status='ignored', decided_by=?, decided_at=? WHERE id=?",
                     (by, _stamp(), rec_id))
        if addr and x['status'] == 'other':
            conn.execute("INSERT OR REPLACE INTO vib_senders (address, kind, vendor) VALUES (?,?,?)",
                         (addr, 'ignore', ''))
    elif action == 'entered':
        conn.execute("UPDATE vib_received SET status='inputted', decided_by=?, decided_at=? WHERE id=?",
                     (by, _stamp(), rec_id))
    elif action == 'undo':
        conn.execute("UPDATE vib_received SET status=CASE WHEN status='ignored' AND ? IN "
                     "(SELECT address FROM vib_senders WHERE kind='ignore') THEN 'other' ELSE 'received' END, "
                     "row_id=NULL, decided_by='', decided_at='' WHERE id=?", (addr, rec_id))
        conn.execute("DELETE FROM vib_senders WHERE address=? AND kind='ignore'", (addr,))
    else:
        conn.close()
        raise ValueError('Unknown action')
    conn.commit()
    conn.close()
    return True


# ── 4. Jobber ────────────────────────────────────────────────────────────────

_INV_FIELDS = 'id invoiceNumber subject invoiceStatus issuedDate jobberWebUri amounts { total invoiceBalance } client { name }'


def _jobber(query, variables=None):
    """Read Jobber through the Jobber app connected in Pumps. Changing anything in Jobber is refused."""
    if re.match(r'\s*mutation\b', query):
        raise RuntimeError('Blocked: Vendor Invoices only reads Jobber.')
    force = False
    for _ in range(4):
        token = _cfg['jobber_token'](force)
        r = http_requests.post(JOBBER_GQL, timeout=60, json={'query': query, 'variables': variables or {}},
                               headers={'Authorization': f'Bearer {token}', 'Content-Type': 'application/json',
                                        'X-JOBBER-GRAPHQL-VERSION': _cfg['jobber_version']})
        if r.status_code == 401 and not force:
            force = True
            continue
        if r.status_code == 429:
            threading.Event().wait(20)
            continue
        if r.status_code >= 400:
            raise RuntimeError(f'Jobber API HTTP {r.status_code}')
        out = r.json()
        errors = out.get('errors') or []
        if errors:
            if any((e.get('extensions') or {}).get('code') == 'THROTTLED' for e in errors):
                threading.Event().wait(15)
                continue
            raise RuntimeError(' | '.join(e.get('message', '?') for e in errors))
        return out.get('data') or {}
    raise RuntimeError('Jobber kept refusing the request (signed out or rate limited)')


def _inv_out(n):
    a = n.get('amounts') or {}
    return {'id': n.get('id'), 'number': str(n.get('invoiceNumber') or ''), 'subject': n.get('subject') or '',
            'status': (n.get('invoiceStatus') or '').lower(), 'issued': n.get('issuedDate') or '',
            'url': n.get('jobberWebUri') or '', 'total': a.get('total'), 'balance': a.get('invoiceBalance'),
            'client': (n.get('client') or {}).get('name') or ''}


def jobber_search(term):
    data = _jobber(f'query($q: String!) {{ invoices(searchTerm: $q, first: 25) {{ nodes {{ {_INV_FIELDS} }} }} }}',
                   {'q': term})
    out = [_inv_out(n) for n in ((data.get('invoices') or {}).get('nodes') or []) if n]
    want = re.sub(r'\D', '', term or '')
    out.sort(key=lambda i: (i['number'] != want, i['number']))
    return out


def jobber_invoice(jobber_id):
    data = _jobber(f'query($id: EncodedId!) {{ invoice(id: $id) {{ {_INV_FIELDS} }} }}', {'id': jobber_id})
    n = data.get('invoice')
    return _inv_out(n) if n else None


def suggest_jobber_number(row):
    """'INV 36437/PAY APP 3 - ...' in the notes -> '36437' (only a suggestion; a person picks the invoice)."""
    m = re.search(r'(?i)\bINV(?:OICE)?\.?\s*#?\s*(\d{4,6})\b', row.get('notes') or '')
    return m.group(1) if m else ''


def _is_paid(inv):
    return (inv.get('status') or '') == 'paid'


def link_jobber(row_id, jobber_id, by='', notify=False):
    inv = jobber_invoice(jobber_id)
    if not inv:
        raise ValueError('Jobber has no such invoice')
    paid_at = _stamp() if _is_paid(inv) else ''
    conn = _conn()
    conn.execute('''UPDATE vib_rows SET jobber_id=?, jobber_number=?, jobber_client=?, jobber_subject=?, jobber_total=?,
                      jobber_status=?, jobber_url=?, jobber_linked_by=?, jobber_linked_at=?, jobber_checked_at=?,
                      jobber_paid_at=?, paid_email_at='', updated_at=?, updated_by=? WHERE id=?''',
                 (inv['id'], inv['number'], inv['client'], inv['subject'], inv['total'], inv['status'], inv['url'],
                  by, _stamp(), _stamp(), paid_at, _stamp(), by, row_id))
    conn.commit()
    conn.close()
    _log('jobber', f"{by} linked a row to Jobber invoice #{inv['number']} ({inv['client']}).")
    sent = None
    if paid_at and notify:
        sent = notify_paid(row_id, by=by)
    return {'invoice': inv, 'already_paid': bool(paid_at), 'emailed': sent}


def unlink_jobber(row_id, by=''):
    conn = _conn()
    conn.execute('''UPDATE vib_rows SET jobber_id='', jobber_number='', jobber_client='', jobber_subject='',
                      jobber_total=NULL, jobber_status='', jobber_url='', jobber_linked_by='', jobber_linked_at='',
                      jobber_checked_at='', jobber_paid_at='', paid_email_at='', updated_at=?, updated_by=?
                    WHERE id=?''', (_stamp(), by, row_id))
    conn.commit()
    conn.close()
    return True


def check_jobber_paid():
    """Every linked Jobber invoice not yet paid: is it paid now? If so, email Christian (once)."""
    out = {'checked': 0, 'paid': 0, 'emailed': 0, 'errors': []}
    if not jobber_ready():
        out['errors'].append('Jobber is not connected (it is connected in the Pumps app)')
        return out
    conn = _conn()
    rows = conn.execute("SELECT id, jobber_id FROM vib_rows WHERE jobber_id<>'' AND jobber_paid_at=''").fetchall()
    conn.close()
    for r in rows:
        try:
            inv = jobber_invoice(r['jobber_id'])
        except Exception as e:
            out['errors'].append(str(e))
            continue
        out['checked'] += 1
        if not inv:
            continue
        paid = _is_paid(inv)
        conn = _conn()
        conn.execute("UPDATE vib_rows SET jobber_status=?, jobber_total=?, jobber_checked_at=?, jobber_paid_at=? "
                     "WHERE id=?", (inv['status'], inv['total'], _stamp(), _stamp() if paid else '', r['id']))
        conn.commit()
        conn.close()
        if paid:
            out['paid'] += 1
            ok, _ = notify_paid(r['id'])
            out['emailed'] += 1 if ok else 0
    _set_setting('last_jobber_check', json.dumps({'at': _stamp(), **out, 'errors': out['errors'][:3]}))
    return out


# ── 5. Emails to Christian ───────────────────────────────────────────────────

def send_mail(to, subject, body_html):
    """Send from the PO@ mailbox through Microsoft 365."""
    if not graph_ready():
        raise RuntimeError('Microsoft 365 is not connected, so email cannot be sent')
    token = _cfg['graph_token']()
    r = http_requests.post(f"{GRAPH}/users/{_cfg['mail_from']}/sendMail", timeout=60,
                           headers={'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'},
                           json={'message': {'subject': subject, 'body': {'contentType': 'HTML', 'content': body_html},
                                             'toRecipients': [{'emailAddress': {'address': a}} for a in to]},
                                 'saveToSentItems': True})
    if r.status_code >= 400:
        raise RuntimeError(f'Microsoft 365 said HTTP {r.status_code}: {r.text[:200]}')
    return True


def _money_text(v):
    n = _money(v)
    return '' if n is None else (f'${n:,.2f}' if n >= 0 else f'(${-n:,.2f})')


def _date_text(v):
    d = _parse_date(v) if re.match(r'^\d{4}-\d{2}-\d{2}', v or '') else None
    return d.strftime('%m/%d/%Y') if d else (v or '')


def _table(pairs):
    cells = ''.join(f'<tr><td style="padding:5px 10px;color:#64748b;font-weight:600;width:190px">{html.escape(k)}</td>'
                    f'<td style="padding:5px 10px">{html.escape(str(v or "—"))}</td></tr>' for k, v in pairs)
    return f'<table style="border-collapse:collapse;font-size:14px">{cells}</table>'


def _app_link(row_id=None):
    base = (_cfg.get('website_url') or '').rstrip('/')
    return f"{base}/vendor-invoices" + (f"?row={row_id}" if row_id else '')


def notify_paid(row_id, by='automatic'):
    """Tell Christian the Jobber invoice behind a vendor invoice is paid, so the vendor can be paid."""
    to = christian_email()
    if not to:
        return False, "Christian's email is not set (Settings)"
    conn = _conn()
    r = conn.execute("SELECT r.*, t.label AS tab_label FROM vib_rows r JOIN vib_tabs t ON t.id=r.tab_id WHERE r.id=?",
                     (row_id,)).fetchone()
    conn.close()
    if not r:
        return False, 'No such row'
    if r['paid_email_at'] and by == 'automatic':
        return False, 'Already emailed'
    subject = (f"Paid in Jobber: {r['vendor'] or 'Vendor'} invoice #{r['inv_no'] or '-'}"
               f" (SEI invoice #{r['jobber_number']})")
    jobber_link = (f' &nbsp; <a href="{html.escape(r["jobber_url"])}">Open in Jobber</a>'
                   if r['jobber_url'] else '')
    body = (f'<div style="font-family:Arial,sans-serif;max-width:640px">'
            f'<h2 style="color:#047857;margin:0 0 6px">Jobber invoice #{html.escape(r["jobber_number"])} is paid</h2>'
            f'<p style="font-size:14px">The SEI invoice this vendor invoice was billed on has been paid in Jobber, '
            f'so the vendor can be paid.</p>'
            + _table([('Sub contractor', r['vendor']), ('Vendor invoice #', r['inv_no']),
                      ('Invoice date', _date_text(r['inv_date'])), ('Invoice total', _money_text(r['inv_total'])),
                      ('SEI proposal #', r['proposal']), ('Date billed', _date_text(r['date_billed'])),
                      ('Amount SEI billed', _money_text(r['amt_billed'])), ('Notes', r['notes']),
                      ('Project mgr', r['pm']), ('Jobber invoice', f"#{r['jobber_number']} — {r['jobber_client']}"),
                      ('Jobber total', _money_text(r['jobber_total'])), ('Sheet tab', r['tab_label'])])
            + f'<p><a href="{html.escape(_app_link(row_id))}" style="background:#0E7490;color:#fff;padding:10px 20px;'
              f'border-radius:6px;text-decoration:none;font-weight:700;display:inline-block">Open Vendor Invoices</a>'
              f'{jobber_link}</p>'
              f'<p style="color:#94a3b8;font-size:12px">Sent by The Office App (Vendor Invoices).</p></div>')
    try:
        send_mail([to], subject, body)
    except Exception as e:
        _log('email', f"Could not email Christian that Jobber #{r['jobber_number']} is paid: {e}")
        return False, str(e)
    conn = _conn()
    conn.execute("UPDATE vib_rows SET paid_email_at=? WHERE id=?", (_stamp(), row_id))
    conn.commit()
    conn.close()
    _log('email', f"Emailed {to}: Jobber #{r['jobber_number']} paid ({r['vendor']} #{r['inv_no']})"
                  + (f' — sent by {by}' if by != 'automatic' else '') + '.')
    return True, to


def alert_not_inputted(force=False):
    """Vendor invoices in PO@ nobody has entered on the sheet after a few days: one email to Christian listing
    them, repeated every few days while they are still not entered."""
    to = christian_email()
    days, again = _int_setting('alert_days', 3), _int_setting('realert_days', 7)
    if not to or days <= 0:
        return {'sent': False, 'why': 'no email set' if not to else 'off', 'count': 0}
    today = _today()
    conn = _conn()
    due = []
    for x in conn.execute("SELECT * FROM vib_received WHERE status='received' AND row_id IS NULL "
                          "ORDER BY received_at"):
        try:
            got = datetime.strptime(x['received_at'][:10], '%Y-%m-%d').date()
        except ValueError:
            continue
        if (today - got).days < days:
            continue
        if x['alerted_at'] and not force:
            last = datetime.strptime(x['alerted_at'][:10], '%Y-%m-%d').date()
            if again <= 0 or (today - last).days < again:
                continue
        due.append((x, (today - got).days))
    conn.close()
    if not due:
        return {'sent': False, 'why': 'nothing waiting', 'count': 0}
    rows = ''.join(
        f'<tr><td style="padding:5px 8px;border-bottom:1px solid #e2e8f0">{html.escape(x["vendor"] or x["sender"])}</td>'
        f'<td style="padding:5px 8px;border-bottom:1px solid #e2e8f0">{html.escape(x["inv_no"] or "—")}</td>'
        f'<td style="padding:5px 8px;border-bottom:1px solid #e2e8f0">{html.escape(_date_text(x["inv_date"]) or "—")}</td>'
        f'<td style="padding:5px 8px;border-bottom:1px solid #e2e8f0;text-align:right">{_money_text(x["inv_total"]) or "—"}</td>'
        f'<td style="padding:5px 8px;border-bottom:1px solid #e2e8f0">{html.escape(x["received_at"][:10])}</td>'
        f'<td style="padding:5px 8px;border-bottom:1px solid #e2e8f0;text-align:right">{n}</td></tr>' for x, n in due)
    body = (f'<div style="font-family:Arial,sans-serif;max-width:760px">'
            f'<h2 style="color:#B45309;margin:0 0 6px">{len(due)} vendor invoice{"s" if len(due) != 1 else ""} '
            f'not entered yet</h2><p style="font-size:14px">These came into {html.escape(_cfg.get("mail_from") or "PO@")} '
            f'but nobody has put them on the vendor invoice sheet.</p>'
            f'<table style="border-collapse:collapse;font-size:13px"><tr style="background:#f1f5f9;text-align:left">'
            f'<th style="padding:5px 8px">Sub contractor</th><th style="padding:5px 8px">Invoice #</th>'
            f'<th style="padding:5px 8px">Invoice date</th><th style="padding:5px 8px">Total</th>'
            f'<th style="padding:5px 8px">Received</th><th style="padding:5px 8px">Days waiting</th></tr>{rows}</table>'
            f'<p><a href="{html.escape(_app_link())}" style="background:#0E7490;color:#fff;padding:10px 20px;'
            f'border-radius:6px;text-decoration:none;font-weight:700;display:inline-block">Open Vendor Invoices</a></p>'
            f'<p style="color:#94a3b8;font-size:12px">Sent by The Office App (Vendor Invoices).</p></div>')
    try:
        send_mail([to], f'{len(due)} vendor invoice{"s" if len(due) != 1 else ""} in PO@ not entered yet', body)
    except Exception as e:
        _log('email', f'Could not email Christian about invoices not entered: {e}')
        return {'sent': False, 'why': str(e), 'count': len(due)}
    conn = _conn()
    for x, _ in due:
        conn.execute("UPDATE vib_received SET alerted_at=?, alert_count=alert_count+1 WHERE id=?", (_stamp(), x['id']))
    conn.commit()
    conn.close()
    _log('email', f'Emailed {to}: {len(due)} vendor invoice(s) in PO@ not entered yet.')
    return {'sent': True, 'count': len(due), 'to': to}


def run_cycle():
    """Every hour: read PO@, check linked Jobber invoices for payment, then chase invoices nobody entered.
    Each gunicorn worker schedules this; a file lock lets only one run it at a time."""
    import fcntl
    with open(os.path.join(_cfg.get('data_dir') or '.', '.vendor_invoices_cycle.lock'), 'w') as fh:
        try:
            fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            return
        try:
            for step in (scan_inbox, check_jobber_paid, alert_not_inputted):
                try:
                    step()
                except Exception as e:
                    print(f'⚠ Vendor Invoices {step.__name__}: {e}')
        finally:
            fcntl.flock(fh, fcntl.LOCK_UN)


# ── Routes ───────────────────────────────────────────────────────────────────

def _deny():
    if 'username' not in session:
        return redirect(url_for('login'))
    return redirect(url_for('dashboard'))


def _forbidden():
    return jsonify({'success': False, 'error': 'Office only'}), 403


def _guard():
    return None if vendor_invoices_allowed() else _forbidden()


def _err(e, code=400):
    return jsonify({'success': False, 'error': str(e)}), code


@bp.route('/vendor-invoices')
def page():
    if not vendor_invoices_allowed():
        return _deny()
    return render_template_string(PAGE_TEMPLATE, full_name=_who(), colors=COLORS, cols=COLS)


@bp.route('/vendor-invoices/api/state')
def api_state():
    if _guard():
        return _guard()
    def _j(key):
        try:
            return json.loads(_setting(key) or 'null')
        except ValueError:
            return None
    return jsonify({'success': True, 'tabs': tabs_summary(), 'inbox': inbox(),
                    'status': {'graph': graph_ready(), 'jobber': jobber_ready(), 'christian': christian_email(),
                               'mailbox': _cfg.get('mail_from') or '', 'scan_ready': bool(_cfg.get('fetch_emails')),
                               'last_scan': _j('last_scan'), 'last_jobber_check': _j('last_jobber_check')}})


@bp.route('/vendor-invoices/api/tab/<int:tab_id>')
def api_tab(tab_id):
    if _guard():
        return _guard()
    rows = tab_rows(tab_id)
    for r in rows:
        r['suggest'] = suggest_jobber_number(r) if not r['jobber'] else ''
    return jsonify({'success': True, 'rows': rows})


@bp.route('/vendor-invoices/api/find-row/<int:row_id>')
def api_find_row(row_id):
    if _guard():
        return _guard()
    conn = _conn()
    r = conn.execute("SELECT tab_id FROM vib_rows WHERE id=?", (row_id,)).fetchone()
    conn.close()
    return jsonify({'success': bool(r), 'tab_id': r['tab_id'] if r else None})


@bp.route('/vendor-invoices/api/search')
def api_search():
    if _guard():
        return _guard()
    return jsonify({'success': True, 'rows': search_rows(request.args.get('q', ''))})


@bp.route('/vendor-invoices/api/cell', methods=['POST'])
def api_cell():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    try:
        return jsonify({'success': True, 'row': set_cell(int(d.get('id')), d.get('key'), d.get('value'), _who())})
    except (TypeError, ValueError) as e:
        return _err(e)


@bp.route('/vendor-invoices/api/style', methods=['POST'])
def api_style():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    keys = d.get('keys') or ([d['key']] if d.get('key') else [])
    try:
        return jsonify({'success': True, 'row': set_style(int(d.get('id')), keys, bg=d.get('bg'), fg=d.get('fg'),
                                                          bold=d.get('bold'), by=_who())})
    except (TypeError, ValueError) as e:
        return _err(e)


@bp.route('/vendor-invoices/api/row/add', methods=['POST'])
def api_row_add():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    try:
        return jsonify({'success': True, 'row': add_row(int(d.get('tab_id')), d.get('after_id'), d.get('values'),
                                                        by=_who())})
    except (TypeError, ValueError) as e:
        return _err(e)


@bp.route('/vendor-invoices/api/row/delete', methods=['POST'])
def api_row_delete():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    return jsonify({'success': delete_row(int(d.get('id') or 0), by=_who())})


@bp.route('/vendor-invoices/api/tab/add', methods=['POST'])
def api_tab_add():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    tid, made = add_tab(d.get('label', ''))
    return jsonify({'success': True, 'tab_id': tid, 'made': made})


@bp.route('/vendor-invoices/api/inbox/<int:rec_id>/enter', methods=['POST'])
def api_inbox_enter(rec_id):
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    try:
        return jsonify({'success': True, 'row': enter_received(rec_id, d.get('tab_id'), by=_who())})
    except ValueError as e:
        return _err(e)


@bp.route('/vendor-invoices/api/inbox/<int:rec_id>/decide', methods=['POST'])
def api_inbox_decide(rec_id):
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    try:
        return jsonify({'success': decide_received(rec_id, d.get('action'), by=_who())})
    except ValueError as e:
        return _err(e)


@bp.route('/vendor-invoices/api/inbox/<int:rec_id>/file')
def api_inbox_file(rec_id):
    if not vendor_invoices_allowed():
        return _deny()
    conn = _conn()
    x = conn.execute("SELECT filename, file_path FROM vib_received WHERE id=?", (rec_id,)).fetchone()
    conn.close()
    if not x or not x['file_path'] or not os.path.exists(x['file_path']):
        return 'Not found', 404
    return send_file(x['file_path'], mimetype='application/pdf', download_name=x['filename'] or 'invoice.pdf')


@bp.route('/vendor-invoices/api/scan', methods=['POST'])
def api_scan():
    if _guard():
        return _guard()
    return jsonify({'success': True, 'summary': scan_inbox()})


@bp.route('/vendor-invoices/api/alert', methods=['POST'])
def api_alert():
    if _guard():
        return _guard()
    return jsonify({'success': True, 'result': alert_not_inputted(force=True)})


@bp.route('/vendor-invoices/api/jobber/search')
def api_jobber_search():
    if _guard():
        return _guard()
    if not jobber_ready():
        return _err('Jobber is not connected (connect it in the Pumps app)')
    try:
        return jsonify({'success': True, 'invoices': jobber_search(request.args.get('q', '').strip())})
    except Exception as e:
        return _err(e, 502)


@bp.route('/vendor-invoices/api/jobber/link', methods=['POST'])
def api_jobber_link():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    try:
        out = link_jobber(int(d.get('id')), str(d.get('jobber_id') or ''), by=_who(), notify=bool(d.get('notify')))
    except (TypeError, ValueError) as e:
        return _err(e)
    except Exception as e:
        return _err(e, 502)
    conn = _conn()
    r = conn.execute("SELECT * FROM vib_rows WHERE id=?", (int(d.get('id')),)).fetchone()
    conn.close()
    return jsonify({'success': True, **out, 'row': _row_out(r)})


@bp.route('/vendor-invoices/api/jobber/unlink', methods=['POST'])
def api_jobber_unlink():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    unlink_jobber(int(d.get('id') or 0), by=_who())
    return jsonify({'success': True})


@bp.route('/vendor-invoices/api/jobber/check', methods=['POST'])
def api_jobber_check():
    if _guard():
        return _guard()
    return jsonify({'success': True, 'result': check_jobber_paid()})


@bp.route('/vendor-invoices/api/notify-paid', methods=['POST'])
def api_notify_paid():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    ok, info = notify_paid(int(d.get('id') or 0), by=_who())
    return jsonify({'success': ok, 'error': '' if ok else info, 'to': info if ok else ''})


@bp.route('/vendor-invoices/api/settings', methods=['GET', 'POST'])
def api_settings():
    if _guard():
        return _guard()
    if request.method == 'POST':
        d = request.get_json(silent=True) or {}
        if 'christian_email' in d:
            ce = (d['christian_email'] or '').strip()
            if ce and not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', ce):
                return _err('That email address does not look right')
            if _cfg.get('set_christian_email'):
                _cfg['set_christian_email'](ce)
        for key in ('alert_days', 'realert_days'):
            if key in d:
                try:
                    _set_setting(key, max(0, int(d[key])))
                except (TypeError, ValueError):
                    return _err(f'{key} must be a number of days')
        if 'pump_words' in d:
            _set_setting('pump_words', ', '.join(w.strip() for w in str(d['pump_words']).split(',') if w.strip()))
        if 'scan_since' in d:
            sd = _parse_date(d['scan_since'])
            if not sd:
                return _err('Pick a date to read PO@ from')
            _set_setting('scan_since', sd.isoformat())
    conn = _conn()
    senders = [dict(r) for r in conn.execute("SELECT * FROM vib_senders ORDER BY kind, address")]
    log = [dict(r) for r in conn.execute("SELECT * FROM vib_log ORDER BY id DESC LIMIT 60")]
    conn.close()
    return jsonify({'success': True, 'christian_email': christian_email(),
                    'alert_days': _int_setting('alert_days', 3), 'realert_days': _int_setting('realert_days', 7),
                    'pump_words': _setting('pump_words', DEFAULT_PUMP_WORDS), 'scan_since': _setting('scan_since'),
                    'senders': senders, 'log': log})


@bp.route('/vendor-invoices/api/senders/delete', methods=['POST'])
def api_sender_delete():
    if _guard():
        return _guard()
    d = request.get_json(silent=True) or {}
    conn = _conn()
    conn.execute("DELETE FROM vib_senders WHERE address=?", ((d.get('address') or '').lower(),))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@bp.route('/vendor-invoices/api/import', methods=['POST'])
def api_import():
    if _guard():
        return _guard()
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm')):
        return _err('Choose the vendor invoices workbook (.xlsx)')
    try:
        tabs = parse_workbook(f.read())
    except Exception as e:
        return _err(f'Could not read that workbook: {e}')
    if not tabs:
        return _err('No vendor invoice tabs in that workbook (row 1 should have SUB CONTRACTOR, INVOICE # ...)')
    out = import_tabs(tabs, by=_who(), replace=request.form.get('replace') == '1')
    return jsonify({'success': True, **out})


@bp.route('/vendor-invoices/api/export')
def api_export():
    if not vendor_invoices_allowed():
        return _deny()
    return send_file(export_workbook(), as_attachment=True,
                     download_name=f'VENDOR INVOICES {_today().isoformat()}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── Hook-up ──────────────────────────────────────────────────────────────────

def init_vendor_invoices(app, db_path, *, data_dir, get_christian_email, set_christian_email, graph_token=None,
                         graph_enabled=False, mail_from='', fetch_emails=None, graph_attachments=None,
                         email_attachments=None, extract_fields=None, jobber_token=None, jobber_connected=None,
                         jobber_version='', scheduler_cls=None, website_url='', seed_file=SEED_FILE):
    """Create the tables, bring in the workbook the first time, register the routes and run the cycle hourly.
    fetch_emails(log_table=, extensions=, since=) is shared_mail.fetch_scan_emails for this database;
    extract_fields(text, sender) reads an invoice PDF's text into sub_contractor / invoice_number / ...;
    jobber_token(force_refresh) hands over a token from the Jobber app connected for Pumps."""
    _cfg.update(db_path=db_path, data_dir=data_dir, get_christian_email=get_christian_email,
                set_christian_email=set_christian_email, graph_token=graph_token, graph_enabled=graph_enabled,
                mail_from=mail_from, fetch_emails=fetch_emails, graph_attachments=graph_attachments,
                email_attachments=email_attachments, extract_fields=extract_fields, jobber_token=jobber_token,
                jobber_connected=jobber_connected or (lambda: False), jobber_version=jobber_version,
                website_url=website_url, seed_file=seed_file)
    init_db()
    try:
        seed_from_workbook()
    except Exception as e:
        print(f'⚠ Vendor Invoices: could not import the workbook: {e}')
    app.register_blueprint(bp)
    if scheduler_cls and os.environ.get('VI_AUTO_RUN', 'true').lower() not in ('0', 'false', 'no', 'off'):
        try:
            sched = scheduler_cls()
            sched.add_job(run_cycle, 'interval', hours=1, id='vendor_invoices_cycle',
                          next_run_time=datetime.now() + timedelta(minutes=6))
            sched.start()
            print('✓ Vendor Invoices: PO@ scan, Jobber paid check and reminders scheduled every hour')
        except Exception as e:
            print(f'⚠ Could not start the Vendor Invoices scheduler: {e}')


from vendor_invoices_page import PAGE_TEMPLATE  # noqa: E402
